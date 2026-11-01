#!/usr/bin/env python3
r"""Capture the English web app from the B2B commercial review showcase workbook.

Unlike capture_viewer.py, this exercises the running, lazy web application.
No model is called unless --generate-doc is explicitly supplied. An existing
English node document is reused when available; otherwise the screenshots show
deterministic calculation. --no-ai also excludes existing documents. --prepare
saves a reusable browser context and project evidence without writing images.

Examples (use a dedicated preview server, never a user's active workspace):
    uv run --extra screenshots python scripts/capture_app.py \
      --base-url http://127.0.0.1:8877/review/ \
      --workbook validation_screenshots/run/fy26_commercial_review.xlsx --no-ai
    uv run --extra screenshots python scripts/capture_app.py \
      --base-url http://127.0.0.1:8877/review/ \
      --workbook validation_screenshots/run/fy26_commercial_review.xlsx \
      --generate-doc --prepare --save-state validation_screenshots/app/state.json
    uv run --extra screenshots python scripts/capture_app.py \
      --base-url http://127.0.0.1:8877/review/ \
      --workbook validation_screenshots/run/fy26_commercial_review.xlsx \
      --reuse-project PROJECT_ID --storage-state validation_screenshots/app/state.json
    python -S scripts/capture_app.py --check

--check uses only the Python standard library and never contacts a server.
The manifest contains hashes and synthetic-fixture provenance, not browser
cookies, source workbook bytes, or generated document text.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import re
import struct
import sys
import time
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = ROOT / "imgs"
MANIFEST_NAME = "app_manifest.json"
SCHEMA = 2
WIDTH, HEIGHT, SCALE = 1440, 900, 2
NODE = "Summary!C10"
PROJECT_NAME = "FY26 B2B Commercial Review"
SOURCES = (
    "src/linexcel/assets/app.html",
    "src/linexcel/assets/cytoscape.min.js",
    "validation_workbooks.py",
    "scripts/capture_app.py",
)
SHOTS = {
    "app_graph.png": "Interactive dependency graph with deterministic calculation and saved-value comparison",
    "app_graph_documented.png": "Wide graph reader with a separate English AI documentation tab",
    "app_node_documented.png": "English node documentation in the full detail panel, with separate input and output token counts",
    "app_captures.png": "Sheet capture gallery and the selected full-resolution image",
    "app_import.png": "Accessible main-workbook and reference-file selection",
}
DETERMINISTIC_SHOTS = {
    **SHOTS,
    # Keep published image URLs stable; their contents and captions describe
    # the selected mode, not the historical filename.
    "app_graph_documented.png": "Wide graph reader with deterministic dependency values; AI documentation skipped",
    "app_node_documented.png": "Full cell detail with source formula, saved value and targeted calculation; AI documentation skipped",
}


def shot_captions(documented: bool) -> dict[str, str]:
    return SHOTS if documented else DETERMINISTIC_SHOTS


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source_digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes().replace(b"\r\n", b"\n")).hexdigest()


def source_hashes() -> dict[str, str]:
    return {name: source_digest(ROOT / name) for name in SOURCES}


def png_dimensions(path: Path) -> tuple[int, int]:
    with path.open("rb") as stream:
        header = stream.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n" or header[12:16] != b"IHDR":
        raise ValueError(f"Not a PNG image: {path}")
    return struct.unpack(">II", header[16:24])


def check(output: Path, workbook: Path | None = None) -> int:
    """Check evidence and image hashes without Playwright or project imports."""
    try:
        manifest = json.loads((output / MANIFEST_NAME).read_text(encoding="utf-8"))
        problems = []
        if manifest.get("schema") != SCHEMA:
            problems.append("unsupported or missing manifest schema")
        for name, expected in source_hashes().items():
            if manifest.get("sources", {}).get(name) != expected:
                problems.append(f"changed source: {name}")
        if manifest.get("uiLanguage") != "en":
            problems.append("screenshots are not recorded as English")
        if manifest.get("viewport") != {
            "width": WIDTH,
            "height": HEIGHT,
            "scale": SCALE,
        }:
            problems.append("unexpected viewport or pixel density")
        fixture = manifest.get("fixture", {})
        if fixture.get("name") != "fy26_commercial_review.xlsx":
            problems.append("source is not the current commercial review fixture")
        if fixture.get("generator") != "validation_workbooks.build_sales_workbook":
            problems.append("source is not the synthetic Sales fixture")
        if fixture.get("generatorSha256") != source_digest(
            ROOT / "validation_workbooks.py"
        ):
            problems.append("synthetic fixture generator changed")
        if not re.fullmatch(r"[0-9a-f]{64}", fixture.get("sha256", "")):
            problems.append("missing fixture hash")
        if workbook is not None and digest(workbook) != fixture.get("sha256"):
            problems.append("supplied workbook differs from the captured fixture")
        documentation = manifest.get("documentation", {})
        documented = documentation.get("status") == "completed"
        doc = manifest.get("document")
        if documented:
            doc = doc or {}
            if (
                manifest.get("documentLanguage") != "en"
                or doc.get("language") != "en"
                or doc.get("nodeId") != NODE
                or not doc.get("taskId")
                or not doc.get("model")
            ):
                problems.append("missing English document provenance")
            if doc.get("providerRequests") != 2:
                problems.append("missing draft/review provider-call evidence")
            if not re.fullmatch(r"[0-9a-f]{64}", doc.get("markdownSha256", "")):
                problems.append("missing document content hash")
            for field in ("inputTokens", "outputTokens"):
                value = doc.get("usage", {}).get(field)
                if (
                    not isinstance(value, (int, float))
                    or isinstance(value, bool)
                    or not math.isfinite(value)
                    or value < 0
                ):
                    problems.append(f"missing document usage: {field}")
        elif (
            documentation.get("status") != "skipped"
            or documentation.get("reason")
            not in ("--no-ai", "no existing English document")
            or documentation.get("requested") is not False
            or doc is not None
            or manifest.get("documentLanguage") is not None
            or manifest.get("modelRequestsDuringCapture") != 0
        ):
            problems.append("missing or inconsistent skipped documentation evidence")
        if type(documentation.get("requested")) is not bool:
            problems.append("missing documentation request selection")
        requests = manifest.get("modelRequestsDuringCapture")
        if type(requests) is not int or requests not in (0, 1):
            problems.append("unexpected model request count")
        elif requests and not documentation.get("requested"):
            problems.append("model request without explicit documentation selection")
        if not manifest.get("capture", {}).get("taskId"):
            problems.append("missing sheet-capture provenance")
        if manifest.get("capture", {}).get("imageCount") != 4:
            problems.append("missing captures for the four showcase sheets")
        expected_shots = shot_captions(documented)
        if set(manifest.get("images", {})) != set(expected_shots):
            problems.append("the five expected app screenshots are not recorded")
        for name, caption in expected_shots.items():
            path = output / name
            record = manifest.get("images", {}).get(name, {})
            if record.get("caption") != caption:
                problems.append(
                    f"image caption does not match the documentation mode: {name}"
                )
            if not path.is_file():
                problems.append(f"missing image: {name}")
                continue
            if digest(path) != record.get("sha256"):
                problems.append(f"image content changed: {name}")
            actual = png_dimensions(path)
            if actual != (WIDTH * SCALE, HEIGHT * SCALE):
                problems.append(f"unexpected PNG dimensions: {name}: {actual}")
            if list(actual) != record.get("pixels"):
                problems.append(f"image dimensions do not match the manifest: {name}")
        if problems:
            print("App screenshots need recapturing:\n- " + "\n- ".join(problems))
            return 1
    except (OSError, ValueError, TypeError, KeyError) as error:
        print(f"Cannot verify app screenshots: {error}")
        return 1
    print(
        "App screenshots match their English UI, synthetic fixture and source hashes."
    )
    return 0


def normalized_formula(value: object) -> object:
    if not isinstance(value, str) or not value.startswith("="):
        return value
    # LibreOffice may add these interoperability prefixes or quote sheet names.
    value = value.replace("_xlfn.", "").replace("_xlws.", "")
    return re.sub(r"'([A-Za-z_][A-Za-z_0-9]*)'!", r"\1!", value)


def verify_fixture(workbook: Path) -> dict:
    """Compare every nonempty cell with the generator; allow recalculated caches."""
    from openpyxl import load_workbook

    sys.path.insert(0, str(ROOT))
    from validation_workbooks import build_sales_workbook

    if workbook.name != "fy26_commercial_review.xlsx":
        raise ValueError(
            "Only fy26_commercial_review.xlsx from build_sales_workbook is accepted."
        )
    with zipfile.ZipFile(workbook) as archive:
        if any(name.startswith("xl/media/") for name in archive.namelist()):
            raise ValueError("The synthetic fixture must not contain embedded images.")
    expected = load_workbook(io.BytesIO(build_sales_workbook()), data_only=False)
    actual = load_workbook(workbook, data_only=False)
    try:
        if actual.sheetnames != expected.sheetnames:
            raise ValueError("Worksheet names differ from the synthetic Sales fixture.")
        for sheet in expected.sheetnames:

            def cells(book, sheet=sheet):
                return {
                    cell.coordinate: normalized_formula(cell.value)
                    for row in book[sheet]
                    for cell in row
                    if cell.value is not None
                }

            wanted, found = cells(expected), cells(actual)
            if wanted != found:
                changed = sorted(
                    key
                    for key in wanted.keys() | found.keys()
                    if wanted.get(key) != found.get(key)
                )
                raise ValueError(
                    f"Not the unmodified synthetic Sales fixture: {sheet}: {changed[:8]}"
                )
    finally:
        expected.close()
        actual.close()
    return {
        "name": workbook.name,
        "sha256": digest(workbook),
        "generator": "validation_workbooks.build_sales_workbook",
        "generatorSha256": source_digest(ROOT / "validation_workbooks.py"),
        "verification": "All nonempty cells and formulas match the synthetic generator; recalculated caches are allowed.",
    }


def private_state_path(path: Path) -> Path:
    """Browser cookies must never be written among repository screenshot assets."""
    path = path.resolve()
    if path.is_relative_to(ROOT) and not path.is_relative_to(
        ROOT / "validation_screenshots"
    ):
        raise ValueError(
            "Store browser state outside the repository or in ignored validation_screenshots/."
        )
    return path


def api(page, base: str, path: str) -> dict:
    response = page.request.get(base + "api/" + path)
    if not response.ok:
        raise RuntimeError(f"API request failed ({response.status}): {path}")
    return response.json()


def wait_tasks(page, base: str, project: str, timeout: int) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        data = api(page, base, "projects/" + project)
        if not any(t["status"] in ("queued", "running") for t in data.get("tasks", [])):
            # Project snapshots intentionally contain only task metadata. Load
            # the two types of evidence needed for these screenshots explicitly.
            data["tasks"] = [
                api(page, base, "tasks/" + task["id"])["task"]
                if task.get("resultAvailable")
                and task.get("result") is None
                and task.get("operation") in ("document", "capture")
                else task
                for task in data.get("tasks", [])
            ]
            return data
        page.wait_for_timeout(400)
    raise RuntimeError("Timed out waiting for the selected project's tasks.")


def task_time(task: dict) -> float:
    return task.get("finishedAt") or task.get("createdAt") or 0


def english_document(data: dict) -> dict | None:
    docs = sorted(
        (
            task
            for task in data.get("tasks", [])
            if task.get("operation") == "document"
            and task.get("nodeId") == NODE
            and task.get("status") == "succeeded"
            and (task.get("result") or {}).get("language") == "en"
        ),
        key=task_time,
        reverse=True,
    )
    if not docs:
        return None
    task = docs[0]
    result = task["result"]
    if not str(result.get("markdown", "")).strip() or not result.get("model"):
        raise ValueError(
            "The English node document is empty or has no model provenance."
        )
    for key in ("inputTokens", "outputTokens"):
        value = result.get("usage", {}).get(key)
        if (
            not isinstance(value, (int, float))
            or isinstance(value, bool)
            or not math.isfinite(value)
            or value < 0
        ):
            raise ValueError(f"The document has no usable separate {key} count.")
    return task


def require_english(page, *, documentation: bool = False) -> None:
    if page.locator("html").get_attribute("lang") != "en":
        raise ValueError("The application did not activate its real English UI.")
    if page.locator("#ui-language").input_value() != "en":
        raise ValueError("English UI selection is missing.")
    if documentation and page.locator("#ai-language").input_value() != "en":
        raise ValueError(
            "Documentation language must be English independently of the UI."
        )


def open_workspace_menu(page) -> None:
    menu = page.locator("#workspace-menu")
    if menu.is_visible() and menu.get_attribute("aria-expanded") != "true":
        menu.click()


def close_workspace_menu(page) -> None:
    close = page.locator("#workspace-close")
    if close.is_visible():
        close.click()


def set_document_language(page) -> None:
    open_workspace_menu(page)
    setting = page.locator("#ai-language")
    if not setting.is_visible():
        page.get_by_text("On-demand AI documentation", exact=True).click()
    setting.select_option("en")
    close_workspace_menu(page)


def document_evidence(document: dict | None, args) -> dict:
    if document is None:
        if args.generate_doc:
            raise ValueError("Requested documentation did not complete.")
        return {
            "documentation": {
                "status": "skipped",
                "requested": False,
                "reason": "--no-ai" if args.no_ai else "no existing English document",
            },
            "documentLanguage": None,
            "document": None,
        }
    result = document["result"]
    return {
        "documentation": {"status": "completed", "requested": args.generate_doc},
        "documentLanguage": "en",
        "document": {
            "taskId": document["id"],
            "nodeId": NODE,
            "language": result["language"],
            "model": result["model"],
            "usage": result["usage"],
            "providerRequests": sum(
                stage.get("usage", {}).get("requests", 0)
                for stage in result["validation"]["specialized_review"]["stages"]
            ),
            "markdownSha256": hashlib.sha256(result["markdown"].encode()).hexdigest(),
        },
    }


def guard_model_request(route, *, allowed: bool, requests: list) -> None:
    request = route.request
    if request.method == "POST" and urlsplit(request.url).path.endswith("/tasks"):
        body = request.post_data_json
        if body.get("operation") in ("document", "describe_capture"):
            requests.append(body)
            if not allowed or body.get("operation") != "document" or len(requests) > 1:
                route.abort()
                return
    route.continue_()


def check_model_requests(requests: list, *, allowed: bool) -> None:
    if requests and (
        not allowed or len(requests) != 1 or requests[0].get("operation") != "document"
    ):
        raise ValueError(
            "Unexpected model request(s); only one explicit English node request is allowed."
        )


def frame(page, selector: str, offset: int = 0) -> None:
    page.locator(selector).evaluate(
        "(element, offset) => window.scrollTo(0, element.getBoundingClientRect().top + scrollY - offset)",
        offset,
    )
    page.wait_for_timeout(150)
    if page.evaluate("document.documentElement.scrollWidth > innerWidth"):
        raise ValueError("The screenshot contains horizontal page overflow.")


def fit_graph(page) -> None:
    """Show the complete loaded neighborhood before photographing its cards."""
    page.locator("#graph-fit").click()
    page.wait_for_function(
        "() => !document.querySelector('#graph-canvas')._cy.animated()"
    )
    page.evaluate(
        "() => new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)))"
    )


def capture(args) -> int:
    from playwright.sync_api import expect, sync_playwright

    workbook = args.workbook.resolve()
    fixture = verify_fixture(workbook)
    watched = source_hashes()
    output = args.output.resolve()
    state_out = private_state_path(args.save_state) if args.save_state else None
    state_in = private_state_path(args.storage_state) if args.storage_state else None
    parts = urlsplit(args.base_url)
    if (
        parts.scheme not in ("http", "https")
        or parts.username
        or parts.password
        or parts.query
        or parts.fragment
    ):
        raise ValueError(
            "--base-url must be an HTTP(S) app root without credentials, query or fragment."
        )
    base = urlunsplit(
        (parts.scheme, parts.netloc, parts.path.rstrip("/") + "/", "", "")
    )
    images = {}
    with sync_playwright() as driver:
        browser = driver.chromium.launch(
            channel=os.environ.get("LINEXCEL_BROWSER_CHANNEL") or None
        )
        context = browser.new_context(
            viewport={"width": WIDTH, "height": HEIGHT},
            device_scale_factor=SCALE,
            storage_state=str(state_in) if state_in else None,
            accept_downloads=True,
        )
        page = context.new_page()
        errors, model_requests = [], []
        page.on("pageerror", lambda error: errors.append(str(error)))

        context.route(
            "**/api/**",
            lambda route: guard_model_request(
                route, allowed=args.generate_doc, requests=model_requests
            ),
        )
        try:
            page.goto(base + "?lang=en")
            expect(page.locator("#connection")).to_have_text("Server connected")
            projects = api(page, base, "projects").get("projects", [])
            open_workspace_menu(page)
            if args.reuse_project:
                if any(p["id"] != args.reuse_project for p in projects):
                    raise ValueError(
                        "Use a dedicated browser context containing only the synthetic screenshot project."
                    )
                page.locator(f'[data-project="{args.reuse_project}"]').click()
                project = args.reuse_project
            else:
                if projects:
                    raise ValueError(
                        "Use a fresh context or explicitly reuse the single synthetic project."
                    )
                page.locator("#workbook").set_input_files(workbook)
                page.locator("#project-name").fill(PROJECT_NAME)
                page.locator("#import-button").click()
                expect(page.locator("#project-title")).to_have_text(PROJECT_NAME)
                project = page.evaluate("localStorage.getItem('linexcel-project')")
            data = wait_tasks(page, base, project, args.timeout)
            if data.get("graph", {}).get("meta", {}).get("sha256") != fixture["sha256"]:
                raise ValueError(
                    "The project was imported from different workbook bytes. Refusing unrelated documentation."
                )
            if data["project"].get("name") != PROJECT_NAME:
                raise ValueError(
                    f"The synthetic project must be named {PROJECT_NAME!r}."
                )
            page.reload()
            expect(page.locator("#project-title")).to_have_text(PROJECT_NAME)
            close_workspace_menu(page)
            if not args.no_ai:
                set_document_language(page)
            require_english(page)
            page.locator("#tab-graph").click()
            page.locator("#graph-search").fill(NODE)
            with page.expect_response(
                lambda response: (
                    response.request.method == "POST"
                    and response.url.endswith("/tasks")
                )
            ):
                page.locator(f'#graph-results [data-related="{NODE}"]').click()
            data = wait_tasks(page, base, project, args.timeout)
            document = None if args.no_ai else english_document(data)
            if document is None and args.generate_doc:
                with page.expect_response(
                    lambda response: (
                        response.request.method == "POST"
                        and response.url.endswith("/tasks")
                    )
                ):
                    page.locator("#graph-tab-ai").click()
                    page.locator("#graph-ai [data-ai-start]").click()
                data = wait_tasks(page, base, project, args.timeout)
                document = english_document(data)
                if document is None:
                    raise ValueError(
                        "English documentation failed or was cancelled; no screenshot will pretend it succeeded."
                    )
            check_model_requests(model_requests, allowed=args.generate_doc)
            captures = sorted(
                (
                    t
                    for t in data["tasks"]
                    if t["operation"] == "capture"
                    and t["status"] == "succeeded"
                    and (t.get("result") or {}).get("screenshots")
                ),
                key=task_time,
            )
            if not captures:
                page.locator("#tab-captures").click()
                with page.expect_response(
                    lambda response: (
                        response.request.method == "POST"
                        and response.url.endswith("/tasks")
                    )
                ):
                    page.locator("#capture").click()
                data = wait_tasks(page, base, project, args.timeout)
                captures = sorted(
                    (
                        t
                        for t in data["tasks"]
                        if t["operation"] == "capture"
                        and t["status"] == "succeeded"
                        and (t.get("result") or {}).get("screenshots")
                    ),
                    key=task_time,
                )
            if not captures:
                raise ValueError(
                    "No successful sheet images exist; capture screenshots cannot be substituted."
                )
            captured = captures[-1]
            check_model_requests(model_requests, allowed=args.generate_doc)
            if state_out:
                state_out.parent.mkdir(parents=True, exist_ok=True)
                context.storage_state(path=str(state_out))
                evidence = {
                    "projectId": project,
                    "fixture": fixture,
                    "documentTaskId": document["id"] if document else None,
                    "documentation": document_evidence(document, args)["documentation"],
                    "captureTaskId": captured["id"],
                    "baseUrl": base,
                    "modelRequests": len(model_requests),
                }
                state_out.with_suffix(".project.json").write_text(
                    json.dumps(evidence, indent=2) + "\n", encoding="utf-8"
                )
            if (
                errors
                or source_hashes() != watched
                or digest(workbook) != fixture["sha256"]
            ):
                raise ValueError(
                    "Browser errors or source changes occurred during preparation."
                )
            if args.prepare:
                status = (
                    f"English document {document['id']}"
                    if document
                    else "AI documentation skipped"
                )
                print(
                    f"Prepared synthetic project {project}; {status}; no repository images written."
                )
                return 0
            page.reload()
            expect(page.locator("#project-title")).to_have_text(PROJECT_NAME)
            close_workspace_menu(page)
            if document:
                set_document_language(page)
            page.locator("#tab-graph").click()
            page.locator("#graph-search").fill(NODE)
            with page.expect_response(
                lambda response: (
                    response.request.method == "POST"
                    and response.url.endswith("/tasks")
                )
            ):
                page.locator(f'#graph-results [data-related="{NODE}"]').click()
            page.locator("#graph-panel-calculation").wait_for()
            wait_tasks(page, base, project, args.timeout)
            expect(page.locator("#graph-recalculate")).to_be_enabled()
            expect(
                page.locator("#graph-preview .value-pair .value-card").nth(1)
            ).to_contain_text("816668.1")
            require_english(page, documentation=bool(document))
            output.mkdir(parents=True, exist_ok=True)
            captions = shot_captions(bool(document))

            def shot(name: str, selector: str, offset: int = 0):
                frame(page, selector, offset)
                path = output / name
                page.screenshot(path=str(path))
                images[name] = {
                    "caption": captions[name],
                    "sha256": digest(path),
                    "pixels": list(png_dimensions(path)),
                }

            fit_graph(page)
            shot("app_graph.png", ".graph-investigation", 70)
            if document:
                page.locator("#graph-tab-ai").click()
                page.locator("#graph-ai .ai-result").wait_for()
                page.locator("#graph-ai .token-usage").wait_for()
            page.locator("#graph-options > summary").click()
            page.locator("#graph-reader-width").focus()
            page.keyboard.press("End")
            expect(page.locator("#graph-reader-width")).to_have_value("62")
            page.locator("#graph-options > summary").click()
            if document:
                page.locator("#graph-panel-ai").evaluate(
                    """panel => {
                        const heading = panel.querySelector('.ai-result h3');
                        if (!heading) throw new Error('Missing rendered document heading');
                        panel.scrollTop += heading.getBoundingClientRect().top
                            - panel.getBoundingClientRect().top - 16;
                    }"""
                )
            else:
                page.locator("#graph-panel-calculation").evaluate(
                    """panel => {
                        const heading = panel.querySelector('.detail-section h3');
                        panel.scrollTop += heading.getBoundingClientRect().top
                            - panel.getBoundingClientRect().top - 16;
                    }"""
                )
            fit_graph(page)
            shot("app_graph_documented.png", ".graph-investigation", 70)
            # Opening the same cell's full reader can reuse the active result.
            # Verify the visible state instead of requiring an extra operation.
            page.locator("#graph-full-detail").click()
            # Lazy neighborhood/result requests can remount the cell details.
            # Wait before framing; otherwise scroll anchoring moves the shot.
            page.wait_for_function(
                "() => document.querySelector('#recalculate')"
                " && !document.querySelector('#recalculate').disabled"
            )
            if document:
                page.locator("#node-ai .token-usage").wait_for()
                shot("app_node_documented.png", "#node-ai", 80)
                usage = page.locator("#node-ai .token-usage").bounding_box()
                if not usage or usage["y"] < 0 or usage["y"] + usage["height"] > HEIGHT:
                    raise ValueError(
                        "Input/output token counts are not visible in the documentation screenshot."
                    )
            else:
                expect(page.locator("#node-detail .value-pair")).to_contain_text(
                    "816668.1"
                )
                shot("app_node_documented.png", "#node-detail", 40)
            page.locator("#tab-captures").click()
            page.wait_for_function(
                "() => document.querySelector('#capture-selected figure img')"
                "?.naturalWidth > 0"
            )
            expect(page.locator(".capture-thumb")).to_have_count(
                min(80, len(captured["result"]["screenshots"]))
            )
            shot("app_captures.png", "#workspace-content")
            # A fresh tab can show a real no-selection workspace without rewriting UI text.
            state = context.storage_state()
            for origin in state.get("origins", []):
                origin["localStorage"] = [
                    entry
                    for entry in origin.get("localStorage", [])
                    if entry["name"] != "linexcel-project"
                ]
            import_context = browser.new_context(
                viewport={"width": WIDTH, "height": HEIGHT},
                device_scale_factor=SCALE,
                storage_state=state,
            )
            import_context.route(
                "**/api/**",
                lambda route: guard_model_request(
                    route, allowed=False, requests=model_requests
                ),
            )
            import_page = import_context.new_page()
            import_page.on("pageerror", lambda error: errors.append(str(error)))
            import_page.goto(base + "?lang=en")
            expect(import_page.locator("#welcome")).to_be_visible()
            open_workspace_menu(import_page)
            import_page.locator("#workbook").set_input_files(workbook)
            import_page.locator("#project-name").fill(PROJECT_NAME)
            import_page.locator("#project-name").blur()
            image_path = output / "app_import.png"
            import_page.screenshot(path=str(image_path))
            images["app_import.png"] = {
                "caption": captions["app_import.png"],
                "sha256": digest(image_path),
                "pixels": list(png_dimensions(image_path)),
            }
            import_context.close()
            check_model_requests(model_requests, allowed=args.generate_doc)
            if errors:
                raise ValueError(f"Browser errors occurred: {errors}")
            if source_hashes() != watched or digest(workbook) != fixture["sha256"]:
                raise ValueError(
                    "Sources or workbook changed during capture; retry after freezing the tree."
                )
            manifest = {
                "schema": SCHEMA,
                "capturedAt": datetime.now(UTC).isoformat(),
                "uiLanguage": "en",
                **document_evidence(document, args),
                "viewport": {"width": WIDTH, "height": HEIGHT, "scale": SCALE},
                "sources": watched,
                "fixture": fixture,
                "images": images,
                "capture": {
                    "taskId": captured["id"],
                    "imageCount": len(captured["result"]["screenshots"]),
                },
                "modelRequestsDuringCapture": len(model_requests),
            }
            (output / MANIFEST_NAME).write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            return check(output, workbook)
        finally:
            context.close()
            browser.close()


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    result.add_argument(
        "--check",
        action="store_true",
        help="Verify the manifest and PNGs using only the standard library.",
    )
    result.add_argument(
        "--base-url",
        default="http://127.0.0.1:8877/review/",
        help="Root URL of the dedicated preview application.",
    )
    result.add_argument(
        "--workbook",
        type=Path,
        help="Generated fy26_commercial_review.xlsx; required except for --check.",
    )
    result.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Screenshot output directory (default: imgs/).",
    )
    result.add_argument(
        "--storage-state",
        type=Path,
        help="Ignored/private Playwright state from an earlier preparation.",
    )
    result.add_argument(
        "--save-state",
        type=Path,
        help="Save reusable state and a small .project.json evidence file outside tracked assets.",
    )
    result.add_argument(
        "--reuse-project",
        help="Reuse the sole synthetic project in the supplied browser context.",
    )
    ai = result.add_mutually_exclusive_group()
    ai.add_argument(
        "--generate-doc",
        action="store_true",
        help="Allow one English documentation operation (draft and review calls).",
    )
    ai.add_argument(
        "--no-ai",
        action="store_true",
        help="Capture deterministic calculation only, ignoring any existing AI document.",
    )
    result.add_argument(
        "--prepare",
        action="store_true",
        help="Prepare/reuse evidence and save state; do not write repository images.",
    )
    result.add_argument(
        "--timeout",
        type=int,
        default=900,
        help="Maximum seconds per server task wait (default: 900).",
    )
    return result


def main(argv: list[str] | None = None) -> int:
    cli = parser()
    args = cli.parse_args(argv)
    if args.check:
        if args.generate_doc or args.no_ai or args.prepare:
            cli.error(
                "--check cannot generate documentation or prepare a server project"
            )
        return check(args.output.resolve(), args.workbook)
    if args.workbook is None:
        cli.error(
            "--workbook is required; only the generated Sales fixture is accepted"
        )
    if args.prepare and args.save_state is None:
        cli.error(
            "--prepare requires --save-state so existing documentation can be reused"
        )
    if args.reuse_project and args.storage_state is None:
        cli.error("--reuse-project requires --storage-state")
    if args.timeout < 1:
        cli.error("--timeout must be positive")
    try:
        return capture(args)
    except (OSError, ImportError, ValueError, RuntimeError) as error:
        print(f"App capture failed: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
