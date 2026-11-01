#!/usr/bin/env python3
"""Prepare, run and compare a synthetic, native Microsoft Excel oracle.

Run from the checkout with ``uv run scripts/excel_oracle.py --help``. Preparation
and comparison work on Linux. Only --run needs Windows, desktop Excel and
Windows PowerShell. No private workbook is accepted by this harness.

An agreement is evidence about the listed cells, engine version and Excel build,
not a certification of Excel compatibility. See the generated README.txt.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import platform
import re
import shutil
import subprocess
import sys
import time
import uuid
import zipfile
from datetime import UTC, datetime
from importlib.metadata import version
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.packaging.relationship import Relationship
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.external_link.external import (
    ExternalBook,
    ExternalLink,
    ExternalSheetNames,
)
from openpyxl.worksheet.table import Table, TableColumn

PROTOCOL = 1
ROOT = Path(__file__).resolve().parent.parent
WORKER = Path(__file__).with_name("excel_oracle_worker.ps1")
TOLERANCE = {"absolute": 1e-9, "relative": 1e-12}


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def json_bytes(value: Any) -> bytes:
    return (
        json.dumps(value, indent=2, sort_keys=True, allow_nan=False) + "\n"
    ).encode()


def write_json(path: Path, value: Any) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(json_bytes(value))
    temporary.replace(path)


def workbook_bytes(book: Workbook) -> bytes:
    """Remove ZIP/core timestamps so source hashes are reproducible."""
    stream = io.BytesIO()
    book.save(stream)
    output = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(output, "w") as dest:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]+",
                    rb"\g<1>2026-01-01T00:00:00Z",
                    data,
                )
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            # Stored entries avoid different zlib builds changing source hashes
            # when a Linux-prepared bundle is regenerated on Windows.
            info.compress_type = zipfile.ZIP_STORED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            dest.writestr(info, data)
    return output.getvalue()


def target(sheet: str, cell: str, family: str, **kwargs: Any) -> dict:
    return {
        "nodeId": f"{sheet}!{cell}",
        "sheet": sheet,
        "range": cell,
        "family": family,
        "tolerance": dict(TOLERANCE),
        "comparison": "exact_type",
        **kwargs,
    }


def corpus() -> tuple[dict[str, bytes], list[dict]]:
    """Curated fixtures only; mutations by Excel are explicit manifest actions."""
    sources: dict[str, bytes] = {}
    cases: list[dict] = []

    def add(case_id: str, book: Workbook, targets: list[dict], **options: Any):
        filename = f"{case_id}.xlsx"
        sources[filename] = workbook_bytes(book)
        cases.append(
            {
                "id": case_id,
                "file": filename,
                "references": [],
                "actions": [],
                "iterate": False,
                "maxIterations": 1000,
                "maxChange": 1e-10,
                "targets": targets,
                **options,
            }
        )

    for epoch in ("1900", "1904"):
        book = Workbook()
        if epoch == "1904":
            book.epoch = CALENDAR_MAC_1904
        sheet = book.active
        assert sheet is not None
        sheet.title = "Scalar"
        cells = {
            "A1": 6,
            "A2": 7,
            "A3": "7",
            "A4": True,
            "A5": "#N/A",
            "A6": "#N/A",
            "B1": "=A1*A2",
            "B2": '=SUMPRODUCT({2,"3",TRUE},{10,10,10})',
            "B3": "=TRUE=1",
            "B4": '=1<"1"',
            "B5": '=IFERROR(A6,"caught")',
            "B6": '=""',
            "B7": '=A5&" kept"',
            "B8": "=YEAR(DATE(2026,2,1))",
            "B9": "=DATE(2026,2,1)",
            "B10": "=1/0",
            "B11": "=NA()",
            "C1": "=B1+5",
            "C2": "=IF(B3,1,2)",
            "C10": "=IFERROR(B10,99)",
        }
        for address, value in cells.items():
            sheet[address] = value
        sheet["A5"].data_type = "s"
        add(
            f"scalar-{epoch}",
            book,
            [
                target(
                    "Scalar",
                    address,
                    "dates" if address in {"B8", "B9"} else "coercion_errors",
                )
                for address in cells
                if address.startswith(("B", "C"))
            ],
        )

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Tables"
    for row in (
        ["Amount", "Qty", "Total"],
        [10, 2, "=[@Amount]*[@Qty]"],
        [20, 3, "=[@Amount]*[@Qty]"],
        ["=SUM(A2:A3)", "=SUM(B2:B3)", "=SUM(C2:C3)"],
    ):
        sheet.append(row)
    table = Table(displayName="SalesTable", ref="A1:C4", totalsRowCount=1)
    table.tableColumns = [
        TableColumn(id=i, name=name, totalsRowFunction="sum")
        for i, name in enumerate(("Amount", "Qty", "Total"), 1)
    ]
    sheet.add_table(table)
    for address, formula in {
        "E1": "=SUM(SalesTable[Amount])",
        "E2": "=SUM(SalesTable[[#Totals],[Amount]])",
        "E3": "=ROWS(SalesTable[#All])",
        "E4": "=INDEX(SalesTable[#Headers],1,1)",
        "E5": "=SUM(SalesTable[[Amount]:[Qty]])",
        "E6": "=E1*2",
    }.items():
        sheet[address] = formula
    # Let Excel serialize its structured-reference syntax. These exact Formula2
    # assignments are part of the curated manifest, just like native spills.
    # openpyxl's bare structured formulas can make Excel reject the whole file.
    table_actions = []
    for row in sheet:
        for cell in row:
            if cell.data_type == "f":
                table_actions.append(
                    {
                        "kind": "formula2",
                        "sheet": sheet.title,
                        "cell": cell.coordinate,
                        "formula": cell.value,
                    }
                )
                cell.value = 1
    add(
        "tables",
        book,
        [
            target("Tables", address, "tables")
            for address in ["C2", "C3", "E1", "E2", "E3", "E4", "E5", "E6"]
        ],
        actions=table_actions,
    )

    book = Workbook()
    book.remove(book.active)
    for name, value in (("Jan", 1), ("Feb", 2), ("Mar", 3)):
        book.create_sheet(name)["A1"] = value
    sheet = book.create_sheet("Outputs")
    book.defined_names.add(DefinedName("Base", attr_text="2"))
    book.defined_names.add(DefinedName("Computed", attr_text="SUM(Jan!$A$1,Feb!$A$1)"))
    sheet.defined_names.add(DefinedName("Base", attr_text="5"))
    sheet["A1"], sheet["A2"], sheet["A3"] = (
        "=SUM(Jan:Mar!A1)",
        "=Computed*Base",
        "=A1+A2",
    )
    add("names-3d", book, [target("Outputs", f"A{i}", "names_3d") for i in range(1, 4)])

    book = Workbook()
    book.remove(book.active)
    targets = []
    for name in ("Hidden", "Filtered"):
        sheet = book.create_sheet(name)
        for row in (["Label", "Amount"], ["a", 10], ["b", 20], ["a", 30]):
            sheet.append(row)
        if name == "Hidden":
            sheet.row_dimensions[3].hidden = True
        for row_number, value in enumerate((2, 2, 1, 1), 1):
            sheet.cell(row_number, 10, value)
        for address, formula in {
            "D1": "=SUBTOTAL(9,B2:B4)",
            "D2": "=SUBTOTAL(109,B2:B4)",
            "D3": "=_xlfn.AGGREGATE(9,5,B2:B4)",
            "D4": "=SUM(B2:B4)",
            "D5": "=D1*2",
            "G1": "=_xlfn.AGGREGATE(12,5,Z1:Z2)",
            "G2": "=_xlfn.AGGREGATE(14,5,B2:B4,1.5)",
            "G3": "=_xlfn.AGGREGATE(15,5,B2:B4,0.2)",
            "G4": "=_xlfn.AGGREGATE(17,5,B2:B4,1.5)",
            "G5": "=_xlfn.AGGREGATE(19,5,B2:B4,1.5)",
            "G6": "=_xlfn.AGGREGATE(13,6,J1:J4)",
            "G7": "=SUBTOTAL(109,B2:B4)",
            "G8": "=D2+G7",
        }.items():
            sheet[address] = formula
            targets.append(
                target(
                    name,
                    address,
                    "aggregate_edges" if address.startswith("G") else "visibility",
                )
            )
    add(
        "visibility",
        book,
        targets,
        actions=[
            {
                "kind": "filter",
                "sheet": "Filtered",
                "range": "A1:B4",
                "field": 1,
                "criteria": "a",
            }
        ],
    )

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Arrays"
    sheet["A1"] = "=1"  # Native Excel writes real spill metadata via Formula2.
    sheet["D1"], sheet["D2"], sheet["D3"] = "=1", "=1", "=D1+D2"
    add(
        "arrays",
        book,
        [
            target("Arrays", "A1", "dynamic_arrays", range="A1:B3", shape=[3, 2]),
            target("Arrays", "B3", "spill_children"),
            *[target("Arrays", f"D{i}", "array_dependents") for i in range(1, 4)],
        ],
        actions=[
            {"kind": "formula2", "sheet": "Arrays", "cell": cell, "formula": formula}
            for cell, formula in {
                "A1": "=SEQUENCE(3,2)",
                "D1": "=SUM(A1#)",
                "D2": "=INDEX(A1#,3,2)",
            }.items()
        ],
    )

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Inputs"
    sheet["A1"], sheet["A2"], sheet["A3"], sheet["B1"] = 2, 3, 4, "Inputs!A2"
    sheet = book.create_sheet("Dynamic")
    formulas = {
        "A1": "=INDIRECT(Inputs!B1)*2",
        "A2": "=SUM(OFFSET(Inputs!A1,0,0,3,1))",
        "A3": "=_xlfn.FORMULATEXT(A1)",
        "A4": "=_xlfn.SHEETS()",
        "A5": "=NOW()",
        "A6": "=RAND()",
        "D1": "=A5-A5",
        "D2": "=A6-A6",
        "D3": "=A1+A2",
    }
    for address, formula in formulas.items():
        sheet[address] = formula
    add(
        "dynamic-context",
        book,
        [
            target(
                "Dynamic",
                address,
                "volatiles"
                if address in {"A5", "A6", "D1", "D2"}
                else "dynamic_context",
                comparison="observation_only"
                if address in {"A5", "A6"}
                else "exact_type",
            )
            for address in formulas
        ],
    )

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Cycles"
    sheet["A1"], sheet["B1"] = "=(A1+10)/2", "=A1*2"
    book.calculation.iterate = True
    book.calculation.iterateCount = 1000
    book.calculation.iterateDelta = 1e-10
    add(
        "cycles-convergent",
        book,
        [target("Cycles", cell, "cycles") for cell in ("A1", "B1")],
        iterate=True,
    )
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Cycles"
    sheet["A1"], sheet["B1"] = "=A1+1", "=A1*2"
    add(
        "cycles-divergent",
        book,
        [
            target("Cycles", cell, "cycles_divergent", comparison="observation_only")
            for cell in ("A1", "B1")
        ],
        iterate=True,
        maxIterations=50,
    )

    source = Workbook()
    sheet = source.active
    assert sheet is not None
    sheet.title = "Data"
    sheet["A1"], sheet["A2"], sheet["B1"], sheet["B2"] = 7, 3, "=A1*A2", 10
    sources["refs/source.xlsx"] = workbook_bytes(source)
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "External"
    sheet["A1"] = "='[1]Data'!B1+1"
    sheet["A2"] = "=SUM('[1]Data'!B1:B2)"
    sheet["B1"] = "=A1+A2"
    link = ExternalLink(
        externalBook=ExternalBook(
            id="rId1", sheetNames=ExternalSheetNames(sheetName=["Data"])
        )
    )
    link.file_link = Relationship(
        type="externalLinkPath", Target="refs/source.xlsx", TargetMode="External"
    )
    book._external_links.append(link)
    add(
        "external-local",
        book,
        [target("External", cell, "external_local") for cell in ("A1", "A2", "B1")],
        references=["refs/source.xlsx"],
    )
    return sources, cases


def specification() -> tuple[dict, dict[str, bytes]]:
    sources, cases = corpus()
    manifest = {
        "format": "linexcel-excel-oracle",
        "version": PROTOCOL,
        "generatorSha256": sha256(Path(__file__).read_text(encoding="utf-8").encode()),
        "workerSha256": sha256(WORKER.read_text(encoding="utf-8").encode()),
        "openpyxl": version("openpyxl"),
        "sources": {
            name: {"sha256": sha256(data), "bytes": len(data)}
            for name, data in sorted(sources.items())
        },
        "cases": cases,
        "limits": {
            "realWorkbooks": False,
            "macros": False,
            "connections": False,
            "externalReferences": "only generated refs/source.xlsx",
            "comparisonScope": "listed cells and dependents; no certification",
        },
    }
    return manifest, sources


README = """Native Excel oracle, protocol 1

This directory contains generated, public synthetic workbooks only.
Copy the directory and the same linexcel checkout to a Windows desktop with
Microsoft Excel and run (same uv.lock/openpyxl version as preparation):

  uv run scripts/excel_oracle.py --run PATH_TO_THIS_DIRECTORY

Transfer the complete directory back if desired, then compare on either OS:

  uv run scripts/excel_oracle.py --compare PATH_TO_THIS_DIRECTORY

Every invocation creates a separate run/comparison receipt. --receipt selects
a particular run receipt instead of latest-run.json. No result is overwritten.
The source files are hashed, regenerated for validation, and never opened by
Excel: only writable run copies are opened. Unlisted or changed inputs fail.
Excel Formula2 and AutoFilter actions are declared in manifest.json. Comparison
uses the Excel-saved copies, including native spill/filter metadata. Oracle
values are typed COM Value2 reads after full calculation and saving, not an
assumption that cached values are correct. Source/output hashes, Excel version,
build, calculation/iteration settings and engine versions are recorded.

AutomationSecurity=3 is set before Workbooks.Open(UpdateLinks=0); events and
alerts are disabled. Generated sources contain no macros, XLM, Power Query or
active connections. The external-reference case opens only the generated local
reference copy. No arbitrary external workbook, refresh or VBA is executed.
The watchdog terminates only an Excel PID whose start time and window handle
identify the new COM instance; it never kills an existing user's Excel process.
If COM activation hangs before ownership can be recorded, cleanup is reported
as unverified and no unknown Excel process is terminated.

Numeric comparison uses abs=1e-9 and rel=1e-12; text, booleans, blanks and Excel
error types are exact and distinct. Array dimensions and every element must
match. Literal text '#N/A' is not an Excel error. NOW/RAND and divergent cycle
values are observations only, never numerical certification. Their deterministic
dependents are checked separately. Missing Excel/Formula2 support, missing
results, timeouts, engine refusals and differences are recorded explicitly.
Historical caches never count as a successful engine recalculation.
Agreement covers only this corpus, these engine versions and this Excel build.
Preparing or testing the protocol on Linux does not constitute an Excel run.

Microsoft API references:
https://learn.microsoft.com/en-us/office/vba/api/excel.application.automationsecurity
https://learn.microsoft.com/en-us/office/vba/api/excel.workbooks.open
https://learn.microsoft.com/en-us/office/vba/api/excel.application.calculatefullrebuild
"""


def prepare(bundle: Path) -> Path:
    manifest, sources = specification()
    bundle.mkdir(parents=True, exist_ok=False)
    (bundle / "sources").mkdir()
    for name, data in sources.items():
        path = bundle / "sources" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        path.chmod(0o444)
    write_json(bundle / "manifest.json", manifest)
    (bundle / "README.txt").write_text(README, encoding="utf-8")
    return bundle / "manifest.json"


def validate_bundle(bundle: Path) -> dict:
    expected, sources = specification()
    actual = json.loads((bundle / "manifest.json").read_text(encoding="utf-8"))
    if actual != expected:
        raise ValueError(
            "Manifest is not this curated corpus; use the same checkout and openpyxl version, or prepare a new bundle."
        )
    directory = bundle / "sources"
    if directory.is_symlink():
        raise ValueError("Source directory must not be a symbolic link")
    found = {
        path.relative_to(directory).as_posix()
        for path in directory.rglob("*")
        if path.is_file()
    }
    if found != set(sources) or any(path.is_symlink() for path in directory.rglob("*")):
        raise ValueError("Source inventory changed or contains symbolic links")
    for name, data in sources.items():
        if (directory / name).read_bytes() != data:
            raise ValueError(f"Source changed: {name}")
    return actual


def windows_powershell() -> str:
    if os.name != "nt":
        raise RuntimeError(
            "Native Excel oracle requires Windows and desktop Excel; no Excel result was produced."
        )
    executable = shutil.which("powershell.exe")
    if not executable:
        raise RuntimeError("Windows PowerShell is unavailable")
    return executable


def owned_cleanup(powershell: str, process_receipt: Path) -> dict:
    if not process_receipt.exists():
        return {
            "status": "unverified",
            "reason": "COM instance ownership was not recorded",
        }
    try:
        result = subprocess.run(
            [
                powershell,
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(WORKER),
                "-CleanupReceipt",
                str(process_receipt),
            ],
            capture_output=True,
            text=True,
            timeout=15,
        )
        return {
            "status": "checked" if result.returncode == 0 else "failed",
            "exitCode": result.returncode,
            "output": result.stdout[-2048:],
            "error": result.stderr[-2048:],
        }
    except (OSError, subprocess.TimeoutExpired) as exc:
        return {"status": "failed", "reason": str(exc)}


def run_excel(bundle: Path, seconds: float) -> Path:
    manifest = validate_bundle(bundle)
    powershell = windows_powershell()
    run_id = f"{time.strftime('%Y%m%dT%H%M%SZ', time.gmtime())}-{uuid.uuid4().hex[:8]}"
    run_dir = bundle / "runs" / run_id
    run_dir.mkdir(parents=True)
    receipt: dict[str, Any] = {
        "format": "linexcel-excel-receipt",
        "version": PROTOCOL,
        "status": "running",
        "manifestSha256": sha256((bundle / "manifest.json").read_bytes()),
        "runId": run_id,
        "startedAt": datetime.now(UTC).isoformat(),
        "platform": platform.platform(),
        "caseTimeoutSeconds": seconds,
        "cases": [],
    }
    receipt_path = run_dir / "receipt.json"
    write_json(receipt_path, receipt)
    try:
        for case in manifest["cases"]:
            case_dir = run_dir / case["id"]
            case_dir.mkdir()
            for name in [case["file"], *case["references"]]:
                dest = case_dir / name
                dest.parent.mkdir(parents=True, exist_ok=True)
                dest.write_bytes((bundle / "sources" / name).read_bytes())
            request_path = case_dir / "request.json"
            result_path = case_dir / "excel-result.json"
            process_path = case_dir / "excel-process.json"
            write_json(
                request_path, {"case": case, "directory": str(case_dir.resolve())}
            )
            with (
                (case_dir / "powershell.stdout.log").open("wb") as stdout,
                (case_dir / "powershell.stderr.log").open("wb") as stderr,
            ):
                child = subprocess.Popen(
                    [
                        powershell,
                        "-NoProfile",
                        "-NonInteractive",
                        "-ExecutionPolicy",
                        "Bypass",
                        "-File",
                        str(WORKER),
                        "-Request",
                        str(request_path),
                        "-Result",
                        str(result_path),
                        "-ProcessReceipt",
                        str(process_path),
                    ],
                    stdout=stdout,
                    stderr=stderr,
                )
                timed_out = False
                try:
                    child.wait(timeout=seconds)
                except subprocess.TimeoutExpired:
                    timed_out = True
                    child.kill()
                    child.wait(timeout=10)
                finally:
                    if child.poll() is None:
                        child.kill()
                        child.wait(timeout=10)
                    cleanup = owned_cleanup(powershell, process_path)
            outcome: dict[str, Any] = {
                "id": case["id"],
                "status": "timed_out" if timed_out else "failed",
                "exitCode": child.returncode,
                "cleanup": cleanup,
            }
            if result_path.exists():
                native = json.loads(result_path.read_text(encoding="utf-8-sig"))
                outcome["excel"] = native
                if (
                    not timed_out
                    and child.returncode == 0
                    and native.get("status") == "completed"
                    and cleanup["status"] == "checked"
                ):
                    try:
                        validate_native(case, native)
                        outcome["status"] = "completed"
                    except ValueError as exc:
                        outcome["protocolError"] = str(exc)
            outcome["outputs"] = {
                name: {
                    "sha256": sha256((case_dir / name).read_bytes()),
                    "path": f"runs/{run_id}/{case['id']}/{name}",
                }
                for name in [case["file"], *case["references"]]
                if (case_dir / name).exists()
            }
            receipt["cases"].append(outcome)
            write_json(receipt_path, receipt)
    finally:
        validate_bundle(bundle)
        receipt["finishedAt"] = datetime.now(UTC).isoformat()
        receipt["status"] = (
            "completed"
            if len(receipt["cases"]) == len(manifest["cases"])
            and all(case["status"] == "completed" for case in receipt["cases"])
            else "incomplete"
        )
        write_json(receipt_path, receipt)
        write_json(
            bundle / "latest-run.json", {"receipt": f"runs/{run_id}/receipt.json"}
        )
    return receipt_path


def typed(value: Any, kind: str = "scalar") -> dict:
    if kind == "error":
        return {"kind": "error", "value": value}
    if isinstance(value, list):
        matrix = value if not value or isinstance(value[0], list) else [value]
        return {
            "kind": "array",
            "rows": len(matrix),
            "columns": len(matrix[0]) if matrix else 0,
            "values": [[typed(cell) for cell in row] for row in matrix],
        }
    if value is None:
        return {"kind": "blank", "value": None}
    if isinstance(value, bool):
        return {"kind": "boolean", "value": value}
    if isinstance(value, (int, float)) and math.isfinite(value):
        return {"kind": "number", "value": value}
    if isinstance(value, str):
        return {"kind": "text", "value": value}
    return {"kind": "uncomparable", "value": str(value)}


def equivalent(actual: dict, expected: dict, tolerance: dict) -> bool:
    if (
        actual.get("kind") != expected.get("kind")
        or actual.get("kind") == "uncomparable"
    ):
        return False
    if actual["kind"] == "array":
        return (
            actual.get("rows") == expected.get("rows")
            and actual.get("columns") == expected.get("columns")
            and len(actual.get("values", [])) == len(expected.get("values", []))
            and all(
                len(a) == len(b)
                and all(equivalent(x, y, tolerance) for x, y in zip(a, b, strict=True))
                for a, b in zip(actual["values"], expected["values"], strict=True)
            )
        )
    if actual["kind"] == "number":
        return math.isclose(
            actual["value"],
            expected["value"],
            rel_tol=tolerance["relative"],
            abs_tol=tolerance["absolute"],
        )
    return type(actual.get("value")) is type(expected.get("value")) and actual.get(
        "value"
    ) == expected.get("value")


def engine_request(request_path: Path) -> None:
    """Private subprocess entry point. Results must prove calculation provenance."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    from linexcel import ExecutionPolicy, analyze
    from linexcel.lazy import evaluate_node

    request = json.loads(request_path.read_text(encoding="utf-8"))
    path = Path(request["file"])
    refs = Path(request["refs"]) if request.get("refs") else None
    sheet, address = request["nodeId"].rsplit("!", 1)
    c1, r1, c2, r2 = range_boundaries(request.get("range", address))
    ids = [
        [f"{sheet}!{get_column_letter(column)}{row}" for column in range(c1, c2 + 1)]
        for row in range(r1, r2 + 1)
    ]
    flat_ids = [node_id for row in ids for node_id in row]
    readings = {}
    if request["mode"] == "lazy":
        for node_id in flat_ids:
            result = evaluate_node(path, node_id, refs_dir=refs)
            readings[node_id] = {
                "calculated": result.get("status") == "completed"
                and result.get("provenance") == "targeted_recalculation",
                "status": result.get("status"),
                "provenance": result.get("provenance"),
                "value": typed(result.get("value"), result.get("valueKind", "scalar")),
                "diagnostics": result.get("diagnostics"),
                "coverage": result.get("coverage"),
                "arrayEvidence": result.get("arrayEvidence"),
                "evaluationContext": result.get("evaluationContext"),
            }
    else:
        result = analyze(
            path,
            refs_dir=refs,
            targets=flat_ids,
            execution=ExecutionPolicy(seconds=request["seconds"], memory_mb=2048),
        )
        for node_id in flat_ids:
            source_sheet, cell = node_id.rsplit("!", 1)
            node = next(
                (
                    node
                    for node in result.nodes
                    if node.get("sheet") == source_sheet and node.get("addr") == cell
                ),
                {},
            )
            readings[node_id] = {
                "calculated": node.get("valueSource") == "engine",
                "status": node.get("valueSource", "missing_target"),
                "provenance": node.get("valueSource"),
                "value": typed(node.get("value"), node.get("valueKind", "scalar")),
                "diagnostics": result.warnings,
                "execution": result.graph.get("meta", {}).get("execution"),
            }
    if len(flat_ids) == 1:
        output = readings[flat_ids[0]]
    else:
        # The public cell API returns an anchor's scalar value. Array evidence is
        # therefore assembled from every fresh cell read, never saved spill
        # caches. Each lazy call checks the native extent against source bounds.
        calculated = all(reading["calculated"] for reading in readings.values())
        output = {
            "calculated": calculated,
            "status": "completed" if calculated else "incomplete_array",
            "provenance": "individually_recalculated_array_cells"
            if request["mode"] == "lazy"
            else "batch_target_cells",
            "value": {
                "kind": "array",
                "rows": len(ids),
                "columns": len(ids[0]),
                "values": [
                    [readings[node_id]["value"] for node_id in row] for row in ids
                ],
            },
            "cells": readings,
            "scope": "Fixed SEQUENCE fixture; lazy array cells are separate recalculations with individually recorded contexts",
        }
    write_json(Path(request["result"]), output)


def safe_output(bundle: Path, relative: str) -> Path:
    path = (bundle / relative).resolve()
    if not path.is_relative_to((bundle / "runs").resolve()) or not relative.startswith(
        "runs/"
    ):
        raise ValueError("Output path escapes the oracle run")
    return path


def validate_typed(value: dict) -> None:
    kind = value.get("kind")
    item = value.get("value")
    if kind == "array":
        rows, columns, matrix = (
            value.get("rows"),
            value.get("columns"),
            value.get("values"),
        )
        if (
            type(rows) is not int
            or type(columns) is not int
            or not 1 <= rows <= 100
            or not 1 <= columns <= 100
            or not isinstance(matrix, list)
            or len(matrix) != rows
        ):
            raise ValueError("Invalid oracle array dimensions")
        for row in matrix:
            if not isinstance(row, list) or len(row) != columns:
                raise ValueError("Invalid oracle array row")
            for cell in row:
                validate_typed(cell)
        return
    valid = (
        (
            kind == "number"
            and isinstance(item, (int, float))
            and not isinstance(item, bool)
            and math.isfinite(item)
        )
        or (kind == "boolean" and type(item) is bool)
        or (kind in {"text", "error", "uncomparable"} and isinstance(item, str))
        or (kind == "blank" and item is None)
    )
    if not valid:
        raise ValueError("Oracle value is missing a valid explicit type")


def comparable(value: dict) -> bool:
    """An unknown native error is missing evidence, not a numerical difference."""
    if value["kind"] == "array":
        return all(comparable(cell) for row in value["values"] for cell in row)
    return value["kind"] != "uncomparable"


def validate_native(spec: dict, native: dict) -> None:
    identity, settings = native.get("excel", {}), native.get("calculation", {})
    if (
        not identity.get("version")
        or not identity.get("build")
        or identity.get("automationSecurity") != 3
    ):
        raise ValueError("Excel identity/security evidence is incomplete")
    if (
        settings.get("mode") != -4135
        or settings.get("iteration") is not spec["iterate"]
        or settings.get("maxIterations") != spec["maxIterations"]
        or settings.get("maxChange") != spec["maxChange"]
        or settings.get("calculateBeforeSave") is not False
    ):
        raise ValueError("Excel calculation settings differ from the declared case")
    if settings.get("date1904") is not (spec["id"] == "scalar-1904"):
        raise ValueError("Excel date epoch differs from the declared case")
    if settings.get("stateAfterRebuild", 0) != 0:
        raise ValueError("Native Excel did not finish the requested calculation")
    observations = native.get("observations", {})
    if set(observations) != {item["nodeId"] for item in spec["targets"]}:
        raise ValueError("Excel observations do not cover the declared targets exactly")
    for selection in spec["targets"]:
        validate_typed(observations[selection["nodeId"]]["value"])
        if selection.get("shape"):
            shape = native.get("arrayShapes", {}).get(selection["nodeId"], {})
            if [shape.get("rows"), shape.get("columns")] != selection["shape"]:
                raise ValueError("Native Excel spill shape evidence is missing")


def validate_receipt(bundle: Path, manifest: dict, path: Path) -> dict:
    receipt = json.loads(path.read_text(encoding="utf-8"))
    if (
        receipt.get("format") != "linexcel-excel-receipt"
        or receipt.get("version") != PROTOCOL
        or receipt.get("manifestSha256")
        != sha256((bundle / "manifest.json").read_bytes())
    ):
        raise ValueError("Excel receipt does not identify this bundle")
    expected = {case["id"]: case for case in manifest["cases"]}
    cases = receipt.get("cases", [])
    if len({case["id"] for case in cases}) != len(cases) or any(
        case["id"] not in expected for case in cases
    ):
        raise ValueError("Excel receipt contains duplicate or unknown cases")
    for case in cases:
        spec = expected[case["id"]]
        if case.get("status") == "completed":
            if case.get("cleanup", {}).get("status") != "checked":
                raise ValueError("Excel process cleanup was not verified")
            if set(case.get("outputs", {})) != {spec["file"], *spec["references"]}:
                raise ValueError(
                    "Completed Excel case has an incomplete output inventory"
                )
            validate_native(spec, case.get("excel", {}))
        for name, output in case.get("outputs", {}).items():
            if output["path"] != f"runs/{receipt.get('runId')}/{case['id']}/{name}":
                raise ValueError(
                    "Excel output does not belong to its declared run/case"
                )
            file = safe_output(bundle, output["path"])
            if sha256(file.read_bytes()) != output["sha256"]:
                raise ValueError(f"Excel output changed: {file.name}")
    return receipt


def compare(bundle: Path, receipt_path: Path | None, seconds: float) -> Path:
    manifest = validate_bundle(bundle)
    if receipt_path is None:
        latest = json.loads((bundle / "latest-run.json").read_text(encoding="utf-8"))
        receipt_path = safe_output(bundle, latest["receipt"])
    receipt = validate_receipt(bundle, manifest, receipt_path)
    comparison_dir = bundle / "comparisons" / uuid.uuid4().hex
    comparison_dir.mkdir(parents=True)
    report: dict[str, Any] = {
        "format": "linexcel-excel-comparison",
        "version": PROTOCOL,
        "receiptSha256": sha256(receipt_path.read_bytes()),
        "manifestSha256": receipt["manifestSha256"],
        "engineVersions": {
            name: version(name) for name in ("linexcel", "formualizer", "openpyxl")
        },
        "engineSourceSha256": {
            path.relative_to(ROOT).as_posix(): sha256(path.read_bytes())
            for path in sorted((ROOT / "src" / "linexcel").rglob("*.py"))
        },
        "python": platform.python_version(),
        "platform": platform.platform(),
        "scope": "Only declared targets; observations do not certify numerical agreement",
        "rows": [],
    }
    by_case = {case["id"]: case for case in receipt["cases"]}
    for spec in manifest["cases"]:
        native = by_case.get(spec["id"], {})
        for selection in spec["targets"]:
            for mode in ("lazy", "batch"):
                row: dict[str, Any] = {
                    "case": spec["id"],
                    "nodeId": selection["nodeId"],
                    "family": selection["family"],
                    "mode": mode,
                    "tolerance": selection["tolerance"],
                    "status": "oracle_unavailable",
                }
                observations = native.get("excel", {}).get("observations", {})
                oracle = observations.get(selection["nodeId"])
                if native.get("status") != "completed":
                    row["oracleFailure"] = {
                        "status": native.get("status", "missing"),
                        **{
                            key: native["excel"][key]
                            for key in ("stage", "error", "calculation")
                            if key in native.get("excel", {})
                        },
                    }
                if native.get("status") == "completed" and oracle is not None:
                    row["oracle"] = oracle
                    workbook = safe_output(
                        bundle, native["outputs"][spec["file"]]["path"]
                    )
                    request_path = (
                        comparison_dir / f"request-{len(report['rows'])}.json"
                    )
                    result_path = comparison_dir / f"result-{len(report['rows'])}.json"
                    write_json(
                        request_path,
                        {
                            "mode": mode,
                            "file": str(workbook),
                            "refs": str(workbook.parent / "refs")
                            if spec["references"]
                            else None,
                            "nodeId": selection["nodeId"],
                            "range": selection["range"],
                            "result": str(result_path),
                            "seconds": seconds,
                        },
                    )
                    with request_path.with_suffix(".log").open("wb") as log:
                        try:
                            child = subprocess.run(
                                [
                                    sys.executable,
                                    str(Path(__file__).resolve()),
                                    "--engine-request",
                                    str(request_path),
                                ],
                                stdout=log,
                                stderr=log,
                                timeout=seconds + 10,
                                cwd=ROOT,
                            )
                            if child.returncode == 0 and result_path.exists():
                                actual = json.loads(
                                    result_path.read_text(encoding="utf-8")
                                )
                                row["engine"] = actual
                                row["status"] = (
                                    (
                                        "agreed"
                                        if equivalent(
                                            actual["value"],
                                            oracle["value"],
                                            selection["tolerance"],
                                        )
                                        else "different"
                                    )
                                    if actual["calculated"]
                                    else "not_recalculated"
                                )
                                if not comparable(oracle["value"]):
                                    row["status"] = "oracle_uncomparable"
                                elif selection["comparison"] == "observation_only":
                                    row["status"] = "observation_only"
                            else:
                                row.update(
                                    status="engine_failed",
                                    exitCode=child.returncode,
                                )
                        except subprocess.TimeoutExpired:
                            row["status"] = "engine_timed_out"
                report["rows"].append(row)
                write_json(comparison_dir / "comparison.json", report)
    statuses: dict[str, int] = {}
    for row in report["rows"]:
        statuses[row["status"]] = statuses.get(row["status"], 0) + 1
    report["counts"] = statuses
    report["status"] = (
        "agreed_for_comparable_targets"
        if report["rows"]
        and all(
            row["status"] in {"agreed", "observation_only"} for row in report["rows"]
        )
        else "differences_or_incomplete"
    )
    write_json(comparison_dir / "comparison.json", report)
    return comparison_dir / "comparison.json"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    operation = parser.add_mutually_exclusive_group(required=True)
    operation.add_argument("--prepare", type=Path, metavar="NEW_DIRECTORY")
    operation.add_argument("--run", type=Path, metavar="BUNDLE")
    operation.add_argument("--compare", type=Path, metavar="BUNDLE")
    operation.add_argument("--engine-request", type=Path, help=argparse.SUPPRESS)
    parser.add_argument(
        "--receipt", type=Path, help="Explicit Excel receipt for --compare"
    )
    parser.add_argument(
        "--seconds",
        type=float,
        default=120,
        help="Native Excel/engine timeout per case or target (default 120)",
    )
    args = parser.parse_args(argv)
    if not math.isfinite(args.seconds) or args.seconds <= 0:
        parser.error("--seconds must be finite and positive")
    if args.receipt and not args.compare:
        parser.error("--receipt requires --compare")
    try:
        if args.prepare:
            output = prepare(args.prepare.resolve())
        elif args.run:
            output = run_excel(args.run.resolve(), args.seconds)
        elif args.compare:
            output = compare(args.compare.resolve(), args.receipt, args.seconds)
        else:
            engine_request(args.engine_request)
            return 0
        print(output)
        if args.run or args.compare:
            result = json.loads(output.read_text(encoding="utf-8"))
            return (
                0
                if result["status"] in {"completed", "agreed_for_comparable_targets"}
                else 2
            )
        return 0
    except (OSError, ValueError, RuntimeError, KeyError, TypeError) as exc:
        print(f"Excel oracle: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
