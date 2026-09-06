# One-shot: build a wide workbook (~180 lineage nodes) to stress the viewer.
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

wb = Workbook()
wb.remove(wb.active)

# Parameters sheet (inputs)
params = wb.create_sheet("Params")
params["A1"] = "rate"
params["B1"] = 0.2
params["A2"] = "discount"
params["B2"] = 0.05
wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Params!$B$1")
wb.defined_names["Discount"] = DefinedName("Discount", attr_text="Params!$B$2")

# 12 data sheets, each with stretched formula groups referencing Params
for s in range(12):
    ws = wb.create_sheet(f"Data{s + 1:02d}")
    ws.append(["qty", "price", "ht", "ttc"])
    for r in range(2, 42):
        ws.cell(row=r, column=1, value=r * 3 + s)
        ws.cell(row=r, column=2, value=10 + s)
        ws.cell(row=r, column=3, value=f"=A{r}*B{r}")
        ws.cell(row=r, column=4, value=f"=C{r}*(1+Rate)")

# 6 summary sheets aggregating the data sheets (cross-sheet refs)
for s in range(6):
    ws = wb.create_sheet(f"Summary{s + 1}")
    a, b = f"Data{2 * s + 1:02d}", f"Data{2 * s + 2:02d}"
    ws["A1"] = "total ht"
    ws["B1"] = f"=SUM({a}!C2:C41)+SUM({b}!C2:C41)"
    ws["A2"] = "total ttc"
    ws["B2"] = f"=SUM({a}!D2:D41)+SUM({b}!D2:D41)"
    ws["A3"] = "net"
    ws["B3"] = "=B2*(1-Discount)"
    for r in range(5, 15):
        ws.cell(row=r, column=1, value=f"part {r}")
        ws.cell(row=r, column=2, value=f"=B$3/{r}")

# Top sheet chaining the summaries
top = wb.create_sheet("Top")
for s in range(6):
    top.cell(row=s + 1, column=1, value=f"Summary{s + 1}")
    top.cell(row=s + 1, column=2, value=f"=Summary{s + 1}!B3")
top["A8"] = "grand total"
top["B8"] = "=SUM(B1:B6)"

wb.save("dense_check.xlsx")
print("dense_check.xlsx written")
