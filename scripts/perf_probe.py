"""Build a stress workbook and time formualizer boot + evaluate_all variants.

openpyxl writes no cached formula values, so this file also exercises the
engine-only path a real recalculated file would not.
"""

from __future__ import annotations

import io
import sys
import time

import formualizer as fz
from openpyxl import Workbook

ROWS = int(sys.argv[1]) if len(sys.argv) > 1 else 20_000


def build() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 1
    ws["B1"] = 2
    # Long dependency chain: each row reads the row above it.
    for r in range(2, ROWS + 1):
        ws.cell(row=r, column=1).value = f"=A{r-1}*1.0001+B{r-1}"
        ws.cell(row=r, column=2).value = f"=B{r-1}+MOD(A{r},7)"
    # Cross-sheet references: every row reads the other sheet.
    ws2 = wb.create_sheet("Cross")
    for r in range(1, ROWS + 1):
        ws2.cell(row=r, column=1).value = f"=Data!A{r}*2"
    ws3 = wb.create_sheet("Lookups")
    for r in range(1, min(ROWS, 2000) + 1):
        ws3.cell(row=r, column=1).value = f"=SUM(Data!A1:A{r})"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def timeit(label, fn):
    t = time.perf_counter()
    out = fn()
    print(f"{label}: {time.perf_counter() - t:.2f}s", flush=True)
    return out


data = build()
print(f"workbook bytes: {len(data) / 1_048_576:.1f} MB, {ROWS} chain rows", flush=True)

wb = timeit("from_bytes (default)", lambda: fz.Workbook.from_bytes(data))
timeit("evaluate_all (default)", lambda: wb.evaluate_all())

cfg = fz.EvaluationConfig()
cfg.enable_parallel = False
wb2 = timeit(
    "from_bytes (parallel off)",
    lambda: fz.Workbook.from_bytes(data, config=fz.WorkbookConfig(eval_config=cfg)),
)
timeit("evaluate_all (parallel off)", lambda: wb2.evaluate_all())

cfg3 = fz.EvaluationConfig()
cfg3.warmup_enabled = True
wb3 = timeit(
    "from_bytes (warmup on)",
    lambda: fz.Workbook.from_bytes(data, config=fz.WorkbookConfig(eval_config=cfg3)),
)
timeit("evaluate_all (warmup on)", lambda: wb3.evaluate_all())
