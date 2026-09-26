"""Deterministic lazy workbook graph and exact, sparse target closure.

These entry points must be called in a resource-limited child process. Import
never initializes or evaluates the native formula engine. Unsupported dependencies fail
closed before any evaluation; no clipped ranges or cached formula substitutions.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import math
import re
import zipfile
from collections import defaultdict
from dataclasses import replace
from pathlib import Path, PureWindowsPath
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import column_index_from_string
from openpyxl.utils.datetime import to_excel

from linexcel.external import _EXTERNAL_REF_RE, read_external_links
from linexcel.refs import (
    a1,
    parse_ref,
    parse_ref_detailed,
    quote_sheet,
    ref_to_r1c1,
    split_sheet_prefix,
)
from linexcel.values import (
    ERROR_KIND_TEXT,
    EXCEL_ERRORS,
    _excel_error_text,
    _is_uncomputed,
)

_INPUT_ERROR_KINDS = {value: kind for kind, value in ERROR_KIND_TEXT.items()}

_VOLATILE = {"RAND", "RANDBETWEEN", "RANDARRAY", "NOW", "TODAY"}
_PATTERN_CELL = r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*"
_PATTERN_REFERENCE = re.compile(
    rf"(?:{_PATTERN_CELL}(?::{_PATTERN_CELL})?|"
    r"\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}|\$?[1-9][0-9]*:\$?[1-9][0-9]*)"
)

_DYNAMIC = {
    "INDIRECT",
    "OFFSET",
    "CELL",
    "INFO",
    "SHEET",
    "SHEETS",
    "SUBTOTAL",
    "AGGREGATE",
    "FORMULATEXT",
    "RTD",
    "WEBSERVICE",
    "STOCKHISTORY",
}


def _json(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _engine_value(value: Any) -> Any:
    """Normalize real Excel results; never stringify a native limitation."""
    if _is_uncomputed(value):
        raise ValueError(f"Engine could not calculate this value: {value.get('kind')}")
    error = _excel_error_text(value)
    if error is not None:
        return error
    if isinstance(value, (list, tuple)):
        return [_engine_value(item) for item in value]
    if value is None or isinstance(value, (str, int, float, bool, dt.date, dt.time)):
        return _json(value)
    raise ValueError(f"Unsupported engine value representation: {type(value).__name__}")


def _id(sheet: str, coordinate: str) -> str:
    return f"{quote_sheet(sheet)}!{coordinate}"


def _pattern_signature(
    node: dict, nodes: dict[str, dict]
) -> tuple[str, list, str | None]:
    """Engine-free R1C1 key; preserve names, literals and operator token types."""
    formula = node["formula"]
    if node.get("diagnostics"):
        return (
            formula,
            [],
            "Formula has unsupported or ambiguous structural diagnostics.",
        )
    try:
        tokens = Tokenizer(formula).items
        identity = []
        readable = []
        for token in tokens:
            value = token.value
            tag = (token.type, token.subtype)
            if token.type == "OPERAND" and token.subtype == "RANGE":
                binding = nodes.get(node.get("reference_bindings", {}).get(value), {})
                if binding.get("kind") == "name":
                    # A name such as TVA or RC is not a column/cell reference.
                    # Include its resolved scope in the identity; never rewrite it.
                    tag = ("DEFINED_NAME", binding["id"])
                    value = binding["name"]
                else:
                    detail = parse_ref_detailed(value)
                    _, body = split_sheet_prefix(value)
                    if (
                        "[" in value
                        or detail is None
                        or not _PATTERN_REFERENCE.fullmatch(body)
                        or binding.get("kind") == "unresolved"
                    ):
                        return (
                            formula,
                            [],
                            f"Reference cannot be canonicalized safely: {value}",
                        )
                    converted = ref_to_r1c1(value, node["row"], node["column"])
                    if converted is None:
                        return (
                            formula,
                            [],
                            f"Reference cannot be canonicalized safely: {value}",
                        )
                    value = converted
                    tag = ("REFERENCE", "R1C1")
            # Do not fold string literals or remove whitespace: a space may be
            # Excel's intersection operator. This prefers false negatives to
            # grouping formulas whose syntax has different meaning.
            identity.append([*tag, value])
            readable.append(value)
        return "=" + "".join(readable), identity, None
    except (ValueError, IndexError) as exc:
        return formula, [], f"Formula cannot be canonicalized safely: {exc}"


def _member_ranges(members: list[dict]) -> list[str]:
    """Tile actual members with exact rectangles; never fill gaps in a group."""
    rows: dict[int, list[int]] = defaultdict(list)
    for node in members:
        rows[node["row"]].append(node["column"])
    rectangles = []
    active: dict[tuple[int, int], tuple[int, int]] = {}
    for row, columns in sorted(rows.items()):
        columns.sort()
        spans = []
        start = end = columns[0]
        for column in columns[1:]:
            if column == end + 1:
                end = column
            else:
                spans.append((start, end))
                start = end = column
        spans.append((start, end))
        next_active = {}
        for span in spans:
            previous = active.pop(span, None)
            if previous and previous[1] == row - 1:
                next_active[span] = (previous[0], row)
            else:
                if previous:
                    rectangles.append((previous[0], span[0], previous[1], span[1]))
                next_active[span] = (row, row)
        for span, (r1, r2) in active.items():
            rectangles.append((r1, span[0], r2, span[1]))
        active = next_active
    for span, (r1, r2) in active.items():
        rectangles.append((r1, span[0], r2, span[1]))
    return [
        a1(r1, c1) if (r1, c1) == (r2, c2) else f"{a1(r1, c1)}:{a1(r2, c2)}"
        for r1, c1, r2, c2 in sorted(rectangles)
    ]


def _formula_patterns(nodes: dict[str, dict]) -> list[dict]:
    groups = {}
    members: dict[str, list[dict]] = defaultdict(list)
    for node in nodes.values():
        if node["kind"] != "cell":
            continue
        signature, tokens, reason = _pattern_signature(node, nodes)
        identity = ["copied-r1c1-v1", node["sheet"], tokens]
        if reason:
            identity += ["isolated", node["id"], node["formula"]]
        key = (
            "formula-sha256:"
            + hashlib.sha256(
                json.dumps(identity, ensure_ascii=False, separators=(",", ":")).encode(
                    "utf-8"
                )
            ).hexdigest()
        )
        node["patternId"] = key
        if key not in groups:
            groups[key] = {
                "id": key,
                "sheet": node["sheet"],
                "signature": signature,
                "status": "isolated" if reason else "canonical",
                "diagnostics": [reason] if reason else [],
            }
        members[key].append(node)
    result = []
    for key, group in groups.items():
        cells = sorted(members[key], key=lambda n: (n["row"], n["column"]))
        group.update(
            memberIds=[node["id"] for node in cells],
            memberCount=len(cells),
            ranges=_member_ranges(cells),
            representativeId=cells[0]["id"],
            representativeFormula=cells[0]["formula"],
        )
        for node in cells:
            node["patternMemberCount"] = len(cells)
        result.append(group)
    return sorted(result, key=lambda group: (group["sheet"], group["representativeId"]))


def _cell_context(sheet, cell, hidden_columns: set[int]) -> dict:
    """Bounded literal nearby text, never an inferred business label or input."""
    nearby = []
    for position in ("left", "above"):
        for distance in range(1, 5):
            row = cell.row - (distance if position == "above" else 0)
            column = cell.column - (distance if position == "left" else 0)
            other = sheet._cells.get((row, column))
            if (
                other is not None
                and isinstance(other.value, str)
                and other.data_type not in {"f", "e"}
                and other.value.strip()
            ):
                nearby.append(
                    {
                        "cell": other.coordinate,
                        "text": other.value[:240],
                        "position": position,
                        "truncated": len(other.value) > 240,
                    }
                )
                break
    row_dimension = sheet.row_dimensions.get(cell.row)
    return {
        "nearbyLabels": nearby,
        "sheetState": sheet.sheet_state,
        "rowHidden": bool(row_dimension and row_dimension.hidden),
        "columnHidden": cell.column in hidden_columns,
    }


def _stored_date_numbers(path: Path, cache) -> dict[tuple[str, str], float]:
    """Retain numeric date inputs/caches before Python calendar conversion.

    openpyxl maps both serials 59 and 60 to February 28, 1900 and rounds time
    fractions. Calculation and cache comparison must use the actual file value.
    ISO-date OOXML cells have no numeric payload and keep their normal conversion.
    """
    from linexcel.loader import _parse_sheet_targets

    wanted = {
        sheet.title: {
            cell.coordinate
            for cell in sheet._cells.values()
            if isinstance(cell.value, (dt.datetime, dt.date, dt.time))
        }
        for sheet in cache
    }
    if not any(wanted.values()):
        return {}
    result = {}
    with zipfile.ZipFile(path) as archive:
        paths = _parse_sheet_targets(
            archive.read("xl/workbook.xml").decode("utf-8"),
            archive.read("xl/_rels/workbook.xml.rels").decode("utf-8"),
        )
        for sheet, member in paths.items():
            addresses = wanted.get(sheet, set())
            if not addresses:
                continue
            with archive.open(member) as source:
                for _, cell in ElementTree.iterparse(source, events=("end",)):
                    tag = cell.tag.rsplit("}", 1)[-1]
                    if tag == "c":
                        address = cell.get("r")
                        if address in addresses and cell.get("t", "n") == "n":
                            try:
                                value = float(cell.findtext("{*}v", ""))
                            except ValueError:
                                pass
                            else:
                                if math.isfinite(value):
                                    result[(sheet, address)] = value
                        cell.clear()
                    elif tag == "row":
                        cell.clear()
    return result


def build_structure(path: Path, refs_dir: Path | None = None) -> dict:
    """Read saved cells and tokenize references, without executing formulas."""
    book = load_workbook(path, data_only=False, keep_links=False)
    cache = load_workbook(path, data_only=True, keep_links=False)
    nodes: dict[str, dict] = {}
    names: dict[tuple[str | None, str], str] = {}
    edges: set[tuple[str, str]] = set()
    array_regions = []
    external_books: dict[str, Any] = {}
    links = read_external_links(path.read_bytes())
    sheet_details = []
    sheet_lookup = {name.casefold(): name for name in book.sheetnames}
    try:
        date_numbers = _stored_date_numbers(path, cache)
        for sheet in book:
            hidden_columns: set[int] = set()
            for key, dimension in sheet.column_dimensions.items():
                if dimension.hidden:
                    first = dimension.min or column_index_from_string(key)
                    last = dimension.max or first
                    hidden_columns.update(range(first, last + 1))
            sheet_counts = {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "formulas": 0,
                "inputs": 0,
            }
            # Sparse stored cells only: do not expand a formatted whole column.
            for cell in sheet._cells.values():
                if cell.value is None and cell.data_type not in {
                    "inlineStr",
                    "s",
                    "str",
                }:
                    continue
                formula = cell.value if cell.data_type == "f" else None
                diagnostics = []
                if formula is not None and not isinstance(formula, str):
                    region = parse_ref(getattr(formula, "ref", ""), sheet.title)
                    if region is not None:
                        array_regions.append(region)
                if cell.data_type == "e" and cell.value not in _INPUT_ERROR_KINDS:
                    diagnostics.append(
                        "Unknown Excel error input is not supported for recalculation."
                    )
                if formula is not None and not isinstance(formula, str):
                    formula = getattr(formula, "text", None) or "=<array/data-table>"
                    diagnostics.append("Array/data-table formula is not supported.")
                nid = _id(sheet.title, cell.coordinate)
                count_key = "formulas" if formula else "inputs"
                sheet_counts[count_key] += 1
                saved = cache[sheet.title][cell.coordinate].value
                stored_serial = date_numbers.get((sheet.title, cell.coordinate))
                # Python calendars cannot represent Excel's fictitious leap
                # day. Keep its exact serial rather than displaying February 28.
                if (
                    book.epoch.year != 1904
                    and stored_serial is not None
                    and 60 <= stored_serial < 61
                ):
                    saved = stored_serial
                if (
                    formula is None
                    and saved is None
                    and cell.data_type in {"inlineStr", "s", "str"}
                ):
                    saved = ""
                nodes[nid] = {
                    "id": nid,
                    "kind": "cell" if formula else "input",
                    "sheet": sheet.title,
                    "cell": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "formula": formula,
                    "numberFormat": cell.number_format,
                    "sourceDataType": cell.data_type,
                    "context": _cell_context(sheet, cell, hidden_columns),
                    "value": _json(saved),
                    "cached_value": _json(saved),
                    "cachedValueKind": (
                        "error"
                        if cache[sheet.title][cell.coordinate].data_type == "e"
                        else "scalar"
                    ),
                    "cached_comparison_value": (
                        stored_serial
                        if stored_serial is not None
                        else to_excel(saved, book.epoch)
                        if isinstance(saved, (dt.datetime, dt.date, dt.time))
                        else _json(saved)
                    ),
                    "cache_status": "unknown" if formula and saved is None else "saved",
                    "provenance": "saved_cache" if formula else "saved_input",
                    "dependencies": [],
                    "diagnostics": diagnostics,
                }
            sheet_details.append(sheet_counts)
        for scope, collection in [
            (None, book.defined_names),
            *[(sheet.title, sheet.defined_names) for sheet in book],
        ]:
            for name, definition in collection.items():
                nid = f"name:{scope or '*'}:{name}"
                names[(scope, name.casefold())] = nid
                nodes[nid] = {
                    "id": nid,
                    "kind": "name",
                    "name": name,
                    "sheet": scope,
                    "formula": "=" + definition.attr_text.lstrip("="),
                    "value": None,
                    "cached_value": None,
                    "cache_status": "unknown",
                    "dependencies": [],
                    "diagnostics": [],
                }
                literal = _name_literal(nodes[nid])
                if literal is not None:
                    _expression, value, kind = literal
                    nodes[nid].update(
                        value=value,
                        valueKind=kind,
                        provenance="defined_constant",
                    )

        def external_input(text: str) -> dict | None:
            match = _EXTERNAL_REF_RE.fullmatch(text)
            if not match or refs_dir is None:
                return None
            key = match.group("book")
            filename = links[key].name if key in links else key
            # Only explicit uploads; never follow workbook-declared directories.
            if PureWindowsPath(filename).name != filename:
                return None
            uploaded = refs_dir / filename
            if not uploaded.is_file() or uploaded.is_symlink():
                return None
            if uploaded.resolve().parent != refs_dir.resolve():
                return None
            sheet = match.group("sheet").replace("''", "'")
            rect = parse_ref(match.group("cell"), default_sheet=sheet)
            if rect is None or rect.ncells != 1:
                return None
            coordinate = a1(rect.r1, rect.c1)
            nid = f"external:{filename}:{_id(sheet, coordinate)}"
            if nid in nodes:
                return nodes[nid]
            node = {
                "id": nid,
                "kind": "input",
                "external": True,
                "sheet": f"[{filename}]{sheet}",
                "sourceSheet": sheet,
                "sourceFile": filename,
                "cell": coordinate,
                "row": rect.r1,
                "column": rect.c1,
                "reference": text,
                "formula": None,
                "value": None,
                "cached_value": None,
                "cache_status": "unknown",
                "provenance": "reference_input",
                "dependencies": [],
                "diagnostics": [],
            }
            try:
                if filename not in external_books:
                    external_books[filename] = load_workbook(
                        uploaded, data_only=False, keep_links=False
                    )
                other = external_books[filename]
                cell = other[sheet][coordinate]
                if cell.data_type == "f":
                    raise ValueError("External formula inputs are not supported.")
                if cell.data_type == "e" and cell.value not in _INPUT_ERROR_KINDS:
                    raise ValueError("Unknown external Excel error input.")
                for member in other[sheet]._cells.values():
                    region_text = getattr(member.value, "ref", None)
                    region = parse_ref(region_text, sheet) if region_text else None
                    if (
                        region
                        and region.r1 <= cell.row <= region.r2
                        and region.c1 <= cell.column <= region.c2
                    ):
                        raise ValueError(
                            "External array result is not a literal input."
                        )
                if isinstance(cell.value, (dt.datetime, dt.date, dt.time)):
                    raise ValueError("External date inputs are not yet supported.")
                literal = cell.value
                if literal is None and cell.data_type in {"inlineStr", "s", "str"}:
                    literal = ""
                node.update(
                    value=_json(literal),
                    cached_value=_json(literal),
                    sourceDataType=cell.data_type,
                    cachedValueKind="error" if cell.data_type == "e" else "scalar",
                    cache_status="saved",
                    sourceSha256=hashlib.sha256(uploaded.read_bytes()).hexdigest(),
                )
            except Exception as exc:
                node["diagnostics"].append(f"External reference unavailable: {exc}")
            nodes[nid] = node
            return node

        def dependency(text: str, sheet: str | None) -> str:
            external = external_input(text)
            if external is not None:
                return external["id"]
            prefix, name_body = split_sheet_prefix(text)
            scope = sheet_lookup.get(prefix.casefold()) if prefix is not None else sheet
            name_id = None
            if prefix is None or scope is not None:
                name_id = names.get((scope, name_body.casefold())) or names.get(
                    (None, name_body.casefold())
                )
            if name_id:
                return name_id
            rect = parse_ref(text, default_sheet=sheet)
            if rect and rect.sheet:
                rect = replace(rect, sheet=sheet_lookup.get(rect.sheet.casefold()))
            if rect and rect.sheet in book.sheetnames and "[" not in text:
                nid = rect.to_a1()
                if rect.ncells == 1:
                    nodes.setdefault(
                        nid,
                        {
                            "id": nid,
                            "kind": "input",
                            "sheet": rect.sheet,
                            "cell": a1(rect.r1, rect.c1),
                            "row": rect.r1,
                            "column": rect.c1,
                            "formula": None,
                            "value": None,
                            "cached_value": None,
                            "cache_status": "blank",
                            "provenance": "saved_blank",
                            "dependencies": [],
                            "diagnostics": [],
                        },
                    )
                else:
                    nodes.setdefault(
                        nid,
                        {
                            "id": nid,
                            "kind": "range",
                            "sheet": rect.sheet,
                            "reference": text,
                            "bounds": [rect.r1, rect.c1, rect.r2, rect.c2],
                            "cell_count": rect.ncells,
                            "dependencies": [],
                            "diagnostics": [],
                        },
                    )
                return nid
            nid = f"unresolved:{sheet or '*'}:{text}"
            nodes.setdefault(
                nid,
                {
                    "id": nid,
                    "kind": "unresolved",
                    "sheet": sheet,
                    "reference": text,
                    "dependencies": [],
                    "diagnostics": [f"Unresolved dependency: {text}"],
                },
            )
            return nid

        for node in list(nodes.values()):
            formula = node.get("formula")
            if not formula:
                continue
            try:
                tokens = Tokenizer(formula).items
                if len(tokens) <= 512:
                    node["formulaTokens"] = [
                        {
                            "value": token.value,
                            "type": token.type,
                            "subtype": token.subtype,
                        }
                        for token in tokens
                    ]
                for token in tokens:
                    if token.type == "FUNC" and token.subtype == "OPEN":
                        function = token.value[:-1].upper().removeprefix("_XLFN.")
                        if function in _VOLATILE:
                            node["volatile"] = True
                        if function in _DYNAMIC:
                            node["diagnostics"].append(
                                f"Dynamic/workbook-sensitive dependency: {function}"
                            )
                    if token.type == "OPERAND" and token.subtype == "RANGE":
                        dep = dependency(token.value, node.get("sheet"))
                        node.setdefault("reference_bindings", {})[token.value] = dep
                        if dep not in node["dependencies"]:
                            node["dependencies"].append(dep)
                        edges.add((dep, node["id"]))
            except Exception as exc:
                node["diagnostics"].append(f"Formula tokenization failed: {exc}")
        references = (
            sorted(p.name for p in refs_dir.iterdir() if p.is_file())
            if refs_dir and refs_dir.is_dir()
            else []
        )
        for node in nodes.values():
            if node["kind"] in {"cell", "input"}:
                for region in array_regions:
                    if (
                        node["sheet"] == region.sheet
                        and region.r1 <= node["row"] <= region.r2
                        and region.c1 <= node["column"] <= region.c2
                    ):
                        node["diagnostics"].append(
                            "Cell belongs to an unsupported array/data-table "
                            "result region."
                        )
            node["label"] = (
                node.get("name")
                or node.get("cell")
                or node.get("reference")
                or node["id"]
            )
            node["cachedValue"] = node.get("cached_value")
            node["valueSource"] = node.get("provenance", "unknown")
        patterns = _formula_patterns(nodes)
        return {
            "nodes": list(nodes.values()),
            "formulaPatterns": patterns,
            "edges": [
                {"source": s, "target": t, "kind": "dep"} for s, t in sorted(edges)
            ],
            "meta": {
                "mode": "structural",
                "evaluated": False,
                "patternSemantics": (
                    "Sheet-local copied-formula signatures: A1 references become "
                    "relative R1C1 offsets with absolute/mixed anchors preserved. "
                    "Names retain their scope, strings and operators retain their "
                    "tokens. Formulas with structural diagnostics or ambiguous "
                    "references remain isolated. "
                    "Groups describe matching templates, not equal values or "
                    "proof of a historical copy action. Member ranges tile actual "
                    "cells without filling gaps. No formula evaluation at import."
                ),
                "formulaPatternCount": len(patterns),
                "copiedPatternCount": sum(p["memberCount"] > 1 for p in patterns),
                "copiedFormulaCells": sum(
                    p["memberCount"] for p in patterns if p["memberCount"] > 1
                ),
                "filename": path.name,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "sheets": book.sheetnames,
                "sheetDetails": sheet_details,
                "contextSemantics": "Literal nearby text snippets, not inferred "
                "labels. At most one text cell within four cells left/above, "
                "each capped at 240 characters. Formula and calculation "
                "inputs are never truncated by this presentation limit.",
                "reference_files": references,
                "limits": [
                    "External single-cell literal inputs are supported from uploads; "
                    "external formulas/ranges, structured and dynamic references "
                    "remain unsupported.",
                    "Range members are expanded sparsely only for "
                    "requested calculations.",
                    "A missing saved formula value is unknown, never zero.",
                ],
            },
        }
    finally:
        book.close()
        cache.close()
        for other in external_books.values():
            other.close()


def _absolute_name(node: dict) -> str | None:
    reference = node["formula"][1:]
    if "[" in reference:
        return None
    rect = parse_ref(reference, default_sheet=node.get("sheet"))
    if rect is None or rect.sheet is None:
        return None
    body = reference.rsplit("!", 1)[-1]
    absolute = r"(?:\$[A-Za-z]{1,3}\$[0-9]+|\$[A-Za-z]{1,3}|\$[0-9]+)"
    if not re.fullmatch(absolute + r"(?::" + absolute + r")?", body):
        return None
    return rect.to_a1()


def _name_literal(node: dict) -> tuple[str, Any, str] | None:
    """A defined scalar literal, never a formula or relative reference.

    Parentheses keep a signed number or error a single operand when substituted
    into a consuming formula. The source name/formula remains in the graph.
    """
    body = node["formula"][1:].strip()
    if re.fullmatch(r'"(?:[^"]|"")*"', body):
        return f"({body})", body[1:-1].replace('""', '"'), "scalar"
    if body.upper() in {"TRUE", "FALSE"}:
        return f"({body.upper()})", body.upper() == "TRUE", "scalar"
    if body in EXCEL_ERRORS:
        return f"({body})", body, "error"
    if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", body):
        value = float(body)
        if math.isfinite(value):
            return f"({body})", value, "scalar"
    return None


def _name_expression(node: dict) -> str | None:
    reference = _absolute_name(node)
    if reference is not None:
        return reference
    literal = _name_literal(node)
    return literal[0] if literal is not None else None


def _cache_comparison(
    node: dict, raw: Any, normalized: Any, epoch: dt.datetime
) -> dict:
    comparison_value = (
        to_excel(raw, epoch)
        if isinstance(raw, (dt.datetime, dt.date, dt.time))
        else normalized
    )
    comparison = "unknown"
    if node.get("cache_status") == "saved":
        saved = node.get("cached_comparison_value", node.get("cached_value"))
        same_boolean_type = isinstance(saved, bool) == isinstance(
            comparison_value, bool
        )
        same_error_type = (node.get("cachedValueKind") == "error") == (
            _excel_error_text(raw) is not None
        )
        comparison = (
            "equal"
            if same_boolean_type and same_error_type and saved == comparison_value
            else "different"
        )
    return {"comparisonValue": comparison_value, "comparison": comparison}


def evaluate_node(path: Path, node_id: str, refs_dir: Path | None = None) -> dict:
    """Evaluate a complete target closure in the caller's isolated process.

    This function itself does not create a process; the API task supervisor
    supplies isolation, cancellation, memory and time budgets.
    """
    graph = build_structure(path, refs_dir)
    nodes = {node["id"]: node for node in graph["nodes"]}
    if node_id not in nodes:
        raise ValueError(f"Unknown node: {node_id}")
    target = nodes[node_id]
    result: dict[str, Any] = {
        "node_id": node_id,
        "nodeId": node_id,
        "patternId": target.get("patternId"),
        "status": "unsupported",
        "value": None,
        "cachedValue": target.get("cached_value"),
        "cached_value": target.get("cached_value"),
        "comparison": "unknown",
        "provenance": None,
        "steps": [],
        "diagnostics": [],
    }
    if target["kind"] not in {"cell", "input"}:
        result["diagnostics"] = ["Select a single cell for recalculation."]
        return result
    closure: set[str] = set()
    pending = [node_id]
    reachable_deps: dict[str, list[str]] = {}
    while pending:
        nid = pending.pop()
        if nid in closure:
            continue
        closure.add(nid)
        node = nodes[nid]
        if node["diagnostics"]:
            result["diagnostics"].extend(
                f"{nid}: {message}" for message in node["diagnostics"]
            )
        reachable_deps[nid] = list(node["dependencies"])
        if node["kind"] == "range":
            r1, c1, r2, c2 = node["bounds"]
            members = [
                key
                for key, cell in nodes.items()
                if cell["kind"] in {"cell", "input"}
                and cell.get("sheet") == node["sheet"]
                and r1 <= cell["row"] <= r2
                and c1 <= cell["column"] <= c2
            ]
            pending.extend(members)
            reachable_deps[nid].extend(members)
        elif node["kind"] == "name":
            if _name_expression(node) is None:
                result["diagnostics"].append(
                    f"{nid}: Only absolute reference or scalar literal defined names "
                    "support recalculation."
                )
        pending.extend(node["dependencies"])
    result["volatile"] = any(nodes[nid].get("volatile", False) for nid in closure)
    result["dependency_count"] = len(closure)
    # Topological removal detects cycles through compact ranges too. Never
    # pretend to reproduce Excel iterative-calculation settings in a new engine.
    indegrees = {nid: 0 for nid in closure}
    for deps in reachable_deps.values():
        for dep in deps:
            indegrees[dep] += 1
    ready = [nid for nid, degree in indegrees.items() if degree == 0]
    visited = 0
    while ready:
        nid = ready.pop()
        visited += 1
        for dep in reachable_deps[nid]:
            indegrees[dep] -= 1
            if indegrees[dep] == 0:
                ready.append(dep)
    if visited != len(closure):
        result["diagnostics"].append(
            "Circular dependencies require unsupported iteration semantics."
        )
    steps: list[dict[str, Any]] = [
        {
            "node_id": nid,
            "nodeId": nid,
            "label": nodes[nid].get("label", nid),
            "formula": nodes[nid].get("formula"),
            "cachedValue": nodes[nid].get("cachedValue"),
            "cachedValueKind": nodes[nid].get("cachedValueKind", "scalar"),
            "valueSource": nodes[nid].get("valueSource", "unknown"),
            "numberFormat": nodes[nid].get("numberFormat"),
            "dependencies": reachable_deps[nid],
            "kind": nodes[nid]["kind"],
            "comparison": "unknown",
            "evaluationStatus": (
                "not_evaluated"
                if nodes[nid]["kind"] == "cell"
                else "input"
                if nodes[nid]["kind"] == "input"
                else "structural"
            ),
            "diagnostics": list(nodes[nid].get("diagnostics", [])),
        }
        for nid in sorted(closure)
    ]
    result["steps"] = steps
    result["coverage"] = {
        "status": "not_evaluated",
        "totalNodes": len(closure),
        "formulaCells": sum(nodes[nid]["kind"] == "cell" for nid in closure),
        "calculatedFormulaCells": 0,
        "unavailableFormulaCells": 0,
        "unsupportedFormulaCells": 0,
        "notEvaluatedFormulaCells": sum(
            nodes[nid]["kind"] == "cell" for nid in closure
        ),
        "inputCells": sum(nodes[nid]["kind"] == "input" for nid in closure),
        "structuralNodes": sum(
            nodes[nid]["kind"] not in {"cell", "input"} for nid in closure
        ),
    }
    result["step_semantics"] = (
        "Static dependency closure, including the target and sparse range members. "
        "Formula values are read from engine memory after one target evaluation, "
        "without evaluating steps again or substituting saved formula caches. "
        "The engine may calculate formulas from unselected IF branches; this is "
        "not an execution trace or proof that every value was used by the target. "
        "Inputs are source values; ranges and names are structural. A missing "
        "engine value is unavailable, not zero or the saved cache."
    )
    if result["diagnostics"]:
        return result

    if target["kind"] == "input":
        # Literal inputs already are deterministic source values. Reading one
        # needs neither a formula engine nor conversion of an error to text.
        value = target.get("value")
        result.update(
            status="completed",
            value=value,
            valueKind=target.get("cachedValueKind", "scalar"),
            comparison="equal",
            provenance="input_read",
        )
        result["coverage"]["status"] = "complete"
        return result

    # Native parser and engine deliberately imported only after closure checks.
    import formualizer as fz

    from linexcel.engine import _register_legacy_normal_functions, is_too_deep
    from linexcel.excel_compat import register_excel_functions, rewrite_formula

    for nid in closure:
        formula = nodes[nid].get("formula")
        if formula and is_too_deep(formula):
            result["diagnostics"].append(
                f"{nid}: Formula exceeds the native engine's safe complexity limit."
            )
    if result["diagnostics"]:
        return result

    source = load_workbook(path, data_only=False, keep_links=False)
    try:
        config = fz.EvaluationConfig()
        config.enable_parallel = False
        config.date_system = "1904" if source.epoch.year == 1904 else "1900"
        engine = fz.Workbook(config=fz.WorkbookConfig(eval_config=config))
        register_excel_functions(engine, config.date_system)
        _register_legacy_normal_functions(engine, config.date_system)
        for sheet in source.sheetnames:
            engine.add_sheet(sheet)
        external_sheets = {}
        occupied = set(source.sheetnames)
        for nid in sorted(closure):
            node = nodes[nid]
            if node.get("external") and node["sheet"] not in external_sheets:
                name = (
                    "_linexcel_"
                    + hashlib.sha256(node["sheet"].encode()).hexdigest()[:16]
                )
                while name in occupied:
                    name += "_"
                occupied.add(name)
                engine.add_sheet(name)
                external_sheets[node["sheet"]] = name
        for nid in sorted(closure):
            node = nodes[nid]
            if node["kind"] not in {"cell", "input"}:
                continue
            sheet = external_sheets.get(node["sheet"], node["sheet"])
            row, column = node["row"], node["column"]
            if node.get("formula"):
                tokens = Tokenizer(node["formula"]).items
                for token in tokens:
                    if token.type == "OPERAND" and token.subtype == "RANGE":
                        dep = node.get("reference_bindings", {}).get(token.value)
                        named = nodes.get(dep, {})
                        if named.get("external"):
                            token.value = _id(
                                external_sheets[named["sheet"]], named["cell"]
                            )
                        elif named.get("kind") == "name":
                            replacement = _name_expression(named)
                            assert replacement is not None
                            token.value = replacement
                        elif named.get("sheet"):
                            # Canonical source sheet spelling also fixes native
                            # case-sensitive resolution of Excel references.
                            rect = parse_ref(token.value, default_sheet=sheet)
                            if rect is not None:
                                token.value = replace(
                                    rect, sheet=named["sheet"]
                                ).to_a1()
                engine.set_formula(
                    sheet,
                    row,
                    column,
                    rewrite_formula("=" + "".join(t.value for t in tokens), sheet),
                )
            else:
                value = (
                    node["value"]
                    if node.get("external")
                    else source[sheet][node["cell"]].value
                )
                if value is None and not node.get("external"):
                    if source[sheet][node["cell"]].data_type in {
                        "inlineStr",
                        "s",
                        "str",
                    }:
                        value = ""
                if isinstance(value, (dt.datetime, dt.date, dt.time)):
                    value = node["cached_comparison_value"]
                if node.get("sourceDataType") == "e":
                    value = {"type": "Error", "kind": _INPUT_ERROR_KINDS[value]}
                if value is not None:
                    engine.set_value(sheet, row, column, value)
        value = engine.evaluate_cell(
            external_sheets.get(target["sheet"], target["sheet"]),
            target["row"],
            target["column"],
        )
        # get_value reads the current engine cache; it does not open another
        # evaluation request. In particular, never call evaluate_cell per step:
        # that would recalculate volatile functions and break snapshot coherence.
        for step in steps:
            if step["kind"] != "cell":
                continue
            node = nodes[step["nodeId"]]
            result["coverage"]["notEvaluatedFormulaCells"] -= 1
            try:
                raw = (
                    value
                    if step["nodeId"] == node_id
                    else engine.get_value(
                        external_sheets.get(node["sheet"], node["sheet"]),
                        node["row"],
                        node["column"],
                    )
                )
            except (ValueError, RuntimeError) as exc:
                step["evaluationStatus"] = "unavailable"
                step["diagnostics"].append(
                    "Could not read this operation's engine value: "
                    f"{type(exc).__name__}: {exc}"
                )
                result["coverage"]["unavailableFormulaCells"] += 1
                continue
            if raw is None and step["nodeId"] != node_id:
                step["evaluationStatus"] = "unavailable"
                step["diagnostics"].append(
                    "The engine did not expose a value after this target evaluation. "
                    "A blank result cannot be distinguished from an uncomputed cell."
                )
                result["coverage"]["unavailableFormulaCells"] += 1
                continue
            try:
                calculated = _engine_value(raw)
            except ValueError as exc:
                step["evaluationStatus"] = "unsupported"
                step["diagnostics"].append(str(exc))
                result["coverage"]["unsupportedFormulaCells"] += 1
                continue
            step.update(
                evaluationStatus="calculated",
                calculatedValue=calculated,
                calculatedValueKind=(
                    "error"
                    if _excel_error_text(raw) is not None
                    else "array"
                    if isinstance(calculated, list)
                    else "scalar"
                ),
                calculatedProvenance="targeted_recalculation",
            )
            step.update(_cache_comparison(node, raw, calculated, source.epoch))
            result["coverage"]["calculatedFormulaCells"] += 1
        result["coverage"]["status"] = (
            "complete"
            if result["coverage"]["calculatedFormulaCells"]
            == result["coverage"]["formulaCells"]
            else "partial"
        )
        try:
            normalized = _engine_value(value)
        except ValueError as exc:
            result["diagnostics"].append(str(exc))
            return result
        result.update(
            status="completed", value=normalized, provenance="targeted_recalculation"
        )
        result.update(_cache_comparison(target, value, normalized, source.epoch))
        result["valueKind"] = (
            "error"
            if _excel_error_text(value) is not None
            else "array"
            if isinstance(normalized, list)
            else "scalar"
        )
        return result
    finally:
        source.close()
