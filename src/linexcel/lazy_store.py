"""Disk-backed, sparse OOXML index for the interactive lazy viewer.

ZIP inflation and XML parsing use the standard library's native zlib/Expat
implementations; SQLite owns the indexes and the bounded page cache. Only one
stored cell is decoded at a time. In particular shared strings, copied formula
members and worksheet row objects never accumulate in Python memory.
"""

from __future__ import annotations

import hashlib
import json
import math
import os
import posixpath
import re
import sqlite3
import tempfile
import zipfile
import zlib
from collections.abc import Mapping
from dataclasses import replace
from pathlib import Path, PureWindowsPath
from typing import Any
from urllib.parse import quote, unquote
from xml.etree import ElementTree as ET

from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.datetime import (
    CALENDAR_MAC_1904,
    CALENDAR_WINDOWS_1900,
    from_excel,
    from_ISO8601,
    to_excel,
)

from linexcel.refs import a1, parse_ref, split_sheet_prefix

SCHEMA_VERSION = 1
MAX_PAGE_SIZE = 200
_NODE_KEYS = (
    "id",
    "kind",
    "sheet",
    "cell",
    "row",
    "column",
    "formula",
    "numberFormat",
    "sourceDataType",
    "context",
    "value",
    "cached_value",
    "cachedValueKind",
    "cached_comparison_value",
    "cache_status",
    "provenance",
    "dependencies",
    "diagnostics",
    "reference_bindings",
    "volatile",
)
_SHORT_KEYS = {key: chr(65 + i) for i, key in enumerate(_NODE_KEYS)}
_LONG_KEYS = {short: key for key, short in _SHORT_KEYS.items()}
_XSTRING_ESCAPE = re.compile(r"_x([0-9A-Fa-f]{4})_")


def _decode_xstring(value: str) -> str:
    """Decode OOXML ST_Xstring once, preserving escaped literal spellings.

    For example ``_x005F_x0041_`` means the literal ``_x0041_``, not ``A``.
    UTF-16 surrogate pairs encode one Unicode character; an unpaired surrogate
    is invalid source data, not a replacement character or a readable value.
    """
    if "_x" not in value:
        return value
    decoded = _XSTRING_ESCAPE.sub(lambda match: chr(int(match[1], 16)), value)
    if any(0xD800 <= ord(character) <= 0xDFFF for character in decoded):
        decoded = decoded.encode("utf-16-le", "surrogatepass").decode("utf-16-le")
    return decoded


def _dump(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def _pack(node):
    # Per-cell native compression and short field names avoid writing several
    # gigabytes of repeated schema, aliases and saved values for a large sheet.
    packed = {}
    for key, value in node.items():
        if key in {
            "label",
            "cachedValue",
            "valueSource",
            "formulaTokens",
            "patternId",
            "patternMemberCount",
        }:
            continue
        if (
            key in {"cached_value", "cached_comparison_value"}
            and value == node.get("value")
            and type(value) is type(node.get("value"))
        ):
            continue
        packed[_SHORT_KEYS.get(key, key)] = value
    return zlib.compress(_dump(packed).encode(), level=1)


def _unpack(data):
    packed = json.loads(zlib.decompress(data))
    node = {_LONG_KEYS.get(key, key): value for key, value in packed.items()}
    if "value" in node:
        node.setdefault("cached_value", node["value"])
        node.setdefault("cached_comparison_value", node["value"])
    return node


def file_sha256(path: Path) -> str:
    with path.open("rb") as source:
        return hashlib.file_digest(source, "sha256").hexdigest()


def _part(part: str, target: str) -> str:
    return posixpath.normpath(
        target.lstrip("/")
        if target.startswith("/")
        else posixpath.join(posixpath.dirname(part), target)
    )


def _relationships(archive, part):
    name = posixpath.join(
        posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels"
    )
    if name not in archive.namelist():
        return {}
    return {
        r.get("Id"): (r.get("Type", ""), r.get("Target", ""), r.get("TargetMode"))
        for r in ET.fromstring(archive.read(name))
    }


def _rid(element):
    return next(
        (v for k, v in element.attrib.items() if k.rsplit("}", 1)[-1] == "id"), None
    )


def _stream_elements(archive, member: str, tag: str):
    """Yield complete records and detach them, including empty worksheet rows."""
    with archive.open(member) as source:
        stack = []
        for event, element in ET.iterparse(source, events=("start", "end")):
            if event == "start":
                stack.append(element)
                continue
            local = element.tag.rsplit("}", 1)[-1]
            if local == tag:
                yield element
                if len(stack) > 1:
                    stack[-2].remove(element)
                element.clear()
            elif local == "row":
                if len(stack) > 1:
                    stack[-2].remove(element)
                element.clear()
            stack.pop()


class WorkbookIndex(Mapping):
    """A read-through node mapping; iteration streams rows, never a workbook."""

    def __init__(self, path: Path, *, writable=False):
        self.path = Path(path)
        uri_path = str(Path(path).resolve())
        if os.name == "nt":
            if not uri_path.startswith("\\\\?\\"):
                uri_path = (
                    "\\\\?\\UNC\\" + uri_path[2:]
                    if uri_path.startswith("\\\\")
                    else "\\\\?\\" + uri_path
                )
            # Path.as_uri interprets \\?\ as a host. Encode the original Windows
            # pathname instead: SQLite needs its extended prefix past MAX_PATH.
            readonly_uri = "file:" + quote(uri_path, safe="/:") + "?mode=ro"
        else:
            readonly_uri = Path(uri_path).as_uri() + "?mode=ro"
        self.db = sqlite3.connect(
            uri_path if writable else readonly_uri,
            uri=not writable,
        )
        self.db.execute("PRAGMA cache_size=-8192")
        self.db.execute("PRAGMA temp_store=FILE")
        self._meta: dict[str, Any] | None = None

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.close()

    def close(self):
        self.db.close()

    @property
    def meta(self) -> dict[str, Any]:
        if self._meta is None:
            row = self.db.execute(
                "SELECT value FROM metadata WHERE key='meta'"
            ).fetchone()
            self._meta = json.loads(row[0]) if row else {}
        return self._meta

    def __len__(self):
        return self.db.execute("SELECT count(*) FROM nodes").fetchone()[0]

    def __iter__(self):
        return (row[0] for row in self.db.execute("SELECT id FROM nodes ORDER BY seq"))

    def __contains__(self, key):
        return (
            self.db.execute("SELECT 1 FROM nodes WHERE id=?", (key,)).fetchone()
            is not None
        )

    def __getitem__(self, key):
        row = self.db.execute(
            "SELECT data,pattern_id FROM nodes WHERE id=?", (key,)
        ).fetchone()
        if row is None:
            raise KeyError(key)
        return self._decode(row)

    def _decode(self, row, *, context=False):
        node = _unpack(row[0])
        node["label"] = (
            node.get("name") or node.get("cell") or node.get("reference") or node["id"]
        )
        node["cachedValue"] = node.get("cached_value")
        node["valueSource"] = node.get("provenance", "unknown")
        if row[1]:
            node["patternId"] = row[1]
            count = self.db.execute(
                "SELECT member_count FROM patterns WHERE id=?", (row[1],)
            ).fetchone()
            if count:
                node["patternMemberCount"] = count[0]
        if context and node.get("cell") and not node.get("external"):
            nearby = []
            for position in ("left", "above"):
                for distance in range(1, 5):
                    r = node["row"] - (distance if position == "above" else 0)
                    c = node["column"] - (distance if position == "left" else 0)
                    hit = self.db.execute(
                        (
                            "SELECT data FROM nodes WHERE sheet=? "
                            "AND row=? AND col=? AND "
                            "kind='input'"
                        ),
                        (node["sheet"], r, c),
                    ).fetchone()
                    if hit:
                        other = _unpack(hit[0])
                        value = other.get("value")
                        if (
                            isinstance(value, str)
                            and value.strip()
                            and other.get("sourceDataType") in {"s", "str", "inlineStr"}
                        ):
                            nearby.append(
                                {
                                    "cell": other["cell"],
                                    "text": value[:240],
                                    "position": position,
                                    "truncated": len(value) > 240,
                                }
                            )
                            break
            node.setdefault("context", {})["nearbyLabels"] = nearby
        if context and node.get("formula"):
            try:
                tokens = Tokenizer(node["formula"]).items
                if len(tokens) <= 512:
                    node["formulaTokens"] = [
                        {"value": t.value, "type": t.type, "subtype": t.subtype}
                        for t in tokens
                    ]
            except Exception:
                pass  # The import diagnostics already describe the failure.
        return node

    def values(self):
        for row in self.db.execute("SELECT data,pattern_id FROM nodes ORDER BY seq"):
            yield self._decode(row)

    def items(self):
        for node in self.values():
            yield node["id"], node

    def put(self, node, *, only_if_absent=False):
        sql = "INSERT OR IGNORE" if only_if_absent else "INSERT"
        suffix = (
            ""
            if only_if_absent
            else (
                " ON CONFLICT(id) DO UPDATE SET "
                "data=excluded.data,pattern_id=excluded.pattern_id"
            )
        )
        self.db.execute(
            sql
            + (
                " INTO nodes(id,kind,sheet,row,col,data,pattern_id,formula,search_val"
                "ue) VALUES (?,?,?,?,?,?,?,?,?)"
            )
            + suffix,
            (
                node["id"],
                node["kind"],
                node.get("sheet"),
                node.get("row"),
                node.get("column"),
                _pack(node),
                node.get("patternId"),
                node.get("formula"),
                str(node["value"]) if node.get("value") is not None else None,
            ),
        )

    def range_members(self, sheet, bounds, *, kind=None):
        r1, c1, r2, c2 = bounds
        clause = "kind=?" if kind else "kind IN ('cell','input')"
        arguments = (sheet, r1, r2, c1, c2) + ((kind,) if kind else ())
        for row in self.db.execute(
            "SELECT id FROM nodes WHERE sheet=? AND row BETWEEN ? AND ? "
            f"AND col BETWEEN ? AND ? AND {clause} ORDER BY row,col",
            arguments,
        ):
            yield row[0]

    def prepare_input_selection(self):
        self.db.execute(
            "CREATE TEMP TABLE IF NOT EXISTS selected_inputs "
            "(id TEXT PRIMARY KEY) WITHOUT ROWID"
        )
        self.db.execute("DELETE FROM selected_inputs")

    def select_input(self, node_id):
        self.db.execute("INSERT OR IGNORE INTO selected_inputs VALUES (?)", (node_id,))

    def select_range_inputs(self, sheet, bounds):
        r1, c1, r2, c2 = bounds
        self.db.execute(
            "INSERT OR IGNORE INTO selected_inputs SELECT id FROM nodes "
            "WHERE sheet=? AND row BETWEEN ? AND ? AND col BETWEEN ? AND ? "
            "AND kind='input'",
            (sheet, r1, r2, c1, c2),
        )

    def selected_inputs(self):
        for (data,) in self.db.execute(
            "SELECT n.data FROM selected_inputs s JOIN nodes n ON n.id=s.id "
            "ORDER BY n.sheet,n.row,n.col"
        ):
            yield _unpack(data)

    def selected_input_count(self):
        return self.db.execute("SELECT count(*) FROM selected_inputs").fetchone()[0]

    def selected_input_sample(self, limit):
        return [
            r[0]
            for r in self.db.execute(
                "SELECT id FROM selected_inputs ORDER BY id LIMIT ?", (limit,)
            )
        ]

    def node(self, node_id):
        row = self.db.execute(
            "SELECT data,pattern_id FROM nodes WHERE id=?", (node_id,)
        ).fetchone()
        return self._decode(row, context=True) if row else None


def _schema(db):
    db.executescript("""
        PRAGMA journal_mode=DELETE;
        PRAGMA synchronous=NORMAL;
        CREATE TABLE metadata(key TEXT PRIMARY KEY,value TEXT NOT NULL);
        CREATE TABLE nodes(seq INTEGER PRIMARY KEY,id TEXT UNIQUE NOT NULL,
            kind TEXT NOT NULL,sheet TEXT,row INTEGER,col INTEGER,
            data BLOB NOT NULL,pattern_id TEXT,formula TEXT,search_value TEXT);
        CREATE INDEX nodes_position ON nodes(sheet,row,col);
        CREATE INDEX nodes_kind ON nodes(kind,seq);
        CREATE INDEX nodes_pattern ON nodes(pattern_id,row,col);
        CREATE TABLE edges(source TEXT NOT NULL,target TEXT NOT NULL,
            PRIMARY KEY(source,target)) WITHOUT ROWID;
        CREATE INDEX edges_target ON edges(target);
        CREATE TABLE patterns(id TEXT PRIMARY KEY,sheet TEXT NOT NULL,
            member_count INTEGER NOT NULL,data TEXT NOT NULL);
        CREATE TABLE strings(id INTEGER PRIMARY KEY,value TEXT NOT NULL);
        CREATE TABLE shared(sheet TEXT NOT NULL,id TEXT NOT NULL,
            origin TEXT NOT NULL,formula TEXT NOT NULL,
            PRIMARY KEY(sheet,id)) WITHOUT ROWID;
        CREATE TABLE arrays(sheet TEXT NOT NULL,r1 INTEGER,c1 INTEGER,
            r2 INTEGER,c2 INTEGER);
    """)


def _read_source(path, store):
    """Index literals and formula text in one forward XML pass per sheet."""
    from linexcel.lazy import _INPUT_ERROR_KINDS, _id, _json, _name_literal

    db = store.db
    names = {}
    links = {}
    diagnostics = []
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = _relationships(archive, "xl/workbook.xml")
        prop = workbook.find("{*}workbookPr")
        epoch1904 = prop is not None and prop.get("date1904", "0") in {"1", "true"}
        epoch = CALENDAR_MAC_1904 if epoch1904 else CALENDAR_WINDOWS_1900
        formats = ["General"]
        style_part = next(
            (
                _part("xl/workbook.xml", v[1])
                for v in rels.values()
                if v[0].endswith("/styles")
            ),
            "xl/styles.xml",
        )
        if style_part in archive.namelist():
            root = ET.fromstring(archive.read(style_part))
            custom = {
                int(e.attrib["numFmtId"]): e.get("formatCode", "General")
                for e in root.findall("{*}numFmts/{*}numFmt")
            }
            formats = [
                custom.get(
                    int(e.get("numFmtId", "0")),
                    BUILTIN_FORMATS.get(int(e.get("numFmtId", "0")), "General"),
                )
                for e in root.findall("{*}cellXfs/{*}xf")
            ] or ["General"]
        format_kinds = [(is_date_format(f), is_timedelta_format(f)) for f in formats]
        string_part = next(
            (
                _part("xl/workbook.xml", v[1])
                for v in rels.values()
                if v[0].endswith("/sharedStrings")
            ),
            "xl/sharedStrings.xml",
        )
        if string_part in archive.namelist():
            try:
                for index, element in enumerate(
                    _stream_elements(archive, string_part, "si")
                ):
                    # Phonetic guides are not displayed cell text.
                    value = "".join(
                        e.text or ""
                        for e in element.findall("{*}t") + element.findall("{*}r/{*}t")
                    )
                    try:
                        value = _decode_xstring(value)
                    except UnicodeError as exc:
                        diagnostics.append(
                            {
                                "part": string_part,
                                "message": "Invalid Unicode shared string "
                                f"{index}: {exc}",
                            }
                        )
                        continue
                    db.execute("INSERT INTO strings VALUES (?,?)", (index, value))
                    if index % 10000 == 0:
                        db.commit()
            except (ET.ParseError, zipfile.BadZipFile, OSError) as exc:
                diagnostics.append(
                    {
                        "part": string_part,
                        "message": "Shared-string table is incomplete; missing strings "
                        "are unavailable, never empty values: "
                        f"{type(exc).__name__}: {exc}",
                    }
                )
        sheets: list[dict[str, Any]] = []
        for item in workbook.findall("{*}sheets/{*}sheet"):
            relation = rels.get(_rid(item))
            sheets.append(
                {
                    "name": item.get("name"),
                    "state": item.get("state", "visible"),
                    "part": _part("xl/workbook.xml", relation[1]) if relation else None,
                    "formulas": 0,
                    "inputs": 0,
                    "status": "complete",
                    "diagnostics": [],
                }
            )
        for index, link in enumerate(
            workbook.findall("{*}externalReferences/{*}externalReference"), 1
        ):
            relation = rels.get(_rid(link))
            if relation:
                part = _part("xl/workbook.xml", relation[1])
                target = next(
                    (
                        v[1]
                        for v in _relationships(archive, part).values()
                        if v[0].endswith("/externalLinkPath")
                    ),
                    None,
                )
                if target:
                    links[str(index)] = PureWindowsPath(unquote(target)).name
        for definition in workbook.findall("{*}definedNames/{*}definedName"):
            local = definition.get("localSheetId")
            scope = (
                sheets[int(local)]["name"]
                if local is not None and int(local) < len(sheets)
                else None
            )
            name = definition.get("name", "")
            if not name:
                diagnostics.append({"message": "Unnamed defined-name entry skipped."})
                continue
            nid = f"name:{scope or '*'}:{name}"
            names[(scope, name.casefold())] = nid
            node = {
                "id": nid,
                "kind": "name",
                "name": name,
                "sheet": scope,
                "formula": "=" + (definition.text or "").lstrip("="),
                "value": None,
                "cached_value": None,
                "cache_status": "unknown",
                "dependencies": [],
                "diagnostics": [],
            }
            literal = _name_literal(node)
            if literal is not None:
                _, value, kind = literal
                node.update(value=value, valueKind=kind, provenance="defined_constant")
            store.put(node)
        for sheet in sheets:
            row_number = 0
            column_number = 0
            previous_address = None
            row_hidden = False
            hidden_columns = []
            try:
                if not sheet["part"]:
                    raise ValueError("Worksheet relationship is missing")
                with archive.open(sheet["part"]) as source:
                    stack = []
                    for event, element in ET.iterparse(source, events=("start", "end")):
                        tag = element.tag.rsplit("}", 1)[-1]
                        if event == "start":
                            stack.append(element)
                            if tag == "row":
                                try:
                                    row_number = int(element.get("r", row_number + 1))
                                except (ValueError, TypeError):
                                    row_number += 1
                                    sheet["status"] = "partial"
                                    message = (
                                        "Invalid row coordinate; omitted cell "
                                        "coordinates may be ambiguous."
                                    )
                                    sheet["diagnostics"].append(message)
                                    diagnostics.append(
                                        {"sheet": sheet["name"], "message": message}
                                    )
                                column_number = 0
                                previous_address = None
                                row_hidden = element.get("hidden", "0") in {"1", "true"}
                            elif tag == "col" and element.get("hidden", "0") in {
                                "1",
                                "true",
                            }:
                                try:
                                    hidden_columns.append(
                                        (
                                            int(element.get("min")),
                                            int(element.get("max")),
                                        )
                                    )
                                except (ValueError, TypeError):
                                    diagnostics.append(
                                        {
                                            "sheet": sheet["name"],
                                            "message": "Invalid hidden-column "
                                            "metadata; visibility is unavailable "
                                            "for that interval.",
                                        }
                                    )
                            continue
                        if tag == "c":
                            address = element.get("r")
                            try:
                                if address is None:
                                    previous_column = (
                                        coordinate_to_tuple(previous_address)[1]
                                        if previous_address
                                        else 0
                                    )
                                    address = a1(row_number, previous_column + 1)
                                previous_address = address
                            except (ValueError, TypeError, KeyError):
                                address = None
                            # Excel often persists millions of formatting-only
                            # cells. They are absent inputs, so detach immediately
                            # without coordinate/style/date allocation or SQL.
                            if len(element) == 0 and element.get("t", "n") == "n":
                                stack[-2].remove(element)
                                element.clear()
                                stack.pop()
                                continue
                            try:
                                if not isinstance(address, str):
                                    raise ValueError("Cell coordinate is unavailable")
                                row, column_number = coordinate_to_tuple(address)
                                if not (
                                    1 <= row <= 1048576 and 1 <= column_number <= 16384
                                ):
                                    raise ValueError(
                                        "Cell coordinate is outside Excel bounds"
                                    )
                                style = int(element.get("s", "0"))
                                if not 0 <= style < len(formats):
                                    raise ValueError("Cell style is missing or invalid")
                            except (ValueError, TypeError, KeyError) as exc:
                                message = (
                                    f"Cell metadata unreadable at {address}: {exc}"
                                )
                                try:
                                    if not isinstance(address, str):
                                        raise ValueError(
                                            "Cell coordinate is unavailable"
                                        )
                                    row, column_number = coordinate_to_tuple(address)
                                    if not (
                                        1 <= row <= 1048576
                                        and 1 <= column_number <= 16384
                                    ):
                                        raise ValueError("Invalid coordinate")
                                except (ValueError, TypeError, KeyError):
                                    sheet["status"] = "partial"
                                    sheet["diagnostics"].append(message)
                                    diagnostics.append(
                                        {"sheet": sheet["name"], "message": message}
                                    )
                                else:
                                    assert isinstance(address, str)
                                    formula_element = element.find("{*}f")
                                    formula = (
                                        "=" + (formula_element.text or "")
                                        if formula_element is not None
                                        else None
                                    )
                                    store.put(
                                        {
                                            "id": _id(sheet["name"], address),
                                            "kind": "cell"
                                            if formula is not None
                                            else "input",
                                            "sheet": sheet["name"],
                                            "cell": address,
                                            "row": row,
                                            "column": column_number,
                                            "formula": formula,
                                            "value": None,
                                            "cache_status": "unknown",
                                            "dependencies": [],
                                            "diagnostics": [message],
                                        }
                                    )
                                    sheet[
                                        "formulas" if formula is not None else "inputs"
                                    ] += 1
                                stack[-2].remove(element)
                                element.clear()
                                stack.pop()
                                continue
                            assert isinstance(address, str)
                            formula_element = element.find("{*}f")
                            formula = None
                            problems = []
                            if formula_element is not None:
                                formula = "=" + (formula_element.text or "")
                                form_type = formula_element.get("t")
                                if form_type == "shared":
                                    shared_id = formula_element.get("si", "")
                                    if formula_element.text:
                                        db.execute(
                                            "INSERT OR REPLACE INTO shared "
                                            "VALUES (?,?,?,?)",
                                            (
                                                sheet["name"],
                                                shared_id,
                                                address,
                                                formula,
                                            ),
                                        )
                                    else:
                                        shared = db.execute(
                                            "SELECT origin,formula FROM shared "
                                            "WHERE sheet=? AND id=?",
                                            (sheet["name"], shared_id),
                                        ).fetchone()
                                        if shared:
                                            try:
                                                formula = Translator(
                                                    shared[1], origin=shared[0]
                                                ).translate_formula(address)
                                            except Exception as exc:
                                                problems.append(
                                                    "Shared formula translation "
                                                    f"failed: {exc}"
                                                )
                                        else:
                                            problems.append(
                                                "Shared formula anchor is missing "
                                                "or occurs after this cell."
                                            )
                                elif form_type in {"array", "dataTable"}:
                                    rect = parse_ref(
                                        formula_element.get("ref", ""), sheet["name"]
                                    )
                                    if rect:
                                        db.execute(
                                            "INSERT INTO arrays VALUES (?,?,?,?,?)",
                                            (
                                                sheet["name"],
                                                rect.r1,
                                                rect.c1,
                                                rect.r2,
                                                rect.c2,
                                            ),
                                        )
                                    problems.append(
                                        "Array/data-table formula is not supported."
                                    )
                                    if not formula_element.text:
                                        formula = "=<array/data-table>"
                            cell_type = element.get("t", "n")
                            if cell_type not in {
                                "n",
                                "s",
                                "str",
                                "inlineStr",
                                "b",
                                "d",
                                "e",
                            }:
                                problems.append(
                                    f"Unsupported stored cell type: {cell_type}"
                                )
                            raw = element.findtext("{*}v")
                            if raw == "" and cell_type not in {"str", "inlineStr"}:
                                raw = None
                            number_format = (
                                formats[style]
                                if 0 <= style < len(formats)
                                else "General"
                            )
                            value = None
                            comparison = None
                            try:
                                if cell_type == "inlineStr":
                                    text_element = element.find("{*}is")
                                    value = (
                                        ""
                                        if text_element is None
                                        else "".join(
                                            e.text or ""
                                            for e in text_element.findall("{*}t")
                                            + text_element.findall("{*}r/{*}t")
                                        )
                                    )
                                    value = _decode_xstring(value)
                                elif cell_type == "s":
                                    if raw is None:
                                        raise ValueError(
                                            "Shared-string index is missing"
                                        )
                                    hit = db.execute(
                                        "SELECT value FROM strings WHERE id=?",
                                        (int(raw),),
                                    ).fetchone()
                                    if hit is None:
                                        raise ValueError(f"Missing shared string {raw}")
                                    value = hit[0]
                                elif cell_type in {"str", "e"}:
                                    value = (
                                        raw
                                        if raw is not None
                                        else ""
                                        if cell_type == "str" and formula is None
                                        else None
                                    )
                                    if cell_type == "str" and value is not None:
                                        value = _decode_xstring(value)
                                elif cell_type == "b" and raw is not None:
                                    if raw not in {"0", "1"}:
                                        raise ValueError(
                                            "Boolean cell value must be 0 or 1"
                                        )
                                    value = raw == "1"
                                elif cell_type == "d" and raw is not None:
                                    value = from_ISO8601(raw)
                                    comparison = to_excel(value, epoch)
                                elif raw is not None:
                                    value = (
                                        float(raw)
                                        if any(t in raw for t in ".eE")
                                        else int(raw)
                                    )
                                    if isinstance(value, float) and not math.isfinite(
                                        value
                                    ):
                                        raise ValueError(
                                            "Non-finite stored numeric value"
                                        )
                                    comparison = value
                                    if (
                                        style < len(format_kinds)
                                        and format_kinds[style][0]
                                    ):
                                        if not (not epoch1904 and 60 <= value < 61):
                                            value = from_excel(
                                                value,
                                                epoch,
                                                timedelta=format_kinds[style][1],
                                            )
                                if (
                                    cell_type == "e"
                                    and formula is None
                                    and value not in _INPUT_ERROR_KINDS
                                ):
                                    problems.append(
                                        "Unknown Excel error input is not "
                                        "supported for recalculation."
                                    )
                            except Exception as exc:
                                problems.append(
                                    "Stored cell value is unreadable: "
                                    f"{type(exc).__name__}: {exc}"
                                )
                                value = None
                                comparison = None
                            if (
                                formula is not None
                                or value is not None
                                or cell_type in {"inlineStr", "s", "str"}
                                or problems
                            ):
                                node = {
                                    "id": _id(sheet["name"], address),
                                    "kind": "cell" if formula is not None else "input",
                                    "sheet": sheet["name"],
                                    "cell": address,
                                    "row": row,
                                    "column": column_number,
                                    "formula": formula,
                                    "numberFormat": number_format,
                                    "sourceDataType": "f"
                                    if formula is not None
                                    else cell_type,
                                    "context": {
                                        "sheetState": sheet["state"],
                                        "rowHidden": row_hidden,
                                        "columnHidden": any(
                                            first <= column_number <= last
                                            for first, last in hidden_columns
                                        ),
                                    },
                                    "value": _json(value),
                                    "cached_value": _json(value),
                                    "cachedValueKind": "error"
                                    if cell_type == "e"
                                    else "scalar",
                                    "cached_comparison_value": comparison
                                    if comparison is not None
                                    else _json(value),
                                    "cache_status": "unknown"
                                    if formula is not None and value is None
                                    else "saved",
                                    "provenance": "saved_cache"
                                    if formula is not None
                                    else "saved_input",
                                    "dependencies": [],
                                    "diagnostics": problems,
                                }
                                store.put(node)
                                sheet[
                                    "formulas" if formula is not None else "inputs"
                                ] += 1
                                if (sheet["formulas"] + sheet["inputs"]) % 10000 == 0:
                                    db.commit()
                            stack[-2].remove(element)
                            element.clear()
                        elif tag == "row":
                            stack[-2].remove(element)
                            element.clear()
                        stack.pop()
            except (
                ET.ParseError,
                ValueError,
                KeyError,
                zipfile.BadZipFile,
                OSError,
            ) as exc:
                message = f"Worksheet import incomplete: {type(exc).__name__}: {exc}"
                sheet["status"] = "partial"
                sheet["diagnostics"].append(message)
                diagnostics.append({"sheet": sheet["name"], "message": message})
            sheet.pop("part", None)
            db.commit()
    return (
        {
            "sheets": [s["name"] for s in sheets],
            "sheetDetails": sheets,
            "dateSystem": "1904" if epoch1904 else "1900",
            "diagnostics": diagnostics,
        },
        names,
        links,
    )


def _index_dependencies(store, metadata, names, links, refs_dir):
    from linexcel.external import _EXTERNAL_REF_RE
    from linexcel.lazy import _DYNAMIC, _VOLATILE, _id, _pattern_signature

    db = store.db
    sheets = metadata["sheets"]
    lookup = {sheet.casefold(): sheet for sheet in sheets}
    partial = {s["name"] for s in metadata["sheetDetails"] if s["status"] != "complete"}
    external_books = {}
    with tempfile.TemporaryDirectory(
        prefix="linexcel-refs-", dir=store.path.parent
    ) as temporary:

        def external_input(text):
            match = _EXTERNAL_REF_RE.fullmatch(text)
            if not match or refs_dir is None:
                return None
            key = match.group("book")
            filename = links.get(key, key)
            if PureWindowsPath(filename).name != filename:
                return None
            uploaded = refs_dir / filename
            if (
                not uploaded.is_file()
                or uploaded.is_symlink()
                or uploaded.resolve().parent != refs_dir.resolve()
            ):
                return None
            sheet = match.group("sheet").replace("''", "'")
            rect = parse_ref(match.group("cell"), default_sheet=sheet)
            if rect is None or rect.ncells != 1:
                return None
            coordinate = a1(rect.r1, rect.c1)
            nid = f"external:{filename}:{_id(sheet, coordinate)}"
            if nid in store:
                return store[nid]
            node = {
                "id": nid,
                "kind": "input",
                "external": True,
                "sheet": f"[{filename}]{sheet}",
                "sourceSheet": sheet,
                "sourceFile": filename,
                "cell": coordinate,
                "row": rect.r1,
                "column": rect.c1,
                "reference": text,
                "formula": None,
                "value": None,
                "cached_value": None,
                "cache_status": "unknown",
                "provenance": "reference_input",
                "dependencies": [],
                "diagnostics": [],
            }
            try:
                if filename not in external_books:
                    external_path = Path(temporary) / f"{len(external_books)}.sqlite"
                    other = WorkbookIndex(external_path, writable=True)
                    _schema(other.db)
                    external_books[filename] = (other, {}, "")
                    detail, _, _ = _read_source(uploaded, other)
                    external_books[filename] = (other, detail, file_sha256(uploaded))
                other, detail, digest = external_books[filename]
                if sheet not in detail.get("sheets", []):
                    raise ValueError("External source sheet is unavailable.")
                if any(
                    s["name"] == sheet and s["status"] != "complete"
                    for s in detail["sheetDetails"]
                ):
                    raise ValueError("External source sheet import is incomplete.")
                member = other.get(_id(sheet, coordinate))
                if member and (member.get("formula") or member.get("diagnostics")):
                    raise ValueError(
                        "External formula or unreadable inputs are not supported."
                    )
                region = other.db.execute(
                    (
                        "SELECT 1 FROM arrays WHERE sheet=? "
                        "AND ? BETWEEN r1 AND r2 AND ? "
                        "BETWEEN c1 AND c2 LIMIT 1"
                    ),
                    (sheet, rect.r1, rect.c1),
                ).fetchone()
                if region:
                    raise ValueError("External array result is not a literal input.")
                if (
                    member
                    and member.get("numberFormat")
                    and is_date_format(member["numberFormat"])
                ):
                    raise ValueError("External date inputs are not yet supported.")
                literal = member.get("value") if member else None
                node.update(
                    value=literal,
                    cached_value=literal,
                    sourceDataType=member.get("sourceDataType", "n") if member else "n",
                    cachedValueKind=member.get("cachedValueKind", "scalar")
                    if member
                    else "scalar",
                    cache_status="saved",
                    sourceSha256=digest,
                )
            except Exception as exc:
                node["diagnostics"].append(f"External reference unavailable: {exc}")
            store.put(node)
            return node

        def dependency(text, sheet):
            external = external_input(text)
            if external is not None:
                return external["id"]
            prefix, body = split_sheet_prefix(text)
            scope = lookup.get(prefix.casefold()) if prefix is not None else sheet
            name_id = None
            if prefix is None or scope is not None:
                name_id = names.get((scope, body.casefold())) or names.get(
                    (None, body.casefold())
                )
            if name_id:
                return name_id
            rect = parse_ref(text, default_sheet=sheet)
            if rect and rect.sheet:
                rect = replace(rect, sheet=lookup.get(rect.sheet.casefold()))
            if rect and rect.sheet in sheets and "[" not in text:
                nid = rect.to_a1()
                if nid in store:
                    return nid
                if rect.ncells == 1:
                    node = {
                        "id": nid,
                        "kind": "input",
                        "sheet": rect.sheet,
                        "cell": a1(rect.r1, rect.c1),
                        "row": rect.r1,
                        "column": rect.c1,
                        "formula": None,
                        "value": None,
                        "cached_value": None,
                        "cache_status": "blank",
                        "provenance": "saved_blank",
                        "dependencies": [],
                        "diagnostics": [],
                    }
                else:
                    node = {
                        "id": nid,
                        "kind": "range",
                        "sheet": rect.sheet,
                        "reference": text,
                        "bounds": [rect.r1, rect.c1, rect.r2, rect.c2],
                        "cell_count": rect.ncells,
                        "dependencies": [],
                        "diagnostics": [],
                    }
                if rect.sheet in partial:
                    node["diagnostics"].append(
                        "Source worksheet is incomplete; missing cells "
                        "cannot be treated as "
                        "blank."
                    )
                store.put(node, only_if_absent=True)
                return nid
            nid = f"unresolved:{sheet or '*'}:{text}"
            store.put(
                {
                    "id": nid,
                    "kind": "unresolved",
                    "sheet": sheet,
                    "reference": text,
                    "dependencies": [],
                    "diagnostics": [f"Unresolved dependency: {text}"],
                },
                only_if_absent=True,
            )
            return nid

        try:
            # New blank/range nodes have no formulas. A keyset cursor stays bounded
            # and remains valid across commits and node updates.
            last = 0
            while True:
                batch = db.execute(
                    (
                        "SELECT seq,data FROM nodes WHERE seq>? "
                        "AND kind IN ('cell','name') "
                        "ORDER BY seq LIMIT 512"
                    ),
                    (last,),
                ).fetchall()
                if not batch:
                    break
                for seq, data in batch:
                    last = seq
                    node = _unpack(data)
                    if node["kind"] == "cell":
                        if db.execute(
                            (
                                "SELECT 1 FROM arrays WHERE sheet=? "
                                "AND ? BETWEEN r1 AND r2 AND ? "
                                "BETWEEN c1 AND c2 LIMIT 1"
                            ),
                            (node["sheet"], node["row"], node["column"]),
                        ).fetchone():
                            node["diagnostics"].append(
                                "Cell belongs to an unsupported array/data-table "
                                "result region."
                            )
                    try:
                        tokens = Tokenizer(node["formula"]).items
                        for token in tokens:
                            if token.type == "FUNC" and token.subtype == "OPEN":
                                function = (
                                    token.value[:-1].upper().removeprefix("_XLFN.")
                                )
                                if function in _VOLATILE:
                                    node["volatile"] = True
                                if function in _DYNAMIC:
                                    node["diagnostics"].append(
                                        "Dynamic/workbook-sensitive dependency: "
                                        f"{function}"
                                    )
                            if token.type == "OPERAND" and token.subtype == "RANGE":
                                dep = dependency(token.value, node.get("sheet"))
                                node.setdefault("reference_bindings", {})[
                                    token.value
                                ] = dep
                                if dep not in node["dependencies"]:
                                    node["dependencies"].append(dep)
                                db.execute(
                                    "INSERT OR IGNORE INTO edges VALUES (?,?)",
                                    (dep, node["id"]),
                                )
                    except Exception as exc:
                        node["diagnostics"].append(
                            f"Formula tokenization failed: {exc}"
                        )
                    if node["kind"] == "cell":
                        signature, tokens, reason = _pattern_signature(node, store)
                        identity = ["copied-r1c1-v1", node["sheet"], tokens]
                        if reason:
                            identity += ["isolated", node["id"], node["formula"]]
                        key = (
                            "formula-sha256:"
                            + hashlib.sha256(_dump(identity).encode()).hexdigest()
                        )
                        node["patternId"] = key
                        group = {
                            "id": key,
                            "sheet": node["sheet"],
                            "signature": signature,
                            "status": "isolated" if reason else "canonical",
                            "diagnostics": [reason] if reason else [],
                            "representativeId": node["id"],
                            "representativeFormula": node["formula"],
                        }
                        db.execute(
                            (
                                "INSERT INTO patterns VALUES (?,?,1,?) "
                                "ON CONFLICT(id) DO UPDATE SET "
                                "member_count=member_count+1"
                            ),
                            (key, node["sheet"], _dump(group)),
                        )
                    store.put(node)
                db.commit()
            # Arrays can include saved literal-looking cells. Mark them in the
            # index so target closures refuse cached array results consistently.
            for sheet, r1, c1, r2, c2 in db.execute("SELECT * FROM arrays"):
                for row in db.execute(
                    (
                        "SELECT data,pattern_id FROM nodes WHERE sheet=? "
                        "AND row BETWEEN ? "
                        "AND ? AND col BETWEEN ? AND ? AND kind='input'"
                    ),
                    (sheet, r1, r2, c1, c2),
                ):
                    node = store._decode(row)
                    node["diagnostics"].append(
                        "Cell belongs to an unsupported array/data-table result region."
                    )
                    store.put(node)
            db.commit()
        finally:
            for other, _, _ in external_books.values():
                other.close()


def build_index(path: Path, index_path: Path, refs_dir: Path | None = None) -> dict:
    "Publish a complete disk index atomically; preserve any previous index on failure."
    path, index_path = Path(path), Path(index_path)
    index_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(
        prefix=index_path.name + ".", suffix=".building", dir=index_path.parent
    )
    os.close(handle)
    temp_path = Path(temporary)
    try:
        with WorkbookIndex(temp_path, writable=True) as store:
            _schema(store.db)
            metadata, names, links = _read_source(path, store)
            _index_dependencies(store, metadata, names, links, refs_dir)
            counts = dict(
                store.db.execute("SELECT kind,count(*) FROM nodes GROUP BY kind")
            )
            per_sheet = {}
            for sheet, kind, count in store.db.execute(
                "SELECT sheet,kind,count(*) FROM nodes GROUP BY sheet,kind"
            ):
                per_sheet.setdefault(sheet or "defined names", {})[kind] = count
            pattern_count, copied, copied_cells = store.db.execute(
                "SELECT count(*),coalesce(sum(member_count>1),0),coalesce(sum(CASE "
                "WHEN member_count>1 THEN member_count ELSE 0 END),0) FROM patterns"
            ).fetchone()
            metadata.update(
                schemaVersion=SCHEMA_VERSION,
                mode="structural",
                evaluated=False,
                lazy=True,
                storage="sqlite",
                filename=path.name,
                sha256=file_sha256(path),
                nodeCount=sum(counts.values()),
                nodeCounts=counts,
                perSheetCounts=per_sheet,
                sheetNodeCounts={s: sum(c.values()) for s, c in per_sheet.items()},
                formulaCount=counts.get("cell", 0),
                inputCount=counts.get("input", 0),
                patternCount=pattern_count,
                formulaPatternCount=pattern_count,
                copiedPatternCount=copied,
                copiedFormulaCells=copied_cells,
                edgeCount=store.db.execute("SELECT count(*) FROM edges").fetchone()[0],
                status="partial" if metadata["diagnostics"] else "complete",
                reference_files=sorted(
                    p.name for p in refs_dir.iterdir() if p.is_file()
                )
                if refs_dir and refs_dir.is_dir()
                else [],
                patternSemantics=(
                    "Sheet-local copied-formula signatures: A1 references "
                    "become relative "
                    "R1C1 offsets with absolute/mixed anchors preserved. Names retain "
                    "their scope, strings and operators retain their tokens. Formulas "
                    "with structural diagnostics or ambiguous references "
                    "remain isolated. Groups describe matching templates, "
                    "not equal values or proof of a "
                    "historical copy action. Member ranges tile actual cells without "
                    "filling gaps. No formula evaluation at import."
                ),
                contextSemantics=(
                    "Literal nearby text snippets, not inferred labels. "
                    "At most one text "
                    "cell within four cells left/above, each capped at 240 characters. "
                    "Formula and calculation inputs are never truncated by this "
                    "presentation limit."
                ),
                limits=[
                    (
                        "External single-cell literal inputs are supported "
                        "from uploads; external formulas/ranges, structured "
                        "and dynamic references remain "
                        "unsupported."
                    ),
                    (
                        "Range members are expanded sparsely only for "
                        "requested calculations."
                    ),
                    "A missing saved formula value is unknown, never zero.",
                    (
                        "Browser pages and graphs are bounded views "
                        "of the complete disk "
                        "index."
                    ),
                ],
            )
            store.db.execute(
                "INSERT INTO metadata VALUES ('meta',?)", (_dump(metadata),)
            )
            store.db.commit()
        os.replace(temp_path, index_path)
    finally:
        temp_path.unlink(missing_ok=True)
    return graph_page(index_path)


def get_node(index_path: Path, node_id: str) -> dict | None:
    with WorkbookIndex(index_path) as store:
        return store.node(node_id)


def _filters(sheet=None, query=None):
    conditions, values = [], []
    if sheet:
        conditions.append("sheet=?")
        values.append(sheet)
    if query:
        # Literal substring search, with SQL metacharacters escaped.
        conditions.append(
            "(id LIKE ? ESCAPE '\\' OR formula LIKE ? ESCAPE '\\' OR search_value "
            "LIKE ? ESCAPE '\\')"
        )
        pattern = (
            "%"
            + query.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
            + "%"
        )
        values += [pattern] * 3
    return " WHERE " + " AND ".join(conditions) if conditions else "", values


def _edges_for_nodes(store, ids):
    if not ids:
        return []
    placeholders = ",".join("?" for _ in ids)
    rows = store.db.execute(
        f"SELECT source,target FROM edges WHERE target IN ({placeholders}) "
        f"AND source IN ({placeholders}) ORDER BY source,target",
        ids + ids,
    )
    return [{"source": s, "target": t, "kind": "dep"} for s, t in rows]


def _pattern(store, row, member_limit=MAX_PAGE_SIZE):
    from linexcel.lazy import _member_ranges

    data, count = row
    group = json.loads(data)
    members = [
        {"id": n, "row": r, "column": c}
        for n, r, c in store.db.execute(
            "SELECT id,row,col FROM nodes WHERE pattern_id=? ORDER BY row,col LIMIT ?",
            (group["id"], member_limit),
        )
    ]
    group.update(
        memberCount=count,
        memberIds=[n["id"] for n in members],
        ranges=_member_ranges(members) if members else [],
        membersTruncated=count > len(members),
    )
    return group


def graph_page(
    index_path: Path, *, offset=0, limit=200, sheet=None, query=None
) -> dict:
    offset, limit = max(0, int(offset)), max(1, min(MAX_PAGE_SIZE, int(limit)))
    with WorkbookIndex(index_path) as store:
        where, args = _filters(sheet, query)
        total = store.db.execute("SELECT count(*) FROM nodes" + where, args).fetchone()[
            0
        ]
        nodes = [
            store._decode(row, context=True)
            for row in store.db.execute(
                "SELECT data,pattern_id FROM nodes"
                + where
                + " ORDER BY seq LIMIT ? OFFSET ?",
                args + [limit, offset],
            )
        ]
        patterns = []
        for key in dict.fromkeys(n["patternId"] for n in nodes if n.get("patternId")):
            row = store.db.execute(
                "SELECT data,member_count FROM patterns WHERE id=?", (key,)
            ).fetchone()
            patterns.append(_pattern(store, row, member_limit=20))
        return {
            "lazy": True,
            "nodes": nodes,
            "edges": _edges_for_nodes(store, [n["id"] for n in nodes]),
            "formulaPatterns": patterns,
            "meta": store.meta,
            "pagination": {
                "offset": offset,
                "limit": limit,
                "total": total,
                "hasMore": offset + len(nodes) < total,
            },
        }


def patterns_page(index_path: Path, *, offset=0, limit=200, sheet=None) -> dict:
    offset, limit = max(0, int(offset)), max(1, min(MAX_PAGE_SIZE, int(limit)))
    with WorkbookIndex(index_path) as store:
        where, args = (" WHERE sheet=?", [sheet]) if sheet else ("", [])
        total = store.db.execute(
            "SELECT count(*) FROM patterns" + where, args
        ).fetchone()[0]
        groups = [
            _pattern(store, row, member_limit=20)
            for row in store.db.execute(
                "SELECT data,member_count FROM patterns"
                + where
                + (
                    " ORDER BY sheet,json_extract(data,'$.representativeId') LIMIT ? "
                    "OFFSET ?"
                ),
                args + [limit, offset],
            )
        ]
        return {
            "formulaPatterns": groups,
            "pagination": {
                "offset": offset,
                "limit": limit,
                "total": total,
                "hasMore": offset + len(groups) < total,
            },
        }


def graph_neighborhood(index_path: Path, node_id: str, limit=200, depth=1) -> dict:
    limit, depth = max(1, min(MAX_PAGE_SIZE, int(limit))), max(0, min(2, int(depth)))
    with WorkbookIndex(index_path) as store:
        if node_id not in store:
            raise ValueError(f"Unknown node: {node_id}")
        ids = {node_id: None}
        frontier = [node_id]
        truncated = False
        for _ in range(depth):
            following = []
            for current in frontier:
                rows = store.db.execute(
                    (
                        "SELECT source FROM edges WHERE target=? "
                        "UNION SELECT target FROM "
                        "edges WHERE source=? LIMIT ?"
                    ),
                    (current, current, limit + 1),
                )
                for row in rows:
                    if row[0] in ids:
                        continue
                    if len(ids) >= limit:
                        truncated = True
                        break
                    ids[row[0]] = None
                    following.append(row[0])
            frontier = following
        nodes = [store.node(nid) for nid in ids]
        return {
            "lazy": True,
            "nodes": nodes,
            "edges": _edges_for_nodes(store, list(ids)),
            "formulaPatterns": [
                _pattern(
                    store,
                    store.db.execute(
                        "SELECT data,member_count FROM patterns WHERE id=?", (key,)
                    ).fetchone(),
                    member_limit=20,
                )
                for key in dict.fromkeys(
                    n["patternId"] for n in nodes if n.get("patternId")
                )
            ],
            "meta": store.meta,
            "pagination": {
                "limit": limit,
                "returned": len(nodes),
                "truncated": truncated,
                "total": store.meta["nodeCount"],
            },
            "nodeId": node_id,
            "depth": depth,
        }


def evidence_graph(index_path: Path, node_id: str | None = None) -> dict:
    """Small deterministic evidence sample, with exact workbook-level counts."""
    if node_id:
        with WorkbookIndex(index_path) as store:
            node = store.node(node_id)
            if node is None:
                raise ValueError(f"Unknown node: {node_id}")
            ids = list(dict.fromkeys([node_id, *node.get("dependencies", [])[:12]]))
            groups = []
            if node.get("patternId"):
                row = store.db.execute(
                    "SELECT data,member_count FROM patterns WHERE id=?",
                    (node["patternId"],),
                ).fetchone()
                if row:
                    groups.append(_pattern(store, row, member_limit=20))
            return {
                "lazy": True,
                "nodes": [store.node(nid) for nid in ids],
                "edges": _edges_for_nodes(store, ids),
                "formulaPatterns": groups,
                "meta": store.meta,
                "pagination": {
                    "returned": len(ids),
                    "total": len(set([node_id, *node.get("dependencies", [])])),
                    "truncated": len(node.get("dependencies", [])) > 12,
                },
            }
    with WorkbookIndex(index_path) as store:
        ids = []
        sheets = store.meta["sheets"]
        # Spread the bounded sample across the workbook, including the last sheet.
        selected = (
            [
                sheets[round(i * (len(sheets) - 1) / min(17, len(sheets) - 1))]
                for i in range(min(18, len(sheets)))
            ]
            if len(sheets) > 1
            else sheets
        )
        for sheet in selected:
            row = store.db.execute(
                (
                    "SELECT id FROM nodes WHERE sheet=? ORDER BY kind='cell' DESC,seq "
                    "LIMIT 1"
                ),
                (sheet,),
            ).fetchone()
            if row:
                ids.append(row[0])
        for row in store.db.execute(
            "SELECT id FROM nodes ORDER BY kind='cell' DESC,seq LIMIT 36"
        ):
            if len(ids) >= 18:
                break
            if row[0] not in ids:
                ids.append(row[0])
        return {
            "lazy": True,
            "nodes": [store.node(nid) for nid in ids],
            "edges": _edges_for_nodes(store, ids),
            "formulaPatterns": [],
            "meta": store.meta,
            "pagination": {
                "returned": len(ids),
                "total": store.meta["nodeCount"],
                "truncated": len(ids) < store.meta["nodeCount"],
            },
        }


def export_structure(index_path: Path) -> dict:
    """Explicit compatibility export. The web API must use bounded pages instead."""
    with WorkbookIndex(index_path) as store:
        nodes = [
            store._decode(row, context=True)
            for row in store.db.execute(
                "SELECT data,pattern_id FROM nodes ORDER BY seq"
            )
        ]
        patterns = [
            _pattern(store, row, member_limit=row[1])
            for row in store.db.execute(
                "SELECT data,member_count FROM patterns ORDER BY "
                "sheet,json_extract(data,'$.representativeId')"
            )
        ]
        edges = [
            {"source": s, "target": t, "kind": "dep"}
            for s, t in store.db.execute(
                "SELECT source,target FROM edges ORDER BY source,target"
            )
        ]
        return {
            "nodes": nodes,
            "formulaPatterns": patterns,
            "edges": edges,
            "meta": store.meta,
        }
