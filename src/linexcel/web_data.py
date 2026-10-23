"""Owner-scoped, read-only Datasette views of the complete imported source index.

The SQLite database is derived from stored values, never from calculation
results or browser samples. Datasette runs in this process behind WebApp's
authentication; it does not bind a second HTTP port.
"""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import sqlite3
import tempfile
import threading
import zipfile
from collections import OrderedDict
from contextlib import closing
from pathlib import Path
from typing import Any
from urllib.parse import quote, unquote, urlencode
from xml.etree import ElementTree as ET

from datasette.app import Datasette, PrefixedUrlString, Urls
from datasette.database import Database
from openpyxl.utils.cell import get_column_letter

from linexcel.graph_store import generation
from linexcel.lazy_store import WorkbookIndex, _unpack

DATA_VERSION = 2
MAX_GRID_COLUMNS = 512
MAX_INSTANCES = 8


def _json(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def _sqlite_value(value):
    # Keep the exact JSON representation in cells.value_json alongside the
    # browsable SQLite scalar. SQLite has no boolean or arbitrary-size integer.
    if isinstance(value, int) and not -(2**63) <= value < 2**63:
        return str(value)
    if isinstance(value, (dict, list)):
        return _json(value)
    return value


def _value_kind(value):
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, str):
        return "string"
    return "json"


def _worksheet_source_cell(node, source_sha256, sheet):
    """Exclude dependency placeholders and inputs from reference workbooks."""
    provenance = node.get("sourceValueProvenance") or {}
    return (
        provenance.get("workbookSha256") == source_sha256
        and provenance.get("sheet") == sheet
        and provenance.get("cell") == node.get("cell")
        and bool(provenance.get("worksheetPart"))
    )


def build_source_database(index: Path, destination: Path, *, max_bytes: int) -> dict:
    """Stream every indexed worksheet cell into a provenance table and grids.

    A sheet wider than SQLite's useful grid width uses the complete cells table;
    it is never silently truncated. Formatting-only blanks and synthetic blank
    dependency placeholders do not become source records.
    """
    descriptor, temporary = tempfile.mkstemp(
        prefix="data-", suffix=".building", dir=destination.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary)
    try:
        with (
            WorkbookIndex(index) as source,
            closing(sqlite3.connect(temporary)) as output,
        ):
            metadata = source.meta
            output.executescript("""
                PRAGMA journal_mode=DELETE;
                PRAGMA cache_size=-8192;
                PRAGMA temp_store=FILE;
                CREATE TABLE cells (
                    sheet TEXT NOT NULL, cell TEXT NOT NULL,
                    row INTEGER NOT NULL, column INTEGER NOT NULL,
                    value, value_kind TEXT NOT NULL, value_json TEXT NOT NULL,
                    source_type TEXT, value_role TEXT, value_status TEXT,
                    number_format TEXT, source_formula TEXT, canonical_formula TEXT,
                    formula_xml_text TEXT, formula_status TEXT,
                    worksheet_part TEXT, workbook_sha256 TEXT,
                    raw_xml_value TEXT, formula_provenance TEXT,
                    value_provenance TEXT, diagnostics TEXT,
                    PRIMARY KEY (sheet, cell)
                );
                CREATE INDEX cells_position ON cells(sheet,row,column);
                CREATE TABLE source_metadata (key TEXT PRIMARY KEY,value TEXT NOT NULL);
            """)
            sheets = []
            for sequence, sheet in enumerate(metadata.get("sheets", []), 1):
                count, rows, maximum, previous_row = 0, 0, 0, None
                for (packed,) in source.db.execute(
                    "SELECT data FROM nodes WHERE sheet=? "
                    "AND kind IN ('cell','input') ORDER BY row,col",
                    (sheet,),
                ):
                    node = _unpack(packed)
                    if not _worksheet_source_cell(node, metadata.get("sha256"), sheet):
                        continue
                    count += 1
                    maximum = max(maximum, node["column"])
                    if node["row"] != previous_row:
                        rows += 1
                        previous_row = node["row"]
                table = (
                    f"sheet_{sequence:03d}" if maximum <= MAX_GRID_COLUMNS else "cells"
                )
                columns = [get_column_letter(c) for c in range(1, maximum + 1)]
                if table != "cells":
                    column_sql = ",".join(f'"{column}"' for column in columns)
                    output.execute(
                        f'CREATE TABLE "{table}" (row INTEGER PRIMARY KEY'
                        + ("," + column_sql if column_sql else "")
                        + ")"
                    )
                sheets.append(
                    {
                        "sheet": sheet,
                        "table": table,
                        "rowCount": rows,
                        "cellCount": count,
                        "columns": columns if table != "cells" else [],
                        "layout": "grid" if table != "cells" else "cells",
                    }
                )
                current_row = None
                row_values = {}

                def flush_row():
                    if table == "cells" or current_row is None:
                        return
                    present = sorted(row_values)
                    names = ",".join(f'"{get_column_letter(c)}"' for c in present)
                    placeholders = ",".join("?" for _ in present)
                    output.execute(
                        f'INSERT INTO "{table}" (row'
                        + ("," + names if names else "")
                        + ") VALUES (?"
                        + ("," + placeholders if placeholders else "")
                        + ")",
                        (current_row, *(row_values[c] for c in present)),
                    )

                for (packed,) in source.db.execute(
                    "SELECT data FROM nodes WHERE sheet=? "
                    "AND kind IN ('cell','input') ORDER BY row,col",
                    (sheet,),
                ):
                    node = _unpack(packed)
                    if not _worksheet_source_cell(node, metadata.get("sha256"), sheet):
                        continue
                    value = node.get("value")
                    provenance = node.get("sourceValueProvenance") or {}
                    formula = node.get("formulaProvenance") or {}
                    output.execute(
                        "INSERT INTO cells VALUES ("
                        + ",".join("?" for _ in range(21))
                        + ")",
                        (
                            sheet,
                            node["cell"],
                            node["row"],
                            node["column"],
                            _sqlite_value(value),
                            _value_kind(value),
                            _json(value),
                            node.get("sourceDataType"),
                            provenance.get("role"),
                            provenance.get("status"),
                            node.get("numberFormat"),
                            node.get("sourceFormula"),
                            node.get("canonicalFormula"),
                            formula.get("xmlText"),
                            formula.get("status"),
                            provenance.get("worksheetPart"),
                            provenance.get("workbookSha256"),
                            provenance.get("rawXmlValue"),
                            _json(formula),
                            _json(provenance),
                            _json(node.get("diagnostics") or []),
                        ),
                    )
                    if node["row"] != current_row:
                        flush_row()
                        current_row = node["row"]
                        row_values = {}
                    if table != "cells":
                        row_values[node["column"]] = _sqlite_value(value)
                    if output.total_changes % 10000 == 0:
                        output.commit()
                        if temporary_path.stat().st_size > max_bytes:
                            raise DataLimitError
                flush_row()
                output.commit()
            result = {
                "sourceSha256": metadata.get("sha256"),
                "sheets": sheets,
                "sourceStatus": metadata.get("status", "unknown"),
                "diagnostics": metadata.get("diagnostics", []),
                "semantics": {
                    "values": "Stored input values and saved formula caches; "
                    "no recalculation.",
                    "formatting": "Excel number formats are recorded in cells; "
                    "grids show decoded values, not Excel display formatting.",
                    "coverage": "Worksheet cells retained by the structural index "
                    "with source evidence from this workbook. Formatting-only "
                    "blanks and synthetic dependency placeholders are excluded; "
                    "index diagnostics remain applicable.",
                    "coordinates": "Column names are Excel letters; row is the "
                    "original Excel row number. No header inference.",
                    "types": "cells.value_json and value_kind retain JSON types; "
                    "grids use SQLite scalars (booleans are 0/1).",
                    "formulas": "source_formula adds presentation '='; "
                    "formula_xml_text is exact decoded <f> text; canonical_formula "
                    "may be translated or normalized.",
                },
            }
            for key, value in result.items():
                output.execute(
                    "INSERT INTO source_metadata VALUES (?,?)", (key, _json(value))
                )
            output.commit()
        if temporary_path.stat().st_size > max_bytes:
            raise DataLimitError
        os.replace(temporary_path, destination)
        return result
    finally:
        temporary_path.unlink(missing_ok=True)


class DataLimitError(Exception):
    """The derived source snapshot would exceed the configured storage limit."""


class SourceDatabase(Database):
    """Keep native paths for metadata, and escape Windows SQLite URI paths."""

    def connect(self, write=False):
        if os.name != "nt":
            return super().connect(write=write)
        assert not write and not self.is_mutable
        connection = sqlite3.connect(
            "file:" + quote(str(self.path), safe="/:") + "?immutable=1",
            uri=True,
            check_same_thread=False,
        )
        self._all_file_connections.append(connection)
        return connection


class MountedUrls(Urls):
    """Datasette 0.65 may hand request.path back to its base_url helper."""

    def path(self, path, format=None):
        if path.startswith(self.ds.setting("base_url")):
            path = PrefixedUrlString(path)
        return super().path(path, format=format)


class EmbeddedDatasette(Datasette):
    @property
    def urls(self):
        return MountedUrls(self)


class DataBrowser:
    """A small in-process Datasette cache, confined to authenticated projects."""

    def __init__(self, store):
        self.store = store
        self._lock = threading.RLock()
        self._instances: OrderedDict[tuple, Any] = OrderedDict()
        self._loop = asyncio.new_event_loop()
        self._thread = threading.Thread(target=self._loop.run_forever, daemon=True)
        self._thread.start()

    def close(self):
        with self._lock:
            for instance in self._instances.values():
                self._close_instance(instance)
            self._instances.clear()
            self._loop.call_soon_threadsafe(self._loop.stop)
            self._thread.join(timeout=10)
            self._loop.close()

    def invalidate(self, owner: str, project_id: str):
        with self._lock:
            for key in list(self._instances):
                if key[:2] == (owner, project_id):
                    self._close_instance(self._instances.pop(key))

    @staticmethod
    def _close_instance(instance):
        instance.executor.shutdown(wait=True)
        for database in instance.databases.values():
            database.close()

    def _snapshot(self, owner: str, project_id: str) -> tuple[Path | None, dict]:
        from linexcel.web import APIError, _read

        path = self.store.project_path(owner, project_id)
        project = _read(path / "project.json")
        if self.store._revision(path) != project["revision"]:
            raise APIError(409, "Sources modifiées ; importez une nouvelle version")
        if project["filename"].lower().endswith((".xla", ".xlam")):
            return None, {
                "available": False,
                "reason": "macro_only",
                "sheets": [],
                "revision": project["revision"],
            }
        index = path / "index.sqlite"
        if not index.is_file():
            raise APIError(409, "Importez le classeur avant de parcourir les données")
        index_generation = generation(index)
        signature = hashlib.sha256(
            f"{DATA_VERSION}:{project['revision']}:{index_generation}".encode()
        ).hexdigest()
        destination = path / "source-data.sqlite"
        stamp_path = path / "source-data.json"
        saved: dict[str, Any] = _read(stamp_path) if stamp_path.exists() else {}
        if saved.get("signature") != signature or not destination.is_file():
            # Instances release handles before Windows replaces a rebuilt file.
            for key in list(self._instances):
                if key[:2] == (owner, project_id):
                    self._close_instance(self._instances.pop(key))
            remaining = (
                self.store.config.max_storage_mb * 1048576
                - self.store._storage_bytes()
                + (destination.stat().st_size if destination.exists() else 0)
            )
            try:
                metadata = build_source_database(
                    index, destination, max_bytes=remaining
                )
            except DataLimitError:
                raise APIError(
                    413, "Stockage insuffisant pour la vue des données"
                ) from None
            if (
                generation(index) != index_generation
                or self.store._revision(path) != project["revision"]
            ):
                # Source ownership was held throughout; source edits still must
                # invalidate the snapshot rather than serve stale provenance.
                raise APIError(409, "Sources modifiées pendant la lecture")
            from linexcel.web import _write

            saved = {"signature": signature, "metadata": metadata}
            _write(stamp_path, saved)
        prefix = f"{self.store.config.mount_path}/api/projects/{project_id}/datasette"
        metadata = dict(saved["metadata"])
        metadata.update(
            available=True,
            presentationAvailable=True,
            revision=project["revision"],
            url=prefix + "/source",
            cellsUrl=prefix + "/source/cells",
        )
        metadata["sheets"] = [
            {
                **sheet,
                "url": prefix
                + "/source/"
                + quote(sheet["table"])
                + (
                    "?" + urlencode({"sheet__exact": sheet["sheet"]})
                    if sheet["layout"] == "cells"
                    else ""
                ),
            }
            for sheet in metadata["sheets"]
        ]
        return destination, {"signature": signature, **metadata}

    def metadata(self, owner: str, project_id: str) -> dict:
        with self._lock, self.store.lock:
            _, metadata = self._snapshot(owner, project_id)
            metadata.pop("signature", None)
            return metadata

    def presentation(
        self,
        owner: str,
        project_id: str,
        sheet: str,
        *,
        row=1,
        column=1,
        rows=60,
        columns=20,
    ) -> dict:
        """Return a static source worksheet window behind the same owner check."""
        from linexcel.excel_presentation import (
            PresentationLimitError,
            sheet_presentation,
        )
        from linexcel.lazy_store import file_sha256
        from linexcel.web import APIError, _read

        with self._lock, self.store.lock:
            path = self.store.project_path(owner, project_id)
            project = _read(path / "project.json")
            if self.store._revision(path) != project["revision"]:
                raise APIError(409, "Sources modifiées ; importez une nouvelle version")
            if project["filename"].lower().endswith((".xla", ".xlam")):
                raise APIError(409, "Cet add-in ne contient pas de feuille de calcul")
            index = path / "index.sqlite"
            if not index.is_file():
                raise APIError(
                    409, "Importez le classeur avant de parcourir les données"
                )
            before = generation(index)
            with WorkbookIndex(index) as source:
                source_sha = source.meta.get("sha256")
            if file_sha256(path / "workbook.xlsx") != source_sha:
                raise APIError(409, "Sources modifiées ; importez une nouvelle version")
            try:
                result = sheet_presentation(
                    path / "workbook.xlsx",
                    path / "index.sqlite",
                    sheet,
                    row=row,
                    column=column,
                    rows=rows,
                    columns=columns,
                )
            except KeyError:
                raise APIError(404, "Feuille introuvable") from None
            except PresentationLimitError as exc:
                raise APIError(413, str(exc)) from None
            except (ValueError, OSError, zipfile.BadZipFile, ET.ParseError):
                raise APIError(400, "Présentation de la feuille indisponible") from None
            if (
                self.store._revision(path) != project["revision"]
                or generation(index) != before
                or result.get("sourceSha256") != source_sha
            ):
                raise APIError(409, "Sources modifiées pendant la lecture")
            return {"available": True, "revision": project["revision"], **result}

    def request(
        self,
        owner: str,
        project_id: str,
        method: str,
        path: str,
        query: str,
        *,
        origin: str = "http://localhost",
    ):
        from linexcel.web import APIError

        # Keep introspection, plugin endpoints, actor management and other
        # database names outside this read-only, per-project surface.
        if method not in {"GET", "HEAD"}:
            raise APIError(405, "La vue des données est en lecture seule")
        decoded = unquote(path)
        if any(part in {".", ".."} for part in decoded.split("/")) or "\\" in decoded:
            raise APIError(404, "Ressource Datasette introuvable")
        if not (
            path in {"", "/", "/source", "/source.json"}
            or path.startswith(("/source/", "/-/static/"))
        ):
            raise APIError(404, "Ressource Datasette introuvable")
        if len(query) > 8192:
            raise APIError(400, "Recherche Datasette trop longue")
        with self._lock, self.store.lock:
            database, metadata = self._snapshot(owner, project_id)
            if database is None:
                raise APIError(
                    409, "Cet add-in contient des sources VBA, sans vue de données"
                )
            key = (owner, project_id, metadata["signature"])
            if key not in self._instances:
                if len(self._instances) >= MAX_INSTANCES:
                    self._close_instance(self._instances.popitem(last=False)[1])
                tables = {
                    sheet["table"]: {
                        "description": f"Feuille {sheet['sheet']} — lignes Excel "
                        "et valeurs sauvegardées. Aucun recalcul."
                    }
                    for sheet in metadata["sheets"]
                    if sheet["layout"] == "grid"
                }
                tables["cells"] = {
                    "description": "Cellules sources et caches sauvegardés, "
                    "avec formules et provenance. Les formats Excel sont décrits, "
                    "pas rendus."
                }
                instance = EmbeddedDatasette(
                    immutables=[str(database)],
                    metadata={
                        "title": "Données Excel",
                        "description": "Données sauvegardées en lecture seule; "
                        "aucune formule ni macro exécutée.",
                        "databases": {
                            "source": {
                                "title": "Données du classeur",
                                "description": "Source SHA-256 : "
                                + metadata["sourceSha256"],
                                "tables": tables,
                            }
                        },
                    },
                    settings={
                        "base_url": metadata["url"].removesuffix("source"),
                        "allow_download": False,
                        "default_allow_sql": False,
                        "default_page_size": 50,
                        "max_returned_rows": 200,
                        "sql_time_limit_ms": 1000,
                        "num_sql_threads": 2,
                        "truncate_cells_html": 80,
                        "default_cache_ttl": 0,
                    },
                    cache_headers=False,
                )
                # A stable public name avoids leaking the file name/path and
                # ensures each instance contains only its project's source DB.
                instance.remove_database("source-data")
                instance.add_database(
                    SourceDatabase(instance, path=str(database), is_mutable=False),
                    name="source",
                    route="source",
                )
                self._instances[key] = instance
            instance = self._instances[key]
            self._instances.move_to_end(key)
            route = origin + instance.setting("base_url").rstrip("/") + (path or "/")
            if query:
                route += "?" + query
            pending = asyncio.run_coroutine_threadsafe(
                instance.client.request(method, route, avoid_path_rewrites=True),
                self._loop,
            )
            try:
                response = pending.result(timeout=30)
            except TimeoutError:
                pending.cancel()
                raise APIError(
                    504, "Délai dépassé pendant la lecture des données"
                ) from None
            headers = [
                (name, value)
                for name, value in response.headers.items()
                if name.lower() in {"content-type", "location", "content-disposition"}
            ]
            return response.status_code, headers, response.content
