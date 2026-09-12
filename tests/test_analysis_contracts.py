"""Regressions for target parsing, chain diagnostics and value comparisons."""

import datetime
import io

import pytest
from openpyxl import Workbook

from linexcel import analyze
from linexcel.analyzer import _longest_dep_chain, _parse_targets
from linexcel.values import readings_agree


@pytest.mark.parametrize(
    ("target", "sheet"),
    [
        ("'Costs, FY26'!$A$1", "Costs, FY26"),
        ("'L''été, FY26'!A1", "L'été, FY26"),
        ("My Sheet!A1", "My Sheet"),
        ("L'été!A1", "L'été"),
        ("'Bang! sheet'!A1", "Bang! sheet"),
    ],
)
def test_targets_preserve_excel_sheet_names(target, sheet):
    assert _parse_targets([target + ",Other!B2", target]) == [
        (sheet, 1, 1),
        ("Other", 2, 2),
    ]


@pytest.mark.parametrize("target", ["", ",", "S!A1,", "'Broken!A1", "S!A1,,S!B1"])
def test_empty_or_malformed_targets_are_rejected(target):
    with pytest.raises(ValueError, match="Invalid target"):
        _parse_targets([target])


def test_quoted_target_is_evaluated_end_to_end():
    wb = Workbook()
    ws = wb.active
    ws.title = "Costs, FY26"
    ws["A1"] = "=2+3"
    buf = io.BytesIO()
    wb.save(buf)
    result = analyze(buf.getvalue(), targets=["'Costs, FY26'!A1"])
    assert next(n for n in result.nodes if n.get("formula"))["value"] == 5


def chain_graph(pairs):
    nodes = {nid: {"kind": "cell"} for pair in pairs for nid in pair}
    edges = {
        str(i): {"source": src, "target": dst, "kind": "dep"}
        for i, (src, dst) in enumerate(pairs)
    }
    return nodes, edges


def test_chain_depth_across_distinct_groups_and_diamond():
    graph = chain_graph([("a", "b"), ("a", "c"), ("b", "d"), ("c", "d")])
    assert _longest_dep_chain(*graph) == 3
    long_chain = chain_graph([(i, i + 1) for i in range(10_000)])
    assert _longest_dep_chain(*long_chain) == 10_001


def test_chain_depth_ignores_cycles_and_non_dependency_edges():
    nodes, edges = chain_graph([("a", "b"), ("b", "a"), ("x", "y")])
    edges["extra"] = {"source": "y", "target": "a", "kind": "call"}
    assert _longest_dep_chain(nodes, edges) == 2
    assert _longest_dep_chain(*chain_graph([("a", "b"), ("b", "a")])) == 0


def test_full_analysis_warns_for_a_chain_across_distinct_patterns():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    for row in range(2, 32):
        # Different constants prevent the whole chain collapsing to one group.
        ws.cell(row, 1, f"=A{row - 1}+{row}")
    buf = io.BytesIO()
    wb.save(buf)
    result = analyze(buf.getvalue())
    assert any("Long dependency chains" in warning for warning in result.warnings)


@pytest.mark.parametrize(
    ("left", "right", "expected"),
    [
        (False, True, "differ"),
        (True, 1, "differ"),
        (0, False, "differ"),
        (True, True, "same"),
        (7, "seven", "differ"),
        ("7", 7, "differ"),
        (7, "7", "differ"),
        (1, 1.0, "same"),
        (1.0, 1.0 + 1e-10, "same"),
        ("6,7 €", "6.7 €", "format"),
    ],
)
def test_value_verdicts_respect_excel_types(left, right, expected):
    assert readings_agree(left, right, None) == expected


def test_date_serial_agreement_keeps_date_semantics():
    assert readings_agree(46023, datetime.date(2026, 1, 1), "2026-01-01") == "same"
    assert readings_agree(46023, "2026-01-01", "2026-01-01") == "same"
    assert readings_agree(46023, "2026-01-02", "2026-01-01") == "differ"
    assert readings_agree({"type": "Error", "kind": "Div"}, "#DIV/0!", None) == "same"
    assert (
        readings_agree(
            {"type": "Error", "kind": "Div"}, datetime.date(2026, 1, 1), None
        )
        == "differ"
    )


@pytest.mark.parametrize("formula", ['=INDIRECT("A2")', "=OFFSET(A1,1,0)"])
def test_targeted_dynamic_dependencies_state_static_trace_limit(formula):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 2
    ws["A2"] = "=30+40"
    ws["B1"] = formula
    buf = io.BytesIO()
    wb.save(buf)
    result = analyze(buf.getvalue(), targets=["Sheet!B1"])
    assert result.engine.get_value("Sheet", 2, 1) == 70
    warnings = " ".join(result.warnings)
    assert "static upstream trace" in warnings
    assert "dynamic references" in warnings
    assert "were not recomputed" not in warnings


def test_truncated_trace_does_not_claim_omitted_cells_were_not_evaluated(monkeypatch):
    monkeypatch.setattr("linexcel.engine.TRACE_MAX_NODES", 1)
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "=2+3"
    ws["B1"] = "=A1*2"
    buf = io.BytesIO()
    wb.save(buf)
    result = analyze(buf.getvalue(), targets=["Sheet!B1"])
    assert any("hit its budget" in warning for warning in result.warnings)
    assert not any("were not recomputed" in warning for warning in result.warnings)
