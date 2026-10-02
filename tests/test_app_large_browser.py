"""Browser contracts for bounded uploads and server-backed workbook exploration."""

import json
import os
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import pytest

playwright = pytest.importorskip("playwright.sync_api")
APP = Path(__file__).parents[1] / "src/linexcel/assets/app.html"


@pytest.fixture(scope="module")
def browser():
    with playwright.sync_playwright() as driver:
        instance = driver.chromium.launch(
            channel=os.environ.get("LINEXCEL_BROWSER_CHANNEL") or None
        )
        yield instance
        instance.close()


def serve_asset(route):
    path = urlsplit(route.request.url).path
    if path == "/mounted/":
        route.fulfill(body=APP.read_text(encoding="utf-8"), content_type="text/html")
    elif path.endswith("assets/cytoscape.min.js"):
        route.fulfill(
            body=(APP.parent / "cytoscape.min.js").read_text(encoding="utf-8"),
            content_type="application/javascript",
        )
    else:
        return False
    return True


@pytest.mark.parametrize("viewport", [(1440, 900), (390, 844)])
def test_upload_chunks_resume_after_lost_acknowledgement(browser, viewport):
    page = browser.new_page(viewport={"width": viewport[0], "height": viewport[1]})
    page.add_init_script(
        "FileReader.prototype.readAsDataURL = () => {throw Error('whole file read')}"
    )
    workbook, reference = b"0123456789abcdefghij", b"ref"
    state = {"id": "upload-1", "chunkBytes": 8, "files": []}
    chunks, manifests = [], []
    completed = False
    project = {"id": "large", "name": "Large workbook"}

    def handle(route):
        nonlocal completed
        if serve_asset(route):
            return
        path = urlsplit(route.request.url).path.removeprefix("/mounted/api/")
        if path == "config":
            body = {"uploads": {"chunkBytes": 8}, "limits": {"uploadMb": 256}}
        elif path == "projects":
            body = {"projects": [project] if completed else []}
        elif path == "uploads":
            manifest = route.request.post_data_json
            manifests.append(manifest)
            state["files"] = [
                {**item, "received": 0}
                for item in [manifest["workbook"], *manifest["references"]]
            ]
            body = {"upload": state}
        elif "/files/" in path:
            index = int(path.rsplit("/", 1)[-1])
            offset = int(route.request.headers["x-upload-offset"])
            data = route.request.post_data_buffer
            assert data == [workbook, reference][index][offset : offset + 8]
            assert len(data) <= 8
            assert offset == state["files"][index]["received"]
            chunks.append((index, offset, len(data)))
            state["files"][index]["received"] += len(data)
            if index == 0 and offset == 8:
                route.abort("connectionreset")
                return
            body = {"upload": state}
        elif path == "uploads/upload-1":
            assert not completed  # Completed sessions are no longer upload resources.
            body = {"upload": state}
        elif path == "uploads/upload-1/complete":
            assert [f["received"] for f in state["files"]] == [20, 3]
            if not completed:
                completed = True
                route.abort("connectionreset")  # The server committed the project.
                return
            completed = True
            body = {"project": project}
        elif path == "projects/large":
            body = {"project": project, "graph": None, "tasks": []}
        else:
            raise AssertionError(path)
        route.fulfill(body=json.dumps(body), content_type="application/json")

    page.route("**/*", handle)
    try:
        page.goto("http://localhost/mounted/")
        page.locator("#workbook").set_input_files(
            {
                "name": "large.xlsx",
                "mimeType": "application/octet-stream",
                "buffer": workbook,
            }
        )
        page.locator("#references").set_input_files(
            {
                "name": "ref.xlsx",
                "mimeType": "application/octet-stream",
                "buffer": reference,
            }
        )
        page.locator("#import-button").click()
        page.get_by_text(
            "Envoi interrompu. Réessayez pour reprendre", exact=False
        ).wait_for()
        with page.expect_request("**/uploads/upload-1/complete"):
            page.locator("#import-button").click()
        page.wait_for_function("!document.querySelector('#import-button').disabled")
        page.locator("#import-button").click()
        page.locator("#project-title").get_by_text("Large workbook").wait_for()
        assert len(manifests) == 1
        assert manifests[0]["workbook"] == {"name": "large.xlsx", "size": 20}
        assert chunks == [(0, 0, 8), (0, 8, 8), (0, 16, 4), (1, 0, 3)]
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    finally:
        page.close()


@pytest.mark.parametrize("viewport", [(1440, 900), (390, 844)])
def test_lazy_pages_search_neighborhood_and_scoped_notices(browser, viewport):
    page = browser.new_page(viewport={"width": viewport[0], "height": viewport[1]})
    errors, requests = [], []
    page.on("pageerror", lambda error: errors.append(str(error)))
    project = {"id": "large", "name": "Large workbook"}
    all_nodes = [
        {
            "id": f"Data!A{i}",
            "sheet": "Data",
            "cell": f"A{i}",
            "kind": "input",
            "value": i,
        }
        for i in range(1, 402)
    ]
    all_nodes[399].update(
        patternId="copied",
        patternMemberCount=250,
        context={
            "nearbyLabels": [
                {"cell": "B400", "text": "0123456789abcdef" * 12, "position": "left"}
            ]
        },
    )

    def graph_response(items, offset=0, limit=200, total=None):
        return {
            "lazy": True,
            "nodes": items,
            "edges": [],
            "formulaPatterns": [
                {
                    "id": "copied",
                    "sheet": "Data",
                    "memberCount": 250,
                    "membersTruncated": True,
                    "ranges": ["A400:A419"],
                    "memberIds": ["Data!A400"],
                }
            ],
            "pagination": {"offset": offset, "limit": limit, "total": total or 401},
            "meta": {
                "sheets": ["Data", "Empty"],
                "nodeCount": 401,
                "edgeCount": 0,
                "sheetNodeCounts": {"Data": 401, "Empty": 0},
            },
        }

    def handle(route):
        if serve_asset(route):
            return
        url = urlsplit(route.request.url)
        path = url.path.removeprefix("/mounted/api/")
        query = parse_qs(url.query)
        requests.append((path, query))
        if path == "config":
            body = {}
        elif path == "projects":
            body = {"projects": [project]}
        elif path == "projects/large":
            body = {
                "project": project,
                "graph": graph_response(all_nodes[:200]),
                "tasks": [],
            }
        elif path == "projects/large/nodes":
            offset, limit = (
                int(query.get("offset", [0])[0]),
                int(query.get("limit", [200])[0]),
            )
            q = query.get("q", [""])[0].lower()
            filtered = [n for n in all_nodes if q in n["id"].lower()]
            body = graph_response(
                filtered[offset : offset + limit], offset, limit, len(filtered)
            )
        elif path == "projects/large/neighborhood":
            node_id = query["nodeId"][0]
            body = graph_response([n for n in all_nodes if n["id"] == node_id], total=1)
        else:
            raise AssertionError(path)
        route.fulfill(body=json.dumps(body), content_type="application/json")

    page.route("**/*", handle)
    try:
        page.goto("http://localhost/mounted/")
        page.locator("[data-project='large']").click()
        page.locator("[data-node='Data!A1']").wait_for()
        assert page.locator("[data-node]").count() == 200
        assert "401 nœuds" in page.locator("#stats").inner_text()
        page.locator("#node-next").click()
        page.locator("[data-node='Data!A201']").wait_for()
        assert page.locator("[data-node='Data!A1']").count() == 0
        page.locator("#node-next").click()
        page.locator("[data-node='Data!A401']").wait_for()
        assert page.locator("[data-node]").count() == 1
        assert page.locator("#node-next").is_disabled()
        page.locator("#search").fill("A400")
        page.locator("[data-node='Data!A400']").click()
        page.locator("#node-detail h2").get_by_text("Data · A400").wait_for()
        assert "Plages de l’échantillon :" in page.locator(".pattern-card").inner_text()
        page.get_by_text("Repères dans la feuille", exact=True).click()
        context_label = page.locator("#node-detail details[open] li span[data-no-i18n]")
        assert context_label.is_visible()
        assert context_label.inner_text() == "0123456789abcdef" * 12
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        label_box = context_label.bounding_box()
        detail_box = page.locator("#node-detail").bounding_box()
        assert label_box["x"] >= detail_box["x"]
        assert label_box["x"] + label_box["width"] <= (
            detail_box["x"] + detail_box["width"]
        )
        page.locator("[data-tab='graph']").click()
        assert "Data!A400" in page.locator("#graph-notice").inner_text()
        assert page.locator("#graph-lazy-notice").is_visible()
        page.locator("#graph-search").fill("A401")
        page.locator("#graph-results [data-related='Data!A401']").click()
        page.locator("#graph-preview h3").get_by_text("Data · A401").wait_for()
        page.locator("#graph-depth").select_option("2")
        page.wait_for_function(
            "document.querySelector('#graph-notice').textContent"
            ".includes('profondeur 2')"
        )
        for name in ["nodes", "sheets", "captures", "tasks", "diagnostics"]:
            page.locator(f"[data-tab='{name}']").click()
            assert not page.locator("#graph-lazy-notice").is_visible()
            assert (
                page.locator("#workspace-content > [role=tabpanel]:visible").count()
                == 1
            )
            assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        assert sum(path == "projects/large" for path, _ in requests) == 1
        assert all(int(q.get("limit", [200])[0]) <= 200 for _, q in requests)
        assert not errors
    finally:
        page.close()


def test_local_path_import_does_not_read_or_upload_file(browser):
    page = browser.new_page(viewport={"width": 390, "height": 844})
    imports = []
    project = {"id": "local", "name": "Server workbook"}

    def handle(route):
        if serve_asset(route):
            return
        path = urlsplit(route.request.url).path.removeprefix("/mounted/api/")
        if path == "config":
            body = {"uploads": {"localImport": True}}
        elif path == "projects":
            body = {"projects": [project] if imports else []}
        elif path == "projects/local" and route.request.method == "POST":
            imports.append(route.request.post_data_json)
            body = {"project": project}
        elif path == "projects/local":
            body = {"project": project, "graph": None, "tasks": []}
        else:
            raise AssertionError(path)
        route.fulfill(body=json.dumps(body), content_type="application/json")

    page.route("**/*", handle)
    try:
        page.goto("http://localhost/mounted/")
        page.locator("#import-source").select_option("local")
        page.locator("#local-path").fill(r"C:\corpus\large.xlsx")
        page.locator("#local-references").fill(
            "C:\\corpus\\ref.xlsx\nC:\\corpus\\other.xlsx"
        )
        page.locator("#import-button").click()
        page.locator("#project-title").get_by_text("Server workbook").wait_for()
        assert imports[0]["path"] == r"C:\corpus\large.xlsx"
        assert imports[0]["references"] == [
            r"C:\corpus\ref.xlsx",
            r"C:\corpus\other.xlsx",
        ]
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    finally:
        page.close()


@pytest.mark.parametrize("viewport", [(1440, 900), (390, 844)])
@pytest.mark.parametrize("sampled", [False, True])
def test_results_and_capture_images_load_on_demand(browser, viewport, sampled):
    page = browser.new_page(viewport={"width": viewport[0], "height": viewport[1]})
    requests, errors = [], []
    page.on("pageerror", lambda error: errors.append(str(error)))
    project = {"id": "p", "name": "Partial workbook"}
    node = {"id": "Data!A1", "sheet": "Data", "kind": "input", "value": 42}
    tasks = [
        {"id": "calc", "operation": "evaluate", "nodeId": "Data!A1"},
        {"id": "capture", "operation": "capture"},
        {"id": "overview", "operation": "document", "options": {"language": "fr"}},
    ]
    for task in tasks:
        task.update(status="succeeded", resultAvailable=True, result=None)
    results = {
        "calc": {
            "status": "completed",
            "value": 42,
            "steps": [{"nodeId": "Data!A1", "evaluationStatus": "input"}],
            "stepsPagination": {
                "offset": 0,
                "limit": 200,
                "total": 10000,
                "hasMore": True,
            },
        },
        "overview": {"markdown": "Aperçu enregistré.", "language": "fr"},
        "capture": {
            "screenshots": [
                {
                    "sheet": "Data",
                    "url": "api/projects/p/captures/capture/0",
                    "thumbnailUrl": "api/projects/p/captures/capture/0?thumbnail=1",
                }
            ],
        },
    }
    if sampled:
        results["calc"].update(stepTotal=10000, omittedSteps=9999)
        results["calc"]["steps"][0]["dependenciesOmitted"] = 9999

    def handle(route):
        if serve_asset(route):
            return
        url = urlsplit(route.request.url)
        path = url.path.removeprefix("/mounted/api/")
        requests.append(path + ("?" + url.query if url.query else ""))
        if path == "config":
            body = {}
        elif path == "projects":
            body = {"projects": [project]}
        elif path == "projects/p":
            body = {
                "project": project,
                "tasks": tasks,
                "graph": {"nodes": [node], "edges": [], "meta": {"status": "partial"}},
            }
        elif path.startswith("tasks/"):
            task_id = path.split("/")[1]
            body = {
                "task": {
                    **next(t for t in tasks if t["id"] == task_id),
                    "result": results[task_id],
                }
            }
        elif path.startswith("projects/p/captures/"):
            route.fulfill(
                body="<svg xmlns='http://www.w3.org/2000/svg' width='10' height='10'/>",
                content_type="image/svg+xml",
            )
            return
        else:
            raise AssertionError(path)
        route.fulfill(body=json.dumps(body), content_type="application/json")

    page.route("**/*", handle)
    try:
        page.goto("http://localhost/mounted/")
        page.locator("[data-project='p']").click()
        page.locator("[data-node='Data!A1']").wait_for()
        assert not any(path.startswith("tasks/") for path in requests)
        assert page.locator("#index-status").is_visible()
        page.locator("[data-node='Data!A1']").click()
        page.locator("#calculation-steps").get_by_text(
            "Dépendances affichées sous forme d’échantillon"
            if sampled
            else "Aperçu des dépendances limité",
            exact=False,
        ).wait_for()
        assert "10000" in page.locator("#calculation-steps").inner_text()
        if sampled:
            assert "9999" in page.locator("#calculation-steps").inner_text()
            assert (
                "toutes les étapes"
                not in page.locator("#calculation-steps").inner_text()
            )
        assert "tasks/capture" not in requests
        page.locator("[data-tab='captures']").click()
        page.locator("#capture-selected img").wait_for()
        assert (
            page.locator("#capture-selected img")
            .get_attribute("src")
            .endswith("/capture/0")
        )
        assert (
            page.locator("#capture-thumb-0")
            .get_attribute("src")
            .endswith("?thumbnail=1")
        )
        assert requests.count("tasks/capture") == 1
        assert "tasks/overview" not in requests
        page.locator("[data-tab='sheets']").click()
        page.locator("#workbook-ai").get_by_text("Aperçu enregistré.").wait_for()
        assert requests.count("tasks/overview") == 1
        page.locator("[data-tab='captures']").click()
        assert requests.count("tasks/capture") == 1
        assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        assert not errors
    finally:
        page.close()


def test_pending_result_cannot_evaluate_same_address_in_another_project(browser):
    page = browser.new_page()
    delayed, operations = [], []
    projects = [{"id": name, "name": name} for name in ("first", "second")]
    node = {"id": "Data!A1", "sheet": "Data", "formula": "=1+1", "kind": "cell"}
    saved = {
        "id": "saved",
        "nodeId": node["id"],
        "operation": "evaluate",
        "status": "succeeded",
        "resultAvailable": True,
        "result": None,
    }

    def handle(route):
        if serve_asset(route):
            return
        path = urlsplit(route.request.url).path.removeprefix("/mounted/api/")
        if path == "config":
            body = {}
        elif path == "projects":
            body = {"projects": projects}
        elif path in ("projects/first", "projects/second"):
            project = next(p for p in projects if path.endswith(p["id"]))
            body = {
                "project": project,
                "graph": {"nodes": [node], "edges": []},
                "tasks": [saved] if project["id"] == "first" else [],
            }
        elif path == "tasks/saved":
            delayed.append(route)
            return
        elif path.endswith("/tasks"):
            operations.append(path)
            body = {
                "task": {
                    "id": "new",
                    "nodeId": node["id"],
                    "operation": "evaluate",
                    "status": "succeeded",
                    "result": {"status": "completed", "value": 2},
                }
            }
        else:
            raise AssertionError(path)
        route.fulfill(body=json.dumps(body), content_type="application/json")

    page.route("**/*", handle)
    try:
        page.goto("http://localhost/mounted/")
        page.locator("[data-project='first']").click()
        with page.expect_request("**/tasks/saved"):
            page.locator("[data-node='Data!A1']").click()
        page.locator("[data-project='second']").click()
        page.get_by_text("second", exact=True).filter(
            has_not=page.locator("small")
        ).first.wait_for()
        with page.expect_response("**/projects/second/tasks"):
            page.locator("[data-node='Data!A1']").click()
        with page.expect_response("**/tasks/saved"):
            delayed[0].fulfill(
                body=json.dumps(
                    {"task": {**saved, "result": {"status": "completed", "value": 2}}}
                ),
                content_type="application/json",
            )
        page.evaluate(
            "()=>new Promise(r=>requestAnimationFrame(()=>requestAnimationFrame(r)))"
        )
        assert operations == ["projects/second/tasks"]
    finally:
        page.close()


@pytest.mark.parametrize("viewport", [(1440, 900), (390, 844)])
def test_server_hierarchy_and_manual_calculation_are_independent(
    browser, viewport, tmp_path
):
    from openpyxl import Workbook

    from linexcel.graph_store import build_graph, generation, read_view
    from linexcel.lazy_store import build_index, graph_neighborhood, graph_page

    book = Workbook()
    book.active.title = "Data"
    for i in range(1, 205):
        book.active.cell(i, 1, i)
    sheet = book.create_sheet("Result")
    sheet["B1"] = "=Data!A1+Data!A204"
    sheet["B2"] = "=B1*2"
    source, index, artifact = [
        tmp_path / name for name in ("book.xlsx", "index.sqlite", "hierarchy.sqlite")
    ]
    book.save(source)
    build_index(source, index)
    build_graph(index, artifact)
    page = browser.new_page(viewport={"width": viewport[0], "height": viewport[1]})
    errors, operations, requests = [], [], []
    page.on("pageerror", lambda error: errors.append(str(error)))
    project = {"id": "p", "name": "Structure without calculation"}
    tasks = []
    fail_view, delay_build = False, False

    def handle(route):
        if serve_asset(route):
            return
        url = urlsplit(route.request.url)
        path, query = url.path.removeprefix("/mounted/api/"), parse_qs(url.query)
        requests.append(path)
        if path == "config":
            body = {}
        elif path == "projects":
            body = {"projects": [project]}
        elif path == "projects/p":
            body = {"project": project, "tasks": tasks, "graph": graph_page(index)}
        elif path == "projects/p/nodes":
            body = graph_page(index, query=query.get("q", [None])[0])
        elif path == "projects/p/graph-view":
            if fail_view:
                route.fulfill(status=409, json={"message": "Structure modifiée"})
                return
            body = read_view(
                artifact,
                generation(index),
                view=int(query.get("view", ["0"])[0]),
                node_id=query.get("nodeId", [None])[0],
            )
        elif path == "projects/p/neighborhood":
            body = graph_neighborhood(index, query["nodeId"][0])
        elif path == "projects/p/tasks":
            request = route.request.post_data_json
            operations.append(request)
            task = {
                "id": str(len(operations)),
                "operation": request["operation"],
                "nodeId": request.get("nodeId"),
                "status": "running"
                if delay_build and request["operation"] == "build_graph"
                else "succeeded",
                "result": {"status": "completed", "value": 205},
            }
            tasks.append(task)
            body = {"task": task}
        elif path.startswith("tasks/"):
            task = next(t for t in tasks if t["id"] == path.split("/")[1])
            if path.endswith("/cancel"):
                task["status"] = "cancelled"
            body = {"task": task}
        else:
            raise AssertionError(path)
        route.fulfill(json=body)

    page.route("**/*", handle)
    try:
        page.goto("http://localhost/mounted/")
        page.locator("[data-project='p']").click()
        # A no-result node search can leave the current page empty before the
        # hierarchy is built. The workbook graph must still be explorable.
        page.locator("#search").fill("no-such-cell")
        page.wait_for_function(
            "document.querySelectorAll('#node-list [data-node]').length===0"
        )
        page.locator("#tab-graph").click()
        page.locator("#graph-build").click()
        page.wait_for_function(
            "document.querySelector('#graph-canvas')._cy?.nodes().length===2"
        )
        assert page.locator("#graph [data-structure-node]").count() == 2
        assert len(operations) == 1 and operations[0]["operation"] == "build_graph"
        assert not any(op["operation"] == "evaluate" for op in operations)
        page.reload()
        page.locator("[data-project='p']").click()
        page.locator("#tab-graph").click()
        page.locator("#graph-search").fill("Data!A4")
        exact_result = page.locator("#graph-results [data-related]")
        playwright.expect(exact_result).to_have_count(1)
        assert exact_result.get_attribute("data-related") == "Data!A4"
        page.locator("#tab-nodes").click()
        requests.clear()
        operations.clear()
        tasks.clear()
        page.locator("#manual-recalculation").check()
        page.locator("#tab-graph").click()
        assert not any("graph-view" in r for r in requests)
        page.locator("#graph-build").click()
        page.wait_for_function(
            "document.querySelector('#graph-canvas')._cy?.nodes().length===2"
        )
        assert len(operations) == 1 and operations[0]["operation"] == "build_graph"
        assert operations[0]["budget"]["seconds"] == 600
        assert not any("/structure" in r for r in requests)
        assert page.locator("#graph-lazy-notice").is_hidden()
        assert page.locator("#graph-depth").is_disabled()
        page.evaluate(
            "document.querySelector('#graph-canvas')._cy.nodes().filter(n=>n.data('label').startsWith('Result')).emit('tap')"
        )
        page.wait_for_function(
            "document.querySelector('#graph-canvas')._cy?.nodes().length===3"
        )
        assert (
            page.locator("#graph-breadcrumbs [aria-current]").inner_text() == "Result"
        )
        assert page.evaluate("""() => {
            const cy=document.querySelector('#graph-canvas')._cy;
            const boxes=cy.nodes().map(n=>n.boundingBox());
            for(let i=0;i<boxes.length;i++)for(let j=i+1;j<boxes.length;j++){
                const a=boxes[i],b=boxes[j];
                if(a.x1<b.x2&&a.x2>b.x1&&a.y1<b.y2&&a.y2>b.y1)return false;
            }return true;
        }""")
        page.evaluate("""() => {const cy=document.querySelector('#graph-canvas')._cy;
            cy.zoom(.8);cy.pan({x:123,y:234});window.savedGraph=cy;
            cy.nodes().filter(n=>n.data('nodeId')==='Result!B1').emit('tap');} """)
        page.locator("#graph-preview h3").get_by_text("Result · B1").wait_for()
        assert len(operations) == 1
        assert page.evaluate(
            "document.querySelector('#graph-canvas')._cy===window.savedGraph"
        )
        assert page.evaluate("window.savedGraph.zoom()") == 0.8
        assert page.evaluate("window.savedGraph.pan()") == {"x": 123, "y": 234}
        page.locator("#graph-recalculate").click()
        page.wait_for_function(
            "document.querySelector('#graph-preview').textContent.includes('205')"
        )
        assert len(operations) == 2 and operations[-1]["force"] is True
        assert operations[-1]["budget"]["seconds"] == 120
        page.locator("#manual-recalculation").uncheck()
        assert len(operations) == 2
        page.evaluate(
            "window.savedGraph.nodes().filter(n=>n.data('nodeId')==='Result!B2').emit('tap')"
        )
        page.wait_for_function(
            "document.querySelector('#graph-preview h3').textContent.includes('B2')"
        )
        page.wait_for_function("!document.querySelector('#graph-recalculate').disabled")
        assert len(operations) == 3 and operations[-1]["force"] is False
        page.locator("#manual-recalculation").check()
        for name in ["nodes", "sheets", "captures", "tasks", "diagnostics"]:
            page.locator(f"#tab-{name}").click()
            assert page.locator("#graph-build-status").is_hidden()
            assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
        page.locator("#tab-graph").click()
        page.locator("#graph-neighborhood").click()
        assert page.locator("#graph-lazy-notice").is_visible()
        page.reload()
        page.locator("[data-project='p']").click()
        assert page.locator("#manual-recalculation").is_checked()
        page.locator("#tab-graph").click()
        fail_view = True
        page.locator("#graph-build").click()
        page.locator("#graph-build-status").get_by_text("Structure modifiée").wait_for()
        playwright.expect(page.locator("#graph-build")).to_be_enabled()
        fail_view, delay_build = False, True
        page.locator("#graph-build").click()
        page.locator("#graph-build-cancel").wait_for()
        page.reload()
        page.locator("#tab-graph").click()
        page.locator("#graph-build-cancel").click()
        playwright.expect(page.locator("#graph-build")).to_be_enabled()
        assert tasks[-1]["status"] == "cancelled"
        assert page.locator("#graph-neighborhood").is_hidden()
        assert not errors
    finally:
        page.close()
