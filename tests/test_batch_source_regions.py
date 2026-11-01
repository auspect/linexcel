"""Excel-saved array caches and supported 3D references recalculate in batch."""

import io

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

from linexcel import ExecutionPolicy, analyze
from linexcel.engine import _open_workbook


def package(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def array_book(formula="=SEQUENCE(3,2)"):
    book = Workbook()
    sheet = book.active
    sheet["A1"] = ArrayFormula(ref="A1:B3", text=formula)
    for cell in ("B1", "A2", "B2", "A3", "B3"):
        sheet[cell] = 999
    # Excel writes the # operator as ANCHORARRAY in the stored package.
    sheet["D1"] = "=SUM(_xlfn.ANCHORARRAY(A1))"
    sheet["D2"] = "=INDEX(_xlfn.ANCHORARRAY(A1),3,2)"
    sheet["D3"] = "=D1+D2"
    sheet["E1"] = "=SUM(A2:B3)"
    return book


@pytest.mark.parametrize("targeted", [False, True])
def test_saved_array_children_are_not_inputs_or_spill_blockers(targeted):
    result = analyze(
        package(array_book()),
        targets=["Sheet!D3"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    nodes = {node.get("addr"): node for node in result.nodes}
    for address, expected in (("D1", 21), ("D2", 6), ("D3", 27)):
        assert nodes[address]["value"] == expected
        assert nodes[address]["valueSource"] == "engine"
    assert nodes["D1"]["sourceFormula"] == "=SUM(_xlfn.ANCHORARRAY(A1))"


def test_target_of_only_array_children_recalculates_the_anchor():
    result = analyze(
        package(array_book()),
        targets=["Sheet!E1"],
        execution=ExecutionPolicy(isolated=False),
    )
    node = next(node for node in result.nodes if node.get("addr") == "E1")
    assert node["value"] == 18
    assert node["valueSource"] == "engine"


@pytest.mark.parametrize("targeted", [False, True])
@pytest.mark.parametrize("target", ["D3", "E1"])
@pytest.mark.parametrize("formula", ["=SEQUENCE(2,2)", "=SEQUENCE(1,1)"])
def test_changed_array_extent_and_dependents_are_not_certified(
    targeted, target, formula
):
    result = analyze(
        package(array_book(formula)),
        targets=[f"Sheet!{target}"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    for node in result.nodes:
        if node.get("addr") in {"D1", "D2", "D3", "E1"}:
            assert node.get("valueSource") != "engine"
            assert node["value"] is None
            assert node["steps"] is None
    assert any("extent" in warning for warning in result.warnings)


@pytest.mark.parametrize("targeted", [False, True])
@pytest.mark.parametrize("function, expected", [("SUM", 36), ("AVERAGE", 6)])
def test_batch_3d_aggregate_recalculates_source_order_references(
    targeted, function, expected
):
    book = Workbook()
    book.active.title = "First month"
    book.create_sheet("Middle")
    book.create_sheet("Last month")
    output = book.create_sheet("Output")
    for index, sheet in enumerate(book.worksheets[:3], 1):
        sheet["A1"] = index
        sheet["A2"] = "=5*2"
    output["A1"] = f"={function}('Last month:First month'!A1:A2)"
    output["A2"] = "=A1*2"
    result = analyze(
        package(book),
        targets=["Output!A2"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    node = next(node for node in result.nodes if node["id"] == "c:Output!A2")
    assert node["value"] == expected * 2
    assert node["valueSource"] == "engine"
    assert result.engine.get_formula("Output", 1, 1).lstrip("=") == output[
        "A1"
    ].value.lstrip("=")


@pytest.mark.parametrize("sheet_name", ["Sheet", "My Data", "O'Brien's Data"])
def test_native_private_copy_leaves_source_array_values_intact(sheet_name):
    book = array_book()
    book.active.title = sheet_name
    data = package(book)
    workbook = _open_workbook(data, parallel=False)
    workbook.evaluate_all()
    assert workbook.get_value(sheet_name, 3, 2) == 6
    from openpyxl import load_workbook

    assert load_workbook(io.BytesIO(data), data_only=True).active["B3"].value == 999
    result = analyze(data, execution=ExecutionPolicy(isolated=False))
    node = next(node for node in result.nodes if node.get("addr") == "D3")
    assert node["value"] == 27
    assert node["valueSource"] == "engine"
