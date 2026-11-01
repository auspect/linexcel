"""Batch source table references reuse the exact indexed selection contract."""

import io

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from linexcel import ExecutionPolicy, analyze
from linexcel.engine import _find_unresolvable


def table_package(formula, *, row=1, column=5):
    book = Workbook()
    sheet = book.active
    sheet.title = "Tables"
    for values in (
        ("Amount", "Qty", "Computed"),
        (10, 2, None),
        (20, 3, None),
        ("=SUM(A2:A3)", None, None),
    ):
        sheet.append(values)
    sheet.add_table(Table(displayName="SalesTable", ref="A1:C4", totalsRowCount=1))
    sheet.cell(row, column, formula)
    sheet["G1"] = f"={sheet.cell(row, column).coordinate}"
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize("targeted", [False, True])
@pytest.mark.parametrize(
    "formula, row, column, expected",
    [
        ("=[@Amount]*[@Qty]", 2, 3, 20),
        ("=SalesTable[@Amount]*SalesTable[@Qty]", 3, 3, 60),
        ("=SUM(SalesTable[Amount])", 1, 5, 30),
        ("=SUM(SalesTable[[#Totals],[Amount]])", 1, 5, 30),
        ("=ROWS(SalesTable[#All])", 1, 5, 4),
        ("=INDEX(SalesTable[#Headers],1,1)", 1, 5, "Amount"),
        ("=SUM(SalesTable[[Amount]:[Qty]])", 1, 5, 35),
        ("=SUM(salestable)", 1, 5, 35),
    ],
)
def test_batch_table_selections_recalculate(formula, row, column, expected, targeted):
    result = analyze(
        table_package(formula, row=row, column=column),
        targets=["Tables!G1"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    node = next(node for node in result.nodes if node.get("addr") == "G1")
    assert node["value"] == expected
    assert node["valueSource"] == "engine"
    source = next(node for node in result.nodes if node.get("sourceFormula") == formula)
    assert source["valueSource"] == "engine"


@pytest.mark.parametrize(
    "formula",
    [
        "=IFERROR(SUM(Missing[Amount]),99)",
        "=SUM(SalesTable[Missing])",
        "=SUM(SalesTable[[#Headers],[#Totals],[Amount]])",
    ],
)
@pytest.mark.parametrize("targeted", [False, True])
def test_unresolved_table_selection_and_dependents_are_not_recalculations(
    formula, targeted
):
    result = analyze(
        table_package(formula),
        targets=["Tables!G1"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    for node in result.nodes:
        if node.get("addr") in {"E1", "G1"}:
            assert node["value"] is None
            assert node.get("valueSource") != "engine"
            assert node["steps"] is None


@pytest.mark.parametrize("targeted", [False, True])
def test_formulatext_and_dependents_are_unverified(targeted):
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "= 1 +  2"
    sheet["A2"] = "=FORMULATEXT(A1)"
    sheet["A3"] = '=IFERROR(A2,"fallback")'
    sheet["A4"] = '=IFERROR(FORMULATEXT(A1),"fallback")'
    sheet["B1"] = '="FORMULATEXT(A1)"'
    stream = io.BytesIO()
    book.save(stream)
    result = analyze(
        stream.getvalue(),
        targets=["Sheet!A3", "Sheet!A4", "Sheet!B1"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    nodes = {node.get("addr"): node for node in result.nodes}
    for cell in ("A2", "A3", "A4"):
        assert nodes[cell]["value"] is None
        assert nodes[cell].get("valueSource") != "engine"
    assert nodes["B1"]["value"] == "FORMULATEXT(A1)"
    assert nodes["B1"]["valueSource"] == "engine"


def test_retry_keeps_supported_3d_and_table_formulas():
    book = Workbook()
    book.active.title = "Jan"
    book.active["A1"] = 1
    book.create_sheet("Feb")["A1"] = 2
    book.create_sheet("Mar")["A1"] = 3
    book.create_sheet("Output")["A1"] = "=SUM(Jan:Mar!A1)"
    stream = io.BytesIO()
    book.save(stream)
    assert _find_unresolvable(stream.getvalue(), set(book.sheetnames)) == {}
    assert (
        _find_unresolvable(table_package("=SUM(SalesTable[Amount])"), {"Tables"}) == {}
    )
