"""Batch calculations honor saved row metadata and keep refusals nonresults."""

import io
import zipfile
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.filters import AutoFilter, FilterColumn, Filters
from openpyxl.worksheet.table import Table

from linexcel import ExecutionPolicy, analyze
from linexcel.engine import _open_workbook
from linexcel.visibility import batch_visibility_resolver


def package(formula, *, filtered=False, table_filter=False, stale_cache=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    for row in (("Category", "Amount"), ("a", 10), ("b", 20), ("a", 30)):
        sheet.append(row)
    sheet.row_dimensions[3].hidden = True
    sheet["D1"] = formula
    sheet["F2"] = formula
    sheet["H1"] = "=D1+F2"
    sheet["J1"] = "=2+3"
    if filtered:
        sheet.auto_filter.ref = "A1:B4"
        sheet.auto_filter.add_filter_column(0, ["a"])
        sheet.sheet_properties.filterMode = True
    if table_filter:
        table = Table(displayName="SourceTable", ref="A1:B4")
        table.autoFilter = AutoFilter(
            ref="A1:B4",
            filterColumn=[FilterColumn(colId=0, filters=Filters(filter=["a"]))],
        )
        sheet.add_table(table)
    stream = io.BytesIO()
    book.save(stream)
    if not stale_cache:
        return stream.getvalue()
    output = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(output, "w") as dest:
        for part in source.infolist():
            content = source.read(part.filename)
            if part.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                for cell in root.findall("{*}sheetData/{*}row/{*}c"):
                    if cell.find("{*}f") is not None:
                        cell.find("{*}v").text = "999"
                content = ElementTree.tostring(root)
            dest.writestr(part, content)
    return output.getvalue()


@pytest.mark.parametrize("targeted", [False, True])
@pytest.mark.parametrize("filtered", [False, True])
@pytest.mark.parametrize("formula", ["=SUBTOTAL(109,B2:B4)", "=AGGREGATE(9,5,B2:B4)"])
def test_batch_hidden_totals_and_repeated_calls_recalculate(
    formula, filtered, targeted
):
    result = analyze(
        package(formula, filtered=filtered, stale_cache=True),
        targets=["Data!H1"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    nodes = {node.get("addr"): node for node in result.nodes}
    assert nodes["H1"]["value"] == 80
    assert nodes["H1"]["valueSource"] == "engine"
    assert nodes["D1"]["value"] == 40
    assert nodes["D1"]["valueSource"] == "engine"
    assert "LINEXCEL.VISIBILITY" not in nodes["D1"]["formula"]
    assert nodes["D1"]["sourceFormula"] == formula
    if nodes["D1"]["steps"] is not None:
        assert nodes["D1"]["steps"]["value"] == 40


def test_batch_nested_subtotal_decomposition_respects_hidden_rows():
    result = analyze(
        package("=SUBTOTAL(109,B2:B4)*2"),
        targets=["Data!D1"],
        execution=ExecutionPolicy(isolated=False),
    )
    node = next(node for node in result.nodes if node.get("addr") == "D1")
    assert node["value"] == 80
    pending = [node["steps"]] if node["steps"] else []
    while pending:
        step = pending.pop()
        if step.get("label") == "SUBTOTAL" and step.get("evaluated"):
            assert step["value"] == 40
        pending.extend(step.get("children", []))


@pytest.mark.parametrize("targeted", [False, True])
@pytest.mark.parametrize("table_filter", [False, True])
def test_batch_ambiguous_filter_and_its_dependent_are_not_recalculated(
    targeted, table_filter
):
    result = analyze(
        package(
            "=SUBTOTAL(9,B2:B4)", filtered=not table_filter, table_filter=table_filter
        ),
        targets=["Data!H1"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    nodes = {node.get("addr"): node for node in result.nodes}
    for address in ("D1", "H1"):
        assert nodes[address]["value"] is None
        assert nodes[address].get("valueSource") != "engine"
        assert nodes[address]["steps"] is None
    if not targeted:
        assert nodes["J1"]["value"] == 5
        assert nodes["J1"]["valueSource"] == "engine"


@pytest.mark.parametrize("targeted", [False, True])
@pytest.mark.parametrize(
    "formula",
    [
        "=IFERROR(AGGREGATE(12,5,Z1:Z2),99)",
        "=COUNT(AGGREGATE(12,5,Z1:Z2))",
    ],
)
def test_inline_aggregate_nonresult_cannot_become_a_fallback(formula, targeted):
    result = analyze(
        package(formula),
        targets=["Data!H1"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    nodes = {node.get("addr"): node for node in result.nodes}
    for address in ("D1", "H1"):
        assert nodes[address]["value"] is None
        assert nodes[address].get("valueSource") != "engine"


def test_failed_global_evaluation_cannot_recover_a_masked_nonresult(monkeypatch):
    import linexcel.engine as engine_module

    real_open = engine_module._open_workbook

    def interrupted(data, parallel):
        engine = real_open(data, parallel)

        def fail():
            raise RuntimeError("simulated evaluation interruption")

        engine.evaluate_all = fail
        return engine

    monkeypatch.setattr(engine_module, "_open_workbook", interrupted)
    result = analyze(
        package("=IFERROR(AGGREGATE(12,5,Z1:Z2),99)"),
        execution=ExecutionPolicy(isolated=False),
    )
    nodes = {node.get("addr"): node for node in result.nodes}
    for address in ("D1", "H1"):
        assert nodes[address]["value"] is None
        assert nodes[address].get("valueSource") != "engine"


def test_table_filter_visibility_applies_to_cells_outside_table_columns():
    data = package("=SUBTOTAL(9,B2:B4)", table_filter=True)
    engine = _open_workbook(data, parallel=False)
    resolve = batch_visibility_resolver(data, engine)
    metadata = resolve("dAtA!D2:D4", "Data")
    assert metadata["activeFilters"]
    assert metadata["hiddenRows"] == {3}
    assert metadata["expression"] == "Data!D2:D4"


def test_batch_visibility_resolution_keeps_empty_hidden_rows():
    book = Workbook()
    book.active.row_dimensions[2].hidden = True
    stream = io.BytesIO()
    book.save(stream)
    data = stream.getvalue()
    engine = _open_workbook(data, parallel=False)
    metadata = batch_visibility_resolver(data, engine)("A1:A3", "Sheet")
    assert metadata["hiddenRows"] == {2}
    assert metadata["visibilityKnown"]
