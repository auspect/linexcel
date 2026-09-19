"""Optional report context must not materialize millions of formatted cells."""

import io
import re
import zipfile

from openpyxl import Workbook
from openpyxl.comments import Comment

from linexcel import insights
from linexcel.aidoc import build_workbook_dossier


def test_large_context_streams_previews_and_reports_omissions(monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet["B2"] = "Revenue"
    sheet["C2"] = "=10+20"
    sheet["B2"].comment = Comment("Private context", "Analyst")
    sheet.freeze_panes = "B2"
    output = io.BytesIO()
    workbook.save(output)
    original = insights.load_workbook
    modes = []

    def load(*args, **kwargs):
        modes.append(kwargs["read_only"])
        return original(*args, **kwargs)

    monkeypatch.setattr(insights, "load_workbook", load)
    monkeypatch.setattr(insights, "MAX_CONTEXT_XML_BYTES", 0)
    context = insights.extract_workbook_context(output.getvalue())
    assert modes == [True]
    assert context["sheets"][0]["preview"][1] == {
        "row": 2,
        "values": [None, "Revenue", "=10+20"],
    }
    assert context["sheets"][0]["comments"] == []
    assert "not scanned" in context["warnings"][0]
    dossier = build_workbook_dossier({"nodes": [], "meta": {}}, context=context)
    assert dossier["warnings"] == context["warnings"]


def test_missing_dimensions_still_reads_bounded_preview(monkeypatch):
    workbook = Workbook()
    workbook.active["C4"] = "Visible without dimensions"
    output = io.BytesIO()
    workbook.save(output)
    patched = io.BytesIO()
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(patched, "w") as target:
        for part in source.infolist():
            data = source.read(part)
            if part.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension\b[^>]*/>", b"", data)
            target.writestr(part, data)
    monkeypatch.setattr(insights, "MAX_CONTEXT_XML_BYTES", 0)
    context = insights.extract_workbook_context(patched.getvalue())
    sheet = context["sheets"][0]
    assert sheet["preview"][3]["values"][2] == "Visible without dimensions"
    assert sheet["dimensions"] == {"rows": None, "columns": None}


def test_truncated_comments_populate_sheet_specific_warnings(monkeypatch):
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first["A1"].comment = Comment("First note", "Analyst")
    first["A2"].comment = Comment("Second note", "Analyst")
    second = workbook.create_sheet("Second")
    second["A1"] = "Clean"
    output = io.BytesIO()
    workbook.save(output)

    monkeypatch.setattr(insights, "MAX_COMMENTS_PER_SHEET", 1)
    context = insights.extract_workbook_context(output.getvalue())

    assert context["warnings"] == ["Comments on 'First' were truncated for inspection"]
    assert context["sheets"][0]["comments"] == [
        {"cell": "A1", "author": "Analyst", "text": "First note"}
    ]
    assert context["sheets"][0]["warnings"] == context["warnings"]
    assert context["sheets"][1]["warnings"] == []
