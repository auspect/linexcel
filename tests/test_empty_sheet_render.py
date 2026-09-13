"""Only demonstrably empty sheets may be excused by the screenshot gate."""

import io
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from linexcel import insights


def workbook_bytes():
    book = Workbook()
    book.active.title = "Empty Sheet"
    for name, value in (("Text", "hello"), ("Formula", '=IF(1,"","")'), ("Zero", 0)):
        book.create_sheet(name)["A1"] = value
    book.create_sheet("Style Only")["D8"].fill = PatternFill("solid", fgColor="FF0000")
    book.create_sheet("Comment Only")["A1"].comment = Comment("Note", "Author")
    source = book.create_sheet("Chart Data")
    source.append([1, 2])
    chart = BarChart()
    chart.add_data(Reference(source, min_col=1, max_col=2, min_row=1))
    book.create_chartsheet("Chart").add_chart(chart)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def patch_part(data, part, transform):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as dest:
        for info in source.infolist():
            content = source.read(info)
            dest.writestr(
                info, transform(content) if info.filename == part else content
            )
    return out.getvalue()


def test_only_plain_empty_sheet_is_exempted():
    exemptions = insights.empty_sheet_render_exemptions(workbook_bytes())
    assert set(exemptions) == {"Empty Sheet"}
    assert "no cells" in exemptions["Empty Sheet"]


@pytest.mark.parametrize(
    "content",
    [
        b'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
        b'relationships" r:id="rId1"/>',
        b"<headerFooter><oddHeader>Visible title</oddHeader></headerFooter>",
        b'<cols><col min="1" max="1" style="1"/></cols>',
        b'<printOptions gridLines="true"/>',
        b'<printOptions headings="1"/>',
        b"<unexpectedFeature/>",
    ],
)
def test_visible_or_unknown_features_prevent_exemption(content):
    data = patch_part(
        workbook_bytes(),
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(b"</worksheet>", content + b"</worksheet>"),
    )
    assert not insights.empty_sheet_render_exemptions(data)


def test_relationship_mapping_is_used_instead_of_sheet_numbers():
    data = patch_part(
        workbook_bytes(),
        "xl/_rels/workbook.xml.rels",
        lambda xml: (
            xml.replace(b"worksheets/sheet1.xml", b"worksheets/swap.xml")
            .replace(b"worksheets/sheet2.xml", b"worksheets/sheet1.xml")
            .replace(b"worksheets/swap.xml", b"worksheets/sheet2.xml")
        ),
    )
    assert set(insights.empty_sheet_render_exemptions(data)) == {"Text"}


def test_libreoffice_empty_print_metadata_does_not_require_a_screenshot():
    # LibreOffice adds these to a truly empty sheet during recalculation.
    # Empty header/footer nodes and page geometry do not add visible content.
    metadata = (
        b'<printOptions headings="false" gridLines="false" gridLinesSet="true" '
        b'horizontalCentered="false" verticalCentered="false"/>'
        b'<pageSetup paperSize="9" scale="100" fitToWidth="1" fitToHeight="1" '
        b'orientation="portrait" cellComments="none" horizontalDpi="300"/>'
        b'<headerFooter differentFirst="false" differentOddEven="false">'
        b"<oddHeader></oddHeader><oddFooter></oddFooter></headerFooter>"
    )
    data = patch_part(
        workbook_bytes(),
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(b"</worksheet>", metadata + b"</worksheet>"),
    )
    assert set(insights.empty_sheet_render_exemptions(data)) == {"Empty Sheet"}


def test_related_content_prevents_exemption_even_without_a_sheet_marker():
    data = workbook_bytes()
    buffer = io.BytesIO(data)
    with zipfile.ZipFile(buffer, "a") as package:
        package.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            '<Relationships><Relationship Id="rId1" '
            'Target="../comments.xml"/></Relationships>',
        )
    assert not insights.empty_sheet_render_exemptions(buffer.getvalue())


def test_large_or_malformed_xml_is_not_evidence_of_emptiness(monkeypatch):
    data = workbook_bytes()
    broken = patch_part(data, "xl/worksheets/sheet1.xml", lambda xml: b"<broken")
    assert not insights.empty_sheet_render_exemptions(broken)
    monkeypatch.setattr(insights, "MAX_CONTEXT_XML_BYTES", 10)
    assert not insights.empty_sheet_render_exemptions(data)
    assert not insights.empty_sheet_render_exemptions(b"not OOXML")
