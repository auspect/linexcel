"""Deterministic Excel oracles, never expected values copied from file caches."""

import io
import zipfile

import formualizer as fz
import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from linexcel.engine import _open_workbook
from linexcel.excel_compat import register_excel_functions, rewrite_formula


def imported(workbook):
    output = io.BytesIO()
    workbook.save(output)
    return _open_workbook(output.getvalue(), parallel=False)


def native(formula, inputs=()):
    engine = fz.Workbook()
    engine.add_sheet("S")
    register_excel_functions(engine)
    for address, value in inputs:
        engine.set_value("S", *address, value)
    engine.set_formula("S", 20, 10, rewrite_formula(formula, "S"))
    return engine.evaluate_cell("S", 20, 10)


@pytest.mark.parametrize(
    "formula,expected",
    [
        ('=SUMPRODUCT({2,"3",TRUE,"x"},{10,10,10,10})', 20),
        ('=SUMPRODUCT({2,"3"}*{10,10})', 50),
        ("=SUMPRODUCT({1,2;3,4},{5,6;7,8})", 70),
        ("=SUMPRODUCT({TRUE,FALSE})", 0),
        ('=SUMPRODUCT({"#VALUE!",4},{5,6})', 24),
        ("=SUMPRODUCT({1,2},{1;2})", {"type": "Error", "kind": "Value"}),
        ("=SUMPRODUCT({0,#N/A},{0,0})", {"type": "Error", "kind": "Na"}),
        ('=SUMPRODUCT({"x",#VALUE!},{0,0})', {"type": "Error", "kind": "Value"}),
    ],
)
def test_sumproduct_excel_typed_arrays(formula, expected):
    assert native(formula) == expected


def test_sumproduct_range_inputs_dependency_and_changed_source():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for r, value in enumerate([2, "3", True, "x", None], 1):
        ws.cell(r, 1, value)
        ws.cell(r, 2, 10)
    ws["C1"] = "=SUMPRODUCT(A1:A5,B1:B5)"
    ws["D1"] = "=C1+7"
    engine = imported(wb)
    assert engine.evaluate_cell("S", 1, 4) == 27
    engine.set_value("S", 1, 1, 5)
    assert engine.evaluate_cell("S", 1, 4) == 57
    assert engine.get_value("S", 1, 3) == 50


@pytest.mark.parametrize(
    "formula,expected",
    [
        ('=#VALUE!&" copied"', {"type": "Error", "kind": "Value"}),
        ('="#VALUE!"&" copied"', "#VALUE! copied"),
        ('=IFERROR(#VALUE!&" copied","handled")', "handled"),
        ('=IF(FALSE,#VALUE!&" copied",1)', 1),
        ('=1&TRUE&" x"', "1TRUE x"),
        ('=1="1"', False),
        ('=1<"1"', True),
        ('="a"="A"', True),
        ("=TRUE=1", False),
        ('=TRUE>"z"', True),
        ("=A1=0", True),
        ('=A1=""', True),
        ("=A1=FALSE", True),
        ('=#N/A="#N/A"', {"type": "Error", "kind": "Na"}),
        ('=SUMPRODUCT(--({1,2}="1"))', 0),
    ],
)
def test_typed_operators(formula, expected):
    assert native(formula) == expected


def test_error_input_propagates_through_dependents_while_equal_text_does_not():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "#VALUE!"
    ws["A2"] = "#VALUE!"
    ws["A2"].data_type = "s"
    ws["B1"] = '=A1&" copied"'
    ws["B2"] = '=A2&" copied"'
    ws["C1"] = '=IFERROR(B1,"handled")'
    engine = imported(wb)
    engine.evaluate_all()
    assert engine.get_value("S", 1, 2) == {"type": "Error", "kind": "Value"}
    assert engine.get_value("S", 2, 2) == "#VALUE! copied"
    assert engine.get_value("S", 1, 3) == "handled"
    assert "LINEXCEL" not in engine.get_formula("S", 1, 2)
    formulas = engine.sheet("S").get_formulas(fz.RangeAddress("S", 1, 2, 2, 2))
    assert all("LINEXCEL" not in row[0] for row in formulas)


def test_names_constants_scope_shadowing_and_lexical_let():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    other = wb.create_sheet("Other Sheet")
    wb.defined_names.add(DefinedName("CompanyName", attr_text='"Contoso Ltd"'))
    wb.defined_names.add(DefinedName("Rate", attr_text="4"))
    other.defined_names.add(DefinedName("Rate", attr_text="9"))
    ws["A1"] = "=CompanyName"
    ws["A2"] = '=A1&" invoice"'
    ws["B1"] = "=Rate*2"
    other["B1"] = "=Rate*2"
    ws["B2"] = "='Other Sheet'!Rate*3"
    ws["B3"] = "=LET(Rate,2,Rate*3)"
    ws["B4"] = "=LET(Rate,Rate+1,Rate*3)"
    engine = imported(wb)
    engine.evaluate_all()
    assert engine.get_value("S", 1, 1) == "Contoso Ltd"
    assert engine.get_value("S", 2, 1) == "Contoso Ltd invoice"
    assert engine.get_value("S", 1, 2) == 8
    assert engine.get_value("Other Sheet", 1, 2) == 18
    assert engine.get_value("S", 2, 2) == 27
    assert engine.get_value("S", 3, 2) == 6
    assert engine.get_value("S", 4, 2) == 15


def test_named_absolute_formula_reads_updated_dependencies():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 7
    wb.defined_names.add(DefinedName("Rate", attr_text="'S'!$A$1"))
    wb.defined_names.add(DefinedName("TwiceRate", attr_text="Rate*2"))
    ws["B1"] = "=TwiceRate+1"
    engine = imported(wb)
    assert engine.evaluate_cell("S", 1, 2) == 15
    engine.set_value("S", 1, 1, 8)
    assert engine.evaluate_cell("S", 1, 2) == 17


def test_name_scope_casefold_and_unknown_sheet_no_global_fallback():
    from linexcel.excel_compat import NameBindings

    names = NameBindings(["S", "Other"])
    names.update({(None, "RATE"): "2", ("Other", "RATE"): "5"})
    assert native(rewrite_formula("=other!Rate+Rate", "Other", names)) == 10
    assert rewrite_formula("=Missing!Rate", "Other", names) == "=Missing!Rate"


@pytest.mark.parametrize(
    "expression", ["S!$A1", "S!A$1", "S!$A$1:A2", "S!$A1+1", "$A$1", "$A$1*2"]
)
def test_relative_and_mixed_name_references_are_not_inlined(expression):
    assert rewrite_formula("=Rate", "S", {(None, "RATE"): expression}) == "=Rate"


def test_unsafe_defined_expression_does_not_reach_the_native_parser():
    expression = "1" + "+1" * 2000
    assert rewrite_formula("=Rate", "S", {(None, "RATE"): expression}) == "=Rate"


def test_safe_names_cannot_combine_into_an_unsafe_expanded_formula():
    from linexcel.engine import is_too_deep

    names = {(None, "INNER"): "1" + "+1" * 499, (None, "OUTER"): "INNER" + "+1" * 499}
    assert not is_too_deep("=" + names[(None, "INNER")])
    assert not is_too_deep("=" + names[(None, "OUTER")])
    assert rewrite_formula("=OUTER", "S", names) == "=OUTER"


def test_explicit_empty_text_is_not_an_absent_cell():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = ""
    ws["B1"] = "=COUNTA(A1:A2)"
    ws["B2"] = "=ISBLANK(A1)"
    ws["B3"] = "=ISBLANK(A2)"
    engine = imported(wb)
    engine.evaluate_all()
    assert [engine.get_value("S", r, 2) for r in [1, 2, 3]] == [1, False, True]


def test_whitespace_text_survives_import_and_comparison():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = " "
    ws["B1"] = '=A1=""'
    ws["B2"] = "=ISBLANK(A1)"
    ws["B3"] = '=IF(A1="","empty","space")'
    engine = imported(wb)
    engine.evaluate_all()
    assert [engine.get_value("S", r, 2) for r in [1, 2, 3]] == [False, False, "space"]


@pytest.mark.parametrize("epoch", ["1900", "1904"])
def test_dates_match_excel_serial_oracles_and_update_dependents(epoch):
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904 if epoch == "1904" else CALENDAR_WINDOWS_1900
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 2026
    ws["B1"] = "=DATE(A1,2,1)"
    ws["B2"] = f"=B1-{44592 if epoch == '1904' else 46054}"
    ws["B3"] = "=B2+1"
    formulas = [
        "=DATE(1900,2,29)",
        "=DATE(1900,3,0)",
        "=DATE(1900,1,0)",
        "=DATE(1900,1,-1)",
        "=DATE(1904,1,1)",
        "=DATE(9999,12,31)",
        "=DATE(9999,12,32)",
        '=DATE("2026","2","1")',
    ]
    for row, formula in enumerate(formulas, 5):
        ws.cell(row, 2, formula)
    engine = imported(wb)
    engine.evaluate_all()
    num = {"type": "Error", "kind": "Num"}
    expected = (
        [60, 60, 0, num, 1462, 2958465, num, 46054]
        if epoch == "1900"
        else [num, num, num, num, 0, 2957003, num, 44592]
    )
    assert [engine.get_value("S", row, 2) for row in range(5, 13)] == expected
    assert engine.get_value("S", 2, 2) == 0
    assert engine.get_value("S", 3, 2) == 1
    engine.set_value("S", 1, 1, 2027)
    assert engine.evaluate_cell("S", 3, 2) == 366


@pytest.mark.parametrize("epoch", ["1900", "1904"])
def test_day_month_year_keep_numeric_serial_60_from_source(epoch):
    from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

    wb = Workbook()
    wb.epoch = CALENDAR_MAC_1904 if epoch == "1904" else CALENDAR_WINDOWS_1900
    ws = wb.active
    ws.title = "S"
    values = [0, -1, 60, 60.75, "60", "2026-02-01"]
    for row, value in enumerate(values, 1):
        ws.cell(row, 1, value).number_format = "yyyy-mm-dd"
        for col, part in enumerate(["DAY", "MONTH", "YEAR"], 2):
            ws.cell(row, col, f"={part}(A{row})")
    engine = imported(wb)
    engine.evaluate_all()
    num = {"type": "Error", "kind": "Num"}
    expected = (
        [[0, 1, 1900], [num] * 3, *([[29, 2, 1900]] * 3), [1, 2, 2026]]
        if epoch == "1900"
        else [[1, 1, 1904], [num] * 3, *([[1, 3, 1904]] * 3), [1, 2, 2026]]
    )
    assert [
        [engine.get_value("S", row, col) for col in (2, 3, 4)] for row in range(1, 7)
    ] == expected
    assert engine.get_value("S", 3, 1) == 60


def test_external_values_preserve_error_type_and_equal_literal_text(tmp_path):
    from linexcel.analyzer import analyze_workbook
    from linexcel.external import _typed, read_workbook_values
    from linexcel.values import SpreadsheetError

    source = Workbook()
    source.active.title = "Data"
    source.active["A1"] = "#N/A"
    source.active["A1"].data_type = "s"
    source.active["A2"] = "#N/A"
    source.save(tmp_path / "ref.xlsx")
    values = read_workbook_values(tmp_path / "ref.xlsx")
    assert type(values[("Data", 1, 1)]) is str
    assert isinstance(values[("Data", 2, 1)], SpreadsheetError)
    assert type(_typed("#N/A", "str")) is str
    assert isinstance(_typed("#N/A", "e"), SpreadsheetError)
    for row, expected in [(1, "prefix #N/A"), (2, "caught")]:
        wb = Workbook()
        wb.active.title = "S"
        wb.active["A1"] = f'=IFERROR("prefix "&\'[ref.xlsx]Data\'!A{row},"caught")'
        stream = io.BytesIO()
        wb.save(stream)
        graph = analyze_workbook(stream.getvalue(), "external.xlsx", refs_dir=tmp_path)[
            "graph"
        ]
        nodes = {node["id"]: node for node in graph["nodes"]}
        assert nodes["c:S!A1"]["value"] == expected


def test_prefixed_xml_formula_still_gets_typed_operator_correction():
    from xml.etree import ElementTree as ET

    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = '=IFERROR(#N/A&"x","caught")'
    stream = io.BytesIO()
    wb.save(stream)
    result = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(result, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = ET.tostring(ET.fromstring(content))
            target.writestr(entry, content)
    engine = _open_workbook(result.getvalue(), parallel=False)
    assert engine.evaluate_cell("S", 1, 1) == "caught"


@pytest.mark.parametrize("reader", ["calamine", "openpyxl"])
def test_cached_phantom_day_keeps_raw_serial_without_fictitious_iso(reader):
    import datetime

    from linexcel.loader import (
        _load_cached_values_calamine,
        _load_cached_values_openpyxl,
    )
    from linexcel.resolver import _Budget, _ValueResolver
    from linexcel.values import serial_to_date_text

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    originals = [59, 60, 60.75, 60.999999999, 61]
    for row, serial in enumerate(originals, 1):
        ws.cell(row, 1, serial).number_format = "yyyy-mm-dd hh:mm:ss"
    stream = io.BytesIO()
    wb.save(stream)
    data = stream.getvalue()
    load = (
        _load_cached_values_calamine
        if reader == "calamine"
        else _load_cached_values_openpyxl
    )
    cached = load(data)
    assert cached.get("S", 1, 1) == datetime.datetime(1900, 2, 28)
    assert cached.get("S", 5, 1) == datetime.datetime(1900, 3, 1)
    engine = _open_workbook(data, parallel=False)
    resolver = _ValueResolver(engine, {"S"}, cached, [], _Budget(0), False)
    for row in (2, 3, 4):
        assert cached.get("S", row, 1) == originals[row - 1]
        fields = resolver.describe("S", row, 1)
        assert fields["value"] == fields["cachedValue"] == originals[row - 1]
        assert fields["cachedAgreement"] == "same"
        assert "valueDate" not in fields
        assert serial_to_date_text(originals[row - 1]) is None


def test_false_formula_cache_is_never_a_calculation_oracle():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = '=SUMPRODUCT({2,"3"},{10,10})'
    ws["A2"] = "=A1+1"
    stream = io.BytesIO()
    wb.save(stream)
    patched = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(patched, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<v></v>", b"<v>999999</v>")
            target.writestr(entry, content)
    engine = _open_workbook(patched.getvalue(), parallel=False)
    assert engine.evaluate_cell("S", 2, 1) == 21


def test_sqrt_negative_keeps_correct_num_error_not_a_foreign_cache():
    assert native("=SQRT(-1)") == {"type": "Error", "kind": "Num"}
    assert native('=IFERROR(SQRT(-1),"ok")') == "ok"


def test_step_literals_keep_error_and_text_distinct():
    from linexcel.decompose import _render_expr, typed_ast_dict

    error_ast = typed_ast_dict(fz.parse('=IFERROR(#VALUE!&"x","handled")'))
    text_ast = typed_ast_dict(fz.parse('=IFERROR("#VALUE!"&"x","handled")'))
    assert native("=" + _render_expr(error_ast)) == "handled"
    assert native("=" + _render_expr(text_ast)) == "#VALUE!x"


@pytest.mark.parametrize("cache_is_error", [True, False])
def test_cache_comparison_distinguishes_error_from_equal_text(cache_is_error):
    from linexcel.loader import load_cached_values
    from linexcel.resolver import _Budget, _ValueResolver

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = '="#VALUE!"' if cache_is_error else "=#VALUE!"
    stream = io.BytesIO()
    wb.save(stream)
    result = io.BytesIO()
    from xml.etree import ElementTree as ET

    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(result, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                cell = root.find('.//{*}c[@r="A1"]')
                cell.set("t", "e" if cache_is_error else "str")
                cell.find("{*}v").text = "#VALUE!"
                content = ET.tostring(root)
            target.writestr(entry, content)
    data = result.getvalue()
    cached = load_cached_values(data)
    assert cached.is_error("S", 1, 1) is cache_is_error
    engine = _open_workbook(data, parallel=False)
    engine.evaluate_all()
    resolver = _ValueResolver(engine, {"S"}, cached, [], _Budget(0), False)
    fields = resolver.describe("S", 1, 1)
    assert fields["value"] == fields["cachedValue"] == "#VALUE!"
    assert fields["cachedAgreement"] == "differ"
    assert fields["cachedValueKind"] == ("error" if cache_is_error else "text")
    assert fields["valueKind"] == ("text" if cache_is_error else "error")
