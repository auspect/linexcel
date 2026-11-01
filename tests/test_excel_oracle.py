"""Protocol/corpus tests; these do not claim a native Excel execution."""

import importlib.util
import io
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "excel_oracle.py"
SPEC = importlib.util.spec_from_file_location("excel_oracle", SCRIPT)
assert SPEC and SPEC.loader
oracle = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(oracle)


def test_curated_corpus_deterministic_typed_and_contains_declared_families():
    sources, cases = oracle.corpus()
    again, again_cases = oracle.corpus()
    assert sources == again and cases == again_cases
    families = {target["family"] for case in cases for target in case["targets"]}
    assert families == {
        "dates",
        "coercion_errors",
        "tables",
        "names_3d",
        "visibility",
        "aggregate_edges",
        "dynamic_arrays",
        "spill_children",
        "array_dependents",
        "dynamic_context",
        "volatiles",
        "cycles",
        "cycles_divergent",
        "external_local",
    }
    for data in sources.values():
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            assert not any(
                "vba" in name.lower()
                or "macrosheet" in name.lower()
                or "connections" in name.lower()
                or "query" in name.lower()
                for name in names
            )
            assert all(
                item.date_time == (2026, 1, 1, 0, 0, 0) for item in archive.infolist()
            )
    book = load_workbook(io.BytesIO(sources["scalar-1900.xlsx"]))
    assert book["Scalar"]["A5"].data_type == "s"
    assert book["Scalar"]["A6"].data_type == "e"
    assert load_workbook(io.BytesIO(sources["scalar-1904.xlsx"])).epoch.year == 1904
    external = load_workbook(io.BytesIO(sources["external-local.xlsx"]))
    assert external._external_links[0].file_link.Target == "refs/source.xlsx"
    arrays = next(case for case in cases if case["id"] == "arrays")
    assert arrays["targets"][0]["shape"] == [3, 2]
    assert arrays["actions"][0]["formula"] == "=SEQUENCE(3,2)"


def test_prepare_is_immutable_and_modified_corpus_cannot_be_run(tmp_path):
    bundle = tmp_path / "bundle"
    oracle.prepare(bundle)
    manifest = oracle.validate_bundle(bundle)
    assert len(manifest["cases"]) == 10
    with pytest.raises(FileExistsError):
        oracle.prepare(bundle)
    file = bundle / "sources" / "scalar-1900.xlsx"
    file.chmod(0o600)
    file.write_bytes(b"not a curated workbook")
    with pytest.raises(ValueError, match="Source changed"):
        oracle.validate_bundle(bundle)
    assert not (bundle / "runs").exists()


@pytest.mark.parametrize("change", ["manifest", "extra", "symlink"])
def test_bundle_rejects_unlisted_files_or_manifest_changes(tmp_path, change):
    bundle = tmp_path / "bundle"
    oracle.prepare(bundle)
    if change == "manifest":
        manifest = json.loads((bundle / "manifest.json").read_text())
        manifest["cases"][0]["actions"] = [{"kind": "malicious"}]
        oracle.write_json(bundle / "manifest.json", manifest)
    elif change == "extra":
        (bundle / "sources" / "extra.xlsx").write_bytes(b"unlisted")
    else:
        if os.name == "nt":
            pytest.skip("symlink privileges vary on Windows")
        (bundle / "sources" / "extra.xlsx").symlink_to(
            bundle / "sources" / "tables.xlsx"
        )
    with pytest.raises(ValueError):
        oracle.validate_bundle(bundle)


def test_typed_comparison_does_not_conflate_errors_text_bool_blank_or_shapes():
    def compare(a, b):
        return oracle.equivalent(a, b, oracle.TOLERANCE)

    assert compare(oracle.typed(1), oracle.typed(1 + 1e-10))
    assert not compare(oracle.typed(1), oracle.typed(1.01))
    assert not compare(oracle.typed(True), oracle.typed(1))
    assert not compare(oracle.typed(None), oracle.typed(""))
    assert not compare(oracle.typed("#N/A"), oracle.typed("#N/A", "error"))
    assert compare(oracle.typed("#N/A", "error"), oracle.typed("#N/A", "error"))
    assert compare(oracle.typed([[1, 2], [3, 4]]), oracle.typed([[1, 2], [3, 4]]))
    assert not compare(oracle.typed([[1, 2], [3, 4]]), oracle.typed([[1, 2, 3, 4]]))
    assert not compare(oracle.typed(float("nan")), oracle.typed(float("nan")))


def native_result(spec):
    return {
        "status": "completed",
        "excel": {"version": "16.0", "build": "test-only", "automationSecurity": 3},
        "calculation": {
            "mode": -4135,
            "iteration": spec["iterate"],
            "maxIterations": spec["maxIterations"],
            "maxChange": spec["maxChange"],
            "calculateBeforeSave": False,
            "date1904": spec["id"] == "scalar-1904",
        },
        "observations": {
            selection["nodeId"]: {"value": oracle.typed(42)}
            for selection in spec["targets"]
        },
    }


@pytest.mark.parametrize("missing", ["excel", "calculation", "observations"])
def test_native_receipt_requires_excel_settings_and_every_target(missing):
    _, cases = oracle.corpus()
    native = native_result(cases[0])
    oracle.validate_native(cases[0], native)
    del native[missing]
    with pytest.raises(ValueError):
        oracle.validate_native(cases[0], native)


def test_nonwindows_run_does_not_fabricate_excel_receipt(tmp_path):
    if os.name == "nt":
        pytest.skip("non-Windows refusal")
    bundle = tmp_path / "bundle"
    oracle.prepare(bundle)
    assert oracle.main(["--run", str(bundle)]) == 1
    assert not (bundle / "runs").exists()
    assert not (bundle / "latest-run.json").exists()


def test_missing_oracle_cases_cannot_be_reported_as_agreement(tmp_path):
    bundle = tmp_path / "bundle"
    oracle.prepare(bundle)
    receipt = bundle / "runs" / "missing" / "receipt.json"
    receipt.parent.mkdir(parents=True)
    oracle.write_json(
        receipt,
        {
            "format": "linexcel-excel-receipt",
            "version": oracle.PROTOCOL,
            "manifestSha256": oracle.sha256((bundle / "manifest.json").read_bytes()),
            "cases": [],
            "status": "incomplete",
        },
    )
    path = oracle.compare(bundle, receipt, 5)
    report = json.loads(path.read_text())
    assert report["status"] == "differences_or_incomplete"
    assert set(report["counts"]) == {"oracle_unavailable"}
    assert len(report["rows"]) == 2 * sum(
        len(case["targets"]) for case in oracle.corpus()[1]
    )
    assert not list(path.parent.glob("request-*"))


@pytest.mark.parametrize("mutation", ["bytes", "outside", "case", "cleanup"])
def test_receipt_rejects_changed_or_misidentified_outputs(tmp_path, mutation):
    bundle = tmp_path / "bundle"
    oracle.prepare(bundle)
    manifest = oracle.validate_bundle(bundle)
    spec = manifest["cases"][0]
    case_dir = bundle / "runs" / "synthetic-protocol-test" / spec["id"]
    case_dir.mkdir(parents=True)
    file = case_dir / spec["file"]
    file.write_bytes((bundle / "sources" / spec["file"]).read_bytes())
    record = {
        "id": spec["id"],
        "status": "completed",
        "cleanup": {"status": "checked"},
        "excel": native_result(spec),
        "outputs": {
            spec["file"]: {
                "path": file.relative_to(bundle).as_posix(),
                "sha256": oracle.sha256(file.read_bytes()),
            }
        },
    }
    receipt = {
        "format": "linexcel-excel-receipt",
        "version": oracle.PROTOCOL,
        "manifestSha256": oracle.sha256((bundle / "manifest.json").read_bytes()),
        "runId": "synthetic-protocol-test",
        "cases": [record],
    }
    path = case_dir.parent / "receipt.json"
    oracle.write_json(path, receipt)
    oracle.validate_receipt(bundle, manifest, path)
    if mutation == "bytes":
        file.write_bytes(b"changed after Excel run")
    elif mutation == "outside":
        record["outputs"][spec["file"]]["path"] = "runs/../sources/" + spec["file"]
    elif mutation == "case":
        record["id"] = manifest["cases"][1]["id"]
    else:
        record["cleanup"]["status"] = "unverified"
    oracle.write_json(path, receipt)
    with pytest.raises(ValueError):
        oracle.validate_receipt(bundle, manifest, path)


def test_run_watchdog_kills_only_recorded_owned_instance_and_keeps_seeds(
    tmp_path, monkeypatch
):
    bundle = tmp_path / "bundle"
    oracle.prepare(bundle)
    monkeypatch.setattr(oracle, "windows_powershell", lambda: "mock-powershell")
    calls = []

    class HungChild:
        returncode = None

        def __init__(self, argv, **kwargs):
            self.argv = argv
            calls.append(self)

        def wait(self, timeout):
            if self.returncode is None:
                raise subprocess.TimeoutExpired(self.argv, timeout)
            return self.returncode

        def kill(self):
            self.returncode = -9

        def poll(self):
            return self.returncode

    monkeypatch.setattr(oracle.subprocess, "Popen", HungChild)

    # No ownership file means no cleanup subprocess is allowed to kill Excel.
    def forbidden(*args, **kwargs):
        pytest.fail("must not invoke cleanup without an ownership receipt")

    monkeypatch.setattr(oracle.subprocess, "run", forbidden)
    receipt = json.loads(oracle.run_excel(bundle, 0.01).read_text())
    assert len(calls) == 10 and all(child.returncode == -9 for child in calls)
    assert receipt["status"] == "incomplete"
    assert all(case["status"] == "timed_out" for case in receipt["cases"])
    assert all(case["cleanup"]["status"] == "unverified" for case in receipt["cases"])
    oracle.validate_bundle(bundle)


@pytest.mark.parametrize("mode", ["lazy", "batch"])
def test_engine_adapter_performs_real_calculation_in_subprocess(tmp_path, mode):
    sources, _ = oracle.corpus()
    file = tmp_path / "probe.xlsx"
    file.write_bytes(sources["scalar-1900.xlsx"])
    output = tmp_path / "result.json"
    request = tmp_path / "request.json"
    oracle.write_json(
        request,
        {
            "file": str(file),
            "nodeId": "Scalar!C1",
            "refs": None,
            "mode": mode,
            "seconds": 10,
            "result": str(output),
        },
    )
    completed = subprocess.run(
        [sys.executable, str(SCRIPT), "--engine-request", str(request)],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert completed.returncode == 0, completed.stderr
    result = json.loads(output.read_text())
    assert result["calculated"] is True
    assert result["value"] == {"kind": "number", "value": 47}


def test_lazy_array_adapter_recalculates_each_spill_cell_with_stale_saved_values(
    tmp_path,
):
    from openpyxl import Workbook
    from openpyxl.worksheet.formula import ArrayFormula

    book = Workbook()
    book.active.title = "Arrays"
    book.active["A1"] = ArrayFormula(ref="A1:B3", text="=SEQUENCE(3,2)")
    for cell in ("B1", "A2", "B2", "A3", "B3"):
        book.active[cell] = 999
    file = tmp_path / "array.xlsx"
    book.save(file)
    request, output = tmp_path / "request.json", tmp_path / "result.json"
    oracle.write_json(
        request,
        {
            "file": str(file),
            "nodeId": "Arrays!A1",
            "range": "A1:B3",
            "refs": None,
            "mode": "lazy",
            "seconds": 10,
            "result": str(output),
        },
    )
    completed = subprocess.run(
        [sys.executable, str(SCRIPT), "--engine-request", str(request)],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert completed.returncode == 0, completed.stderr
    result = json.loads(output.read_text())
    assert result["calculated"] is True
    assert oracle.equivalent(
        result["value"], oracle.typed([[1, 2], [3, 4], [5, 6]]), oracle.TOLERANCE
    )
    assert len(result["cells"]) == 6
    assert all(
        cell["arrayEvidence"][0]["observed"]["extent"] == "Arrays!A1:B3"
        for cell in result["cells"].values()
    )
