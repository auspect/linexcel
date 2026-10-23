"""Static presentation is bounded and never changes stored source values."""

import datetime as dt
import io
import zipfile
import zlib
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from linexcel.excel_presentation import (
    PresentationLimitError,
    _bounded_node,
    _color,
    _SafeXML,
    format_display,
    sheet_presentation,
)
from linexcel.lazy_store import WorkbookIndex, _pack, build_index


@pytest.mark.parametrize(
    ("value", "code", "display"),
    [
        (1234.567, "#,##0.00", "1,234.57"),
        (0.125, "0.0%", "12.5%"),
        (2.675, "0.00", "2.68"),
        (-2.5, "0", "-3"),
        (-12, "0.00;[Red](0.00)", "(12.00)"),
        (12, '"€ "#,##0.00', "€ 12.00"),
    ],
)
def test_common_number_formats_keep_source_scalar(value, code, display):
    assert format_display(value, code) == (display, "formatted")


@pytest.mark.parametrize(
    "code", ["0,", "# ??/??", "0.00E+00", "[>=10]0", "[$€-40C]0", "0_)"]
)
def test_complex_formats_fail_to_labelled_raw_values(code):
    assert format_display(1234.5, code) == ("1234.5", "unsupported")


def test_date_epochs_times_and_fictitious_1900_day_are_explicit():
    assert format_display("2024-01-01", "yyyy-mm-dd", raw="45292") == (
        "2024-01-01",
        "formatted",
    )
    assert format_display("1904-01-01", "dd/mm/yyyy", raw="0", epoch1904=True) == (
        "01/01/1904",
        "formatted",
    )
    assert format_display("12:34:56", "h:mm:ss AM/PM", raw="0.524259259259259") == (
        "12:34:56 PM",
        "formatted",
    )
    for raw in ("60", "-1"):
        assert format_display(float(raw), "yyyy-mm-dd", raw=raw)[1] == "unsupported"
    assert (
        format_display("2024-01-01", "dddd, mmmm d, yyyy", raw="45292")[1]
        == "unsupported"
    )
    assert format_display("12:00", "[h]:mm", raw="1.5")[1] == "unsupported"
    assert format_display("14:30", "h:mm A/P", raw="0.6041666667")[1] == "unsupported"
    assert format_display(-12, "$0.00") == ("-12", "unsupported")


def test_static_window_keeps_formats_geometry_hidden_cells_and_missing_caches(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "A & B"
    sheet["A1"] = "literal <script>alert(1)</script>"
    sheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    sheet["A1"].fill = PatternFill("solid", fgColor=Color(theme=4, tint=0.2))
    sheet["A1"].alignment = Alignment(horizontal="center", wrap_text=True)
    sheet["A1"].border = Border(bottom=Side(style="thick", color="123456"))
    sheet.merge_cells("A1:C1")
    sheet["A2"] = 1234.567
    sheet["A2"].number_format = "#,##0.00"
    sheet["B2"] = "=A2*2"
    sheet["C2"] = True
    sheet["D2"] = "#DIV/0!"
    sheet["E4"].fill = PatternFill("solid", fgColor="ABCDEF")
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].hidden = True
    sheet.column_dimensions["F"].fill = PatternFill("solid", fgColor="FEDCBA")
    sheet.row_dimensions[5].font = Font(italic=True)
    sheet.freeze_panes = "B2"
    source, index = tmp_path / "book.xlsx", tmp_path / "index.sqlite"
    book.save(source)
    original = source.read_bytes()
    build_index(source, index)
    view = sheet_presentation(source, index, sheet.title, rows=6, columns=6)
    cells = {cell["address"]: cell for cell in view["cells"]}
    assert source.read_bytes() == original
    assert cells["A1"]["display"] == "literal <script>alert(1)</script>"
    heading = view["styles"][str(cells["A1"]["styleId"])]["css"]
    assert heading["fontWeight"] == "bold" and heading["color"] == "#FFFFFF"
    assert heading["backgroundColor"].startswith("#")
    assert heading["borderBottom"] == "3px solid #123456"
    assert heading["textAlign"] == "center" and heading["whiteSpace"] == "pre-wrap"
    assert cells["A2"]["value"] == 1234.567
    assert cells["A2"]["display"] == "1,234.57"
    assert cells["B2"]["value"] is None and cells["B2"]["display"] == ""
    assert cells["B2"]["valueStatus"] == "missing"
    assert cells["B2"]["formula"] == "=A2*2"
    assert cells["C2"]["display"] == "TRUE"
    assert cells["D2"]["display"] == "#DIV/0!"
    assert cells["E4"]["valueStatus"] == "blank"
    assert (
        view["styles"][str(cells["E4"]["styleId"])]["css"]["backgroundColor"]
        == "#ABCDEF"
    )
    assert (
        view["styles"][str(cells["F6"]["styleId"])]["css"]["backgroundColor"]
        == "#FEDCBA"
    )
    assert view["styles"][str(cells["A5"]["styleId"])]["css"]["fontStyle"] == "italic"
    assert "styleId" not in view["rowDimensions"]["1"]
    assert view["rowDimensions"]["1"]["height"] == 40
    assert view["rowDimensions"]["3"]["hidden"]
    assert view["columnDimensions"]["2"]["width"] == 159
    assert view["columnDimensions"]["3"]["hidden"]
    assert view["freezePane"]["topLeftCell"] == "B2"
    assert view["merges"] == [
        {"range": "A1:C1", "row": 1, "column": 1, "lastRow": 1, "lastColumn": 3}
    ]
    clipped = sheet_presentation(
        source, index, sheet.title, row=1, column=2, rows=2, columns=2
    )
    assert clipped["merges"][0]["column"] == 1
    assert all(cell["address"] != "A1" for cell in clipped["cells"])


def test_1904_workbook_and_budget_failures(tmp_path, monkeypatch):
    book = Workbook()
    book.epoch = CALENDAR_MAC_1904
    book.active["A1"] = dt.datetime(2025, 6, 7)
    book.active["A1"].number_format = "dd/mm/yyyy"
    source, index = tmp_path / "book.xlsx", tmp_path / "index.sqlite"
    book.save(source)
    build_index(source, index)
    view = sheet_presentation(source, index, book.active.title, rows=999, columns=999)
    assert view["cells"][0]["display"] == "07/06/2025"
    assert view["window"]["rows"] == 100 and view["window"]["columns"] == 40
    with pytest.raises(ValueError):
        sheet_presentation(source, index, book.active.title, row=0)
    monkeypatch.setattr("linexcel.excel_presentation.MAX_WORKSHEET_BYTES", 1)
    with pytest.raises(PresentationLimitError):
        sheet_presentation(source, index, book.active.title)


def test_theme_indexed_colors_and_xml_entities_are_bounded():
    assert _color(ET.fromstring('<color indexed="2"/>'), []) == "#FF0000"
    assert _color(ET.fromstring('<color rgb="00AABBCC"/>'), []) == "#AABBCC"
    assert _color(ET.fromstring('<color rgb="url(evil)"/>'), []) is None
    stream = _SafeXML(io.BytesIO(b'<!DOCTYPE x [<!ENTITY e "huge">]><x>&e;</x>'))
    with pytest.raises(ValueError):
        while stream.read(3):
            pass
    stream = _SafeXML(
        io.BytesIO('<!DOCTYPE x [<!ENTITY e "huge">]><x/>'.encode("utf-16"))
    )
    with pytest.raises(ValueError):
        stream.read()


def test_metadata_and_worksheet_entity_declarations_are_rejected(tmp_path):
    book = Workbook()
    book.active["A1"] = 1
    source, index = tmp_path / "book.xlsx", tmp_path / "index.sqlite"
    book.save(source)
    build_index(source, index)
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["xl/worksheets/sheet1.xml"] = (
        b'<!DOCTYPE x [<!ENTITY e "huge">]>' + entries["xl/worksheets/sheet1.xml"]
    )
    with zipfile.ZipFile(source, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    with pytest.raises(ValueError):
        sheet_presentation(source, index, book.active.title)


def test_shared_follower_never_labels_translated_formula_as_source(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet["A1"], sheet["A2"], sheet["B1"], sheet["B2"] = "=B1*2", "=B2*2", 1, 2
    source, index = tmp_path / "book.xlsx", tmp_path / "index.sqlite"
    book.save(source)
    with zipfile.ZipFile(source) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
    for address in ("A1", "A2"):
        formula = root.find(f"{{*}}sheetData/{{*}}row/{{*}}c[@r='{address}']/{{*}}f")
        formula.set("t", "shared")
        formula.set("si", "0")
        if address == "A1":
            formula.set("ref", "A1:A2")
        else:
            formula.text = None
    entries["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with zipfile.ZipFile(source, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    build_index(source, index)
    view = sheet_presentation(source, index, sheet.title)
    cells = {cell["address"]: cell for cell in view["cells"]}
    assert cells["A1"]["formulaRole"] == "source"
    assert cells["A1"]["sourceFormula"] == "=B1*2"
    assert cells["A2"]["formulaRole"] == "canonical"
    assert cells["A2"]["sourceFormula"] is None
    assert cells["A2"]["canonicalFormula"] == "=B2*2"


@pytest.mark.parametrize("oversize_format", [False, True])
def test_presentation_rejects_excessive_format_or_response_payload(
    tmp_path, monkeypatch, oversize_format
):
    book = Workbook()
    book.active["A1"] = "a source value"
    if oversize_format:
        book.active["A1"].number_format = '"' + "x" * 1025 + '"0'
    else:
        monkeypatch.setattr("linexcel.excel_presentation.MAX_PRESENTATION_BYTES", 10)
    source, index = tmp_path / "book.xlsx", tmp_path / "index.sqlite"
    book.save(source)
    build_index(source, index)
    with pytest.raises(PresentationLimitError):
        sheet_presentation(source, index, book.active.title)


def test_compressed_node_limit_precedes_unbounded_decoder(monkeypatch):
    calls = []
    monkeypatch.setattr("linexcel.excel_presentation._unpack", calls.append)
    with pytest.raises(PresentationLimitError):
        _bounded_node(_pack({"value": "x" * 100000}), 1000)
    assert not calls
    with pytest.raises(ValueError):
        _bounded_node(b"invalid", 1000)
    with pytest.raises(ValueError):
        _bounded_node(zlib.compress(b"{}")[0:-2], 1000)


def test_cumulative_node_limit_precedes_retaining_the_whole_window(
    tmp_path, monkeypatch
):
    import linexcel.excel_presentation as presentation

    book = Workbook()
    for row in range(1, 9):
        book.active.cell(row, 1, "repeated source text" * 10)
    source, index = tmp_path / "book.xlsx", tmp_path / "index.sqlite"
    book.save(source)
    build_index(source, index)
    with WorkbookIndex(index) as store:
        sizes = [
            len(zlib.decompress(packed))
            for (packed,) in store.db.execute("SELECT data FROM nodes ORDER BY row")
        ]
    assert len(sizes) == 8
    monkeypatch.setattr(presentation, "MAX_PRESENTATION_BYTES", sum(sizes[:2]))
    decoded, original = [], presentation._unpack

    def count(packed):
        decoded.append(packed)
        return original(packed)

    monkeypatch.setattr(presentation, "_unpack", count)
    with pytest.raises(PresentationLimitError):
        sheet_presentation(source, index, book.active.title, rows=8, columns=1)
    assert len(decoded) == 2
