"""Independent engine and linexcel integration contracts for Formualizer."""

import io
import math

import formualizer as fz
import pytest
from openpyxl import Workbook

from linexcel.engine import _open_workbook


def imported(book):
    stream = io.BytesIO()
    book.save(stream)
    return _open_workbook(stream.getvalue(), parallel=False)


def native_formula(formula):
    engine = fz.Workbook()
    engine.add_sheet("S")
    engine.set_formula("S", 1, 1, formula)
    return engine.evaluate_cell("S", 1, 1)


def raw_workbook(formulas, values=()):
    engine = fz.Workbook()
    engine.add_sheet("S")
    for (row, column), value in values:
        engine.set_value("S", row, column, value)
    for (row, column), formula in formulas.items():
        engine.set_formula("S", row, column, formula)
    engine.evaluate_all()
    return engine


def test_missing_sheet_error_is_catchable_and_uncaught_reference_stays_ref():
    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet["A1"] = "=IFERROR(NoSheet!A1,456)"
    sheet["A2"] = "=NoSheet!A1"

    engine = imported(book)
    engine.evaluate_all()

    assert engine.get_value("S", 1, 1) == 456
    assert engine.get_value("S", 2, 1) == {"type": "Error", "kind": "Ref"}


def test_iferror_applies_fallback_elementwise_to_array_errors():
    # The second argument supplies [8, 9, 10] as the fallback array. Only the
    # middle element of the first array is an Excel error.
    engine = raw_workbook({(1, 1): "=IFERROR({1,#N/A,3},{8,9,10})"})

    assert [engine.get_value("S", 1, column) for column in range(1, 4)] == [1, 9, 3]


def test_power_and_exp_overflow_are_catchable_num_errors():
    assert native_formula("=IFERROR(POWER(1E308,2),456)") == 456
    assert native_formula("=IFERROR(EXP(1000),789)") == 789
    assert native_formula("=POWER(1E308,2)") == {"type": "Error", "kind": "Num"}
    assert native_formula("=EXP(1000)") == {"type": "Error", "kind": "Num"}


def test_cumulative_loan_interest_and_principal_match_an_amortization_oracle():
    rate, periods, principal = 0.01, 4, 1000
    payment = rate * principal / (1 - (1 + rate) ** -periods)
    balance = principal
    interest_total = 0
    principal_total = 0
    for _ in range(periods):
        interest = balance * rate
        paid_principal = payment - interest
        interest_total += interest
        principal_total += paid_principal
        balance -= paid_principal

    assert native_formula("=CUMIPMT(0.01,4,1000,1,4,0)") == pytest.approx(
        -interest_total
    )
    assert native_formula("=CUMPRINC(0.01,4,1000,1,4,0)") == pytest.approx(
        -principal_total
    )


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        # 2x2: chi-square 20/3, df=1, survival function erfc(sqrt(x / 2)).
        (
            "=CHISQ.TEST({10,20;20,10},{15,15;15,15})",
            math.erfc(math.sqrt(10 / 3)),
        ),
        # 2x3: chi-square 20/3, df=2, survival function exp(-x / 2).
        (
            "=CHISQ.TEST({10,20,30;20,10,30},{15,15,30;15,15,30})",
            math.exp(-10 / 3),
        ),
        # 1x3: chi-square 10, df=2, survival function exp(-x / 2).
        (
            "=CHISQ.TEST({10,20,30},{20,20,20})",
            math.exp(-5),
        ),
    ],
)
def test_chisq_test_degrees_of_freedom_match_analytic_oracle(formula, expected):
    assert native_formula(formula) == pytest.approx(expected)


def test_sumif_tilde_escaped_wildcard_matches_a_literal_star():
    engine = raw_workbook(
        {(1, 3): '=SUMIF(A1:A3,"~*",B1:B3)'},
        values=[
            ((1, 1), "*"),
            ((1, 2), 7),
            ((2, 1), "x"),
            ((2, 2), 20),
            ((3, 1), "~*"),
            ((3, 2), 30),
        ],
    )

    # Only A1 is the literal '*'; A2 and A3 are not exact matches.
    assert engine.get_value("S", 1, 3) == 7


def test_spill_reader_and_dependent_recalculate_after_shape_change():
    book = Workbook()
    sheet = book.active
    sheet.title = "S"
    sheet["B1"] = 2
    sheet["A1"] = "=SEQUENCE(B1)"
    sheet["C1"] = "=A3"
    sheet["D1"] = "=C1+10"

    engine = imported(book)
    engine.evaluate_all()
    assert engine.get_value("S", 1, 3) == 0
    assert engine.get_value("S", 1, 4) == 10

    engine.set_value("S", 1, 2, 3)
    engine.evaluate_all()
    assert engine.get_value("S", 1, 3) == 3
    assert engine.get_value("S", 1, 4) == 13
