"""Depth and display budgets remain independent in lazy graph navigation."""

from openpyxl import Workbook

from linexcel.lazy_store import (
    MAX_NEIGHBORHOOD_DEPTH,
    build_index,
    graph_neighborhood,
    graph_page,
)


def chain_index(tmp_path, count=30):
    book = Workbook()
    book.active["A1"] = 1
    for row in range(2, count + 1):
        book.active.cell(row, 1, f"=A{row - 1}+1")
    source = tmp_path / "chain.xlsx"
    book.save(source)
    index = tmp_path / "chain.sqlite"
    build_index(source, index)
    return index


def test_neighborhood_reaches_beyond_two_levels_and_caps_requested_depth(tmp_path):
    index = chain_index(tmp_path)
    third = graph_neighborhood(index, "Sheet!A1", depth=3)
    assert [node["id"] for node in third["nodes"]] == [
        f"Sheet!A{row}" for row in range(1, 5)
    ]
    deepest = graph_neighborhood(index, "Sheet!A1", depth=100)
    assert deepest["depth"] == MAX_NEIGHBORHOOD_DEPTH == 20
    assert [node["id"] for node in deepest["nodes"]] == [
        f"Sheet!A{row}" for row in range(1, 22)
    ]
    # Nodes beyond the selected depth are not display-budget truncation.
    assert not deepest["pagination"]["truncated"]


def test_neighborhood_reports_only_omissions_within_requested_depth(tmp_path):
    index = chain_index(tmp_path)
    exact = graph_neighborhood(index, "Sheet!A1", limit=3, depth=2)
    cut = graph_neighborhood(index, "Sheet!A1", limit=3, depth=20)
    assert exact["nodes"] == cut["nodes"]
    assert exact["edges"] == cut["edges"]
    assert not exact["pagination"]["truncated"]
    assert cut["pagination"]["truncated"]
    assert cut["pagination"]["returned"] == 3
    assert cut["pagination"]["limit"] == 3


def test_neighborhood_terminates_for_a_cycle_and_keeps_all_induced_edges(tmp_path):
    book = Workbook()
    book.active["A1"] = "=A2"
    book.active["A2"] = "=A1"
    source = tmp_path / "cycle.xlsx"
    book.save(source)
    index = tmp_path / "cycle.sqlite"
    build_index(source, index)
    around = graph_neighborhood(index, "Sheet!A1", depth=20)
    assert [node["id"] for node in around["nodes"]] == ["Sheet!A1", "Sheet!A2"]
    assert {(edge["source"], edge["target"]) for edge in around["edges"]} == {
        ("Sheet!A1", "Sheet!A2"),
        ("Sheet!A2", "Sheet!A1"),
    }
    assert not around["pagination"]["truncated"]


def test_deep_neighborhood_cannot_expand_the_two_hundred_node_budget(tmp_path):
    book = Workbook()
    for row in range(1, 501):
        book.active.cell(row, 1, row)
    book.active["B1"] = "=SUM(" + ",".join(f"A{row}" for row in range(1, 501)) + ")"
    source = tmp_path / "wide.xlsx"
    book.save(source)
    index = tmp_path / "wide.sqlite"
    build_index(source, index)
    around = graph_neighborhood(index, "Sheet!B1", depth=20, limit=10000)
    assert around["pagination"]["limit"] == 200
    assert around["pagination"]["returned"] == len(around["nodes"]) == 200
    assert around["pagination"]["truncated"]
    ids = {node["id"] for node in around["nodes"]}
    assert all(
        edge["source"] in ids and edge["target"] in ids for edge in around["edges"]
    )


def test_input_only_sheet_can_be_explored_beyond_the_initial_node_page(tmp_path):
    book = Workbook()
    book.active.title = "First"
    for row in range(1, 251):
        book.active.cell(row, 1, row)
    inputs = book.create_sheet("Public values")
    inputs["J9"] = 27
    book.create_sheet("Empty")
    source = tmp_path / "values.xlsx"
    book.save(source)
    index = tmp_path / "values.sqlite"
    initial = build_index(source, index)
    assert initial["meta"]["sheets"] == ["First", "Public values", "Empty"]
    assert "'Public values'!J9" not in {node["id"] for node in initial["nodes"]}
    assert graph_page(index, sheet="Public values", kind="cell")["nodes"] == []
    values = graph_page(index, sheet="Public values", kind="input", limit=1)
    assert values["pagination"]["total"] == 1
    assert values["nodes"][0]["id"] == "'Public values'!J9"
    around = graph_neighborhood(index, values["nodes"][0]["id"], depth=20)
    assert around["nodes"][0]["value"] == 27
    assert not around["edges"]
    assert not around["pagination"]["truncated"]
    assert graph_page(index, sheet="Empty")["nodes"] == []
