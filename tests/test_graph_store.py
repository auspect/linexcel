"""Whole-workbook coverage remains exact in every bounded hierarchy view."""

import os
import sqlite3
from contextlib import closing
from pathlib import Path

import pytest
from openpyxl import Workbook

from linexcel.graph_store import build_graph, generation, read_view
from linexcel.lazy_store import build_index


def test_hierarchy_conserves_topology_and_can_locate_every_cell(tmp_path):
    book = Workbook()
    book.active.title = "Data"
    for i in range(1, 1100):
        book.active.cell(i, 1, i)
        book.active.cell(i, 2, f"=A{i}+A1")
    for i in range(34):
        sheet = book.create_sheet(f"Sheet{i}")
        sheet["A1"] = "=Data!B1099"
    book.active["C1"] = "=C2"
    book.active["C2"] = "=C1"
    book.active["C3"] = "=SUM(A1:A1099)"
    source, index, artifact = [
        tmp_path / x for x in ("book.xlsx", "index.sqlite", "hierarchy.sqlite")
    ]
    book.save(source)
    build_index(source, index)
    original = index.read_bytes()
    meta = build_graph(index, artifact)
    assert index.read_bytes() == original
    with closing(sqlite3.connect(index)) as db:
        assert (
            meta["edgeCount"] == db.execute("SELECT count(*) FROM edges").fetchone()[0]
        )
        all_ids = {row[0] for row in db.execute("SELECT id FROM nodes")}
    seen, pending, visited = set(), [0], set()
    while pending:
        view_id = pending.pop()
        assert view_id not in visited
        visited.add(view_id)
        view = read_view(artifact, generation(index), view=view_id)
        assert len(view["nodes"]) <= 33
        assert len(view["edges"]) <= 33 * 32
        assert sum(n["nodeCount"] for n in view["nodes"]) == len(all_ids)
        assert (
            sum(n["internalEdges"] for n in view["nodes"])
            + sum(e["weight"] for e in view["edges"])
        ) == meta["edgeCount"]
        for node in view["nodes"]:
            if node["kind"] == "group":
                assert node["label"].count(" … ") <= 1
                pending.append(node["viewId"])
            elif node["kind"] != "outside":
                seen.add(node["nodeId"])
    assert seen == all_ids
    view = read_view(artifact, generation(index), node_id="Data!B1099")
    assert any(n.get("nodeId") == "Data!B1099" for n in view["nodes"])
    with pytest.raises(ValueError, match="stale"):
        read_view(artifact, "different index")
    with pytest.raises(KeyError):
        read_view(artifact, generation(index), node_id="missing")


def test_hierarchy_long_paths_close_handles_and_empty_workbook(tmp_path):
    root = tmp_path / ("long-directory-" * 8) / ("deep-directory-" * 8)
    if os.name == "nt":
        root = Path("\\\\?\\" + str(root))
    root.mkdir(parents=True)
    book, index, artifact = [
        root / name for name in ("book.xlsx", "index.sqlite", "hierarchy.sqlite")
    ]
    Workbook().save(book)
    build_index(book, index)
    build_graph(index, artifact)
    assert read_view(artifact, generation(index))["nodes"] == []
    # Windows refuses these replacements if any SQLite connection is left open.
    artifact.rename(root / "old.sqlite")
    build_graph(index, artifact)
    assert read_view(artifact, generation(index))["meta"]["nodeCount"] == 0
