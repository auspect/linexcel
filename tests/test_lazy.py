import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from linexcel.lazy import build_structure, evaluate_node


def save(tmp_path, setup):
    book = Workbook()
    setup(book)
    path = tmp_path / "test.xlsx"
    book.save(path)
    return path


def test_import_is_structural_and_whole_columns_stay_compact(tmp_path, monkeypatch):
    def setup(book):
        sheet = book.active
        sheet["A1"] = 2
        sheet["A1048576"] = 7
        sheet["B1"] = "=SUM(A:A)"

    path = save(tmp_path, setup)
    import builtins

    original = builtins.__import__

    def guarded(name, *args, **kwargs):
        assert name != "formualizer"
        return original(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", guarded)
    graph = build_structure(path)
    assert graph["meta"]["evaluated"] is False
    assert len(graph["nodes"]) == 4
    formula = next(n for n in graph["nodes"] if n["id"] == "Sheet!B1")
    assert formula["cached_value"] is None
    assert formula["cache_status"] == "unknown"
    assert len(formula["dependencies"]) == 1


def test_source_context_is_bounded_and_preserves_calculation_inputs(tmp_path):
    def setup(book):
        sheet = book.active
        sheet["A3"] = "Long row label " * 40
        sheet["E1"] = "Revenue"
        sheet["B3"] = 2
        sheet["C3"] = 3
        sheet["E3"] = "=B3*C3"
        sheet["E3"].number_format = '#,##0.00 "€"'
        sheet.column_dimensions.group("B", "D", hidden=True)
        sheet.row_dimensions[3].hidden = True
        sheet.sheet_state = "hidden"
        book.create_sheet("Visible")

    path = save(tmp_path, setup)
    graph = build_structure(path)
    nodes = {n["id"]: n for n in graph["nodes"]}
    cell = nodes["Sheet!E3"]
    assert cell["numberFormat"] == '#,##0.00 "€"'
    assert cell["context"]["sheetState"] == "hidden"
    assert cell["context"]["rowHidden"] is True
    assert cell["context"]["columnHidden"] is False
    assert nodes["Sheet!C3"]["context"]["columnHidden"] is True
    labels = cell["context"]["nearbyLabels"]
    assert [(x["cell"], x["position"]) for x in labels] == [
        ("A3", "left"),
        ("E1", "above"),
    ]
    assert labels[0]["truncated"] and len(labels[0]["text"]) == 240
    assert len(nodes["Sheet!A3"]["value"]) > 240
    assert "=" + "".join(t["value"] for t in cell["formulaTokens"]) == "=B3*C3"
    assert graph["meta"]["sheetDetails"][0]["formulas"] == 1
    result = evaluate_node(path, "Sheet!E3")
    assert result["value"] == 6
    step = next(s for s in result["steps"] if s["nodeId"] == "Sheet!B3")
    assert step["cachedValue"] == 2 and step["valueSource"] == "saved_input"


def test_target_closure_excludes_unrelated_invalid_formula(tmp_path):
    def setup(book):
        sheet = book.active
        sheet["A1"] = 2
        sheet["A2"] = "=A1*3"
        sheet["B1"] = "=SUM(A1:A2)"
        sheet["Z1"] = '=INDIRECT("Missing!A1")'

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "completed"
    assert result["value"] == 8
    assert result["comparison"] == "unknown"
    assert "Sheet!Z1" not in {s["node_id"] for s in result["steps"]}


def test_unresolved_transitive_dependency_refuses_calculation(tmp_path):
    def setup(book):
        sheet = book.active
        sheet["A1"] = "='[1]Other'!A1"
        sheet["B1"] = "=SUM(A:A)"

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "unsupported"
    assert result["value"] is None
    assert any("Unresolved" in d for d in result["diagnostics"])


def test_dynamic_reference_refuses_calculation(tmp_path):
    path = save(tmp_path, lambda b: setattr(b.active["A1"], "value", '=INDIRECT("B1")'))
    result = evaluate_node(path, "Sheet!A1")
    assert result["status"] == "unsupported"
    assert any("INDIRECT" in d for d in result["diagnostics"])


def test_quoted_sheet_and_blank_inputs(tmp_path):
    def setup(book):
        book.active.title = "L'été"
        book.active["A1"] = 4
        book.create_sheet("Result")["B1"] = "='L''été'!A1+A2"

    result = evaluate_node(save(tmp_path, setup), "Result!B1")
    assert result["value"] == 4


def test_absolute_names_are_visible_and_evaluated(tmp_path):
    def setup(book):
        book.active["A1"] = 3
        book.defined_names.add(DefinedName("Total", attr_text="Sheet!$A$1"))
        book.active["B1"] = "=Total*2"

    path = save(tmp_path, setup)
    graph = build_structure(path)
    assert any(n["kind"] == "name" for n in graph["nodes"])
    result = evaluate_node(path, "Sheet!B1")
    assert result["status"] == "completed"
    assert result["value"] == 6


def test_reference_inventory_is_explicit(tmp_path):
    path = save(tmp_path, lambda b: setattr(b.active["A1"], "value", 1))
    refs = tmp_path / "refs"
    refs.mkdir()
    (refs / "other.xlsx").write_bytes(b"uploaded")
    assert build_structure(path, refs)["meta"]["reference_files"] == ["other.xlsx"]


def test_relative_name_refuses_incorrect_binding(tmp_path):
    def setup(book):
        book.active["A1"] = 3
        book.defined_names.add(DefinedName("Total", attr_text="Sheet!A1"))
        book.active["B1"] = "=Total*2"

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "unsupported"
    assert any("absolute reference" in d for d in result["diagnostics"])


def test_sparse_far_range_does_not_clip_actual_inputs(tmp_path):
    def setup(book):
        book.active["A1"] = 1
        book.active["A100001"] = 9
        book.active["B1"] = "=SUM(A1:A100001)"

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "completed"
    assert result["value"] == 10


def test_array_result_cannot_be_used_as_fresh_input(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula

    def setup(book):
        book.active["A1"] = ArrayFormula(ref="A1:A2", text="=SEQUENCE(2)")
        book.active["A2"] = 2
        book.active["B1"] = "=A2*3"

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "unsupported"
    assert any("result region" in d for d in result["diagnostics"])


def test_excel_error_inputs_are_not_silently_treated_as_text(tmp_path):
    def setup(book):
        book.active["A1"] = "#N/A"
        book.active["B1"] = "=COUNTA(A1)"

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "completed"
    assert result["value"] == 1  # COUNTA counts an error; it is still not text.


def test_external_literal_input_recalculates_and_preserves_coordinates(tmp_path):
    refs = tmp_path / "refs"
    refs.mkdir()
    other = Workbook()
    other.active.title = "Data"
    other.active["C7"] = 12
    other.save(refs / "source.xlsx")

    def setup(book):
        book.active["A1"] = "='[source.xlsx]Data'!C7*2"
        book.active["A2"] = "=ROW('[source.xlsx]Data'!C7)"
        book.active["A3"] = "=COLUMN('[source.xlsx]Data'!C7)"

    path = save(tmp_path, setup)
    assert evaluate_node(path, "Sheet!A1", refs)["value"] == 24
    assert evaluate_node(path, "Sheet!A2", refs)["value"] == 7
    assert evaluate_node(path, "Sheet!A3", refs)["value"] == 3
    external = next(
        n for n in build_structure(path, refs)["nodes"] if n.get("external")
    )
    assert external["valueSource"] == "reference_input"
    assert external["sourceFile"] == "source.xlsx"


def test_external_formula_never_substitutes_saved_value(tmp_path):
    refs = tmp_path / "refs"
    refs.mkdir()
    other = Workbook()
    other.active["A1"] = "=2*3"
    other.save(refs / "source.xlsx")
    path = save(
        tmp_path, lambda b: setattr(b.active["A1"], "value", "='[source.xlsx]Sheet'!A1")
    )
    result = evaluate_node(path, "Sheet!A1", refs)
    assert result["status"] == "unsupported"
    assert any("External formula" in d for d in result["diagnostics"])


def test_metadata_sensitive_functions_fail_closed(tmp_path):
    def setup(book):
        book.active["A1"] = 1
        book.active["A2"] = 100
        book.active.row_dimensions[2].hidden = True
        book.active["B1"] = "=SUBTOTAL(109,A1:A2)"

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "unsupported"
    assert any("SUBTOTAL" in d for d in result["diagnostics"])


def test_cycle_via_range_refuses_iteration_semantics(tmp_path):
    path = save(tmp_path, lambda b: setattr(b.active["A1"], "value", "=SUM(A1:A2)"))
    result = evaluate_node(path, "Sheet!A1")
    assert result["status"] == "unsupported"
    assert any("Circular" in d for d in result["diagnostics"])


def test_sparse_blank_and_coordinate_semantics(tmp_path):
    def setup(book):
        book.active["A1"] = 1
        book.active["A3"] = 100
        book.active["B1"] = "=COUNTBLANK(A1:A5)"
        book.active["B2"] = "=COUNTA(A1:A5)"
        book.active["B3"] = "=ROW(A3)+COLUMN(C1)"

    path = save(tmp_path, setup)
    assert evaluate_node(path, "Sheet!B1")["value"] == 3
    assert evaluate_node(path, "Sheet!B2")["value"] == 2
    assert evaluate_node(path, "Sheet!B3")["value"] == 6


def test_volatile_dependencies_mark_result_uncacheable(tmp_path):
    def setup(book):
        book.active["A1"] = "=RAND()"
        book.active["B1"] = "=A1*2"

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "completed"
    assert result["volatile"] is True


def test_literal_empty_string_is_not_absent_cell(tmp_path):
    def setup(book):
        book.active["A1"] = ""
        book.active["B1"] = "=COUNTA(A1:A2)"
        book.active["B2"] = "=COUNTBLANK(A1:A2)"
        book.active["B3"] = "=ISBLANK(A1)"

    path = save(tmp_path, setup)
    assert evaluate_node(path, "Sheet!B1")["value"] == 1
    assert evaluate_node(path, "Sheet!B2")["value"] == 2
    assert evaluate_node(path, "Sheet!B3")["value"] is False


def test_external_numeric_link_uses_upload_not_recorded_disk_path(tmp_path):
    from openpyxl.packaging.relationship import Relationship
    from openpyxl.workbook.external_link.external import (
        ExternalBook,
        ExternalLink,
        ExternalSheetNames,
    )

    refs = tmp_path / "refs"
    refs.mkdir()
    other = Workbook()
    other.active["A1"] = 42
    other.save(refs / "source.xlsx")

    def setup(book):
        link = ExternalLink(
            externalBook=ExternalBook(
                sheetNames=ExternalSheetNames(sheetName=["Sheet"])
            )
        )
        link.file_link = Relationship(
            type="externalLinkPath",
            Target="file:///C:/unavailable/source.xlsx",
            TargetMode="External",
        )
        book._external_links.append(link)
        book.active["A1"] = "='[1]Sheet'!A1"

    result = evaluate_node(save(tmp_path, setup), "Sheet!A1", refs)
    assert result["status"] == "completed"
    assert result["value"] == 42


def test_formula_pattern_identity_is_exact_stable_and_propagated(tmp_path):
    def setup(book):
        book.active["A1"] = "=1+2"
        book.active["A2"] = "=1+2"
        book.active["A3"] = "=1+3"
        book.active["A4"] = '=INDIRECT("Z1")'

    path = save(tmp_path, setup)
    first = {n["id"]: n.get("patternId") for n in build_structure(path)["nodes"]}
    second = {n["id"]: n.get("patternId") for n in build_structure(path)["nodes"]}
    assert first == second
    assert first["Sheet!A1"] == first["Sheet!A2"]
    assert first["Sheet!A1"] != first["Sheet!A3"]
    assert evaluate_node(path, "Sheet!A1")["patternId"] == first["Sheet!A1"]
    unsupported = evaluate_node(path, "Sheet!A4")
    assert unsupported["status"] == "unsupported"
    assert unsupported["patternId"] == first["Sheet!A4"]


def test_recalculated_excel_error_matches_same_saved_error(tmp_path):
    import zipfile
    from xml.etree import ElementTree

    path = save(tmp_path, lambda b: setattr(b.active["A1"], "value", "=#N/A"))
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ElementTree.fromstring(parts["xl/worksheets/sheet1.xml"])
    cell = root.find("{*}sheetData/{*}row/{*}c")
    cell.set("t", "e")
    cell.find("{*}v").text = "#N/A"
    parts["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    result = evaluate_node(path, "Sheet!A1")
    assert result["status"] == "completed"
    assert result["value"] == "#N/A"
    assert result["cachedValue"] == "#N/A"
    assert result["comparison"] == "equal"


def test_engine_arrays_preserve_values_and_limitations_are_not_results():
    import pytest

    from linexcel.lazy import _engine_value

    assert _engine_value([[1, {"type": "Error", "kind": "Div"}]]) == [[1, "#DIV/0!"]]
    for kind in ("NImpl", "Circ", "Spill", "Unknown"):
        with pytest.raises(ValueError, match="could not calculate"):
            _engine_value({"type": "Error", "kind": kind})
    with pytest.raises(ValueError, match="Unsupported engine value"):
        _engine_value({"unknown": "shape"})


def test_cache_comparison_keeps_booleans_distinct_from_numbers(tmp_path):
    import zipfile
    from xml.etree import ElementTree

    def setup(book):
        book.active["A1"] = "=1"
        book.active["A2"] = "=TRUE()"
        book.active["A3"] = "=TRUE()"
        book.active["A4"] = "=1"

    path = save(tmp_path, setup)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ElementTree.fromstring(parts["xl/worksheets/sheet1.xml"])
    cells = root.findall("{*}sheetData/{*}row/{*}c")
    for cell, cache_type in zip(cells, ["b", "n", "b", "n"], strict=True):
        cell.set("t", cache_type)
        cell.find("{*}v").text = "1"
    parts["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    assert evaluate_node(path, "Sheet!A1")["comparison"] == "different"
    assert evaluate_node(path, "Sheet!A2")["comparison"] == "different"
    assert evaluate_node(path, "Sheet!A3")["comparison"] == "equal"
    assert evaluate_node(path, "Sheet!A4")["comparison"] == "equal"
    for row, comparison in [
        (1, "different"),
        (2, "different"),
        (3, "equal"),
        (4, "equal"),
    ]:
        step = evaluate_node(path, f"Sheet!A{row}")["steps"][0]
        assert step["comparison"] == comparison


@pytest.mark.parametrize("epoch_year", [1900, 1904])
def test_date_and_time_caches_compare_using_workbook_epoch(tmp_path, epoch_year):
    import datetime as dt
    import zipfile
    from xml.etree import ElementTree

    from openpyxl.utils.datetime import (
        CALENDAR_MAC_1904,
        CALENDAR_WINDOWS_1900,
        to_excel,
    )

    epoch = CALENDAR_MAC_1904 if epoch_year == 1904 else CALENDAR_WINDOWS_1900
    dates = [
        dt.datetime(2001, 11, 20),
        dt.time(13, 14, 15),
        dt.datetime(2001, 11, 20, 13, 14, 15),
    ]

    def setup(book):
        book.epoch = epoch
        for row, formula in enumerate(
            [
                "=DATE(2001,11,20)",
                "=TIME(13,14,15)",
                "=DATE(2001,11,20)+TIME(13,14,15)",
            ],
            1,
        ):
            book.active.cell(row, 1, formula).number_format = "yyyy-mm-dd hh:mm:ss"

    path = save(tmp_path, setup)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ElementTree.fromstring(parts["xl/worksheets/sheet1.xml"])
    for cell, date in zip(root.findall("{*}sheetData/{*}row/{*}c"), dates, strict=True):
        cell.find("{*}v").text = str(to_excel(date, epoch))
    parts["xl/worksheets/sheet1.xml"] = ElementTree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    for row, date in enumerate(dates, 1):
        result = evaluate_node(path, f"Sheet!A{row}")
        assert result["status"] == "completed"
        assert result["comparisonValue"] == pytest.approx(to_excel(date, epoch))
        assert result["comparison"] == "equal"
        step = next(s for s in result["steps"] if s["nodeId"] == f"Sheet!A{row}")
        assert step["comparison"] == "equal"
        assert step["comparisonValue"] == result["comparisonValue"]
    # Dates remain serial numbers at the calculation boundary. Native temporal
    # tags must not turn a date difference into a calendar date in JSON.
    assert evaluate_node(path, "Sheet!A1")["value"] == to_excel(dates[0], epoch)


def _replace_formula_caches(path, cache_text):
    import zipfile
    from xml.etree import ElementTree

    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name in parts:
        if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
            continue
        root = ElementTree.fromstring(parts[name])
        for cell in root.findall("{*}sheetData/{*}row/{*}c"):
            if cell.find("{*}f") is not None:
                cell.find("{*}v").text = cache_text
        parts[name] = ElementTree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


@pytest.mark.parametrize("saved_cache", [None, "999"])
def test_dependency_results_use_one_fresh_engine_across_sheets_and_ranges(
    tmp_path, saved_cache
):
    def setup(book):
        book.active.title = "Inputs"
        book.active["A1"] = 2
        calc = book.create_sheet("Calc")
        calc["A1"] = "=Inputs!A1*3"
        calc["B1"] = "=A1+1"
        book.defined_names.add(DefinedName("Totals", attr_text="'Calc'!$A$1:$B$1"))
        book.create_sheet("Summary")["A1"] = "=SUM(Totals)"

    path = save(tmp_path, setup)
    if saved_cache:
        _replace_formula_caches(path, saved_cache)
    result = evaluate_node(path, "Summary!A1")
    steps = {s["nodeId"]: s for s in result["steps"]}
    assert result["value"] == 13
    for node, value in {"Calc!A1": 6, "Calc!B1": 7, "Summary!A1": 13}.items():
        step = steps[node]
        assert step["evaluationStatus"] == "calculated"
        assert step["calculatedValue"] == value
        assert step["calculatedProvenance"] == "targeted_recalculation"
        assert step["cachedValue"] == (999 if saved_cache else None)
        assert step["comparison"] == ("different" if saved_cache else "unknown")
    assert steps["Inputs!A1"]["evaluationStatus"] == "input"
    assert "calculatedValue" not in steps["Inputs!A1"]
    ranges = [s for s in steps.values() if s["kind"] == "range"]
    assert ranges and {"Calc!A1", "Calc!B1"} <= set(ranges[0]["dependencies"])
    assert all(
        s["evaluationStatus"] == "structural"
        for s in steps.values()
        if s["kind"] in {"range", "name"}
    )
    assert result["coverage"]["status"] == "complete"
    assert (
        result["coverage"]["formulaCells"]
        == result["coverage"]["calculatedFormulaCells"]
        == 3
    )
    assert result["coverage"]["totalNodes"] == len(steps)


def test_if_snapshot_does_not_claim_a_branch_execution_trace(tmp_path):
    def setup(book):
        book.active["A1"] = "=1/0"
        book.active["B1"] = "=IF(TRUE(),42,A1)"

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    steps = {s["nodeId"]: s for s in result["steps"]}
    assert result["value"] == 42
    assert steps["Sheet!A1"]["calculatedValue"] == "#DIV/0!"
    assert steps["Sheet!A1"]["calculatedValueKind"] == "error"
    assert "not an execution trace" in result["step_semantics"]
    assert "unselected IF branches" in result["step_semantics"]


def test_volatile_dependencies_share_exactly_one_evaluation_request(
    tmp_path, monkeypatch
):
    import formualizer as fz

    original = fz.Workbook
    calls = []

    class ObservedWorkbook:
        def __init__(self, **kwargs):
            self.engine = original(**kwargs)

        def __getattr__(self, name):
            return getattr(self.engine, name)

        def evaluate_cell(self, *args):
            calls.append(args)
            return self.engine.evaluate_cell(*args)

    monkeypatch.setattr(fz, "Workbook", ObservedWorkbook)

    def setup(book):
        book.active["A1"] = "=RAND()"
        book.active["B1"] = "=A1*2"
        book.active["C1"] = "=A1+B1"

    result = evaluate_node(save(tmp_path, setup), "Sheet!C1")
    steps = {s["nodeId"]: s for s in result["steps"]}
    random_value = steps["Sheet!A1"]["calculatedValue"]
    assert 0 <= random_value <= 1
    assert steps["Sheet!B1"]["calculatedValue"] == 2 * random_value
    assert result["value"] == random_value + 2 * random_value
    assert steps["Sheet!C1"]["calculatedValue"] == result["value"]
    assert calls == [("Sheet", 1, 3)]


@pytest.mark.parametrize(
    "missing_value,status",
    [
        (None, "unavailable"),
        ({"type": "Error", "kind": "NImpl"}, "unsupported"),
        (ValueError("No readable cell"), "unavailable"),
        (RuntimeError("Cannot read cell"), "unavailable"),
    ],
)
def test_partial_engine_snapshot_keeps_target_and_never_substitutes_cache(
    tmp_path, monkeypatch, missing_value, status
):
    import formualizer as fz

    original = fz.Workbook

    class PartialWorkbook:
        def __init__(self, **kwargs):
            self.engine = original(**kwargs)

        def __getattr__(self, name):
            return getattr(self.engine, name)

        def get_value(self, *args):
            if isinstance(missing_value, Exception):
                raise missing_value
            return missing_value

    monkeypatch.setattr(fz, "Workbook", PartialWorkbook)

    def setup(book):
        book.active["A1"] = "=2+3"
        book.active["B1"] = "=A1+1"

    path = save(tmp_path, setup)
    _replace_formula_caches(path, "999")
    result = evaluate_node(path, "Sheet!B1")
    steps = {s["nodeId"]: s for s in result["steps"]}
    assert result["status"] == "completed" and result["value"] == 6
    assert steps["Sheet!A1"]["evaluationStatus"] == status
    assert "calculatedValue" not in steps["Sheet!A1"]
    assert steps["Sheet!A1"]["cachedValue"] == 999
    assert steps["Sheet!B1"]["calculatedValue"] == 6
    assert result["coverage"]["status"] == "partial"
    assert result["coverage"]["calculatedFormulaCells"] == 1
    assert result["coverage"][status + "FormulaCells"] == 1


def test_refused_closure_keeps_structure_without_partially_calculating(tmp_path):
    def setup(book):
        book.active["A1"] = '=INDIRECT("C1")'
        book.active["B1"] = "=A1+1"
        book.active["C1"] = 2

    result = evaluate_node(save(tmp_path, setup), "Sheet!B1")
    assert result["status"] == "unsupported"
    assert result["coverage"]["status"] == "not_evaluated"
    assert result["coverage"]["calculatedFormulaCells"] == 0
    assert result["coverage"]["notEvaluatedFormulaCells"] == 2
    assert all(
        s["evaluationStatus"] == "not_evaluated"
        for s in result["steps"]
        if s["kind"] == "cell"
    )
    assert all("calculatedValue" not in s for s in result["steps"])
