"""3D aggregates expand reference arguments in source sheet order."""

import pytest
from openpyxl import Workbook

from linexcel.lazy import build_structure, evaluate_node


@pytest.mark.parametrize(
    "function,expected",
    [
        ("SUM", 36),
        ("AVERAGE", 6),
        ("COUNT", 6),
        ("COUNTA", 6),
        ("MIN", 1),
        ("MAX", 10),
        ("PRODUCT", 6000),
    ],
)
def test_three_dimensional_aggregate_preserves_range_semantics(
    tmp_path, function, expected
):
    book = Workbook()
    book.active.title = "Jan"
    book.create_sheet("Middle month")
    book.create_sheet("Mar")
    book.create_sheet("Output")
    for position, sheet in enumerate(book.worksheets[:3], 1):
        sheet["A1"] = position
        sheet["A2"] = "=5*2"
    book["Output"]["A1"] = f"={function}(Jan:Mar!A1:A2)"
    book["Output"]["A2"] = "=A1*2"
    path = tmp_path / "3d.xlsx"
    book.save(path)
    result = evaluate_node(path, "Output!A2")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == expected * 2
    assert result["coverage"]["formulaCells"] == 5
    span = next(n for n in build_structure(path)["nodes"] if n["kind"] == "sheet_span")
    assert span["sheets"] == ["Jan", "Middle month", "Mar"]
    assert span["cell_count"] == 6


def test_three_dimensional_quoted_and_reversed_span(tmp_path):
    book = Workbook()
    book.active.title = "First month"
    book.create_sheet("Middle")
    book.create_sheet("Last month")
    book.create_sheet("Output")
    for position, sheet in enumerate(book.worksheets[:3], 1):
        sheet["A1"] = position
    book["Output"]["A1"] = "=SUM('Last month:First month'!A1,4)"
    path = tmp_path / "3d.xlsx"
    book.save(path)
    assert evaluate_node(path, "Output!A1")["value"] == 10


@pytest.mark.parametrize(
    "formula",
    [
        "=SUM(First:Last!A1+2)",
        "=ABS(First:Last!A1)",
        "=SUM(First:Last!A1)+ABS(First:Last!A1)",
        "=SUM(First:Missing!A1)",
    ],
)
def test_three_dimensional_unsupported_context_is_never_silently_rewritten(
    tmp_path, formula
):
    book = Workbook()
    book.active.title = "First"
    book.active["A1"] = 1
    book.create_sheet("Last")["A1"] = 2
    book.create_sheet("Output")["A1"] = formula
    path = tmp_path / "3d.xlsx"
    book.save(path)
    result = evaluate_node(path, "Output!A1")
    assert result["status"] == "unsupported"
    assert result["diagnostics"]


def test_three_dimensional_text_booleans_and_errors_keep_reference_types(tmp_path):
    book = Workbook()
    book.active.title = "First"
    book.active["A1"] = "10"
    book.active["A2"] = True
    book.create_sheet("Last")["A1"] = 4
    book["Last"]["A2"] = "#N/A"
    output = book.create_sheet("Output")
    output["A1"] = "=SUM(First:Last!A1)"
    output["A2"] = "=SUM(First:Last!A1:A2)"
    output["A3"] = "=COUNTA(First:Last!A1:A2)"
    path = tmp_path / "3d.xlsx"
    book.save(path)
    assert evaluate_node(path, "Output!A1")["value"] == 4
    assert evaluate_node(path, "Output!A2")["value"] == "#N/A"
    assert evaluate_node(path, "Output!A3")["value"] == 4
