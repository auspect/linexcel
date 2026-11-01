"""Source array boundaries and fresh spill values remain part of exact closure."""

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

from linexcel.lazy import build_structure, evaluate_node
from linexcel.structured_refs import formula_tokens


@pytest.mark.parametrize(
    ("formula", "target", "expected"),
    [
        ("=A2*10", "Sheet!C1", 20),
        ("=SUM(A2:A3)", "Sheet!C1", 5),
        ("=SUM(A1#)", "Sheet!C1", 6),
        ("=INDEX(A1#,3,1)", "Sheet!C1", 3),
        ("=A2", "Sheet!A2", 2),
    ],
)
def test_declared_array_has_fresh_children_and_exact_source(
    tmp_path, formula, target, expected
):
    book = Workbook()
    book.active["A1"] = ArrayFormula(ref="A1:A3", text="=SEQUENCE(3,1)")
    book.active["A2"] = 999
    book.active["C1"] = formula
    path = tmp_path / "spill.xlsx"
    book.save(path)
    graph = build_structure(path)
    nodes = {n["id"]: n for n in graph["nodes"]}
    assert nodes["Sheet!A2"]["kind"] == "array_result"
    assert nodes["Sheet!A2"]["dependencies"] == ["Sheet!A1"]
    assert nodes["Sheet!A1"]["sourceFormula"] == "=SEQUENCE(3,1)"
    result = evaluate_node(path, target)
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == expected
    assert result["arrayEvidence"][0]["observed"]["extent"] == "Sheet!A1:A3"


def test_range_covering_only_unsaved_spill_children_includes_anchor(tmp_path):
    book = Workbook()
    book.active["A1"] = ArrayFormula(ref="A1:A3", text="=SEQUENCE(3)")
    book.active["C1"] = "=SUM(A2:A3)"
    path = tmp_path / "empty-children.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!C1")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == 5
    assert "Sheet!A1" in [s["nodeId"] for s in result["steps"]]


def test_extent_change_is_explicit_without_reusing_source_cache(tmp_path):
    book = Workbook()
    book.active["A1"] = ArrayFormula(ref="A1:A3", text="=SEQUENCE(2)")
    book.active["A3"] = 999
    book.active["C1"] = "=A3"
    path = tmp_path / "changed-extent.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!C1")
    assert result["status"] == "unsupported"
    assert result["value"] is None
    assert "extent" in " ".join(result["diagnostics"])


def test_matrix_formula_recalculates_source_inputs(tmp_path):
    book = Workbook()
    book.active["A1"], book.active["A2"] = 2, 3
    book.active["B1"], book.active["B2"] = 4, 5
    book.active["C1"] = ArrayFormula(ref="C1:C2", text="=A1:A2*B1:B2")
    book.active["C2"] = 999
    book.active["D1"] = "=SUM(C1:C2)"
    path = tmp_path / "matrix.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!D1")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == 23


def test_spill_tokenization_keeps_errors_strings_and_exact_formula():
    formula = '=IFERROR(SUM(A1#),"# and A2#")+#N/A'
    assert "=" + "".join(t.value for t in formula_tokens(formula)) == formula
