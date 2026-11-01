"""One replayable operation context and honest iterative stopping criteria."""

import pytest
from openpyxl import Workbook

from linexcel.lazy import evaluate_node

CONTEXT = {"timestampUtc": "2026-10-30T12:30:00+00:00", "timezone": "utc", "seed": 42}


def test_volatile_calculation_replays_clock_seed_and_dependency_snapshot(tmp_path):
    book = Workbook()
    book.active["A1"] = "=RAND()"
    book.active["A2"] = "=NOW()"
    book.active["A3"] = "=A1+A2"
    path = tmp_path / "volatile.xlsx"
    book.save(path)
    first = evaluate_node(path, "Sheet!A3", calculation_context=CONTEXT)
    second = evaluate_node(path, "Sheet!A3", calculation_context=CONTEXT)
    assert first["value"] == second["value"]
    assert first["evaluationContext"] == CONTEXT
    assert first["volatile"]
    steps = {s["nodeId"]: s for s in first["steps"]}
    assert (
        first["value"]
        == steps["Sheet!A1"]["calculatedValue"] + steps["Sheet!A2"]["calculatedValue"]
    )
    assert steps["Sheet!A2"]["calculatedValue"] == pytest.approx(46325 + 12.5 / 24)


def test_iteration_uses_source_limits_and_records_convergence(tmp_path):
    book = Workbook()
    book.calculation.iterate = True
    book.calculation.iterateCount = 200
    book.calculation.iterateDelta = 1e-9
    book.active["A1"] = "=(A1+2)/2"
    book.active["B1"] = "=A1*3"
    path = tmp_path / "cycle.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!B1")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == pytest.approx(6, abs=1e-8)
    assert result["iteration"]["convergedComponents"] == 1
    assert result["iteration"]["cappedComponents"] == 0
    assert result["iteration"]["maxPasses"] <= 200


def test_capped_iterations_never_become_successful_values(tmp_path):
    book = Workbook()
    book.calculation.iterate = True
    book.calculation.iterateCount = 5
    book.calculation.iterateDelta = 0
    book.active["A1"] = "=1-A1"
    book.active["B1"] = "=A1+10"
    path = tmp_path / "cycle.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!B1")
    assert result["status"] == "unsupported"
    assert result["value"] is None
    assert result["iteration"]["cappedComponents"] == 1
    assert "convergence" in " ".join(result["diagnostics"])


@pytest.mark.parametrize(
    "context",
    [
        {**CONTEXT, "seed": -1},
        {**CONTEXT, "seed": True},
        {**CONTEXT, "timestampUtc": "2026-10-30"},
        {**CONTEXT, "timezone": "local"},
    ],
)
def test_invalid_replay_context_is_rejected(tmp_path, context):
    book = Workbook()
    book.active["A1"] = "=RAND()"
    path = tmp_path / "volatile.xlsx"
    book.save(path)
    with pytest.raises(ValueError):
        evaluate_node(path, "Sheet!A1", calculation_context=context)
