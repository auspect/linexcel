"""Source-typed oracles for the isolated interactive calculation path."""

import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from linexcel.lazy import build_structure, evaluate_node


def save(tmp_path, book):
    path = tmp_path / "semantics.xlsx"
    book.save(path)
    return path


def test_literal_names_scopes_quotes_signs_and_case_insensitive_sheets(tmp_path):
    book = Workbook()
    book.active.title = "Data"
    other = book.create_sheet("O'Brien")
    for name, definition in {
        "CompanyName": '"A ""quoted"" company"',
        "Rate": "-0.5",
        "Flag": "TRUE",
        "Problem": "#N/A",
    }.items():
        book.defined_names.add(DefinedName(name, attr_text=definition))
    other.defined_names.add(DefinedName("Rate", attr_text="2"))
    book.active["A1"] = 7
    formulas = {
        "A2": ("=CompanyName", 'A "quoted" company'),
        "A3": ("=Rate^2", 0.25),
        "A4": ("=IF(Flag,3,4)", 3),
        "A5": ('=IFERROR(Problem,"caught")', "caught"),
        "A6": ("='o''brien'!rAtE+Rate", 1.5),
        "A7": ("=data!A1+1", 8),
    }
    for address, (formula, _expected) in formulas.items():
        book.active[address] = formula
    other["A1"] = "=Rate+1"
    path = save(tmp_path, book)
    for address, (_formula, expected) in formulas.items():
        result = evaluate_node(path, f"Data!{address}")
        assert result["status"] == "completed", result["diagnostics"]
        assert result["value"] == expected
    assert evaluate_node(path, "'O''Brien'!A1")["value"] == 3
    name = next(
        n for n in build_structure(path)["nodes"] if n["id"] == "name:*:CompanyName"
    )
    assert name["value"] == 'A "quoted" company'
    assert name["valueSource"] == "defined_constant"


def test_error_input_propagates_while_error_spelling_stays_text(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "#N/A"
    sheet["A2"] = "#N/A"
    sheet["A2"].data_type = "s"
    expected = {
        "B1": ('=IFERROR("prefix "&A1,"caught")', "caught", "scalar"),
        "B2": ('=IFERROR("prefix "&A2,"caught")', "prefix #N/A", "scalar"),
        "B3": ("=SUM(A1:A2)", "#N/A", "error"),
        "B4": ("=COUNTA(A1:A2)", 2, "scalar"),
    }
    for address, (formula, _value, _kind) in expected.items():
        sheet[address] = formula
    path = save(tmp_path, book)
    for address, (_formula, value, kind) in expected.items():
        result = evaluate_node(path, f"Sheet!{address}")
        assert result["status"] == "completed"
        assert result["value"] == value
        assert result["valueKind"] == kind


def test_external_error_input_is_typed_and_not_a_saved_formula(tmp_path):
    other = Workbook()
    other.active.title = "Data"
    other.active["A1"] = "#DIV/0!"
    refs = tmp_path / "refs"
    refs.mkdir()
    other.save(refs / "source.xlsx")
    book = Workbook()
    book.active["A1"] = "=IFERROR('[source.xlsx]Data'!A1,\"caught\")"
    result = evaluate_node(save(tmp_path, book), "Sheet!A1", refs)
    assert result["status"] == "completed"
    assert result["value"] == "caught"


def test_sumproduct_correction_feeds_downstream_dependencies(tmp_path):
    book = Workbook()
    sheet = book.active
    for row, value in enumerate([2, "3", True, None], 1):
        sheet.cell(row, 1, value)
        sheet.cell(row, 2, 10)
    sheet["C1"] = "=SUMPRODUCT(A1:A4,B1:B4)"
    sheet["C2"] = "=SUMPRODUCT(A1:A2*B1:B2)"
    sheet["D1"] = "=C1+C2"
    result = evaluate_node(save(tmp_path, book), "Sheet!D1")
    assert result["value"] == 70
    steps = {s["nodeId"]: s for s in result["steps"]}
    assert steps["Sheet!C1"]["calculatedValue"] == 20
    assert steps["Sheet!C2"]["calculatedValue"] == 50
    assert result["coverage"]["status"] == "complete"


@pytest.mark.parametrize(
    "source_value,kind", [("", "scalar"), ("#N/A", "error"), (3, "scalar")]
)
def test_literal_read_never_imports_native_engine(
    tmp_path, monkeypatch, source_value, kind
):
    import builtins

    book = Workbook()
    book.active["A1"] = source_value
    path = save(tmp_path, book)
    original = builtins.__import__

    def guarded(name, *args, **kwargs):
        assert not name.startswith("formualizer")
        return original(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", guarded)
    result = evaluate_node(path, "Sheet!A1")
    assert result["value"] == source_value
    assert result["valueKind"] == kind
    assert result["provenance"] == "input_read"


def test_error_cache_is_not_equal_to_a_text_result_with_same_spelling(tmp_path):
    book = Workbook()
    book.active["A1"] = '="#N/A"'
    path = save(tmp_path, book)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    cell = root.find("{*}sheetData/{*}row/{*}c")
    cell.set("t", "e")
    cell.find("{*}v").text = "#N/A"
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    result = evaluate_node(path, "Sheet!A1")
    assert result["value"] == "#N/A"
    assert result["valueKind"] == "scalar"
    assert result["comparison"] == "different"


def test_legacy_normal_aliases_use_the_same_supported_functions(tmp_path):
    book = Workbook()
    book.active["A1"] = "=NORMSDIST(0)"
    book.active["A2"] = "=NORMDIST(0,0,1,TRUE)"
    path = save(tmp_path, book)
    assert evaluate_node(path, "Sheet!A1")["value"] == 0.5
    assert evaluate_node(path, "Sheet!A2")["value"] == 0.5


def test_overdeep_target_is_refused_before_native_parsing(tmp_path):
    book = Workbook()
    book.active["A1"] = "=" + "+".join(["1"] * 1500)
    result = evaluate_node(save(tmp_path, book), "Sheet!A1")
    assert result["status"] == "unsupported"
    assert any("complexity" in issue for issue in result["diagnostics"])


def test_stored_date_serials_keep_excel_leap_day_and_submillisecond_precision(tmp_path):
    book = Workbook()
    for row, serial in enumerate([59, 60, 61, 60.5, 46000.123456789], 1):
        book.active.cell(row, 1, serial).number_format = "yyyy-mm-dd hh:mm:ss.000"
        book.active.cell(row, 2, f"=DAY(A{row})")
        book.active.cell(row, 3, f"=A{row}")
    path = save(tmp_path, book)
    graph = {node["id"]: node for node in build_structure(path)["nodes"]}
    for row, serial in enumerate([59, 60, 61, 60.5, 46000.123456789], 1):
        assert graph[f"Sheet!A{row}"]["cached_comparison_value"] == serial
        assert evaluate_node(path, f"Sheet!C{row}")["value"] == serial
    assert graph["Sheet!A2"]["cachedValue"] == 60
    for row, day in enumerate([28, 29, 1, 29], 1):
        assert evaluate_node(path, f"Sheet!B{row}")["value"] == day
