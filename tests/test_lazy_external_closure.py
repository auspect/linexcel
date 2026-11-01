"""Supplied-file calculations retain independent identities and exact closure."""

import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from linexcel.lazy import evaluate_node


def main_book(tmp_path, formula):
    book = Workbook()
    book.active["A1"] = formula
    book.active["A2"] = "=A1*2"
    path = tmp_path / "main.xlsx"
    book.save(path)
    refs = tmp_path / "refs"
    refs.mkdir()
    return path, refs


def test_external_range_streams_literals_and_calculates_formula_members(tmp_path):
    path, refs = main_book(tmp_path, "=SUM('[ref.xlsx]Data'!A1:A1001)")
    book = Workbook()
    book.active.title = "Data"
    for row in range(1, 1001):
        book.active.cell(row, 1, row)
    book.active["A1001"] = "=A1+A2"
    book.save(refs / "ref.xlsx")
    result = evaluate_node(path, "Sheet!A2", refs)
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == (500500 + 3) * 2
    assert result["coverage"]["inputCells"] == 1000
    assert result["coverage"]["formulaCells"] == 3
    assert result["omittedSteps"] > 0


def test_external_formulas_use_their_source_scope_and_tables(tmp_path):
    path, refs = main_book(tmp_path, "='[ref.xlsx]Result'!A1")
    book = Workbook()
    book.active.title = "Data"
    book.active.append(["Amount"])
    book.active.append([10])
    book.active.append([20])
    book.active.add_table(Table(displayName="SalesTable", ref="A1:A3"))
    result_sheet = book.create_sheet("Result")
    book.defined_names.add(DefinedName("Rate", attr_text="2"))
    result_sheet.defined_names.add(DefinedName("Rate", attr_text="3"))
    result_sheet["A1"] = "=SUM(SalesTable[Amount])*Rate"
    book.save(refs / "ref.xlsx")
    result = evaluate_node(path, "Sheet!A2", refs)
    assert result["value"] == 180, result
    formula_step = next(
        s for s in result["steps"] if s["nodeId"] == "external:ref.xlsx:Result!A1"
    )
    assert formula_step["sourceFormula"] == "=SUM(SalesTable[Amount])*Rate"
    assert formula_step["formulaProvenance"]["sheet"] == "Result"


def test_external_transitive_files_and_same_named_sheets_stay_distinct(tmp_path):
    path, refs = main_book(tmp_path, "='[one.xlsx]Sheet'!A1+'[two.xlsx]Sheet'!A1")
    first = Workbook()
    first.active["A1"] = "='[two.xlsx]Sheet'!A2+10"
    first.save(refs / "one.xlsx")
    second = Workbook()
    second.active["A1"] = 3
    second.active["A2"] = 7
    second.save(refs / "two.xlsx")
    result = evaluate_node(path, "Sheet!A2", refs)
    assert result["value"] == 40, result


def test_external_bad_formula_cache_is_observation_not_input(tmp_path):
    path, refs = main_book(tmp_path, "='[ref.xlsx]Sheet'!A1")
    book = Workbook()
    book.active["A1"] = "=6*7"
    source = refs / "ref.xlsx"
    book.save(source)
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    xml = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
    xml.find("{*}sheetData/{*}row/{*}c/{*}v").text = "999"
    parts["xl/worksheets/sheet1.xml"] = ET.tostring(xml)
    with zipfile.ZipFile(source, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    result = evaluate_node(path, "Sheet!A1", refs)
    assert result["value"] == 42
    step = next(s for s in result["steps"] if s["nodeId"].startswith("external:"))
    assert step["cachedValue"] == 999 and step["calculatedValue"] == 42
    assert step["comparison"] == "different"


def test_external_cycle_refuses_instead_of_using_cache(tmp_path):
    path, refs = main_book(tmp_path, "='[one.xlsx]Sheet'!A1")
    first = Workbook()
    first.active["A1"] = "='[two.xlsx]Sheet'!A1+1"
    first.save(refs / "one.xlsx")
    second = Workbook()
    second.active["A1"] = "='[one.xlsx]Sheet'!A1+1"
    second.save(refs / "two.xlsx")
    result = evaluate_node(path, "Sheet!A1", refs)
    assert result["status"] == "unsupported"
    assert any("Circular" in diagnostic for diagnostic in result["diagnostics"])


def test_external_computed_name_reads_its_own_absolute_reference(tmp_path):
    path, refs = main_book(tmp_path, "='[ref.xlsx]Sheet'!A1")
    book = Workbook()
    book.active["B1"] = 7
    book.defined_names.add(DefinedName("Amount", attr_text="Sheet!$B$1*3"))
    book.active["A1"] = "=Amount+1"
    book.save(refs / "ref.xlsx")
    result = evaluate_node(path, "Sheet!A1", refs)
    assert result["value"] == 22, result


@pytest.mark.parametrize(
    "formula",
    [
        "='[missing.xlsx]Sheet'!A1",
        "='[../outside.xlsx]Sheet'!A1",
        "='[ref.xlsx]Missing'!A1",
    ],
)
def test_external_missing_or_unsupplied_path_is_visible_and_refused(tmp_path, formula):
    path, refs = main_book(tmp_path, formula)
    Workbook().save(refs / "ref.xlsx")
    result = evaluate_node(path, "Sheet!A1", refs)
    assert result["status"] == "unsupported"
    assert result["diagnostics"]
