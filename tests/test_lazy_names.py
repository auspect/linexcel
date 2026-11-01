"""Static names retain scoped dependency evidence when expanded for calculation."""

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from linexcel.lazy import evaluate_node


def test_nested_computed_names_preserve_scope_and_precedence(tmp_path):
    book = Workbook()
    book.active.title = "Inputs"
    book.active["A1"] = 4
    calc = book.create_sheet("Calc")
    book.defined_names.add(DefinedName("Rate", attr_text="2"))
    calc.defined_names.add(DefinedName("Rate", attr_text="3"))
    book.defined_names.add(DefinedName("Base", attr_text="Inputs!$A$1+1"))
    calc.defined_names.add(DefinedName("Amount", attr_text="Base*Rate"))
    calc["A1"] = "=Amount^2"
    path = tmp_path / "names.xlsx"
    book.save(path)
    result = evaluate_node(path, "Calc!A1")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == 225
    assert {"name:Calc:Amount", "name:Calc:Rate", "name:*:Base", "Inputs!A1"} <= {
        step["nodeId"] for step in result["steps"]
    }


def test_computed_name_with_absolute_range_includes_formula_members(tmp_path):
    book = Workbook()
    book.active["A1"] = 10
    book.active["A2"] = "=A1*2"
    book.defined_names.add(DefinedName("Total", attr_text="SUM(Sheet!$A$1:$A$2)"))
    book.active["B1"] = "=Total/3"
    path = tmp_path / "names.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!B1")
    assert result["value"] == 10, result
    assert result["coverage"]["calculatedFormulaCells"] == 2


@pytest.mark.parametrize("definition", ["Sheet!A1+2", 'INDIRECT("Sheet!A1")', "Loop+1"])
def test_ambiguous_relative_dynamic_and_cyclic_names_refuse(tmp_path, definition):
    book = Workbook()
    book.active["A1"] = 10
    book.defined_names.add(DefinedName("Loop", attr_text=definition))
    book.active["B1"] = "=Loop"
    path = tmp_path / "names.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!B1")
    assert result["status"] == "unsupported"
    assert result["diagnostics"]
