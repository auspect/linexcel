"""Copied templates are metadata, never an import-time evaluation request."""

import builtins

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from linexcel.lazy import build_structure, evaluate_node


def graph_for(tmp_path, formulas, setup=None):
    book = Workbook()
    for address, formula in formulas.items():
        book.active[address] = formula
    if setup:
        setup(book)
    path = tmp_path / "patterns.xlsx"
    book.save(path)
    graph = build_structure(path)
    return path, graph, {node["id"]: node for node in graph["nodes"]}


def test_vertical_horizontal_and_mixed_anchor_copies_keep_individual_nodes(tmp_path):
    formulas = {
        "D2": "=A2+$B2+C$1+$B$1",
        "D3": "=A3+$B3+C$1+$B$1",
        "E2": "=B2+$B2+D$1+$B$1",
        "E3": "=B3+$B3+D$1+$B$1",
        "D5": "=A5+$B5+C$1+$B$1",
        "F2": "=C2+$C2+E$1+$B$1",  # Absolute column changed.
        "D6": "=A5+$B6+C$1+$B$1",  # Wrong relative offset.
    }
    _, graph, nodes = graph_for(tmp_path, formulas)
    copied = [nodes[f"Sheet!{cell}"] for cell in ("D2", "D3", "E2", "E3", "D5")]
    assert len({node["patternId"] for node in copied}) == 1
    assert nodes["Sheet!F2"]["patternId"] != copied[0]["patternId"]
    assert nodes["Sheet!D6"]["patternId"] != copied[0]["patternId"]
    group = next(
        p for p in graph["formulaPatterns"] if p["id"] == copied[0]["patternId"]
    )
    assert group["memberCount"] == 5
    assert group["ranges"] == ["D2:E3", "D5"]  # Do not fill the gap at row 4.
    assert all(node["formula"] == formulas[node["cell"]] for node in copied)
    assert all(node["patternMemberCount"] == 5 for node in copied)
    assert graph["meta"]["copiedPatternCount"] == 1
    assert graph["meta"]["copiedFormulaCells"] == 5


def test_ranges_whole_axes_and_quoted_cross_sheet_references(tmp_path):
    _, graph, nodes = graph_for(
        tmp_path,
        {
            "D2": "=SUM('O''Brien Data'!A2:B3,$A2:B$7,A:A,1:1,$C:$C,$4:$4)",
            "E3": "=SUM('O''Brien Data'!B3:C4,$A3:C$7,B:B,2:2,$C:$C,$4:$4)",
            "D3": "=SUM('Other Data'!A3:B4,$A3:B$7,A:A,2:2,$C:$C,$4:$4)",
        },
        lambda book: (
            book.create_sheet("O'Brien Data"),
            book.create_sheet("Other Data"),
        ),
    )
    assert nodes["Sheet!D2"]["patternId"] == nodes["Sheet!E3"]["patternId"]
    assert nodes["Sheet!D2"]["patternId"] != nodes["Sheet!D3"]["patternId"]
    group = next(p for p in graph["formulaPatterns"] if p["memberCount"] == 2)
    assert "'O''Brien Data'!" in group["signature"]


def test_defined_names_and_strings_are_not_mistaken_for_references(tmp_path):
    def setup(book):
        book.defined_names.add(DefinedName("FOO", attr_text="Sheet!$Z$1"))
        book.defined_names.add(DefinedName("RC", attr_text="Sheet!$Z$2"))

    _, graph, nodes = graph_for(
        tmp_path,
        {
            "C2": '=FOO+A2+"A2 FOO"',
            "C3": '=foo+A3+"A2 FOO"',
            "C4": '=FOO+A4+"A3 FOO"',
            "C5": '=FOO+A5+"A2 foo"',
            "D1": "=RC",
            "D2": "=D2",
        },
        setup,
    )
    assert nodes["Sheet!C2"]["patternId"] == nodes["Sheet!C3"]["patternId"]
    assert len({nodes[f"Sheet!C{i}"]["patternId"] for i in (2, 4, 5)}) == 3
    # Pretty signatures may both read =RC, but typed identities cannot collide.
    assert nodes["Sheet!D1"]["patternId"] != nodes["Sheet!D2"]["patternId"]
    group = next(p for p in graph["formulaPatterns"] if p["memberCount"] == 2)
    assert '"A2 FOO"' in group["signature"]


def test_groups_are_sheet_local_and_respect_local_names(tmp_path):
    def setup(book):
        other = book.create_sheet("Other")
        other["B1"] = "=FOO+A1"
        book.active.defined_names.add(DefinedName("FOO", attr_text="Sheet!$Z$1"))
        other.defined_names.add(DefinedName("FOO", attr_text="Other!$Z$2"))

    _, graph, nodes = graph_for(tmp_path, {"B1": "=FOO+A1"}, setup)
    assert nodes["Sheet!B1"]["patternId"] != nodes["Other!B1"]["patternId"]
    assert graph["meta"]["copiedPatternCount"] == 0


def test_unsupported_or_malformed_references_are_isolated(tmp_path):
    _, graph, _ = graph_for(
        tmp_path,
        {
            "B1": "=A1:A",
            "B2": "=A2:A",
            "C1": "=A0",
            "C2": "=A0",
            "D1": "=SUM(A1:XFE2)",
            "D2": "=SUM(A2:XFE3)",
            "E1": '=INDIRECT("A1")',
            "E2": '=INDIRECT("A1")',
            "F1": "=SUM(Sheet:Other!A1)",
            "F2": "=SUM(Sheet:Other!A2)",
            "G1": "=Table1[Value]",
            "G2": "=Table1[Value]",
        },
        lambda book: book.create_sheet("Other"),
    )
    assert all(
        p["status"] == "isolated" and p["memberCount"] == 1
        for p in graph["formulaPatterns"]
    )
    assert graph["meta"]["copiedFormulaCells"] == 0


def test_many_copies_import_without_engine_and_store_members_once(
    tmp_path, monkeypatch
):
    original = builtins.__import__

    def guarded(name, *args, **kwargs):
        assert name != "formualizer" and not name.startswith("linexcel.rewrite")
        return original(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", guarded)
    path, graph, nodes = graph_for(
        tmp_path, {f"B{i}": f"=A{i}*2" for i in range(1, 2001)}
    )
    assert graph["meta"]["evaluated"] is False
    assert graph["meta"]["formulaPatternCount"] == 1
    group = graph["formulaPatterns"][0]
    assert group["memberCount"] == 2000
    assert group["ranges"] == ["B1:B2000"]
    assert all("memberIds" not in node for node in nodes.values())
    assert build_structure(path)["formulaPatterns"] == graph["formulaPatterns"]


def test_copied_group_does_not_replace_individual_target_calculation(tmp_path):
    path, graph, nodes = graph_for(
        tmp_path, {"A1": 2, "A2": 5, "B1": "=A1*3", "B2": "=A2*3"}
    )
    assert graph["formulaPatterns"][0]["memberCount"] == 2
    first = evaluate_node(path, "Sheet!B1")
    second = evaluate_node(path, "Sheet!B2")
    assert first["patternId"] == second["patternId"] == nodes["Sheet!B1"]["patternId"]
    assert first["value"] == 6 and second["value"] == 15
    assert "Sheet!B2" not in {s["nodeId"] for s in first["steps"]}
