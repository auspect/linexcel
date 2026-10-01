"""Regression coverage for sparse disk indexing and exact bounded presentation."""

import json
import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook

from linexcel.lazy import evaluate_node
from linexcel.lazy_store import (
    WorkbookIndex,
    build_index,
    evidence_graph,
    get_node,
    graph_neighborhood,
    graph_page,
    patterns_page,
)


def workbook(tmp_path, setup):
    book = Workbook()
    setup(book)
    path = tmp_path / "source.xlsx"
    book.save(path)
    return path


def rewrite(path, change):
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    change(parts)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def test_disk_index_pages_find_off_page_nodes_without_reopening_workbook(
    tmp_path, monkeypatch
):
    def setup(book):
        for row in range(1, 451):
            book.active.cell(row, 1, row)
            book.active.cell(row, 2, f"=A{row}*2")

    path = workbook(tmp_path, setup)
    monkeypatch.setattr(
        "openpyxl.load_workbook",
        lambda *_args, **_kwargs: pytest.fail("Lazy indexing must not open openpyxl"),
    )
    index = tmp_path / "index.sqlite"
    first = build_index(path, index)
    assert len(first["nodes"]) == 200
    assert first["meta"]["nodeCount"] == 900
    assert first["meta"]["formulaCount"] == 450
    assert first["meta"]["sheetNodeCounts"] == {"Sheet": 900}
    assert first["pagination"]["hasMore"]
    assert graph_page(index, offset=890)["pagination"]["hasMore"] is False
    assert graph_page(index, limit=10000)["pagination"]["limit"] == 200
    found = graph_page(index, query="B450")
    assert [n["id"] for n in found["nodes"]] == ["Sheet!B450"]
    assert get_node(index, "Sheet!B450")["formula"] == "=A450*2"
    assert patterns_page(index)["formulaPatterns"][0]["memberCount"] == 450
    neighborhood = graph_neighborhood(index, "Sheet!B450")
    assert {n["id"] for n in neighborhood["nodes"]} == {"Sheet!A450", "Sheet!B450"}
    assert neighborhood["formulaPatterns"][0]["membersTruncated"]
    path.unlink()
    result = evaluate_node(path, "Sheet!B450", index_path=index)
    assert result["status"] == "completed"
    assert result["value"] == 900
    assert len(evidence_graph(index)["nodes"]) <= 18
    assert evidence_graph(index, "Sheet!B450")["formulaPatterns"]


def test_large_overlapping_ranges_stream_all_inputs_and_bound_only_steps(tmp_path):
    book = Workbook(write_only=True)
    sheet = book.create_sheet()
    for row in range(1, 40001):
        sheet.append([1, "=SUM(A1:A30000)+SUM(A15000:A40000)" if row == 1 else None])
    path = tmp_path / "large.xlsx"
    book.save(path)
    index = tmp_path / "index.sqlite"
    build_index(path, index)
    result = evaluate_node(path, "Sheet!B1", index_path=index)
    assert result["status"] == "completed"
    assert result["value"] == 55001
    assert result["coverage"]["inputCells"] == 40000
    assert result["coverage"]["totalNodes"] == 40003
    assert result["coverage"]["closureComplete"]
    assert len(result["steps"]) == 200
    assert result["stepTotal"] == 40003
    assert result["omittedSteps"] == 39803
    assert all(len(step["dependencies"]) <= 200 for step in result["steps"])
    assert any(step.get("dependenciesOmitted", 0) > 20000 for step in result["steps"])
    assert "Sheet!B1" in {step["nodeId"] for step in result["steps"]}


@pytest.mark.parametrize("damage", ["value", "style"])
def test_bad_cell_does_not_prevent_later_cells_or_independent_calculations(
    tmp_path, damage
):
    def setup(book):
        book.active["A1"] = 2
        book.active["B1"] = "=SUM(A1:A2)"
        book.active["C1"] = "=40+2"
        book.active["A2"] = 7

    path = workbook(tmp_path, setup)

    def modify(parts):
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        cell = root.find("{*}sheetData/{*}row/{*}c")
        if damage == "value":
            cell.find("{*}v").text = "not-a-number"
        else:
            cell.set("s", "invalid-style")
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)

    rewrite(path, modify)
    index = tmp_path / "index.sqlite"
    build_index(path, index)
    assert get_node(index, "Sheet!A1")["diagnostics"]
    assert get_node(index, "Sheet!A2")["value"] == 7
    assert evaluate_node(path, "Sheet!B1", index_path=index)["status"] == "unsupported"
    assert evaluate_node(path, "Sheet!C1", index_path=index)["value"] == 42


def test_corrupt_sheet_keeps_other_sheets_and_refuses_missing_range_cells(tmp_path):
    def setup(book):
        book.active["A1"] = 2
        book.active["A2"] = 3
        other = book.create_sheet("Other")
        other["A1"] = "=1+2"
        other["B1"] = "=SUM(Sheet!A1:A5)"

    path = workbook(tmp_path, setup)

    def modify(parts):
        xml = parts["xl/worksheets/sheet1.xml"]
        parts["xl/worksheets/sheet1.xml"] = xml[: xml.index(b'<row r="2"')] + b"<broken"

    rewrite(path, modify)
    index = tmp_path / "index.sqlite"
    result = build_index(path, index)
    assert result["meta"]["status"] == "partial"
    assert result["meta"]["sheetDetails"][0]["status"] == "partial"
    assert get_node(index, "Sheet!A1")["value"] == 2
    assert evaluate_node(path, "Other!A1", index_path=index)["value"] == 3
    assert evaluate_node(path, "Other!B1", index_path=index)["status"] == "unsupported"


def test_shared_formula_translation_and_implicit_cell_coordinates(tmp_path):
    def setup(book):
        book.active["A1"] = 2
        book.active["B1"] = "=A1*3"
        book.active["A2"] = 5
        book.active["B2"] = "=A2*3"

    path = workbook(tmp_path, setup)

    def modify(parts):
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        cells = root.findall("{*}sheetData/{*}row/{*}c")
        first = cells[1].find("{*}f")
        first.set("t", "shared")
        first.set("si", "1")
        first.set("ref", "B1:B2")
        second = cells[3].find("{*}f")
        second.set("t", "shared")
        second.set("si", "1")
        second.text = None
        for cell in cells:
            cell.attrib.pop("r")
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)

    rewrite(path, modify)
    index = tmp_path / "index.sqlite"
    build_index(path, index)
    assert get_node(index, "Sheet!B2")["formula"] == "=A2*3"
    assert evaluate_node(path, "Sheet!B2", index_path=index)["value"] == 15


def test_broken_shared_string_table_retains_numeric_work_and_resolved_prefix(tmp_path):
    def setup(book):
        book.active["A1"] = 1
        book.active["A2"] = 2
        book.active["B1"] = "=40+2"

    path = workbook(tmp_path, setup)

    def modify(parts):
        parts["xl/sharedStrings.xml"] = (
            b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b"<si><t>Valid label</t></si><broken"
        )
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        for cell in root.findall("{*}sheetData/{*}row/{*}c"):
            if cell.get("r") in {"A1", "A2"}:
                cell.set("t", "s")
                cell.find("{*}v").text = "0" if cell.get("r") == "A1" else "99"
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)

    rewrite(path, modify)
    index = tmp_path / "index.sqlite"
    result = build_index(path, index)
    assert result["meta"]["status"] == "partial"
    assert get_node(index, "Sheet!A1")["value"] == "Valid label"
    assert get_node(index, "Sheet!A2")["diagnostics"]
    assert evaluate_node(path, "Sheet!B1", index_path=index)["value"] == 42


def test_failed_replacement_preserves_published_index(tmp_path):
    path = workbook(tmp_path, lambda book: setattr(book.active["A1"], "value", 42))
    index = tmp_path / "index.sqlite"
    first = build_index(path, index)
    path.write_bytes(b"not a zip")
    with pytest.raises(zipfile.BadZipFile):
        build_index(path, index)
    with WorkbookIndex(index) as store:
        assert store["Sheet!A1"]["value"] == 42
        assert store.meta["sha256"] == first["meta"]["sha256"]
    assert not list(tmp_path.glob("*.building"))
    assert len(json.dumps(graph_page(index))) < 10000


@pytest.mark.parametrize("storage", ["s", "inlineStr", "str"])
def test_ooxml_escaped_text_rich_runs_and_formula_caches_are_exact(tmp_path, storage):
    encoded = "  _x005F_x0041_ CR_x000D_LF\n_x0009__xD83D__xDE00_ &  "
    expected = "  _x0041_ CR\rLF\n\t😀 &  "
    path = workbook(tmp_path, lambda book: setattr(book.active["A1"], "value", "text"))

    def modify(parts):
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        cell = root.find("{*}sheetData/{*}row/{*}c")
        cell.clear()
        cell.set("r", "A1")
        cell.set("t", storage)
        if storage == "str":
            ET.SubElement(cell, "f").text = '"constant"'
            ET.SubElement(cell, "v").text = encoded
        else:
            container = (
                ET.Element("si") if storage == "s" else ET.SubElement(cell, "is")
            )
            # Decoding individual rich-text runs would corrupt an escape split
            # across their boundary. Phonetic annotations are not value text.
            for text in (encoded[:7], encoded[7:]):
                ET.SubElement(ET.SubElement(container, "r"), "t").text = text
            ET.SubElement(ET.SubElement(container, "rPh"), "t").text = "phonetic guide"
            if storage == "s":
                ET.SubElement(cell, "v").text = "0"
                table = ET.Element("sst")
                table.append(container)
                parts["xl/sharedStrings.xml"] = ET.tostring(table)
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)

    rewrite(path, modify)
    index = tmp_path / "index.sqlite"
    build_index(path, index)
    node = get_node(index, "Sheet!A1")
    assert node["value"] == expected
    assert node["cachedValue"] == expected
    assert not node["diagnostics"]
    assert node["cachedValueKind"] == "scalar"


@pytest.mark.parametrize("cell_type,raw", [("b", "2"), ("str", "_xD800_")])
def test_invalid_boolean_and_unpaired_surrogate_fail_only_their_cells(
    tmp_path, cell_type, raw
):
    def setup(book):
        book.active["A1"] = 1
        book.active["B1"] = "=40+2"

    path = workbook(tmp_path, setup)

    def modify(parts):
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        cell = root.find("{*}sheetData/{*}row/{*}c")
        cell.set("t", cell_type)
        cell.find("{*}v").text = raw
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)

    rewrite(path, modify)
    index = tmp_path / "index.sqlite"
    build_index(path, index)
    assert get_node(index, "Sheet!A1")["diagnostics"]
    assert evaluate_node(path, "Sheet!A1", index_path=index)["status"] == "unsupported"
    assert evaluate_node(path, "Sheet!B1", index_path=index)["value"] == 42


def test_implicit_coordinate_after_formatting_only_cell_is_preserved(tmp_path):
    def setup(book):
        book.active["A1"].number_format = "0.00"
        book.active["B1"] = 42

    path = workbook(tmp_path, setup)

    def modify(parts):
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        cells = root.findall("{*}sheetData/{*}row/{*}c")
        for cell in cells:
            cell.attrib.pop("r")
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(root)

    rewrite(path, modify)
    index = tmp_path / "index.sqlite"
    build_index(path, index)
    assert get_node(index, "Sheet!A1") is None
    assert get_node(index, "Sheet!B1")["value"] == 42
