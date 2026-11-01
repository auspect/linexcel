"""Visibility adapters operate within the fresh sparse target closure."""

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

from linexcel.lazy import build_structure, evaluate_node


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ("=SUBTOTAL(9,A1:A3)", 9),
        ("=SUBTOTAL(109,A1:A3)", 6),
        ("=AGGREGATE(9,1,A1:A3)", 6),
        ("=AGGREGATE(9,4,A1:A3)", 9),
        ("=AGGREGATE(14,7,A1:A3,1)", 4),
    ],
)
def test_hidden_rows_recalculate_as_source_metadata(tmp_path, formula, expected):
    book = Workbook()
    book.active["A1"], book.active["A2"], book.active["A3"] = 2, 3, 4
    book.active.row_dimensions[2].hidden = True
    book.active["B1"] = formula
    book.active["C1"] = "=B1*2"
    path = tmp_path / "visibility.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!C1")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == expected * 2
    assert result["visibilityEvidence"][0]["nodeId"] == "Sheet!B1"


def test_filter_and_manual_hidden_cause_ambiguity_is_not_guessed(tmp_path):
    book = Workbook()
    book.active.append(["Amount"])
    book.active.append([2])
    book.active.append([3])
    book.active.auto_filter.ref = "A1:A3"
    book.active.auto_filter.add_filter_column(0, ["2"])
    book.active.row_dimensions[3].hidden = True
    book.active["B1"] = "=SUBTOTAL(9,A2:A3)"
    book.active["C1"] = "=SUBTOTAL(109,A2:A3)"
    path = tmp_path / "filter.xlsx"
    book.save(path)
    graph = build_structure(path)
    assert graph["meta"]["sheetDetails"][0]["activeFilters"]
    refused = evaluate_node(path, "Sheet!B1")
    assert refused["status"] == "unsupported"
    assert refused["capabilityDiagnostics"][0]["feature"] == "visibility_aggregates"
    result = evaluate_node(path, "Sheet!C1")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == 2


def test_hidden_unsaved_spill_child_is_excluded(tmp_path):
    book = Workbook()
    book.active["A1"] = ArrayFormula(ref="A1:A3", text="=SEQUENCE(3)")
    book.active.row_dimensions[2].hidden = True
    book.active["C1"] = "=SUBTOTAL(109,A1:A3)"
    path = tmp_path / "hidden-spill.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!C1")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == 4


def test_repeated_visibility_functions_share_registration_safely(tmp_path):
    book = Workbook()
    book.active["A1"], book.active["A2"], book.active["A3"] = 2, 3, 4
    book.active.row_dimensions[2].hidden = True
    book.active["B1"] = "=SUBTOTAL(109,A1:A3)"
    book.active["B2"] = "=SUBTOTAL(109,A1:A3)"
    book.active["C1"] = "=B1+B2"
    path = tmp_path / "repeated-subtotal.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!C1")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == 12
    assert len(result["visibilityEvidence"]) == 2
    steps = {step["nodeId"]: step for step in result["steps"]}
    assert steps["Sheet!B1"]["calculatedValue"] == 6
    assert steps["Sheet!B2"]["calculatedValue"] == 6


@pytest.mark.parametrize(
    "formula",
    [
        "=AGGREGATE(9,6,A1:A2)",
        "=SUBTOTAL(2,A1:A2)",
        "=SUBTOTAL(3,A1:A2)",
        "=COUNT(A1:A2)",
        "=IFERROR(A1,5)",
    ],
)
def test_engine_nonresult_is_not_an_ignorable_excel_error(tmp_path, formula):
    book = Workbook()
    book.active["A1"] = "=AGGREGATE(12,5,Z1:Z2)"
    book.active["A2"] = 5
    book.active["B1"] = formula
    path = tmp_path / "nonresult.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!B1")
    assert result["status"] == "unsupported"
    assert result["value"] is None


def test_inline_fallback_cannot_mask_an_uncomputed_aggregate(tmp_path):
    book = Workbook()
    book.active["A1"] = "=IFERROR(AGGREGATE(12,5,Z1:Z2),99)"
    path = tmp_path / "inline-nonresult.xlsx"
    book.save(path)
    result = evaluate_node(path, "Sheet!A1")
    assert result["status"] == "unsupported"
    assert result["value"] is None
