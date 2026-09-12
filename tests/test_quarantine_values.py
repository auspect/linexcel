"""Quarantined formulas cannot masquerade as blanks or recalculated values."""

import io
import re
import zipfile

import pytest
from openpyxl import Workbook

from linexcel import analyze


def deep_workbook(cached=None, guard="=IFERROR(B1,42)", root=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 3
    ws["B1"] = root or "=" + "+".join(["A1"] * 1000)
    ws["C1"] = "=B1+5"
    ws["D1"] = "=SUM(B1:C1)"
    ws["E1"] = guard
    ws["F1"] = "=2+3"
    output = io.BytesIO()
    wb.save(output)
    if cached is None:
        return output.getvalue()
    patched = io.BytesIO()
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(patched, "w") as dest:
        for part in source.infolist():
            payload = source.read(part.filename)
            if part.filename == "xl/worksheets/sheet1.xml":
                payload = re.sub(
                    rb"</f><v(?:\s*/>|></v>)", b"</f><v>3000</v>", payload, count=1
                )
            dest.writestr(part, payload)
    return patched.getvalue()


@pytest.mark.parametrize("targets", [None, ["S!D1", "S!E1"]])
@pytest.mark.parametrize(
    "guard",
    [
        "=IFERROR(B1,42)",
        "=ISERROR(B1)",
        '=IFERROR(INDIRECT("B1"),42)',
        "=IFERROR(OFFSET(A1,0,1),42)",
    ],
)
def test_missing_quarantined_value_propagates_as_unknown(targets, guard):
    result = analyze(deep_workbook(guard=guard), targets=targets)
    by_addr = {node.get("addr"): node for node in result.nodes}
    for address in ("B1", "C1", "D1"):
        assert by_addr[address]["value"] is None
        assert by_addr[address].get("valueSource") != "engine"
    assert by_addr["E1"]["value"] is None
    assert by_addr["E1"]["steps"] is None
    if targets is None:
        assert by_addr["F1"]["value"] == 5
        assert by_addr["F1"]["valueSource"] == "engine"


@pytest.mark.parametrize("targets", [None, ["S!C1"]])
def test_quarantined_cache_is_attributed_to_the_file(targets):
    result = analyze(deep_workbook(cached=3000), targets=targets)
    by_addr = {node.get("addr"): node for node in result.nodes}
    assert by_addr["B1"]["value"] == 3000
    assert by_addr["B1"]["valueSource"] == "file"
    assert by_addr["C1"]["value"] == 3005


def test_unresolved_external_formula_does_not_claim_its_cache_as_recalculated():
    result = analyze(deep_workbook(cached=3000, root="='[missing.xlsx]Sheet1'!A1"))
    node = next(node for node in result.nodes if node.get("addr") == "B1")
    assert node["value"] == 3000
    assert node["valueSource"] == "file"


def test_incomplete_dependent_trace_keeps_formula_readings_conservative():
    from linexcel.engine import _uncached_dependents

    class MissingTrace:
        def get_value(self, *cell):
            return None

        def trace(self, *args, **kwargs):
            raise ValueError("trace unavailable")

    warnings = []
    unavailable = _uncached_dependents(
        MissingTrace(), deep_workbook(), {("S", 1, 2): "=deep"}, warnings
    )
    assert unavailable == {("S", 1, col) for col in range(2, 7)}
    assert "could not be fully traced" in warnings[0]


def test_failed_evaluation_rebuild_does_not_restore_an_unsafe_formula(monkeypatch):
    import linexcel.engine as engine_module

    real_open = engine_module._open_workbook

    class FailedEvaluation:
        def __init__(self, data, parallel):
            self.engine = real_open(data, parallel)

        def __getattr__(self, name):
            return getattr(self.engine, name)

        def evaluate_all(self):
            raise ValueError("evaluation failed")

    monkeypatch.setattr(engine_module, "_open_workbook", FailedEvaluation)
    monkeypatch.setattr(engine_module, "_formulas_gone", lambda *args: True)
    session = engine_module.boot_engine(deep_workbook(), [])
    assert ("S", 1, 2) in session.quarantined
    assert not session.engine.get_formula("S", 1, 2)
    assert session.engine.get_value("S", 1, 2) == {"type": "Error", "kind": "NImpl"}
