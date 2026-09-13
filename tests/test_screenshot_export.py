"""Preserve screenshot identities across rendering, vision and HTML embedding."""

import io
import json
import re
import struct
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from linexcel import analyze
from linexcel.insights import render_workbook_screenshots


def chart_workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Category", "Value"])
    sheet.append(["A", 5])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2))
    book.create_chartsheet("Chart").add_chart(chart)
    book.create_sheet("After")["A1"] = "After chart"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_chartsheet_pages_keep_workbook_order(monkeypatch, tmp_path):
    monkeypatch.setattr("linexcel.insights.find_libreoffice", lambda: "office")
    monkeypatch.setattr("linexcel.insights.find_pdftoppm", lambda: "poppler")
    # Distinct dimensions let the assertion detect a shifted page assignment.
    dimensions = [(100, 50), (800, 600), (150, 80)]

    def render(command, **kwargs):
        if command[0] == "office":
            folder = Path(command[command.index("--outdir") + 1])
            (folder / "workbook.pdf").write_bytes(b"pdf")
        else:
            assert command[command.index("-r") + 1] == "200"
            for idx, (width, height) in enumerate(dimensions, start=1):
                Path(f"{command[-1]}-{idx}.png").write_bytes(
                    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
                    + struct.pack(">II", width, height)
                )

    monkeypatch.setattr("linexcel.insights.subprocess.run", render)
    shots = render_workbook_screenshots(
        chart_workbook(), "chart.xlsx", tmp_path, dpi=200
    )
    assert isinstance(shots, dict)
    assert list(shots) == ["Data", "Chart", "After"]
    assert [struct.unpack(">II", p[0].read_bytes()[16:24]) for p in shots.values()] == (
        dimensions
    )


def test_page_names_survive_embedding_with_partial_out_of_order_docs(tmp_path):
    paths = [tmp_path / "finance-01.png", tmp_path / "finance-02.png"]
    for path in paths:
        path.write_bytes(b"image bytes")
    result = analyze(chart_workbook(), filename="chart.xlsx")
    html = result.to_html(
        screenshots=paths,
        screenshot_docs={"finance-02": "Second page only."},
    )
    start = re.search(r"var\s+GRAPH\s*=\s*", html).end()
    graph = json.JSONDecoder().raw_decode(html, start)[0]
    assert graph["meta"]["screenshotNames"] == ["finance-01", "finance-02"]
    assert all(s.startswith("data:image/png;") for s in graph["meta"]["screenshots"])
    assert graph["meta"]["screenshotDocs"] == {"finance-02": "Second page only."}


def test_embedded_images_do_not_become_giant_page_names():
    result = analyze(chart_workbook(), filename="chart.xlsx")
    image = "data:image/png;base64,aW1hZ2U="
    html = result.to_html(screenshots=[image])
    start = re.search(r"var\s+GRAPH\s*=\s*", html).end()
    graph = json.JSONDecoder().raw_decode(html, start)[0]
    assert graph["meta"]["screenshotNames"] == [None]
