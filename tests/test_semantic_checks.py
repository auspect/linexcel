"""Capability mismatches annotate risk; they do not rewrite workbook results."""

import copy
import io

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.workbook.defined_name import DefinedName

from linexcel.doc_evidence import read_defined_name_evidence
from linexcel.engine import _open_workbook
from linexcel.semantic_checks import annotate_semantic_risks, engine_capabilities


def test_real_probe_matrix_includes_both_epochs_and_known_comparison_gap():
    report = engine_capabilities()
    assert len(report["checks"]) == 40
    assert {check["epoch"] for check in report["checks"]} == {"1900", "1904"}
    assert all(
        "expected" in check and "observed" in check for check in report["checks"]
    )
    # This is diagnostic, so a future fixed engine should naturally stop
    # reporting a mismatch instead of being required to retain its bug.
    for check in report["checks"]:
        if check["id"].startswith("number_text_eq"):
            assert check["expected"] is False
            assert check["status"] == (
                "passed" if check["observed"] is False else "known_mismatch"
            )
        if check["id"] == "date_serial_subtraction:1900":
            assert check["formula"] == "=DATE(2026,2,1)-46054"
            assert check["expected"] == 0
            if check["observedType"] == "date":
                assert check["observed"]["type"] == "date"
                assert check["status"] == "known_mismatch"


def test_risks_propagate_through_guards_but_preserve_results_and_sources():
    graph = {
        "meta": {},
        "nodes": [
            {"id": "a", "formula": '=10<" "', "value": False, "valueSource": "engine"},
            {
                "id": "b",
                "formula": "=IFERROR(A1,5)",
                "value": False,
                "valueSource": "engine",
            },
            {"id": "c", "formula": "=2+3", "value": 5, "valueSource": "engine"},
        ],
        "edges": [{"source": "a", "target": "b", "kind": "dep"}],
    }
    original = copy.deepcopy(graph)
    annotate_semantic_risks(
        graph,
        capabilities={
            "status": "known_mismatch",
            "checks": [
                {"feature": "mixed_type_comparison", "status": "known_mismatch"},
            ],
        },
    )
    assert graph["nodes"][0]["verification"] == "unverified_engine_semantics"
    assert graph["nodes"][1]["semanticRisks"][0]["origin"] == "dependency"
    assert "verification" not in graph["nodes"][2]
    for before, after in zip(original["nodes"], graph["nodes"]):
        assert (before["value"], before["valueSource"]) == (
            after["value"],
            after["valueSource"],
        )


def test_passed_probes_do_not_certify_a_graph_or_invent_risks():
    graph = {"meta": {}, "nodes": [{"id": "a", "formula": "=A1>2"}], "edges": []}
    annotate_semantic_risks(
        graph, capabilities={"status": "probes_passed", "checks": []}
    )
    assert "verification" not in graph["nodes"][0]
    assert not graph["meta"].get("warnings")


def test_date_risk_requires_a_failed_date_probe_and_date_provenance():
    graph = {
        "nodes": [
            {"id": "date", "valueDate": "2024-01-01"},
            {"id": "calc", "formula": "=A1+1"},
            {"id": "numeric", "formula": "=1+1"},
        ],
        "edges": [{"source": "date", "target": "calc", "kind": "dep"}],
    }
    annotate_semantic_risks(
        graph,
        capabilities={
            "checks": [
                {"feature": "typed_date_arithmetic", "status": "known_mismatch"},
            ]
        },
    )
    assert graph["nodes"][1]["semanticRisks"][0]["feature"] == "typed_date_arithmetic"
    assert "verification" not in graph["nodes"][2]


def test_mixed_type_mismatch_does_not_flag_proven_same_type_literals():
    graph = {
        "nodes": [
            {"id": "numeric", "formula": "=IF(2<3,4,5)"},
            {"id": "text", "formula": '=IF("a"="b",4,5)'},
            {"id": "unknown", "formula": "=IF(A1<3,4,5)"},
        ]
    }
    annotate_semantic_risks(
        graph,
        capabilities={
            "checks": [
                {"feature": "mixed_type_comparison", "status": "known_mismatch"},
            ]
        },
    )
    assert "verification" not in graph["nodes"][0]
    assert "verification" not in graph["nodes"][1]
    assert graph["nodes"][2]["verification"] == "unverified_engine_semantics"


def test_long_unsafe_formula_is_not_parsed_by_risk_annotation():
    graph = {"nodes": [{"id": "deep", "formula": "=" + "+".join(["1"] * 3000)}]}
    annotate_semantic_risks(
        graph,
        capabilities={
            "checks": [
                {"feature": "mixed_type_comparison", "status": "known_mismatch"},
            ]
        },
    )
    assert "verification" not in graph["nodes"][0]


def test_date_serial_arithmetic_is_flagged_without_date_input_nodes():
    graph = {"nodes": [{"id": "a", "formula": "=DATE(2026,2,1)-46054"}]}
    annotate_semantic_risks(
        graph,
        capabilities={
            "checks": [
                {"feature": "typed_date_arithmetic", "status": "known_mismatch"},
            ]
        },
    )
    assert graph["nodes"][0]["semanticRisks"][0]["feature"] == "typed_date_arithmetic"


def test_each_risk_keeps_its_own_formula_or_dependency_origin():
    graph = {
        "nodes": [
            {"id": "date", "formula": "=DATE(2026,2,1)-46054"},
            {"id": "comparison", "formula": '=A1<" "'},
        ],
        "edges": [{"source": "date", "target": "comparison", "kind": "dep"}],
    }
    annotate_semantic_risks(
        graph,
        capabilities={
            "checks": [
                {"feature": "typed_date_arithmetic", "status": "known_mismatch"},
                {"feature": "mixed_type_comparison", "status": "known_mismatch"},
            ]
        },
    )
    reasons = {
        risk["feature"]: risk["origin"] for risk in graph["nodes"][1]["semanticRisks"]
    }
    assert reasons == {
        "mixed_type_comparison": "formula",
        "typed_date_arithmetic": "dependency",
    }


# Versioned synthetic observations: LibreOffice 26.2.4.2, formualizer 0.9.3,
# 2026-09-19. Public inputs contain no corpus data. These are independent
# observations, not an assertion that LibreOffice is an Excel oracle in all
# circumstances. Error kinds remain explicit, never compared as arbitrary text.
ORACLE_EXTENSION_VERSION = 1
ORACLE_EXTENSION_CASES = (
    ("blank_numeric", "=A1=0", True, "confirmed"),
    ("blank_empty_string", '=A1=""', True, "confirmed"),
    ("blank_isblank", "=ISBLANK(A1)", True, "confirmed"),
    ("empty_result_isblank", "=ISBLANK(A2)", False, "confirmed"),
    ("empty_result_numeric", "=A2=0", False, "confirmed"),
    ("true_numeric", "=TRUE=1", False, "unresolved_expectation"),
    ("false_numeric", "=FALSE=0", False, "unresolved_expectation"),
    ("boolean_order", "=TRUE>1", True, "unresolved_expectation"),
    ("guard_division", "=IFERROR(1/0,7)", 7, "confirmed"),
    ("detect_division", "=ISERROR(1/0)", True, "confirmed"),
    ("division_error", "=1/0", {"type": "Error", "kind": "Div"}, "confirmed"),
    ("lazy_branch", "=IF(FALSE,1/0,7)", 7, "confirmed"),
    ("indirect_other_sheet", '=INDIRECT("Data!A1")+1', 11, "confirmed"),
    ("offset_other_sheet", "=OFFSET(Data!A1,1,0)", 20, "confirmed"),
    ("global_constant", "=GlobalRate*10", 20, "known_name_limitation"),
    ("local_shadows_global", "=ScopedRate*10", 30, "known_name_limitation"),
    ("global_reference", "=InputAmount+1", 11, "confirmed"),
    (
        "global_text_constant",
        "=CompanyName",
        "Synthetic Company",
        "known_name_limitation",
    ),
    (
        "undefined_name",
        "=NotDefined+1",
        {"type": "Error", "kind": "Name"},
        "confirmed",
    ),
    (
        "let_local_binding",
        "=_xlfn.LET(total,Data!A1+Data!A2,total/2)",
        15,
        "confirmed",
    ),
    (
        "indirect_missing_guard",
        '=IFERROR(INDIRECT("Missing!A1"),7)',
        7,
        "oracle_disagreement",
    ),
)
# All other v1 native/LO observations matched the stated expectation. These
# exceptions retain the three-way distinction instead of silently changing
# provisional expectations to match whichever engine was used as an oracle.
ORACLE_EXTENSION_DISAGREEMENTS = {
    "true_numeric": {"native": True, "libreoffice": True},
    "false_numeric": {"native": True, "libreoffice": True},
    "boolean_order": {"native": False, "libreoffice": False},
    "global_constant": {
        "native": {"type": "Error", "kind": "Name"},
        "libreoffice": 20,
    },
    "local_shadows_global": {
        "native": {"type": "Error", "kind": "Name"},
        "libreoffice": 30,
    },
    "global_text_constant": {
        "native": {"type": "Error", "kind": "Name"},
        "libreoffice": "Synthetic Company",
    },
    "indirect_missing_guard": {
        "native": 7,
        "libreoffice": {"type": "Error", "kind": "Ref"},
    },
}


def _name_workbook(epoch="1900"):
    workbook = Workbook()
    if epoch == "1904":
        workbook.epoch = CALENDAR_MAC_1904
    sheet = workbook.active
    sheet.title = "Probe"
    sheet["A2"] = '=""'
    data = workbook.create_sheet("Data")
    data["A1"], data["A2"] = 10, 20
    workbook.defined_names.add(DefinedName("GlobalRate", attr_text="2"))
    workbook.defined_names.add(
        DefinedName("CompanyName", attr_text='"Synthetic Company"')
    )
    workbook.defined_names.add(DefinedName("ScopedRate", attr_text="9"))
    sheet.defined_names.add(DefinedName("ScopedRate", attr_text="3"))
    workbook.defined_names.add(DefinedName("InputAmount", attr_text="Data!$A$1"))
    return workbook, sheet


@pytest.mark.parametrize("epoch", ["1900", "1904"])
def test_versioned_oracle_extension_preserves_confirmed_and_unresolved_cases(epoch):
    workbook, sheet = _name_workbook(epoch)
    for row, (_key, formula, _expected, _status) in enumerate(
        ORACLE_EXTENSION_CASES, 1
    ):
        sheet.cell(row, 3, formula)
    output = io.BytesIO()
    workbook.save(output)
    engine = _open_workbook(output.getvalue(), parallel=False)
    engine.evaluate_all()
    for row, (key, _formula, expected, status) in enumerate(ORACLE_EXTENSION_CASES, 1):
        assert (key in ORACLE_EXTENSION_DISAGREEMENTS) == (status != "confirmed")
        value = engine.get_value("Probe", row, 3)
        if status == "unresolved_expectation":
            # LO and fz both returned the opposite boolean to the provisional
            # expectation. Do not enforce either as Excel conformity.
            assert isinstance(value, bool), key
        elif status == "known_name_limitation":
            # A future engine may resolve the limitation. Never require a bug
            # to persist, but reject unrelated failures or wrong numeric values.
            assert value == expected or value == {"type": "Error", "kind": "Name"}
        else:
            assert value == expected, key
            if isinstance(expected, bool):
                assert isinstance(value, bool), key
            elif isinstance(expected, (int, float)):
                assert not isinstance(value, bool), key


def _constant_name_capabilities():
    return {
        "checks": [{"feature": "defined_name_constant", "status": "known_mismatch"}]
    }


def test_name_risk_uses_source_definitions_and_propagates_without_changing_values():
    workbook, _sheet = _name_workbook()
    output = io.BytesIO()
    workbook.save(output)
    graph = {
        "meta": {"definedNameEvidence": read_defined_name_evidence(output.getvalue())},
        "nodes": [
            {
                "id": "constant",
                "sheet": "Probe",
                "formula": "=CompanyName",
                "value": "#NAME?",
                "valueSource": "engine",
                "cachedValue": "Synthetic Company",
                "cachedAgreement": "differ",
                "steps": [{"value": "#NAME?", "evaluated": True}],
            },
            {
                "id": "guard",
                "sheet": "Probe",
                "formula": '=IFERROR(C1,"fallback")',
                "value": "fallback",
                "valueSource": "engine",
            },
            {"id": "reference", "sheet": "Probe", "formula": "=InputAmount+1"},
            {"id": "missing", "sheet": "Probe", "formula": "=NotDefined+1"},
        ],
        "edges": [{"source": "constant", "target": "guard", "kind": "dep"}],
    }
    before = copy.deepcopy(graph)
    annotate_semantic_risks(graph, capabilities=_constant_name_capabilities())
    assert graph["nodes"][0]["semanticRisks"][0]["origin"] == "formula"
    assert graph["nodes"][1]["semanticRisks"][0]["origin"] == "dependency"
    assert "verification" not in graph["nodes"][2]
    assert "verification" not in graph["nodes"][3]
    for original, annotated in zip(before["nodes"], graph["nodes"]):
        assert {key: annotated[key] for key in original} == original


def test_name_risk_respects_sheet_shadowing_and_lexical_bindings():
    workbook, sheet = _name_workbook()
    # A local reference must shadow the global constant of the same name.
    sheet.defined_names.add(DefinedName("GlobalRate", attr_text="Data!$A$1"))
    output = io.BytesIO()
    workbook.save(output)
    formulas = [
        ("Probe", "=GlobalRate", False),
        ("Data", "=GlobalRate", True),
        ("Probe", "=ScopedRate", True),
        ("Data", "='Probe'!GlobalRate", False),
        ("Probe", "=LET(CompanyName,7,CompanyName)", False),
        ("Probe", "=LAMBDA(CompanyName,CompanyName)", False),
        ("Probe", "=LET(x,CompanyName,x)", True),
        ("Probe", "=LET(CompanyName,CompanyName,CompanyName)", True),
        ("Probe", '=LET(CompanyName,"local",CompanyName)&CompanyName', True),
    ]
    graph = {
        "meta": {"definedNameEvidence": read_defined_name_evidence(output.getvalue())},
        "nodes": [
            {"id": str(i), "sheet": sheet_name, "formula": formula}
            for i, (sheet_name, formula, _) in enumerate(formulas)
        ],
    }
    annotate_semantic_risks(graph, capabilities=_constant_name_capabilities())
    for node, (_, formula, expected) in zip(graph["nodes"], formulas):
        assert ("verification" in node) == expected, formula


def test_partial_name_metadata_does_not_assert_unseen_scope_or_absence():
    graph = {
        "meta": {
            "definedNameEvidence": {
                "status": "partial",
                "definitions": [
                    {"name": "CompanyName", "scope": "workbook", "expression": '"x"'}
                ],
            }
        },
        "nodes": [{"id": "name", "sheet": "Probe", "formula": "=CompanyName"}],
    }
    annotate_semantic_risks(graph, capabilities=_constant_name_capabilities())
    assert "verification" not in graph["nodes"][0]


def test_runtime_name_probes_report_raw_errors_or_future_supported_results():
    checks = [
        check
        for check in engine_capabilities()["checks"]
        if check["feature"] == "defined_name_constant"
    ]
    assert len(checks) == 6
    for check in checks:
        if check["observedType"] == "dict":
            assert check["observed"] == {"type": "Error", "kind": "Name"}
            assert check["status"] == "known_mismatch"
        else:
            assert check["observed"] == check["expected"]
            assert check["status"] == "passed"


@pytest.mark.parametrize("targeted", [False, True])
def test_pipeline_exposes_constant_name_risk_for_full_and_targeted_graphs(targeted):
    from linexcel import ExecutionPolicy
    from linexcel.analyzer import analyze_workbook

    workbook, sheet = _name_workbook()
    sheet["C1"] = "=CompanyName"
    sheet["C2"] = '=IFERROR(C1,"fallback")'
    output = io.BytesIO()
    workbook.save(output)
    data = output.getvalue()
    result = analyze_workbook(
        data,
        targets=["Probe!C2"] if targeted else None,
        execution=ExecutionPolicy(isolated=False),
    )
    formula_nodes = [
        node
        for node in result["graph"]["nodes"]
        if node["id"] in {"c:Probe!C1", "c:Probe!C2"}
    ]
    assert len(formula_nodes) == 2
    unsupported = any(
        check["feature"] == "defined_name_constant"
        and check["status"] == "known_mismatch"
        for check in result["graph"]["meta"]["engineCapabilities"]["checks"]
    )
    for node in formula_nodes:
        features = {risk["feature"] for risk in node.get("semanticRisks", [])}
        assert ("defined_name_constant" in features) == unsupported
        if unsupported:
            assert node["verification"] == "unverified_engine_semantics"
    assert data == output.getvalue()
