"""Regression tests for the B2B commercial & margin review showcase workbook."""

from __future__ import annotations

import importlib.util
import io
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

import linexcel
from linexcel.aidoc import build_workbook_dossier
from linexcel.excel_presentation import sheet_presentation
from linexcel.insights import extract_workbook_context
from linexcel.lazy_store import build_index

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def showcase_module():
    spec = importlib.util.spec_from_file_location(
        "showcase_workbooks_under_test", ROOT / "validation_workbooks.py"
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def showcase_bytes(showcase_module) -> bytes:
    return showcase_module.build_sales_workbook()


def _q(val: Decimal, places: str) -> Decimal:
    return val.quantize(Decimal(places), rounding=ROUND_HALF_UP)


def test_showcase_workbook_corporate_structure_and_context(showcase_bytes: bytes):
    wb = load_workbook(io.BytesIO(showcase_bytes), data_only=False)
    try:
        assert wb.sheetnames == ["Summary", "Sales", "Catalog", "Assumptions"]
        for name in wb.sheetnames:
            ws = wb[name]
            assert ws.freeze_panes == "B6"
            assert ws.sheet_properties.tabColor is not None
            assert len(ws.merged_cells.ranges) == 2

        assert wb["Summary"].column_dimensions["G"].hidden is True
        assert wb["Sales"].column_dimensions["H"].hidden is True
        assert len(wb["Summary"]._charts) == 1
        assert len(wb["Catalog"]._charts) == 1

        assert wb["Assumptions"]["D16"].value == "=SUM(D12:D15)"
        defined = {dn.name: dn.attr_text for dn in wb.defined_names.values()}
        assert defined == {"RebateRate": "Assumptions!$D$7"}

        order_ids = [wb["Sales"][f"B{r}"].value for r in range(6, 46)]
        close_dates = [wb["Sales"][f"C{r}"].value for r in range(6, 46)]
        assert order_ids[0] == "SO-26301" and order_ids[-1] == "SO-26340"
        assert len(set(order_ids)) == 40
        assert close_dates == sorted(close_dates)

        skus = [wb["Catalog"][f"B{r}"].value for r in range(6, 14)]
        assert len(skus) == 8
        assert all(isinstance(s, str) and s.startswith("NS-") for s in skus)

        # Ensure legacy synthetic toy placeholders never creep back in.
        legacy_tokens = {"P1", "P2", "P100", "Synthèse", "Ventes"}
        all_strings = {
            cell.value
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and not cell.value.startswith("=")
        }
        assert not (all_strings & legacy_tokens)
    finally:
        wb.close()

    ctx = extract_workbook_context(
        showcase_bytes, filename="fy26_commercial_review.xlsx"
    )
    assert ctx["warnings"] == []
    by_sheet = {s["name"]: s for s in ctx["sheets"]}
    assert by_sheet["Summary"]["hidden_columns"] == ["G"]
    assert by_sheet["Sales"]["hidden_columns"] == ["H"]
    assert [c["cell"] for c in by_sheet["Summary"]["comments"]] == ["C5"]
    assert [c["cell"] for c in by_sheet["Sales"]["comments"]] == ["I5"]
    assert [c["cell"] for c in by_sheet["Catalog"]["comments"]] == ["F5"]
    assert [c["cell"] for c in by_sheet["Assumptions"]["comments"]] == ["D6"]


def test_showcase_workbook_lineage_reconciles_with_decimal_accounting(
    showcase_module, showcase_bytes: bytes
):
    prices = {
        sku: (Decimal(price), Decimal(cost))
        for sku, _name, _fam, price, cost, _bm in showcase_module.CATALOG_ITEMS
    }
    rev_by_region: dict[str, Decimal] = {}
    gp_by_region: dict[str, Decimal] = {}
    for _oid, _dt, _acct, region, sku, qty, disc in showcase_module.SALES_ORDERS:
        price, cost = prices[sku]
        q = Decimal(qty)
        d = Decimal(str(disc))
        direct_cost = _q(q * cost, "0.01")
        net_rev = _q(q * price * (Decimal(1) - d), "0.01")
        gp = net_rev - direct_cost
        rev_by_region[region] = rev_by_region.get(region, Decimal(0)) + net_rev
        gp_by_region[region] = gp_by_region.get(region, Decimal(0)) + gp

    total_rev = sum(rev_by_region.values(), Decimal(0))
    total_gp = sum(gp_by_region.values(), Decimal(0))
    total_plan = sum(
        (
            Decimal(plan)
            for _reg, _vp, plan, _seg, _fy25 in showcase_module.REGIONAL_PLANS
        ),
        Decimal(0),
    )
    variance = _q((total_rev - total_plan) / total_plan, "0.0001")
    gross_margin = _q(total_gp / total_rev, "0.0001")
    net_contrib = _q(
        total_gp - total_rev * Decimal("0.035") - total_gp * Decimal("0.065"),
        "0.01",
    )

    result = linexcel.analyze(showcase_bytes, filename="fy26_commercial_review.xlsx")
    assert result.warnings == []
    assert result.stats["totalFormulas"] == 50
    assert result.stats["totalNodes"] == 20
    assert result.stats["totalEdges"] == 21
    assert result.stats["vbaProcs"] == 0
    assert result.stats["definedNames"] == 1
    assert result.stats["groupedPatterns"] == 2

    by_id = {n["id"]: n for n in result.nodes}
    by_label = {n["label"]: n for n in result.nodes}
    assert by_label["Assumptions!D16"]["value"] == pytest.approx(float(total_plan))
    assert by_label["Summary!C6"]["value"] == pytest.approx(float(total_rev))
    assert by_label["Summary!C7"]["value"] == pytest.approx(float(variance))
    assert by_label["Summary!C8"]["value"] == pytest.approx(float(total_gp))
    assert by_label["Summary!C9"]["value"] == pytest.approx(float(gross_margin))
    assert by_id["g:Summary!C15#4"]["value"] == pytest.approx(
        float(rev_by_region["North America"])
    )
    assert by_label["Summary!C10"]["value"] == pytest.approx(float(net_contrib))
    assert by_label["Summary!C10"]["value"] == pytest.approx(816668.1)
    steps = by_label["Summary!C10"]["steps"]
    assert steps["evaluated"] is True
    assert steps["value"] == pytest.approx(816668.1)
    assert len(steps["children"][0]["children"]) == 2

    dossier = build_workbook_dossier(result.graph)
    coverage = dossier["formula_pattern_coverage"]
    assert coverage["scope"] == "graph_formula_nodes"
    assert coverage["total"] == coverage["shown"] == 8
    assert coverage["omitted"] == 0


def test_showcase_workbook_presentation_formats_and_capture_contract(
    showcase_bytes: bytes, tmp_path: Path
):
    workbook_path = tmp_path / "fy26_commercial_review.xlsx"
    index_path = tmp_path / "index.sqlite"
    workbook_path.write_bytes(showcase_bytes)
    build_index(workbook_path, index_path)

    for sheet_name, hidden_col_idx in (
        ("Summary", "7"),
        ("Sales", "8"),
        ("Catalog", None),
        ("Assumptions", None),
    ):
        view = sheet_presentation(
            workbook_path, index_path, sheet_name, rows=60, columns=15
        )
        if hidden_col_idx is not None:
            assert view["columnDimensions"][hidden_col_idx]["hidden"] is True
        else:
            assert not any(
                dim.get("hidden") for dim in view["columnDimensions"].values()
            )
        unsupported = [
            (cell["address"], cell.get("numberFormat"))
            for cell in view["cells"]
            if cell.get("formatStatus") == "unsupported"
        ]
        assert unsupported == [], f"Unsupported formats on {sheet_name}: {unsupported}"

    spec = importlib.util.spec_from_file_location(
        "capture_app_under_test", ROOT / "scripts" / "capture_app.py"
    )
    assert spec is not None and spec.loader is not None
    capture_app = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = capture_app
    spec.loader.exec_module(capture_app)

    assert capture_app.NODE == "Summary!C10"
    receipt = capture_app.verify_fixture(workbook_path)
    assert receipt["name"] == "fy26_commercial_review.xlsx"
    assert receipt["generator"] == "validation_workbooks.build_sales_workbook"

    wrong_name = tmp_path / "renamed.xlsx"
    wrong_name.write_bytes(showcase_bytes)
    with pytest.raises(ValueError, match="fy26_commercial_review.xlsx"):
        capture_app.verify_fixture(wrong_name)
