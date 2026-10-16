"""Stress fixtures must exercise their intended failure without hidden cycles."""

import runpy
from graphlib import TopologicalSorter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries


def test_dense_broken_reference_fixture_has_no_incidental_cycle(tmp_path):
    script = Path(__file__).resolve().parents[1] / "scripts" / "stress_scenarios.py"
    build = runpy.run_path(str(script))["dense_chain_broken_ref"]
    path, refs = build(tmp_path)
    assert refs is None
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["S"]
        assert sheet["A1"].value == 1
        assert sheet["Z1"].value == "=Ghost!A1"
        formulas = {
            cell.coordinate: cell.value
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        }
        assert len(formulas) == 601

        # Use openpyxl's tokenizer and Python's independent cycle detector;
        # do not rely on linexcel grouping or native iteration diagnostics.
        predecessors = {}
        for address, formula in formulas.items():
            references = set()
            for token in Tokenizer(formula).items:
                if token.type != "OPERAND" or token.subtype != "RANGE":
                    continue
                if "!" in token.value:
                    continue  # The deliberately missing sheet is not this chain.
                c1, r1, c2, r2 = range_boundaries(token.value)
                references.update(
                    sheet.cell(row, column).coordinate
                    for row in range(r1, r2 + 1)
                    for column in range(c1, c2 + 1)
                )
            predecessors[address] = references
        order = tuple(TopologicalSorter(predecessors).static_order())
        assert len(order) == 602  # 601 formulas and the literal input.
    finally:
        workbook.close()
