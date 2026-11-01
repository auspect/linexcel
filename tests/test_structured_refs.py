"""Table calculation uses declared geometry, never cached totals or samples."""

import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from linexcel.lazy import build_structure, evaluate_node
from linexcel.structured_refs import formula_tokens


def workbook(tmp_path, *, headers=None, totals=True):
    book = Workbook()
    data = book.active
    data.title = "Data"
    data.append(headers or ["Item", "Amount", "Quantity"])
    data.append(["A", 10, 2])
    data.append(["B", 20, 3])
    if totals:
        data.append(["Total", 999, 888])
    data.add_table(
        Table(
            displayName="SalesTable",
            ref=f"A1:C{4 if totals else 3}",
            totalsRowCount=int(totals),
        )
    )
    book.create_sheet("Summary")
    return book, tmp_path / "tables.xlsx"


@pytest.mark.parametrize(
    "formula,expected",
    [
        ("=SUM(SalesTable[Amount])", 30),
        ("=SUM(SalesTable[[#Data],[Amount]])", 30),
        ("=SUM(SalesTable[[#Totals],[Amount]])", 999),
        ("=SalesTable[[#Headers],[Amount]]", "Amount"),
        ("=SUM(SalesTable[[#All],[Amount]])", 1029),
        ("=SUM(SalesTable[[#Data],[Amount]:[Quantity]])", 35),
        ("=SUM(SalesTable)", 35),
        ("=SUM(salestable[amount])", 30),
        ("=SUM(Data!SalesTable[Amount])", 30),
    ],
)
def test_table_selectors_are_exact_and_source_formula_is_retained(
    tmp_path, formula, expected
):
    book, path = workbook(tmp_path)
    book["Summary"]["A1"] = formula
    book["Summary"]["A2"] = "=IF(ISNUMBER(A1),A1*2,A1)"
    book.save(path)
    result = evaluate_node(path, "Summary!A2")
    assert result["status"] == "completed", result["diagnostics"]
    assert result["value"] == (expected * 2 if isinstance(expected, int) else expected)
    source = next(s for s in result["steps"] if s["nodeId"] == "Summary!A1")
    assert source["sourceFormula"] == formula
    assert result["coverage"]["closureComplete"]


@pytest.mark.parametrize(
    "reference",
    ["[@Amount]", "SalesTable[@Amount]", "SalesTable[[#This Row],[Amount]]"],
)
def test_table_row_references_bind_to_each_formula_row(tmp_path, reference):
    book, path = workbook(tmp_path, totals=False)
    book["Data"]["C2"] = f"={reference}*2"
    book["Data"]["C3"] = f"={reference}*3"
    book["Summary"]["A1"] = "=SUM(SalesTable[Quantity])"
    book.save(path)
    result = evaluate_node(path, "Summary!A1")
    assert result["value"] == 80, result
    assert result["coverage"]["calculatedFormulaCells"] == 3


@pytest.mark.parametrize(
    "header,reference",
    [
        ("A]B", "SalesTable[A']B]"),
        ("A[B", "SalesTable[A'[B]"),
        ("#Amount", "SalesTable['#Amount]"),
        ("@Amount", "SalesTable['@Amount]"),
        ("O'Brien", "SalesTable[O''Brien]"),
        ("Comma, space", "SalesTable[Comma, space]"),
    ],
)
def test_escaped_table_column_names(tmp_path, header, reference):
    book, path = workbook(tmp_path, headers=["Item", header, "Quantity"])
    book["Summary"]["A1"] = f"=SUM({reference})"
    book.save(path)
    result = evaluate_node(path, "Summary!A1")
    assert result["value"] == 30, result["diagnostics"]


@pytest.mark.parametrize(
    "formula,code",
    [
        ("=SUM(Missing[Amount])", "table_missing"),
        ("=SUM(SalesTable[Unknown])", "table_column_missing"),
        ("=SUM([Amount])", "table_missing"),
        ("=SUM(SalesTable[[#Headers],[#Totals],[Amount]])", "table_disjoint_rows"),
    ],
)
def test_table_failure_stops_transitive_calculation(tmp_path, formula, code):
    book, path = workbook(tmp_path)
    book["Summary"]["A1"] = formula
    book["Summary"]["A2"] = "=A1+1"
    book.save(path)
    result = evaluate_node(path, "Summary!A2")
    assert result["status"] == "unsupported"
    assert any(d["code"] == code for d in result["capabilityDiagnostics"])


def test_bad_declared_table_keeps_source_but_refuses_calculation(tmp_path):
    book, path = workbook(tmp_path)
    book["Summary"]["A1"] = "=SUM(SalesTable[Amount])"
    book.save(path)
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    table = ET.fromstring(parts["xl/tables/table1.xml"])
    table.find("{*}tableColumns").remove(table.find("{*}tableColumns/{*}tableColumn"))
    parts["xl/tables/table1.xml"] = ET.tostring(table)
    with zipfile.ZipFile(path, "w") as target:
        for name, content in parts.items():
            target.writestr(name, content)
    graph = build_structure(path)
    assert any(
        d.get("code") == "table_metadata_invalid" for d in graph["meta"]["diagnostics"]
    )
    result = evaluate_node(path, "Summary!A1")
    assert result["status"] == "unsupported"


def test_tokenization_preserves_quoted_external_sheets_and_literals():
    formulas = [
        "='[source.xlsx]Data'!A1",
        "='[Odd sheet'!A1",
        "=IF(A1=\"a'[b\",SUM(SalesTable[A']B]),0)",
    ]
    for formula in formulas:
        assert "=" + "".join(t.value for t in formula_tokens(formula)) == formula
