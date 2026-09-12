"""Targeted evaluation (--target), the cheaper quarantine retry, and the
chain-depth risk indicator.

Targeted mode boots the engine without ``evaluate_all``, traces the upstream
subgraph of the requested cells, evaluates only it, and leaves the rest of
the workbook out of the lineage. The quarantine change probes the engine
after a failed retry instead of paying a third ``from_bytes`` blindly.
"""

import io

import pytest
from openpyxl import Workbook

from linexcel import analyze
from linexcel.analyzer import analyze_workbook


def workbook(cells: dict[str, object], sheet: str = "S") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for address, value in cells.items():
        ws[address] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def two_sheet_workbook() -> bytes:
    """Data feeds Out!C1; Out!D1 is a separate branch the target never touches."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["B1"] = "=A1*2"
    ws["B2"] = "=B1+A2"
    out = wb.create_sheet("Out")
    out["C1"] = "=Data!B2*3"
    out["D1"] = "=SUM(Data!A1:A2)"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestTargetedAnalysis:
    """--target: only the upstream subgraph is evaluated and graphed."""

    def test_the_graph_is_the_upstream_subgraph(self):
        result = analyze(two_sheet_workbook(), filename="t.xlsx", targets=["Out!C1"])
        ids = {n["id"] for n in result.nodes}
        assert ids == {
            "c:Out!C1",
            "c:Data!B1",
            "c:Data!B2",
            "i:Data!A1",
            "i:Data!A2",
        }
        edges = {(e["source"], e["target"]) for e in result.edges}
        assert ("c:Data!B2", "c:Out!C1") in edges
        assert ("c:Data!B1", "c:Data!B2") in edges
        # The off-target branch is omitted, not just unevaluated.
        assert all("D1" not in i for i in ids)

    def test_off_target_cells_are_not_evaluated(self):
        result = analyze(two_sheet_workbook(), filename="t.xlsx", targets=["Out!C1"])
        # The target path carries recomputed values...
        c1 = next(n for n in result.nodes if n["id"] == "c:Out!C1")
        assert c1["value"] == 120
        assert c1["valueSource"] == "engine"
        # ...and the engine never computed the off-target formula.
        assert result.engine.get_value("Out", 1, 4) is None

    def test_several_targets_union_their_subgraphs(self):
        result = analyze(
            two_sheet_workbook(), filename="t.xlsx", targets=["Out!C1,Out!D1"]
        )
        ids = {n["id"] for n in result.nodes}
        assert "c:Out!C1" in ids and "c:Out!D1" in ids
        d1 = next(n for n in result.nodes if n["id"] == "c:Out!D1")
        assert d1["value"] == 30

    def test_a_constant_target_still_gets_a_node(self):
        result = analyze(two_sheet_workbook(), filename="t.xlsx", targets=["Data!A1"])
        node = next(n for n in result.nodes if n["id"] == "i:Data!A1")
        assert node["value"] == 10

    def test_the_mode_is_announced_and_recorded(self):
        result = analyze(two_sheet_workbook(), filename="t.xlsx", targets=["Out!C1"])
        assert result.graph["targets"] == ["Out!C1"]
        (warning,) = [w for w in result.warnings if w.startswith("Targeted analysis")]
        assert "Out!C1" in warning
        assert "omitted" in warning

    def test_no_target_keeps_the_whole_workbook(self):
        """Default behaviour is untouched: everything is evaluated."""
        result = analyze(two_sheet_workbook(), filename="t.xlsx")
        assert "targets" not in result.graph
        assert not [w for w in result.warnings if w.startswith("Targeted")]
        assert result.engine.get_value("Out", 1, 4) == 30

    def test_a_sheet_with_a_space_in_its_name(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "My Sheet"
        ws["A1"] = 5
        ws["B1"] = "=A1*2"
        buf = io.BytesIO()
        wb.save(buf)
        result = analyze(buf.getvalue(), filename="t.xlsx", targets=["My Sheet!B1"])
        ids = {n["id"] for n in result.nodes}
        assert "c:My Sheet!B1" in ids
        # Input node ids keep the formula spelling, sheet name quoted.
        assert "i:'My Sheet'!A1" in ids

    def test_an_unqualified_target_is_rejected(self):
        with pytest.raises(ValueError, match="sheet-qualified"):
            analyze(two_sheet_workbook(), filename="t.xlsx", targets=["C1"])

    def test_a_range_is_not_a_target(self):
        with pytest.raises(ValueError, match="Invalid target"):
            analyze(two_sheet_workbook(), filename="t.xlsx", targets=["Out!C1:D1"])

    def test_an_unknown_sheet_is_rejected(self):
        with pytest.raises(ValueError, match="does not have"):
            analyze(two_sheet_workbook(), filename="t.xlsx", targets=["Gone!A1"])

    def test_duplicate_targets_are_collapsed(self):
        result = analyze(
            two_sheet_workbook(),
            filename="t.xlsx",
            targets=["Out!C1", "Out!C1"],
        )
        assert result.graph["targets"] == ["Out!C1"]


class TestTargetedWithABrokenPrecedent:
    """A target whose subgraph holds a reference the engine cannot resolve.

    The one-batch evaluation is all-or-nothing, so it fails whole; the
    fallback evaluates the requested targets one at a time, and the warning
    says which of them did not make it.
    """

    @staticmethod
    def _workbook() -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = 3
        ws["B1"] = "=A1*2"
        ws["B2"] = "=B1+NoSheet!A1"  # the blocker, on the path of C1
        ws["C1"] = "=B2*10"
        ws["D1"] = "=A1+100"  # a clean second target
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_the_clean_target_is_evaluated_and_the_broken_one_named(self):
        result = analyze(
            self._workbook(), filename="t.xlsx", targets=["Data!D1", "Data!C1"]
        )
        d1 = next(n for n in result.nodes if n["id"] == "c:Data!D1")
        assert d1["value"] == 103
        joined = " ".join(result.warnings)
        assert "Targeted evaluation" in joined
        assert "Data!C1" in joined

    def test_the_broken_cell_stays_on_the_graph(self):
        result = analyze(self._workbook(), filename="t.xlsx", targets=["Data!C1"])
        ids = {n["id"] for n in result.nodes}
        assert "c:Data!B2" in ids


class TestTruncatedTrace:
    def test_a_truncated_trace_says_the_subgraph_is_partial(self, monkeypatch):
        monkeypatch.setattr("linexcel.engine.TRACE_MAX_NODES", 1)
        result = analyze(two_sheet_workbook(), filename="t.xlsx", targets=["Out!C1"])
        assert any("trace of the target" in w for w in result.warnings)


class TestQuarantineRetryCosts:
    """A failed retry no longer pays a third ``from_bytes`` to get formulas back.

    formualizer 0.9.3 keeps the formula map readable after a failed
    ``evaluate_all``, so the engine the run already holds is kept and the
    rebuild is paid only when the probe says the formulas really are gone.
    """

    @staticmethod
    def _workbook() -> bytes:
        return workbook({"A1": 2, "B1": "=A1*3", "C1": "=NoSheet!A1+1"})

    def test_a_failed_retry_does_not_rebuild(self, monkeypatch):
        import linexcel.engine as eng

        # Make the quarantine scan miss the real blocker, so the retry fails
        # the way the first pass did.
        monkeypatch.setattr(
            eng, "_find_unresolvable", lambda data, sheets: {("S", 1, 2): "=A1*3"}
        )
        boots = 0
        real_open = eng._open_workbook

        def counting_open(data, parallel):
            nonlocal boots
            boots += 1
            return real_open(data, parallel)

        monkeypatch.setattr(eng, "_open_workbook", counting_open)
        warnings: list[str] = []
        session = eng.boot_engine(self._workbook(), warnings)
        # First boot + sanitized retry; the old code paid a third boot here.
        assert boots == 2
        assert not session.engine_alive
        # The decoy stays quarantined so the sweep re-injects its formula.
        assert session.quarantined == {("S", 1, 2): "=A1*3"}

    def test_the_graph_is_still_complete_after_a_failed_retry(self, monkeypatch):
        import linexcel.engine as eng

        monkeypatch.setattr(
            eng, "_find_unresolvable", lambda data, sheets: {("S", 1, 2): "=A1*3"}
        )
        graph = analyze_workbook(self._workbook(), "t.xlsx")["graph"]
        formulas = {n.get("formula") for n in graph["nodes"] if "formula" in n}
        assert "=NoSheet!A1+1" in formulas
        assert "=A1*3" in formulas
        meta_warnings = graph["meta"]["warnings"]
        assert any("Global evaluation incomplete" in w for w in meta_warnings)

    def test_the_probe_rebuilds_only_when_formulas_are_gone(self):
        from linexcel.engine import _formulas_gone

        data = self._workbook()

        class Alive:
            def get_formula(self, sheet, row, col):
                return "A1*3"

        class Empty:
            def get_formula(self, sheet, row, col):
                return None

        class Broken:
            def get_formula(self, sheet, row, col):
                raise RuntimeError("engine gone")

        assert not _formulas_gone(Alive(), data, {"S"}, {})
        assert _formulas_gone(Empty(), data, {"S"}, {})
        assert _formulas_gone(Broken(), data, {"S"}, {})

    def test_a_workbook_without_formulas_needs_no_rebuild(self):
        from linexcel.engine import _formulas_gone

        assert not _formulas_gone(object(), workbook({"A1": 1}), {"S"}, {})


class TestChainDepthWarning:
    """Long dependency chains are flagged as a risk, not timed as a promise."""

    @staticmethod
    def _chain_workbook(length: int) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "S"
        ws["A1"] = 1
        for row in range(2, length + 1):
            ws.cell(row=row, column=1, value=f"=A{row - 1}+1")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_a_long_chain_is_flagged(self):
        result = analyze(self._chain_workbook(40), filename="chain.xlsx")
        (warning,) = [w for w in result.warnings if "Long dependency chains" in w]
        assert "linked calculation steps" in warning
        assert "not a duration estimate" in warning

    def test_a_short_chain_is_not_flagged(self):
        result = analyze(self._chain_workbook(5), filename="chain.xlsx")
        assert not [w for w in result.warnings if "Long dependency chains" in w]

    def test_a_long_chain_is_flagged_in_targeted_mode(self):
        """Targeted mode reads the engine's evaluation plan, exactly."""
        result = analyze(
            self._chain_workbook(40), filename="chain.xlsx", targets=["S!A40"]
        )
        (warning,) = [w for w in result.warnings if "Long dependency chains" in w]
        assert "evaluation layers" in warning

    def test_a_workbook_without_formulas_is_not_flagged(self):
        result = analyze(workbook({"A1": 1, "B1": "text"}), filename="plain.xlsx")
        assert not [w for w in result.warnings if "Long dependency chains" in w]
