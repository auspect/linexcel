"""Targeted evaluation must keep its cost and side effects inside the closure."""

import io

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from linexcel import analyze
from linexcel.progress import Reporter
from linexcel.sweep import sweep_sheets


def test_unrelated_defined_name_does_not_recalculate_its_formula():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 2
    ws["B1"] = "=A1*3"
    ws["C1"] = "=A1*100"
    wb.defined_names.add(DefinedName("Unrelated", attr_text="'S'!$C$1"))
    output = io.BytesIO()
    wb.save(output)

    result = analyze(output.getvalue(), targets=["S!B1"])

    assert result.engine.get_value("S", 1, 3) is None
    assert "n:Unrelated" not in {node["id"] for node in result.nodes}


def test_sparse_targets_do_not_spend_the_scan_budget_on_empty_cells():
    calls = []

    class Engine:
        def sheet(self, sheet):
            return self

        def get_formulas(self, rect):
            calls.append(rect)
            # A sparse scan should request only one occupied cell per call.
            assert rect.start_row == rect.end_row
            assert rect.start_col == rect.end_col
            return [["=1+1"]]

    wanted = {("S", row, 1) for row in (1, 1001, 2001, 3001, 4001)}
    warnings = []
    result = sweep_sheets(
        Engine(),
        {"S": (5000, 16384)},
        {"S"},
        {},
        warnings,
        Reporter(),
        reachable=wanted,
    )

    assert result.formula_count == 5
    assert len(calls) == 5
    assert not warnings


def test_a_sparse_group_does_not_invent_blank_inputs_between_targets():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for address, value in {
        "A1": 2,
        "A2": "=30+40",
        "A3": 5,
        "B1": "=A1*2",
        "B2": "=A2*2",
        "B3": "=A3*2",
    }.items():
        ws[address] = value
    output = io.BytesIO()
    wb.save(output)

    result = analyze(output.getvalue(), targets=["S!B1", "S!B3"])
    ids = {node["id"] for node in result.nodes}
    assert "i:S!A1" in ids and "i:S!A3" in ids
    assert "i:S!A1:A3" not in ids
    assert result.engine.get_value("S", 2, 1) is None


def test_targeted_query_provenance_keeps_only_the_relevant_query():
    from pathlib import Path

    result = analyze(Path("tests/fixtures/power_query.xlsx"), targets=["Loaded!A2"])
    queries = {node["label"] for node in result.nodes if node["kind"] == "query"}
    assert queries == {"BusyProducts"}


def test_a_used_defined_name_keeps_its_precedents():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 2
    ws["B1"] = "=Rate*3"
    wb.defined_names.add(DefinedName("Rate", attr_text="'S'!$A$1"))
    output = io.BytesIO()
    wb.save(output)
    result = analyze(output.getvalue(), targets=["S!B1"])
    ids = {node["id"] for node in result.nodes}
    assert {"n:Rate", "i:S!A1", "c:S!B1"} <= ids
    assert next(node for node in result.nodes if node["id"] == "c:S!B1")["value"] == 6


def test_dense_target_columns_stay_batched():
    from linexcel.sweep import _target_ranges

    assert list(_target_ranges([(row, 7) for row in range(1, 10001)])) == [
        (1, 7, 10000, 7)
    ]
