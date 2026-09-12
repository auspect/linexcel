"""Synthetic Excel time caches retain their meaning across JSON serialization."""

import datetime
import io
import re
import zipfile

import pytest
from openpyxl import Workbook

from linexcel.analyzer import analyze_workbook
from linexcel.loader import load_cached_values
from linexcel.values import readings_agree, serial_to_time_text


@pytest.mark.parametrize(
    ("serial", "stored"),
    [
        (0, datetime.time(0)),
        (0.5, datetime.time(12)),
        (0.25, datetime.time(6)),
        (0.5 + 0.123 / 86400, datetime.time(12, 0, 0, 123000)),
    ],
)
def test_raw_and_serialized_time_caches_agree(serial, stored):
    assert readings_agree(serial, stored, None) == "same"
    assert (
        readings_agree(serial, stored.isoformat(), serial_to_time_text(serial))
        == "same"
    )


@pytest.mark.parametrize("serial", [-1, 1, 1.5, float("nan"), float("inf"), True])
def test_a_time_of_day_does_not_discard_days_or_coerce_booleans(serial):
    assert serial_to_time_text(serial) is None
    assert readings_agree(serial, datetime.time(12), None) == "differ"


@pytest.mark.parametrize("stored", ["0", "00:00:00", "12:00:00"])
def test_time_shaped_text_without_temporal_metadata_stays_text(stored):
    assert readings_agree(0, stored, None) == "differ"


def test_a_true_time_difference_is_not_hidden():
    assert readings_agree(0.5, datetime.time(12, 1), None) == "differ"
    assert readings_agree(0.5, "12:01:00", "12:00:00") == "differ"


def test_no_timezone_is_assumed_for_excel_serials():
    assert readings_agree(0.5, datetime.time(12, tzinfo=datetime.UTC), None) == "differ"
    assert readings_agree(0.5, "12:00:00+00:00", "12:00:00") == "differ"


def _time_workbook(serial: float, cached_serial: float, epoch_1904: bool) -> bytes:
    wb = Workbook()
    if epoch_1904:
        wb.epoch = datetime.datetime(1904, 1, 1)
    ws = wb.active
    ws.title = "Clock"
    ws["A1"] = serial
    ws["B1"] = "=A1"
    ws["B1"].number_format = "hh:mm:ss.000"
    original = io.BytesIO()
    wb.save(original)
    output = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(output, "w") as dest:
        for part in source.infolist():
            payload = source.read(part.filename)
            if part.filename == "xl/worksheets/sheet1.xml":
                payload, count = re.subn(
                    rb"</f><v(?:\s*/>|></v>)",
                    f"</f><v>{cached_serial}</v>".encode(),
                    payload,
                )
                assert count == 1
            dest.writestr(part, payload)
    return output.getvalue()


@pytest.mark.parametrize("epoch_1904", [False, True])
@pytest.mark.parametrize("serial", [0, 0.5, 0.5 + 0.123 / 86400])
@pytest.mark.parametrize("slow_reader", [False, True])
def test_real_in_memory_formula_time_cache_is_not_a_divergence(
    serial, epoch_1904, slow_reader, monkeypatch
):
    if slow_reader:
        monkeypatch.setattr("linexcel.loader.MAX_DENSE_CELLS", 0)
    data = _time_workbook(serial, serial, epoch_1904)
    assert isinstance(load_cached_values(data).get("Clock", 1, 2), datetime.time)
    graph = analyze_workbook(data)["graph"]
    node = next(node for node in graph["nodes"] if node["id"] == "c:Clock!B1")
    assert node["value"] == pytest.approx(serial)
    assert node["valueSource"] == "engine"
    assert node["cachedAgreement"] == "same"
    assert node["valueDate"] == serial_to_time_text(serial)
    assert not any("differs from file" in w for w in graph["meta"]["warnings"])


def test_real_in_memory_time_cache_detects_a_wrong_minute():
    graph = analyze_workbook(_time_workbook(0.5, 0.5 + 1 / 1440, False))["graph"]
    node = next(node for node in graph["nodes"] if node["id"] == "c:Clock!B1")
    assert node["cachedAgreement"] == "differ"
    assert any("differs from file" in w for w in graph["meta"]["warnings"])
