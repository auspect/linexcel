"""Web contracts: isolated work, private ownership, restart and honest failures."""

from __future__ import annotations

import base64
import io
import json
import os
import sys
import threading
import time

import pytest
from openpyxl import Workbook

from linexcel.web import APIError, TaskStore, WebConfig, create_app


def workbook():
    book = Workbook()
    book.active["A1"] = 6
    book.active["B1"] = "=A1*7"
    stream = io.BytesIO()
    book.save(stream)
    return {
        "name": "example.xlsx",
        "data": base64.b64encode(stream.getvalue()).decode(),
    }


def request(app, method, path, body=None, user="alice", **extra):
    data = json.dumps(body).encode() if body is not None else b""
    env = {
        "REQUEST_METHOD": method,
        "PATH_INFO": path,
        "SCRIPT_NAME": "",
        "CONTENT_TYPE": "application/json",
        "CONTENT_LENGTH": str(len(data)),
        "wsgi.input": io.BytesIO(data),
        "wsgi.url_scheme": "http",
        "HTTP_HOST": "127.0.0.1:8765",
        "user": user,
    }
    env.update(extra)
    response = []
    payload = b"".join(
        app(env, lambda status, headers: response.extend([status, headers]))
    )
    return int(response[0].split()[0]), payload, dict(response[1])


def wait(store, task_id):
    deadline = time.monotonic() + 20
    while time.monotonic() < deadline:
        with store.lock:
            task = store.tasks[task_id]
            if task["status"] in {"succeeded", "failed", "cancelled"}:
                return store.get_task(task["owner"], task_id)
        time.sleep(0.02)
    raise AssertionError("Task did not terminate")


@pytest.fixture
def app(tmp_path):
    app = create_app(tmp_path, resolve_user=lambda env: env.get("user"))
    yield app
    app.close()


@pytest.mark.parametrize("arena_limit", [None, "7", "0"])
def test_worker_arena_default_preserves_parent_settings(
    app, tmp_path, monkeypatch, arena_limit
):
    if arena_limit is None:
        monkeypatch.delenv("MALLOC_ARENA_MAX", raising=False)
    else:
        monkeypatch.setenv("MALLOC_ARENA_MAX", arena_limit)
    root = tmp_path / "worker"
    root.mkdir()
    seen = []

    def capture_worker(command, *, env, **kwargs):
        seen.append(env.get("MALLOC_ARENA_MAX"))
        raise RuntimeError("worker launch intercepted")

    monkeypatch.setattr("linexcel.web.subprocess.Popen", capture_worker)
    task = {
        "operation": "import",
        "nodeId": None,
        "budget": {"memoryMb": 512, "seconds": 60, "tokens": 16000},
    }
    with pytest.raises(RuntimeError, match="worker launch intercepted"):
        app.store._worker(task, root, threading.Event())
    expected = "2" if arena_limit is None and sys.platform == "linux" else arena_limit
    assert seen == [expected]
    assert os.environ.get("MALLOC_ARENA_MAX") == arena_limit


def test_full_import_evaluate_reconnect_export_and_delete(app):
    code, body, _ = request(app, "POST", "/api/projects", {"workbook": workbook()})
    assert code == 200
    imported = json.loads(body)
    project = imported["project"]["id"]
    task = wait(app.store, imported["task"]["id"])
    assert task["status"] == "succeeded", task
    code, body, _ = request(app, "GET", f"/api/projects/{project}")
    graph = json.loads(body)["graph"]
    assert graph["meta"]["evaluated"] is False
    assert next(n for n in graph["nodes"] if n["id"] == "Sheet!B1")["value"] is None
    code, body, _ = request(
        app,
        "POST",
        f"/api/projects/{project}/tasks",
        {
            "operation": "evaluate",
            "nodeId": "Sheet!B1",
        },
    )
    assert code == 200
    task_id = json.loads(body)["task"]["id"]
    task = wait(app.store, task_id)
    assert task["status"] == "succeeded", task
    assert task["result"]["value"] == 42
    assert task["result"]["comparison"] == "unknown"
    _, body, _ = request(
        app,
        "POST",
        f"/api/projects/{project}/tasks",
        {
            "operation": "evaluate",
            "nodeId": "Sheet!B1",
        },
    )
    assert json.loads(body)["task"]["id"] == task_id
    _, body, _ = request(
        app,
        "POST",
        f"/api/projects/{project}/tasks",
        {"operation": "evaluate", "nodeId": "Sheet!B1", "force": True},
    )
    forced_id = json.loads(body)["task"]["id"]
    assert forced_id != task_id
    assert wait(app.store, forced_id)["result"]["value"] == 42
    _, body, _ = request(app, "GET", f"/api/projects/{project}/export")
    exported = json.loads(body)
    assert exported["export"]["scope"] == "complete"
    assert exported["graph"]["nodes"] == graph["nodes"]
    assert exported["graph"]["edges"] == graph["edges"]
    assert exported["graph"]["meta"] == graph["meta"]
    assert (
        next(t for t in exported["tasks"] if t["id"] == task_id)["result"]["value"]
        == 42
    )
    assert request(app, "DELETE", f"/api/projects/{project}")[0] == 200
    assert request(app, "GET", f"/api/tasks/{task_id}")[0] == 404


def test_users_cannot_access_projects_tasks_or_exports(app):
    _, body, _ = request(app, "POST", "/api/projects", {"workbook": workbook()})
    imported = json.loads(body)
    project, task_id = imported["project"]["id"], imported["task"]["id"]
    wait(app.store, task_id)
    for method, path in [
        ("GET", f"/api/projects/{project}"),
        ("GET", f"/api/projects/{project}/export"),
        ("DELETE", f"/api/projects/{project}"),
        ("GET", f"/api/tasks/{task_id}"),
        ("POST", f"/api/tasks/{task_id}/cancel"),
    ]:
        assert request(app, method, path, user="bob")[0] == 404
    assert request(app, "GET", "/api/projects", user=None)[0] == 401
    assert json.loads(request(app, "GET", "/api/projects", user="bob")[1]) == {
        "projects": []
    }


def test_restart_reuses_latest_calculation_regardless_of_directory_order(
    tmp_path, monkeypatch
):
    from pathlib import Path

    config = WebConfig(tmp_path)
    first = TaskStore(config)
    try:
        created = first.create("alice", {"workbook": workbook()})
        wait(first, created["task"]["id"])
        project = created["project"]["id"]
        older = first.submit("alice", project, "evaluate", "Sheet!B1", {})
        assert wait(first, older["id"])["status"] == "succeeded"
        newer = first.submit("alice", project, "evaluate", "Sheet!B1", {}, force=True)
        assert wait(first, newer["id"])["status"] == "succeeded"
    finally:
        first.close()

    original_glob = Path.glob

    def newest_files_first(directory, pattern, **kwargs):
        paths = list(original_glob(directory, pattern, **kwargs))
        if pattern == "tasks/*/task.json":
            paths.sort(
                key=lambda path: json.loads(path.read_text("utf-8"))["createdAt"],
                reverse=True,
            )
        return iter(paths)

    monkeypatch.setattr(Path, "glob", newest_files_first)
    recovered = TaskStore(config)
    try:
        reused = recovered.submit("alice", project, "evaluate", "Sheet!B1", {})
        assert reused["id"] == newer["id"]
        assert reused["createdAt"] > older["createdAt"]
    finally:
        recovered.close()


def test_unsupported_evaluation_is_retryable_without_force(app, monkeypatch):
    imported = app.store.create("alice", {"workbook": workbook()})
    assert wait(app.store, imported["task"]["id"])["status"] == "succeeded"
    project = imported["project"]["id"]
    outcomes = iter(
        [
            {
                "status": "unsupported",
                "value": None,
                "diagnostics": ["closure refused"],
            },
            {"status": "completed", "value": 42, "steps": []},
        ]
    )
    monkeypatch.setattr(
        app.store, "_worker", lambda task, root, event: {"result": next(outcomes)}
    )
    refused = app.store.submit("alice", project, "evaluate", "Sheet!B1", {})
    assert wait(app.store, refused["id"])["result"]["status"] == "unsupported"
    retried = app.store.submit("alice", project, "evaluate", "Sheet!B1", {})
    assert retried["id"] != refused["id"]
    assert wait(app.store, retried["id"])["result"]["value"] == 42
    assert (
        app.store.submit("alice", project, "evaluate", "Sheet!B1", {})["id"]
        == retried["id"]
    )


@pytest.mark.parametrize("latest_outcome", ["failed", "unsupported"])
def test_latest_unsuccessful_evaluation_never_falls_back_to_an_older_value(
    app, monkeypatch, latest_outcome
):
    imported = app.store.create("alice", {"workbook": workbook()})
    assert wait(app.store, imported["task"]["id"])["status"] == "succeeded"
    project = imported["project"]["id"]
    latest = (
        {"error": {"kind": "timeout", "message": "target budget exceeded"}}
        if latest_outcome == "failed"
        else {"result": {"status": "unsupported", "value": None}}
    )
    outcomes = iter(
        [
            {"result": {"status": "completed", "value": 11}},
            latest,
            {"result": {"status": "completed", "value": 42}},
        ]
    )
    monkeypatch.setattr(app.store, "_worker", lambda task, root, event: next(outcomes))
    older = app.store.submit("alice", project, "evaluate", "Sheet!B1", {})
    assert wait(app.store, older["id"])["result"]["value"] == 11
    unsuccessful = app.store.submit(
        "alice", project, "evaluate", "Sheet!B1", {}, force=True
    )
    finished = wait(app.store, unsuccessful["id"])
    assert (
        finished["status"] == "failed"
        if latest_outcome == "failed"
        else finished["result"]["status"] == "unsupported"
    )
    retried = app.store.submit("alice", project, "evaluate", "Sheet!B1", {})
    assert retried["id"] not in {older["id"], unsuccessful["id"]}
    assert wait(app.store, retried["id"])["result"]["value"] == 42
    assert app.store.get_task("alice", older["id"])["result"]["value"] == 11


def test_force_joins_active_work_but_keeps_nodes_and_projects_independent(
    tmp_path, monkeypatch
):
    from concurrent.futures import ThreadPoolExecutor

    store = TaskStore(WebConfig(tmp_path, max_workers=1))
    entered, release = threading.Event(), threading.Event()
    try:
        created = [store.create("alice", {"workbook": workbook()}) for _ in range(2)]
        for item in created:
            assert wait(store, item["task"]["id"])["status"] == "succeeded"
        project, other_project = [item["project"]["id"] for item in created]
        calls = []

        def worker(task, root, event):
            calls.append((task["projectId"], task["nodeId"]))
            entered.set()
            assert release.wait(5)
            return {
                "result": {"nodeId": task["nodeId"], "status": "completed", "value": 42}
            }

        monkeypatch.setattr(store, "_worker", worker)
        running = store.submit("alice", project, "evaluate", "Sheet!B1", {}, force=True)
        assert entered.wait(2)
        with ThreadPoolExecutor(max_workers=4) as callers:
            joined = list(
                callers.map(
                    lambda _: store.submit(
                        "alice", project, "evaluate", "Sheet!B1", {}, force=True
                    ),
                    range(4),
                )
            )
        assert {task["id"] for task in joined} == {running["id"]}
        queued = store.submit("alice", project, "evaluate", "Sheet!A1", {}, force=True)
        assert queued["status"] == "queued"
        assert (
            store.submit("alice", project, "evaluate", "Sheet!A1", {}, force=True)["id"]
            == queued["id"]
        )
        other = store.submit(
            "alice", other_project, "evaluate", "Sheet!B1", {}, force=True
        )
        assert len({running["id"], queued["id"], other["id"]}) == 3
        release.set()
        for task in (running, queued, other):
            assert wait(store, task["id"])["status"] == "succeeded"
        assert calls == [
            (project, "Sheet!B1"),
            (project, "Sheet!A1"),
            (other_project, "Sheet!B1"),
        ]
        with pytest.raises(APIError):
            store.get_task("bob", running["id"])
        fresh = store.submit("alice", project, "evaluate", "Sheet!B1", {}, force=True)
        assert fresh["id"] != running["id"]
        assert wait(store, fresh["id"])["status"] == "succeeded"
    finally:
        release.set()
        store.close()


def test_an_evaluation_with_cancellation_requested_is_not_joined(app, monkeypatch):
    imported = app.store.create("alice", {"workbook": workbook()})
    assert wait(app.store, imported["task"]["id"])["status"] == "succeeded"
    project = imported["project"]["id"]
    entered, release = threading.Event(), threading.Event()

    def worker(task, root, event):
        entered.set()
        assert release.wait(5)
        return {"result": {"status": "completed", "value": 42}}

    monkeypatch.setattr(app.store, "_worker", worker)
    try:
        old = app.store.submit("alice", project, "evaluate", "Sheet!B1", {}, force=True)
        assert entered.wait(2)
        app.store.cancel("alice", old["id"])
        new = app.store.submit("alice", project, "evaluate", "Sheet!B1", {}, force=True)
        assert new["id"] != old["id"]
    finally:
        release.set()
    assert wait(app.store, old["id"])["status"] == "cancelled"
    assert wait(app.store, new["id"])["result"]["value"] == 42


def test_refused_formula_and_excel_error_do_not_block_other_nodes(app):
    book = Workbook()
    book.active["A1"] = "=1/0"
    book.active["B1"] = "=IFERROR(A1,42)"
    book.active["C1"] = '=INDIRECT("D"&RAND())'
    book.active["D1"] = 9
    book.active["E1"] = "=2+3"
    stream = io.BytesIO()
    book.save(stream)
    imported = app.store.create(
        "alice",
        {
            "workbook": {
                "name": "mixed.xlsx",
                "data": base64.b64encode(stream.getvalue()).decode(),
            }
        },
    )
    assert wait(app.store, imported["task"]["id"])["status"] == "succeeded"
    project = imported["project"]["id"]
    original = app.store.snapshot("alice", project)["graph"]

    def evaluate(node):
        task = app.store.submit("alice", project, "evaluate", node, {}, force=True)
        finished = wait(app.store, task["id"])
        assert finished["status"] == "succeeded", finished
        return finished["result"]

    error = evaluate("Sheet!A1")
    assert error["status"] == "completed"
    assert error["value"] == "#DIV/0!" and error["valueKind"] == "error"
    guarded = evaluate("Sheet!B1")
    assert guarded["status"] == "completed" and guarded["value"] == 42
    assert {step["nodeId"] for step in guarded["steps"]} == {"Sheet!A1", "Sheet!B1"}
    assert all(
        step["calculatedProvenance"] == "targeted_recalculation"
        for step in guarded["steps"]
    )
    refused = evaluate("Sheet!C1")
    assert refused["status"] == "unsupported"
    assert refused["coverage"]["status"] == "not_evaluated"
    healthy = evaluate("Sheet!E1")
    assert healthy["status"] == "completed" and healthy["value"] == 5
    assert [step["nodeId"] for step in healthy["steps"]] == ["Sheet!E1"]
    assert healthy["coverage"]["formulaCells"] == 1
    assert app.store.snapshot("alice", project)["graph"] == original


def test_graph_asset_is_packaged_mounted_and_allowlisted(tmp_path):
    from pathlib import Path

    from linexcel import web

    app = create_app(
        tmp_path, mount_path="/tools/linexcel", resolve_user=lambda env: env.get("user")
    )
    expected = (Path(web.__file__).parent / "assets" / "cytoscape.min.js").read_bytes()
    try:
        for path, script_name in [
            ("/tools/linexcel/assets/cytoscape.min.js", ""),
            ("/assets/cytoscape.min.js", "/tools/linexcel"),
        ]:
            code, body, headers = request(app, "GET", path, SCRIPT_NAME=script_name)
            assert code == 200 and body == expected
            assert headers["Content-Type"].startswith("text/javascript")
            assert int(headers["Content-Length"]) == len(expected)
            code, body, headers = request(app, "HEAD", path, SCRIPT_NAME=script_name)
            assert code == 200 and body == b""
            assert int(headers["Content-Length"]) == len(expected)
        for path in [
            "/tools/linexcel/assets/../web.py",
            "/tools/linexcel/assets/app.html",
            "/tools/linexcel/assets/unknown.js",
            "/assets/cytoscape.min.js",
        ]:
            assert request(app, "GET", path)[0] == 404
        assert (
            request(app, "GET", "/tools/linexcel/assets/cytoscape.min.js", user=None)[0]
            == 401
        )
        assert (
            request(app, "POST", "/tools/linexcel/assets/cytoscape.min.js", {})[0]
            == 404
        )
    finally:
        app.close()


def test_uploaded_reference_change_never_reuses_old_calculation(app):
    main = Workbook()
    main.active["B1"] = "='[reference.xlsx]Sheet'!A1*2"
    stream = io.BytesIO()
    main.save(stream)
    primary = {
        "name": "main.xlsx",
        "data": base64.b64encode(stream.getvalue()).decode(),
    }
    revisions = []
    for value in (4, 9):
        reference = Workbook()
        reference.active["A1"] = value
        stream = io.BytesIO()
        reference.save(stream)
        body = {
            "workbook": primary,
            "references": [
                {
                    "name": "reference.xlsx",
                    "data": base64.b64encode(stream.getvalue()).decode(),
                }
            ],
        }
        imported = app.store.create("owner", body)
        assert wait(app.store, imported["task"]["id"])["status"] == "succeeded"
        project = imported["project"]
        assert project["name"] == "main.xlsx"
        revisions.append(project["revision"])
        task = app.store.submit("owner", project["id"], "evaluate", "Sheet!B1", {})
        completed = wait(app.store, task["id"])
        assert completed["status"] == "succeeded", completed
        assert completed["result"]["value"] == value * 2
        assert completed["patternId"].startswith("formula-sha256:")
    assert revisions[0] != revisions[1]


def test_mount_and_local_cookie(tmp_path):
    app = create_app(tmp_path, mount_path="/tools/linexcel")
    try:
        assert request(app, "GET", "/tools/linexcel")[0] == 303
        assert request(app, "GET", "", SCRIPT_NAME="/tools/linexcel")[0] == 303
        code, page, headers = request(app, "GET", "/tools/linexcel/")
        assert code == 200 and b"Linexcel" in page
        assert "HttpOnly; SameSite=Strict" in headers["Set-Cookie"]
        assert (
            request(app, "GET", "/api/config", SCRIPT_NAME="/tools/linexcel")[0] == 200
        )
        assert request(app, "GET", "/elsewhere/")[0] == 404
        assert (
            request(
                app,
                "POST",
                "/tools/linexcel/api/projects",
                {},
                HTTP_ORIGIN="https://attacker.example",
            )[0]
            == 403
        )
        assert (
            request(
                app,
                "POST",
                "/tools/linexcel/api/projects",
                {},
                CONTENT_TYPE="text/plain",
            )[0]
            == 415
        )
    finally:
        app.close()


def test_invalid_inputs_are_rejected_before_task_submission(app):
    for body in [
        [],
        {"workbook": {"name": "../bad.xlsx", "data": "UEs="}},
        {"workbook": workbook(), "budget": {"seconds": float("nan")}},
        {"workbook": workbook(), "budget": {"memoryMb": 4096}},
        {"workbook": workbook(), "references": [workbook()]},
    ]:
        assert request(app, "POST", "/api/projects", body)[0] == 400
    assert app.store.tasks == {}


def test_timeout_preserves_previous_graph_and_diagnostic(app):
    imported = app.store.create("owner", {"workbook": workbook()})
    project = imported["project"]["id"]
    wait(app.store, imported["task"]["id"])
    before = app.store.snapshot("owner", project)["graph"]
    task = app.store.submit(
        "owner", project, "evaluate", "Sheet!B1", {"seconds": 0.001}
    )
    failed = wait(app.store, task["id"])
    assert failed["status"] == "failed"
    assert failed["error"]["kind"] == "timeout"
    assert app.store.snapshot("owner", project)["graph"] == before


def test_global_memory_admission_and_cancellation(tmp_path, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    store = TaskStore(WebConfig(tmp_path, total_memory_mb=512, max_workers=2))

    def worker(task, root, event):
        entered.set()
        while not release.wait(0.01) and not event.is_set():
            pass
        return {"result": {"nodes": [], "edges": [], "meta": {}}}

    monkeypatch.setattr(store, "_worker", worker)
    try:
        first = store.create("owner", {"workbook": workbook()})["task"]
        assert entered.wait(2)
        second = store.create("owner", {"workbook": workbook()})["task"]
        assert store.tasks[second["id"]]["status"] == "queued"
        assert store.used_memory == 512
        store.cancel("owner", second["id"])
        assert wait(store, second["id"])["status"] == "cancelled"
        release.set()
        assert wait(store, first["id"])["status"] == "succeeded"
    finally:
        release.set()
        store.close()


def test_persistent_results_and_single_coordinator(tmp_path):
    store = TaskStore(WebConfig(tmp_path))
    try:
        with pytest.raises(RuntimeError, match="coordinator"):
            TaskStore(WebConfig(tmp_path))
        created = store.create("owner", {"workbook": workbook()})
        task_id = created["task"]["id"]
        assert wait(store, task_id)["status"] == "succeeded"
    finally:
        store.close()
    restored = TaskStore(WebConfig(tmp_path))
    try:
        assert restored.get_task("owner", task_id)["status"] == "succeeded"
        path = restored.project_path("owner", created["project"]["id"])
        (path / "workbook.xlsx").write_bytes(b"modified")
        with pytest.raises(APIError, match="Sources modifiées"):
            restored.submit("owner", created["project"]["id"], "import", None, {})
    finally:
        restored.close()


def test_failed_result_write_releases_worker_capacity(tmp_path, monkeypatch):
    import linexcel.web as web

    original = web._write

    def fail_final_task(path, value):
        if path.name == "task.json" and value.get("status") == "succeeded":
            raise OSError("Disk full")
        return original(path, value)

    monkeypatch.setattr(web, "_write", fail_final_task)
    store = TaskStore(WebConfig(tmp_path))
    try:
        created = store.create("owner", {"workbook": workbook()})
        task = wait(store, created["task"]["id"])
        assert task["status"] == "failed"
        assert task["error"]["kind"] == "persistence_error"
        assert store.used_memory == 0
        assert store.snapshot("owner", created["project"]["id"])["graph"]
    finally:
        store.close()


def test_terminal_status_waits_for_final_publication(app, monkeypatch):
    import linexcel.web as web

    imported = app.store.create("alice", {"workbook": workbook()})
    assert wait(app.store, imported["task"]["id"])["status"] == "succeeded"
    project = imported["project"]["id"]
    finalizing, release = threading.Event(), threading.Event()
    condition = app.store.capacity

    class PublicationBarrier:
        """Pause the worker at its final publication, with the lock free."""

        entries = 0

        def __enter__(self):
            if threading.current_thread() is not threading.main_thread():
                self.entries += 1
                if self.entries == 2:
                    finalizing.set()
                    assert release.wait(5)
            return condition.__enter__()

        def __exit__(self, *args):
            return condition.__exit__(*args)

        def __getattr__(self, name):
            return getattr(condition, name)

    original_write = web._write

    def fail_terminal_write(path, value):
        if path.name == "task.json" and value.get("status") == "succeeded":
            raise OSError("Disk full at final publication")
        return original_write(path, value)

    monkeypatch.setattr(app.store, "capacity", PublicationBarrier())
    monkeypatch.setattr(web, "_write", fail_terminal_write)
    monkeypatch.setattr(
        app.store,
        "_worker",
        lambda task, root, event: {"result": {"status": "completed", "value": 42}},
    )
    try:
        task = app.store.submit("alice", project, "evaluate", "Sheet!B1", {})
        assert finalizing.wait(2)
        assert app.store.get_task("alice", task["id"])["status"] == "running"
    finally:
        release.set()
    finished = wait(app.store, task["id"])
    assert finished["status"] == "failed"
    assert finished["error"]["kind"] == "persistence_error"
    assert app.store.used_memory == 0


@pytest.mark.skipif(os.name != "nt", reason="Windows legacy MAX_PATH regression")
def test_nested_windows_data_directory(tmp_path):
    store = TaskStore(WebConfig(tmp_path / ("nested-" * 12)))
    try:
        imported = store.create("a" * 64, {"workbook": workbook()})
        task = wait(store, imported["task"]["id"])
        assert task["status"] == "succeeded", task
        assert store.snapshot("a" * 64, imported["project"]["id"])["graph"]
    finally:
        store.close()
