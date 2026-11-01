"""Datasette serves complete source cells behind the existing owner boundary."""

import asyncio
import gc
import hashlib
import html
import io
import json
import re
import shutil
import sqlite3
import threading
import weakref
import zipfile
from pathlib import Path
from urllib.parse import urlsplit
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook

from linexcel.lazy_store import build_index
from linexcel.web import _write, create_app
from linexcel.web_data import MAX_INSTANCES, DataLimitError, build_source_database


def request(app, path, *, user="alice", query="", method="GET"):
    status_headers = []
    body = b"".join(
        app(
            {
                "REQUEST_METHOD": method,
                "PATH_INFO": path,
                "QUERY_STRING": query,
                "CONTENT_TYPE": "application/json",
                "CONTENT_LENGTH": "2",
                "wsgi.input": io.BytesIO(b"{}"),
                "HTTP_HOST": "127.0.0.1:8765",
                "wsgi.url_scheme": "http",
                "user": user,
            },
            lambda status, headers: status_headers.append((status, dict(headers))),
        )
    )
    status, headers = status_headers[0]
    return int(status.split()[0]), headers, body


@pytest.fixture
def data_app(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data & Source"
    sheet["A1"] = "header (not inferred)"
    sheet["B1"] = "=2*3"
    sheet["C1"] = True
    sheet["D1"] = "literal <script>"
    sheet["A301"] = 301
    sheet["B301"] = "=A301*2"
    book.create_sheet("Empty")
    book.create_sheet("Very wide")["AAA9"] = "far right"
    workbook = io.BytesIO()
    book.save(workbook)
    # Preserve one real saved formula cache and one absent cache distinctly.
    original = zipfile.ZipFile(io.BytesIO(workbook.getvalue()))
    revised = io.BytesIO()
    with zipfile.ZipFile(revised, "w") as archive:
        for name in original.namelist():
            value = original.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(value)
                root.find("{*}sheetData/{*}row/{*}c[@r='B1']/{*}v").text = "6"
                value = ET.tostring(root)
            archive.writestr(name, value)
    app = create_app(
        tmp_path / "app", resolve_user=lambda env: env.get("user"), mount_path="/tools"
    )
    owner = hashlib.sha256(b"alice").hexdigest()
    project = "a" * 32
    path = app.store.root / owner / "projects" / project
    (path / "references").mkdir(parents=True)
    (path / "workbook.xlsx").write_bytes(revised.getvalue())
    revision = app.store._revision(path)
    _write(
        path / "project.json",
        {"id": project, "filename": "sources.xlsx", "revision": revision},
    )
    build_index(path / "workbook.xlsx", path / "index.sqlite", filename="sources.xlsx")
    try:
        yield app, project, path
    finally:
        app.close()


def test_metadata_covers_complete_rows_empty_sheets_and_wide_fallback(data_app):
    app, project, path = data_app
    before = hashlib.sha256((path / "workbook.xlsx").read_bytes()).hexdigest()
    status, headers, body = request(app, f"/tools/api/projects/{project}/data")
    metadata = json.loads(body)
    assert status == 200 and headers["Cache-Control"] == "no-store"
    assert metadata["available"] and metadata["sourceSha256"] == before
    assert metadata["url"] == f"/tools/api/projects/{project}/datasette/source"
    assert metadata["sheets"][0] == {
        "sheet": "Data & Source",
        "table": "sheet_001",
        "rowCount": 2,
        "cellCount": 6,
        "columns": ["A", "B", "C", "D"],
        "layout": "grid",
        "url": metadata["url"] + "/sheet_001",
    }
    assert metadata["sheets"][1]["rowCount"] == 0
    assert metadata["sheets"][2]["layout"] == "cells"
    assert metadata["sheets"][2]["url"].endswith("cells?sheet__exact=Very+wide")
    assert hashlib.sha256((path / "workbook.xlsx").read_bytes()).hexdigest() == before
    assert "signature" not in metadata


def test_real_datasette_filters_sort_and_source_formula_cache_proofs(data_app):
    app, project, _ = data_app
    prefix = f"/tools/api/projects/{project}/datasette"
    status, _, body = request(
        app,
        prefix + "/source/sheet_001.json",
        query="_shape=objects&row__gte=200&_sort_desc=row",
    )
    assert status == 200
    assert json.loads(body)["rows"] == [
        {"row": 301, "A": 301, "B": None, "C": None, "D": None}
    ]
    status, _, body = request(
        app, prefix + "/source/cells.json", query="_shape=objects&cell__exact=B1"
    )
    assert status == 200
    cell = json.loads(body)["rows"][0]
    assert cell["value"] == 6 and cell["value_json"] == "6"
    assert cell["value_role"] == "formula_cache" and cell["value_status"] == "present"
    assert cell["source_formula"] == "=2*3" and cell["formula_xml_text"] == "2*3"
    assert cell["worksheet_part"] == "xl/worksheets/sheet1.xml"
    assert len(cell["workbook_sha256"]) == 64
    status, _, body = request(
        app, prefix + "/source/cells.json", query="_shape=objects&cell__exact=C1"
    )
    assert status == 200
    assert json.loads(body)["rows"][0]["value_kind"] == "boolean"
    assert json.loads(body)["rows"][0]["value_json"] == "true"


def test_every_datasette_request_is_owner_scoped_and_readonly(data_app):
    app, project, _ = data_app
    prefix = f"/tools/api/projects/{project}/datasette"
    for route in ("/source", "/source/cells.json", "/-/static/app.css"):
        assert request(app, prefix + route, user="bob")[0] == 404
    assert request(app, f"/tools/api/projects/{project}/data", user="bob")[0] == 404
    assert request(app, prefix + "/source", method="POST")[0] == 405
    assert request(app, prefix + "/source", query="sql=select+*+from+cells")[0] == 403
    for route in (
        "/-/metadata",
        "/_memory",
        "/source.db",
        "/source/../-/versions",
        "/source/%2e%2e/-/versions",
    ):
        assert request(app, prefix + route)[0] == 404


def test_datasette_html_and_static_assets_keep_mount_and_escape_source(data_app):
    app, project, _ = data_app
    prefix = f"/tools/api/projects/{project}/datasette"
    status, headers, body = request(app, prefix + "/source/sheet_001")
    html = body.decode()
    assert status == 200 and headers["Content-Type"].startswith("text/html")
    assert "literal &lt;script&gt;" in html
    assert '<script src="https://' not in html
    assert prefix + "/-/static/app.css" in html
    asset = re.search(r'href="([^\"]*/-/static/app.css[^\"]*)"', html)[1]
    route, _, query = asset.partition("?")
    status, headers, body = request(app, route, query=query)
    assert status == 200 and "text/css" in headers["Content-Type"] and body
    status, headers, body = request(app, prefix + "/source/sheet_001", method="HEAD")
    assert status == 200 and body == b""


def test_native_generated_links_remain_scoped_and_functional(data_app):
    app, project, _ = data_app
    prefix = f"/tools/api/projects/{project}/datasette"
    status, _, body = request(app, prefix + "/source/sheet_001", query="_facet=A")
    assert status == 200
    links = [
        html.unescape(link.decode()) for link in re.findall(rb'href="([^\"]+)"', body)
    ]
    relevant = []
    for link in links:
        parsed = urlsplit(link)
        if parsed.netloc and parsed.netloc != "127.0.0.1:8765":
            assert parsed.netloc != "localhost"
            continue
        if parsed.path.startswith(prefix + "/source"):
            assert parsed.path.count(prefix) == 1
            relevant.append(parsed)
            assert request(app, parsed.path, query=parsed.query)[0] == 200
    assert any(link.path.endswith(".json") for link in relevant)
    assert any(link.path.endswith(".csv") for link in relevant)
    assert any("_sort=" in link.query for link in relevant)
    assert any("A=" in link.query for link in relevant)
    assert any(link.path.endswith("/301") for link in relevant)


def test_modified_source_is_rejected_and_macro_only_is_explicit(data_app):
    app, project, path = data_app
    route = f"/tools/api/projects/{project}/data"
    assert request(app, route)[0] == 200
    metadata = json.loads((path / "project.json").read_text())
    metadata["filename"] = "tools.xla"
    _write(path / "project.json", metadata)
    status, _, body = request(app, route)
    assert status == 200
    assert json.loads(body) == {
        "available": False,
        "reason": "macro_only",
        "sheets": [],
        "revision": metadata["revision"],
    }
    (path / "workbook.xlsx").write_bytes(b"changed")
    assert request(app, route)[0] == 409


def test_failed_storage_limit_does_not_publish_partial_database(data_app):
    _, _, path = data_app
    output = path / "limited.sqlite"
    with pytest.raises(DataLimitError):
        build_source_database(path / "index.sqlite", output, max_bytes=1)
    assert not output.exists()
    assert not list(Path(path).glob("data-*.building"))


def test_rebuilt_source_invalidates_datasette_and_pages_remain_complete(data_app):
    app, project, path = data_app
    prefix = f"/tools/api/projects/{project}"
    # Open an immutable database first: replacement must release its handles on
    # Windows, rather than keep showing the preceding source generation.
    assert request(app, prefix + "/datasette/source/sheet_001.json")[0] == 200
    book = Workbook()
    for row in range(1, 152):
        book.active.cell(row, 1, row)
    book.save(path / "workbook.xlsx")
    metadata = json.loads((path / "project.json").read_text())
    metadata["revision"] = app.store._revision(path)
    _write(path / "project.json", metadata)
    build_index(path / "workbook.xlsx", path / "index.sqlite", filename="sources.xlsx")
    status, _, body = request(app, prefix + "/data")
    assert status == 200
    assert json.loads(body)["sheets"][0]["rowCount"] == 151
    rows = []
    query = "_shape=objects&_size=50"
    while True:
        status, _, body = request(
            app, prefix + "/datasette/source/sheet_001.json", query=query
        )
        assert status == 200
        page = json.loads(body)
        rows.extend(page["rows"])
        if not page["next"]:
            break
        query = "_shape=objects&_size=50&_next=" + page["next"]
    assert rows == [{"row": row, "A": row} for row in range(1, 152)]


def test_deleting_a_project_releases_datasette_database_handles(data_app):
    app, project, path = data_app
    prefix = f"/tools/api/projects/{project}"
    assert request(app, prefix + "/datasette/source/sheet_001")[0] == 200
    assert request(app, prefix, method="DELETE")[0] == 200
    assert not path.exists()
    assert request(app, prefix + "/datasette/-/static/app.css")[0] == 404


def test_blank_dependency_placeholders_are_not_source_cells(data_app):
    app, project, path = data_app
    book = Workbook()
    book.active["A1"] = "=C14"
    book.save(path / "workbook.xlsx")
    metadata = json.loads((path / "project.json").read_text())
    metadata["revision"] = app.store._revision(path)
    _write(path / "project.json", metadata)
    build_index(path / "workbook.xlsx", path / "index.sqlite", filename="sources.xlsx")
    prefix = f"/tools/api/projects/{project}"
    status, _, body = request(app, prefix + "/data")
    assert status == 200
    sheet = json.loads(body)["sheets"][0]
    assert sheet["rowCount"] == 1 and sheet["cellCount"] == 1
    assert sheet["columns"] == ["A"]
    status, _, body = request(
        app, prefix + "/datasette/source/cells.json", query="_shape=objects"
    )
    assert status == 200
    rows = json.loads(body)["rows"]
    assert [row["cell"] for row in rows] == ["A1"]
    assert rows[0]["value_status"] == "missing"
    assert rows[0]["value_role"] == "formula_cache"


def test_timed_out_read_cancels_coroutine_and_next_request_remains_usable(
    data_app, monkeypatch
):
    import linexcel.web_data as data

    app, project, _ = data_app
    route = f"/tools/api/projects/{project}/datasette/source/sheet_001.json"
    assert request(app, route)[0] == 200
    instance = next(iter(app.data_browser._instances.values()))
    started, cancelled = threading.Event(), threading.Event()
    deadlines = []
    original_run = asyncio.run_coroutine_threadsafe

    async def suspended_request(*args, **kwargs):
        started.set()
        try:
            await asyncio.Event().wait()
        finally:
            cancelled.set()

    def short_deadline(coroutine, loop):
        pending = original_run(coroutine, loop)

        class Deadline:
            def result(self, timeout):
                deadlines.append(timeout)
                return pending.result(timeout=0.05)

            def cancel(self):
                return pending.cancel()

        return Deadline()

    with monkeypatch.context() as patch:
        patch.setattr(instance.client, "request", suspended_request)
        patch.setattr(data.asyncio, "run_coroutine_threadsafe", short_deadline)
        status, _, body = request(app, route)
        assert started.is_set()
        assert status == 504
        assert "Délai dépassé" in json.loads(body)["error"]["message"]
        assert deadlines == [30]
        assert cancelled.wait(2), "The actual event-loop coroutine must stop"

    assert request(app, route)[0] == 200
    assert request(app, f"/tools/api/projects/{project}/data")[0] == 200


def test_real_datasette_eviction_bounds_writers_and_releases_connections(data_app):
    app, project, source_path = data_app
    writers, connections, references, memory_names = [], [], [], set()
    for number in range(12):
        current = project if number == 0 else f"{number:032x}"
        if number:
            path = source_path.with_name(current)
            (path / "references").mkdir(parents=True)
            for filename in ("workbook.xlsx", "index.sqlite"):
                shutil.copyfile(source_path / filename, path / filename)
            _write(
                path / "project.json",
                {
                    "id": current,
                    "filename": "sources.xlsx",
                    "revision": app.store._revision(path),
                },
            )
        status, _, body = request(
            app, f"/tools/api/projects/{current}/datasette/source/sheet_001.json"
        )
        assert status == 200 and json.loads(body)["rows"][0][2] == 6
        instance = next(reversed(app.data_browser._instances.values()))
        internal = instance.databases["_internal"]
        memory_names.add(internal.memory_name)
        writers.append(internal._write_thread)
        connections.extend(internal._connections)
        connections.extend(instance.databases["source"]._all_file_connections)
        references.append(weakref.ref(instance))
        assert len(app.data_browser._instances) == min(number + 1, MAX_INSTANCES)
        assert sum(writer.is_alive() for writer in writers) == min(
            number + 1, MAX_INSTANCES
        )
    assert len(memory_names) == 12
    del instance, internal
    app.close()
    app.close()
    assert not any(writer.is_alive() for writer in writers)
    assert not app.data_browser._thread.is_alive()
    assert app.data_browser._loop.is_closed()
    for connection in connections:
        with pytest.raises(sqlite3.ProgrammingError, match="closed"):
            connection.execute("select 1")
    gc.collect()
    assert not any(reference() is not None for reference in references)


def test_internal_writer_drains_fifo_despite_cancelled_caller_and_failed_job(
    data_app, monkeypatch
):
    app, project, _ = data_app
    route = f"/tools/api/projects/{project}/datasette/source/sheet_001.json"
    assert request(app, route)[0] == 200
    instance = next(iter(app.data_browser._instances.values()))
    internal = instance.databases["_internal"]
    loop = app.data_browser._loop
    started, release, closing = threading.Event(), threading.Event(), threading.Event()
    completed, errors = [], []

    def first(connection):
        started.set()
        assert release.wait(5)
        connection.execute("create table lifecycle (value)")
        connection.execute("insert into lifecycle values (1)")
        completed.append("first")

    def failed(connection):
        assert connection.execute("select value from lifecycle").fetchone()[0] == 1
        completed.append("failed")
        raise ValueError("queued failure")

    def last(connection):
        assert connection.execute("select value from lifecycle").fetchone()[0] == 1
        completed.append("last")

    pending = asyncio.run_coroutine_threadsafe(internal.execute_write_fn(first), loop)
    assert started.wait(2)
    pending.cancel()
    failing = asyncio.run_coroutine_threadsafe(internal.execute_write_fn(failed), loop)
    # block=False returns only after admission, so all three jobs precede STOP.
    asyncio.run_coroutine_threadsafe(
        internal.execute_write_fn(last, block=False), loop
    ).result(2)

    original_join = internal._write_thread.join

    def joining(timeout=None):
        closing.set()
        original_join(timeout)

    monkeypatch.setattr(internal._write_thread, "join", joining)

    def close():
        try:
            internal.close()
        except Exception as error:
            errors.append(error)

    closer = threading.Thread(target=close)
    closer.start()
    assert closing.wait(2)
    try:
        # The state lock makes this observation deterministic without a sleep.
        with internal._state_lock:
            assert internal._closing
        rejected = asyncio.run_coroutine_threadsafe(
            internal.execute_write_fn(last), loop
        )
        with pytest.raises(RuntimeError, match="closing"):
            rejected.result(2)
        assert closer.is_alive()
    finally:
        release.set()
        closer.join(5)
    assert not closer.is_alive() and not errors
    with pytest.raises(ValueError, match="queued failure"):
        failing.result(2)
    assert completed == ["first", "failed", "last"]
    assert pending.cancelled()
    assert not internal._write_thread.is_alive()
    internal.close()


def test_slow_writer_shutdown_is_retryable_and_other_resources_are_closed(
    data_app, monkeypatch
):
    app, project, _ = data_app
    route = f"/tools/api/projects/{project}/datasette/source/sheet_001.json"
    assert request(app, route)[0] == 200
    instance = next(iter(app.data_browser._instances.values()))
    internal = instance.databases["_internal"]
    source_connections = list(instance.databases["source"]._all_file_connections)
    started, release = threading.Event(), threading.Event()

    def slow(connection):
        started.set()
        assert release.wait(5)
        assert connection.execute("select 1").fetchone()[0] == 1

    # No awaiting coroutine remains when the browser's event loop is shut down.
    asyncio.run_coroutine_threadsafe(
        internal.execute_write_fn(slow, block=False), app.data_browser._loop
    ).result(2)
    assert started.wait(2)
    try:
        with monkeypatch.context() as patch:
            patch.setattr(internal._write_thread, "join", lambda timeout=None: None)
            with pytest.raises(ExceptionGroup, match="data browser"):
                app.close()
        assert app.store._lock_file.closed
        assert app.data_browser._loop.is_closed()
        assert len(app.data_browser._instances) == 1
        assert internal._write_thread.is_alive()
        for connection in source_connections:
            with pytest.raises(sqlite3.ProgrammingError, match="closed"):
                connection.execute("select 1")
    finally:
        release.set()
        internal._write_thread.join(5)
        app.close()
    assert not app.data_browser._instances
    assert not internal._connections


def test_failed_invalidation_is_retained_then_reopened_before_next_request(
    data_app, monkeypatch
):
    app, project, _ = data_app
    route = f"/tools/api/projects/{project}/datasette/source/sheet_001.json"
    assert request(app, route)[0] == 200
    instance = next(iter(app.data_browser._instances.values()))
    internal = instance.databases["_internal"]
    started, release = threading.Event(), threading.Event()

    def slow(connection):
        started.set()
        assert release.wait(5)
        assert connection.execute("select 1").fetchone()[0] == 1

    asyncio.run_coroutine_threadsafe(
        internal.execute_write_fn(slow, block=False), app.data_browser._loop
    ).result(2)
    assert started.wait(2)
    try:
        with monkeypatch.context() as patch:
            patch.setattr(internal._write_thread, "join", lambda timeout=None: None)
            with pytest.raises(ExceptionGroup, match="Datasette"):
                app.data_browser.invalidate(
                    hashlib.sha256(b"alice").hexdigest(), project
                )
        assert list(app.data_browser._instances.values()) == [instance]
    finally:
        release.set()
        internal._write_thread.join(5)
    assert request(app, route)[0] == 200
    assert next(iter(app.data_browser._instances.values())) is not instance
    assert not internal._connections
