"""The HTTP coordinator serves bounded disk-index views, never a full graph."""

import base64
import io
import json
import time

import pytest
from openpyxl import Workbook

from linexcel.web import create_app


def call(app, method, path, body=None, query="", user="alice"):
    data = json.dumps(body or {}).encode()
    response = []
    result = b"".join(
        app(
            {
                "REQUEST_METHOD": method,
                "PATH_INFO": path,
                "QUERY_STRING": query,
                "CONTENT_TYPE": "application/json",
                "CONTENT_LENGTH": str(len(data)),
                "wsgi.input": io.BytesIO(data),
                "user": user,
            },
            lambda status, headers: response.append(status),
        )
    )
    return int(response[0].split()[0]), json.loads(result)


def finish(store, task_id):
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        task = store.get_task(task_owner(store, task_id), task_id)
        if task["status"] in {"failed", "succeeded", "cancelled"}:
            assert task["status"] == "succeeded", task
            return task
        time.sleep(0.02)
    pytest.fail("worker did not finish")


def task_owner(store, task_id):
    return store.tasks[task_id]["owner"]


@pytest.fixture
def indexed(tmp_path):
    book = Workbook()
    for row in range(1, 451):
        book.active.cell(row, 1, row)
    book.active["B450"] = "=A450*2"
    stream = io.BytesIO()
    book.save(stream)
    app = create_app(tmp_path, resolve_user=lambda env: env.get("user"))
    try:
        code, imported = call(
            app,
            "POST",
            "/api/projects",
            {
                "workbook": {
                    "name": "pages.xlsx",
                    "data": base64.b64encode(stream.getvalue()).decode(),
                }
            },
        )
        assert code == 200
        finish(app.store, imported["task"]["id"])
        yield app, imported["project"]["id"]
    finally:
        app.close()


def test_complete_structure_pages_are_read_only_and_snapshot_scoped(indexed):
    app, project = indexed
    base = f"/api/projects/{project}/structure"
    original_tasks = set(app.store.tasks)
    nodes, snapshot = [], ""
    for offset in range(0, 451, 200):
        code, page = call(
            app, "GET", base, query=f"offset={offset}&snapshot={snapshot}"
        )
        assert code == 200
        snapshot = page["snapshot"]
        nodes.extend(page["nodes"])
        assert len(page["nodes"]) <= 200
        assert page["pagination"]["hasMore"] == (offset < 400)
    assert len({node["id"] for node in nodes}) == 451
    assert all(set(node) == {"id", "kind", "sheet"} for node in nodes)
    code, page = call(app, "GET", base, query=f"part=edges&snapshot={snapshot}")
    assert code == 200
    assert page["edges"] == [
        {"source": "Sheet!A450", "target": "Sheet!B450", "kind": "dep"}
    ]
    assert not page["pagination"]["hasMore"]
    assert set(app.store.tasks) == original_tasks  # No calculation was submitted.
    assert call(app, "GET", base, query="snapshot=stale")[0] == 409
    assert call(app, "GET", base, query="part=invalid")[0] == 400
    assert call(app, "GET", base, user="bob")[0] == 404


def test_pages_search_detail_neighborhood_and_unloaded_evaluation(indexed):
    app, project = indexed
    base = f"/api/projects/{project}"
    code, snapshot = call(app, "GET", base)
    assert code == 200
    graph = snapshot["graph"]
    assert graph["lazy"] and graph["pagination"]["total"] == 451
    assert len(graph["nodes"]) <= 200
    assert all(n["id"] != "Sheet!B450" for n in graph["nodes"])
    code, page = call(app, "GET", base + "/nodes", query="offset=400&limit=100")
    assert code == 200 and len(page["nodes"]) == 51
    code, found = call(app, "GET", base + "/nodes", query="q=B450&limit=20")
    assert code == 200 and [n["id"] for n in found["nodes"]] == ["Sheet!B450"]
    code, node = call(app, "GET", base + "/node", query="nodeId=Sheet%21B450")
    assert code == 200 and node["node"]["formula"] == "=A450*2"
    code, around = call(app, "GET", base + "/neighborhood", query="nodeId=Sheet%21B450")
    assert code == 200
    assert {n["id"] for n in around["nodes"]} == {"Sheet!B450", "Sheet!A450"}
    code, submitted = call(
        app,
        "POST",
        base + "/tasks",
        {
            "operation": "evaluate",
            "nodeId": "Sheet!B450",
        },
    )
    assert code == 200
    assert finish(app.store, submitted["task"]["id"])["result"]["value"] == 900
    owner = task_owner(app.store, submitted["task"]["id"])
    root = app.store.project_path(owner, project)
    assert (root / "index.sqlite").is_file() and not (root / "graph.json").exists()
    assert app.store.tasks[submitted["task"]["id"]]["result"] is None
    _, reconnected = call(app, "GET", base)
    assert all(t["result"] is None for t in reconnected["tasks"])
    assert any(t.get("resultAvailable") for t in reconnected["tasks"])
    code, exported = call(app, "GET", f"/api/tasks/{submitted['task']['id']}/export")
    assert code == 200 and exported["value"] == 900
    assert (
        call(app, "GET", f"/api/tasks/{submitted['task']['id']}/export", user="bob")[0]
        == 404
    )


def test_disk_queries_are_owner_scoped_and_validate_bounds(indexed):
    app, project = indexed
    for endpoint in ("nodes", "node", "neighborhood", "patterns"):
        path = f"/api/projects/{project}/{endpoint}"
        assert call(app, "GET", path, user="bob")[0] == 404
    for query in ("limit=201", "limit=0", "offset=-1", "depth=8", "limit=x", "q=a&q=b"):
        assert call(app, "GET", f"/api/projects/{project}/nodes", query=query)[0] == 400


def test_evidence_counts_describe_full_index_not_sample(indexed):
    from linexcel.lazy_store import evidence_graph
    from linexcel.web_ai import evidence

    app, project = indexed
    owner = next(iter(app.store.tasks.values()))["owner"]
    graph = evidence_graph(app.store.project_path(owner, project) / "index.sqlite")
    dossier = evidence(graph, None, [])
    assert dossier["nodeCounts"] == {"cell": 1, "input": 450}
    assert dossier["omittedNodes"] + len(dossier["sample"]) == 451
    assert dossier["perSheetCounts"]["Sheet"]["input"] == 450


def test_task_history_is_small_and_full_steps_remain_downloadable(indexed):
    app, project = indexed
    original = next(iter(app.store.tasks.values()))
    task = {**original, "operation": "evaluate", "nodeId": "Sheet!B450"}
    task_id = task["id"]
    path = app.store.project_path(task["owner"], project) / "tasks" / task_id
    result = {"value": 900, "steps": [{"nodeId": f"Sheet!A{i}"} for i in range(600)]}
    app.store._save_result(path, task, result)
    app.store.tasks[task_id] = task
    assert task["result"] is None
    code, observed = call(app, "GET", f"/api/tasks/{task_id}")
    assert code == 200
    preview = observed["task"]["result"]
    assert len(preview["steps"]) == 200
    assert preview["stepsPagination"]["total"] == 600
    code, downloaded = call(app, "GET", f"/api/tasks/{task_id}/export")
    assert code == 200 and downloaded == result


def test_index_disk_quota_is_independent_of_compressed_upload_limit(tmp_path):
    book = Workbook(write_only=True)
    sheet = book.create_sheet()
    for row in range(10000):
        sheet.append([row, row + 1, row + 2])
    stream = io.BytesIO()
    book.save(stream)
    assert len(stream.getvalue()) < 1048576
    app = create_app(
        tmp_path,
        resolve_user=lambda env: env.get("user"),
        max_upload_mb=1,
        max_storage_mb=64,
    )
    try:
        code, imported = call(
            app,
            "POST",
            "/api/projects",
            {
                "workbook": {
                    "name": "compressed.xlsx",
                    "data": base64.b64encode(stream.getvalue()).decode(),
                }
            },
        )
        assert code == 200
        task = finish(app.store, imported["task"]["id"])
        root = app.store.project_path(task["owner"], task["projectId"])
        assert (root / "index.sqlite").stat().st_size > 4 * 1048576
    finally:
        app.close()


def test_hierarchy_task_persists_and_import_invalidates_it(indexed):
    app, project = indexed
    base = f"/api/projects/{project}"
    assert call(app, "GET", base + "/graph-view")[0] == 409
    code, built = call(app, "POST", base + "/tasks", {"operation": "build_graph"})
    assert code == 200
    assert built["task"]["budget"]["seconds"] == 600
    finish(app.store, built["task"]["id"])
    count = len(app.store.tasks)
    code, view = call(app, "GET", base + "/graph-view")
    assert code == 200
    assert view["meta"]["nodeCount"] == 451
    assert view["meta"]["edgeCount"] == 1
    assert len(view["nodes"]) <= 33
    code, located = call(app, "GET", base + "/graph-view", query="nodeId=Sheet!B450")
    assert code == 200
    assert any(n.get("nodeId") == "Sheet!B450" for n in located["nodes"])
    assert len(app.store.tasks) == count  # Reading a view never evaluates.
    assert call(app, "GET", base + "/graph-view", user="bob")[0] == 404
    assert call(app, "GET", base + "/graph-view", query="view=bogus")[0] == 400
    assert call(app, "GET", base + "/graph-view", query="view=123456")[0] == 404
    assert call(app, "GET", base + "/graph-view", query="nodeId=missing")[0] == 404
    code, cached = call(app, "POST", base + "/tasks", {"operation": "build_graph"})
    assert code == 200 and cached["task"]["id"] == built["task"]["id"]
    code, imported = call(
        app, "POST", base + "/tasks", {"operation": "import", "force": True}
    )
    assert code == 200 and imported["task"]["budget"]["seconds"] == 600
    finish(app.store, imported["task"]["id"])
    assert call(app, "GET", base + "/graph-view")[0] == 409
    code, rebuilt = call(app, "POST", base + "/tasks", {"operation": "build_graph"})
    assert code == 200 and rebuilt["task"]["id"] != built["task"]["id"]
    finish(app.store, rebuilt["task"]["id"])
    assert call(app, "GET", base + "/graph-view")[0] == 200


def test_hierarchy_and_import_exclude_each_other_and_failure_keeps_old_graph(
    indexed, monkeypatch
):
    import threading

    app, project = indexed
    base = f"/api/projects/{project}"
    _, built = call(app, "POST", base + "/tasks", {"operation": "build_graph"})
    finish(app.store, built["task"]["id"])
    entered, release = threading.Event(), threading.Event()
    original_worker = app.store._worker

    def wait_worker(task, root, event):
        entered.set()
        release.wait(10)
        return {"error": {"kind": "test_failure", "message": "forced failure"}}

    monkeypatch.setattr(app.store, "_worker", wait_worker)
    try:
        _, failed = call(
            app, "POST", base + "/tasks", {"operation": "build_graph", "force": True}
        )
        assert entered.wait(5)
        assert (
            call(app, "POST", base + "/tasks", {"operation": "import", "force": True})[
                0
            ]
            == 409
        )
    finally:
        release.set()
    deadline = time.monotonic() + 10
    while app.store.tasks[failed["task"]["id"]]["status"] not in {
        "failed",
        "succeeded",
    }:
        assert time.monotonic() < deadline
        time.sleep(0.02)
    assert app.store.tasks[failed["task"]["id"]]["status"] == "failed"
    assert call(app, "GET", base + "/graph-view")[0] == 200
    monkeypatch.setattr(app.store, "_worker", original_worker)


def test_exact_node_search_does_not_include_prefix_matches(indexed):
    app, project = indexed
    code, found = call(app, "GET", f"/api/projects/{project}/nodes", query="q=Sheet!A4")
    assert code == 200
    assert found["pagination"]["total"] == 1
    assert [n["id"] for n in found["nodes"]] == ["Sheet!A4"]
